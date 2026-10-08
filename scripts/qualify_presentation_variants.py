"""Opt-in real Brain/Core test. Isolated URI accounts/state; no Claude access."""
import argparse
import io
import json
import os
from pathlib import Path
import secrets
import shutil
import sys
import tempfile
import time
from contextlib import contextmanager

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


@contextmanager
def isolated_state():
    previous = Path.cwd()
    with tempfile.TemporaryDirectory(prefix='uri-presentation-qualification-') as tmp:
        try:
            yield Path(tmp)
        finally:
            # Windows cannot remove the current working directory.
            os.chdir(previous)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--real-brain', action='store_true', required=True)
    args = parser.parse_args()
    evidence = {'scope': 'real Core and local Ollama Brain; native UI not tested by this script', 'turns': []}
    artifact_dir = ROOT / 'scratch/document_layout_qualification' / time.strftime('%Y%m%d-%H%M%S')
    artifact_dir.mkdir(parents=True, exist_ok=False)
    with isolated_state() as state:
        (state / 'uri_workspace').mkdir()
        # Public capability/policy data only. No accounts, sessions or keys copied.
        for name in ('capabilities_registry.json', 'policy.txt', 'soul.md'):
            source = ROOT / 'uri_workspace' / name
            if source.is_file(): shutil.copyfile(source, state / 'uri_workspace' / name)
        roles = {role: {'provider': 'ollama', 'model': 'gemma4:12b', 'num_ctx': 16384}
                 for role in ('reasoning','drafting','document_composition','semantic_interpretation')}
        (state / 'uri_workspace/model_roles.json').write_text(json.dumps(roles), encoding='utf-8')
        os.chdir(state)
        os.environ['URI_EXTERNAL_CREDENTIAL_SECRET'] = secrets.token_urlsafe(48)
        os.environ['URI_PROVIDER_KEY_SECRET'] = secrets.token_urlsafe(48)
        os.environ['OLLAMA_MODEL'] = 'gemma4:12b'
        os.environ['OLLAMA_NUM_CTX'] = '16384'
        os.environ['URI_INT_ARN_1_ENABLED'] = '0'
        os.environ['GRAPHIFY_HINT_ENABLED'] = '0'
        from fastapi.testclient import TestClient
        from pptx import Presentation
        from uri_core.app.server import app
        session = 'presentation-qualification-' + secrets.token_hex(8)
        with TestClient(app) as client:
            signup = client.post('/auth/signup', json={'username': 'presentation_qualification', 'password': secrets.token_urlsafe(24)})
            signup.raise_for_status()
            headers = {'Authorization': 'Bearer ' + signup.json()['token']}
            requests = [
                'For a disposable design test, our presentation subject is Student Welfare at NIT Sikkim. '
                'The supplied topics are student information, wellbeing, and campus activities; treat these as '
                'suggested topics, not verified institutional services. I want academic, modern and warm visual '
                'drafts with licensed illustrative photos. Just acknowledge this brief now; do not create anything yet.',
                'Generate all three design templates so I can have a visual idea of the presentation. '
                'Create three separate PPTX files, each with exactly two subject-content slides. '
                'Combine student information and wellbeing on the first slide and campus activities on the second; '
                'keep the original subject and photo request.',
                "I don't see three PPTs. Can you check again? Check the local outputs from this conversation.",
            ]
            for request in requests:
                before_files = client.get('/files', headers=headers, params={'session_id': session}).json().get('files', [])
                started = time.monotonic()
                reply = client.post('/ask', headers=headers, json={'text': request, 'session_id': session,
                    'model_override': {'provider_id': 'ollama', 'model': 'gemma4:12b'}})
                duration = round((time.monotonic() - started) * 1000, 1)
                reply.raise_for_status(); data = reply.json()
                result = data.get('result', data)
                response = result.get('response') or {}; execution = result.get('execution') or {}
                row = {'status': result.get('status'), 'duration_ms': duration,
                       'tier': result.get('tier'),
                       'execution_status': execution.get('status'),
                       'branches': [{'capability': b.get('capability'), 'status': b.get('status')}
                                    for b in execution.get('branch_results', [])],
                       'capability': execution.get('capability') or execution.get('tool'),
                       'requested_count': response.get('requested_count'), 'created_count': response.get('created_count'),
                       'images_requested': response.get('images_requested'), 'images_added': response.get('images_added'),
                       'image_note': response.get('image_note'),
                       'reported_slide_count': response.get('slide_count'),
                       'quality_review': response.get('quality_review'),
                       'narrative': result.get('narrative') or response.get('message')}
                after_files = client.get('/files', headers=headers, params={'session_id': session}).json().get('files', [])
                row['files_before'] = len(before_files)
                row['files_after'] = len(after_files)
                row['returned_file_count'] = len(response.get('files') or ([response['file']] if response.get('file') else []))
                evidence['turns'].append(row)
                print(json.dumps(row), flush=True)
            refs = client.get('/files', headers=headers, params={'session_id': session}).json().get('files', [])
            evidence['files'] = []
            for ref in refs:
                if not ref.get('filename', '').endswith('.pptx'): continue
                content = client.get('/files/' + ref['file_id'] + '/content', headers=headers)
                content.raise_for_status()
                from uri_core.services.office_document_writer import inspect_pptx_layout
                layout = inspect_pptx_layout(content.content)
                (artifact_dir / Path(ref['filename']).name).write_bytes(content.content)
                deck = Presentation(io.BytesIO(content.content))
                titles = [next((shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()), '')
                          for slide in deck.slides]
                evidence['files'].append({'filename': ref['filename'], 'size_bytes': len(content.content),
                    'layout': {key: value for key, value in layout.items() if key != 'extracted_text'},
                    'slide_count': len(deck.slides),
                    'content_slide_count': len([title for title in titles if title != 'Image credits']),
                    'slide_titles': titles,
                    'picture_count': sum(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes),
                    'first_slide_text': [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame],
                    'theme_color': str(deck.slides[0].background.fill.fore_color.rgb)})
            evidence['passed'] = (len(evidence['files']) == 3
                and len({row['theme_color'] for row in evidence['files']}) == 3
                and all('student welfare' in ' '.join(row['first_slide_text']).lower() for row in evidence['files'])
                and all(row['content_slide_count'] == 2 for row in evidence['files'])
                and all(row['layout']['status'] == 'estimated_pass' for row in evidence['files'])
                and (evidence['turns'][1].get('quality_review') or {}).get('status') == 'reviewed_estimated'
                and evidence['turns'][1]['returned_file_count'] == 3
                and evidence['turns'][-1]['files_before'] == evidence['turns'][-1]['files_after'] == 3
                and evidence['turns'][-1]['capability'] != 'drive_search')
            trace_path = state / 'uri_workspace/canonical_execution_log.jsonl'
            if trace_path.is_file():
                evidence['dispatch_trace'] = [json.loads(line) for line in trace_path.read_text(encoding='utf-8').splitlines()]
        (ROOT / 'docs/plans/DOCUMENT_LAYOUT_REVIEW_LIVE_CORE_RESULT.json').write_text(json.dumps(evidence, indent=2), encoding='utf-8')
        (artifact_dir / 'result.json').write_text(json.dumps(evidence, indent=2), encoding='utf-8')
        print('LIVE_CORE_QUALIFICATION=' + str(evidence['passed']), flush=True)
        if not evidence['passed']: sys.exit(1)


if __name__ == '__main__':
    main()
