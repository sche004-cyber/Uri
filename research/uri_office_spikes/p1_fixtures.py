"""Spike P1 fixtures: workbooks with chart, image, shared/inline strings, formulas, styles, merges."""
import io
import sys
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill


def tiny_png() -> bytes:
    import struct
    import zlib

    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 4 for _ in range(4))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def build_rich(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount", "Note"])
    for r, (reg, amt) in enumerate([("North", 10), ("South", 20), ("East", 30)], start=2):
        ws.append([reg, amt, f"row {r}"])
    ws["B5"] = "=SUM(B2:B4)"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.merge_cells("D1:E1")
    ws["D1"] = "Merged header"
    ws.column_dimensions["A"].width = 22
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "G2")
    img = Image(io.BytesIO(tiny_png()))
    ws.add_image(img, "G20")
    other = wb.create_sheet("Other")
    other["A1"] = "untouched sheet"
    other["A2"] = "=Sales!B5*2"
    wb.save(path)


def add_inline_string(path: Path):
    """openpyxl writes shared strings only; inject one inlineStr cell (C6) to cover that branch."""
    src = zipfile.ZipFile(path)
    parts = {n: src.read(n) for n in src.namelist()}
    infos = {n: src.getinfo(n) for n in src.namelist()}
    src.close()
    name = "xl/worksheets/sheet1.xml"
    xml = parts[name].decode()
    cell = '<c r="C6" t="inlineStr"><is><t>inline text</t></is></c>'
    if '<row r="6"' in xml:
        raise SystemExit("unexpected row 6")
    xml = xml.replace("</sheetData>", f'<row r="6">{cell}</row></sheetData>')
    parts[name] = xml.encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in parts.items():
            z.writestr(infos[n], data)


if __name__ == "__main__":
    out = Path(sys.argv[1])
    out.mkdir(parents=True, exist_ok=True)
    p = out / "rich.xlsx"
    build_rich(p)
    add_inline_string(p)
    print(p)
