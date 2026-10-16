"""Spike P1 runner: patch fixtures and prove only allowed parts changed (unpacked part hashes)."""
import hashlib
import json
import shutil
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))
import p1_fixtures  # noqa: E402
import xlsx_patch_spike as xp  # noqa: E402


def hashes(p: Path) -> dict:
    with zipfile.ZipFile(p) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


def inject_calc_chain(p: Path):
    z = zipfile.ZipFile(p)
    infos = z.infolist()
    parts = {i.filename: z.read(i.filename) for i in infos}
    z.close()
    parts["xl/calcChain.xml"] = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                                 b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                                 b'<c r="B5" i="1"/><c r="A2" i="2"/></calcChain>')
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>", b'<Override PartName="/xl/calcChain.xml" ContentType="application/'
        b'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
    parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
        b"</Relationships>", b'<Relationship Id="rIdCalc" Type="http://schemas.openxmlformats.org/'
        b'officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as out:
        for i in infos:
            out.writestr(i, parts[i.filename])
        out.writestr("xl/calcChain.xml", parts["xl/calcChain.xml"])


def strip_full_calc(p: Path):
    """Excel-saved files carry calcPr without fullCalcOnLoad; reproduce that."""
    z = zipfile.ZipFile(p)
    infos = z.infolist()
    parts = {i.filename: z.read(i.filename) for i in infos}
    z.close()
    wbx = parts["xl/workbook.xml"]
    assert b'fullCalcOnLoad="1"' in wbx, "fixture changed"
    parts["xl/workbook.xml"] = wbx.replace(b' fullCalcOnLoad="1"', b"")
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as out:
        for i in infos:
            out.writestr(i, parts[i.filename])


CASES = {
    "excel_like_calcpr": {
        "excel_like_calcpr": True,
        "edits": [{"sheet": "Sales", "cell": "B2", "value": 11}],
        "allowed": {"xl/worksheets/sheet1.xml", "xl/workbook.xml"},
        "must_change": {"xl/workbook.xml"},
        "expect": {"B2": 11},
    },
    "values_only": {
        "edits": [{"sheet": "Sales", "cell": "B3", "value": 25},
                  {"sheet": "Sales", "cell": "A4", "value": "West"},
                  {"sheet": "Sales", "cell": "C6", "value": "inline edited"},
                  {"sheet": "Sales", "cell": "F40", "value": True}],
        "allowed": {"xl/worksheets/sheet1.xml", "xl/workbook.xml"},
        "expect": {"B3": 25, "A4": "West", "C6": "inline edited", "F40": True, "B5": "=SUM(B2:B4)"},
    },
    "formula_with_calcchain": {
        "calc_chain": True,
        "edits": [{"sheet": "Sales", "cell": "B6", "formula": "=B5*2"},
                  {"sheet": "Sales", "cell": "B5", "value": 99}],
        "allowed": {"xl/worksheets/sheet1.xml", "xl/workbook.xml", "xl/calcChain.xml",
                    "[Content_Types].xml", "xl/_rels/workbook.xml.rels"},
        "expect": {"B6": "=B5*2", "B5": 99},
    },
}


def run(out: Path) -> dict:
    out.mkdir(parents=True, exist_ok=True)
    report = {"spike": "P1", "cases": {}}
    base = out / "rich.xlsx"
    p1_fixtures.build_rich(base)
    p1_fixtures.add_inline_string(base)
    for name, case in CASES.items():
        src = out / f"{name}_src.xlsx"
        shutil.copy(base, src)
        if case.get("calc_chain"):
            inject_calc_chain(src)
        if case.get("excel_like_calcpr"):
            strip_full_calc(src)
        dst = out / f"{name}_patched.xlsx"
        before = hashes(src)
        res = xp.patch(src, dst, case["edits"])
        after = hashes(dst)
        changed = {n for n in before if after.get(n) != before[n]} | (set(after) - set(before))
        unexpected = sorted(changed - case["allowed"])
        missing = sorted(case.get("must_change", set()) - changed)
        full_calc = b'fullCalcOnLoad="1"' in zipfile.ZipFile(dst).read("xl/workbook.xml")
        wb = load_workbook(dst)
        ws = wb["Sales"]
        values_ok = {c: ws[c].value == v for c, v in case["expect"].items()}
        other_ok = wb["Other"]["A1"].value == "untouched sheet" and wb["Other"]["A2"].value == "=Sales!B5*2"
        untouched_media = {n: before[n] == after.get(n) for n in before
                           if n.startswith(("xl/charts/", "xl/media/", "xl/drawings/"))}
        report["cases"][name] = {
            "changed_parts": sorted(changed), "patcher_reported": res["changed_parts"],
            "unexpected_changes": unexpected, "values_ok": values_ok, "other_sheet_ok": other_ok,
            "chart_image_parts_identical": untouched_media,
            "calc_chain_removed": ("xl/calcChain.xml" not in after) if case.get("calc_chain") else None,
            "missing_required_changes": missing, "full_calc_on_load": full_calc,
            "pass": (not unexpected and not missing and full_calc and all(values_ok.values()) and other_ok and untouched_media
                     and all(untouched_media.values())),
        }
    # refusal checks
    refusals = {}
    for label, edit in {"merged_non_anchor": {"sheet": "Sales", "cell": "E1", "value": "x"},
                        "unknown_sheet": {"sheet": "Nope", "cell": "A1", "value": 1},
                        "bad_ref": {"sheet": "Sales", "cell": "1A", "value": 1}}.items():
        try:
            xp.patch(base, out / "refuse.xlsx", [edit])
            refusals[label] = "NOT_REFUSED"
        except xp.PatchError as e:
            refusals[label] = f"refused: {e}"
    report["refusals"] = refusals
    report["result"] = ("COMPLETED_WITH_RESULT"
                        if all(c["pass"] for c in report["cases"].values())
                        and all(v.startswith("refused") for v in refusals.values()) else "FAILED")
    return report


if __name__ == "__main__":
    rep = run(Path(sys.argv[1]))
    Path(sys.argv[2]).write_text(json.dumps(rep, indent=2))
    print(json.dumps({"result": rep["result"], **{k: v["pass"] for k, v in rep["cases"].items()},
                      "refusals": rep["refusals"]}, indent=1))
