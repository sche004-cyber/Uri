"""Production generator/dispatch/context regression, with no model or network."""
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from pptx import Presentation
from uri_core.core.file_store import FileStore, FileValidationError
from uri_core.tools.generate_document import GenerateDocumentTool
from uri_core.core.canonical_execution import _execute_legacy_capability
from uri_core.core.decision_engine import build_turn_state_and_directory, build_decision_request

VARIANTS = [
    {"label": "Institutional", "theme": "academic"},
    {"label": "Dynamic", "theme": "modern"},
    {"label": "Community", "theme": "warm"},
]


class Composer:
    calls = 0
    def compose(self, **kwargs):
        self.calls += 1
        self.request = kwargs["request_text"]
        return {"body": "# Student Welfare at NIT Sikkim\n- Student support\n# Campus life\n- Counselling and wellbeing",
                "status": "success", "composed_by": "mock", "detail": None}


def generate(tmp_path, **kwargs):
    composer = Composer()
    store = FileStore(str(tmp_path))
    result = GenerateDocumentTool(composer, store).generate(
        request_text="Generate all the design drafts", output_format="pptx",
        title="Student Welfare at NIT Sikkim", session_id="mine", **kwargs)
    return result, composer, store


def test_real_rendered_variants_have_distinct_designs_and_subject(tmp_path):
    result, composer, store = generate(tmp_path, variants=VARIANTS)
    assert composer.calls == 1
    assert "actual presentation slides" in composer.request
    assert result["created_count"] == result["requested_count"] == 3
    assert result["complete"] is True
    records = store.list_for_session("mine")
    assert len(records) == 3
    assert len({record.file_id for record in records}) == 3
    assert len({record.filename for record in records}) == 3
    colors = set()
    for record in records:
        deck = Presentation(tmp_path / record.stored_name)
        assert len(deck.slides) == 2
        assert "Student Welfare at NIT Sikkim" in " ".join(s.text for s in deck.slides[0].shapes if s.has_text_frame)
        colors.add(str(deck.slides[0].background.fill.fore_color.rgb))
    assert len(colors) == 3
    assert store.list_for_session("another") == []


def test_single_file_contract_unchanged(tmp_path):
    result, composer, store = generate(tmp_path)
    assert composer.calls == 1 and len(store.list_for_session("mine")) == 1
    assert "file" in result and "files" not in result


@pytest.mark.parametrize("variants", [
    [], VARIANTS * 2, "three", True,
    [{"label": "Test", "theme": "unknown"}],
    [{"label": "Test", "theme": "modern", "principal": "admin"}],
    [{"label": "Test", "theme": "modern", "session_id": "other"}],
    [{"label": "", "theme": "modern"}],
    [{"label": "x" * 81, "theme": "modern"}],
    [{"label": "A", "theme": "modern"}, {"label": "B", "theme": "modern"}],
    [{"label": "dir/A", "theme": "modern"}, {"label": "A", "theme": "warm"}],
])
def test_invalid_variants_fail_before_model_or_storage(tmp_path, variants):
    result, composer, store = generate(tmp_path, variants=variants)
    assert result["status"] == "error"
    assert composer.calls == 0 and store.list_for_session("mine") == []


def test_docx_never_accepts_presentation_variants(tmp_path):
    composer = Composer()
    result = GenerateDocumentTool(composer, FileStore(str(tmp_path))).generate(
        request_text="Report", output_format="docx", variants=VARIANTS)
    assert result["status"] == "error" and composer.calls == 0


def test_partial_storage_is_not_reported_complete(tmp_path):
    class FailingStore(FileStore):
        def save(self, **kwargs):
            if len(self.list_for_session("mine")) == 2:
                raise FileValidationError("test storage failure")
            return super().save(**kwargs)
    store = FailingStore(str(tmp_path))
    result = GenerateDocumentTool(Composer(), store).generate(
        request_text="Actual topic", output_format="pptx", variants=VARIANTS, session_id="mine")
    assert result["status"] == "error" and result["complete"] is False
    assert result["requested_count"] == 3 and result["created_count"] == len(result["files"]) == 2


def test_images_are_retrieved_once_for_the_variant_set(tmp_path):
    from PIL import Image
    data = io.BytesIO(); Image.new("RGB", (30, 30), "red").save(data, format="PNG")
    image = {"data": data.getvalue(), "credit": "Test, CC0, https://example.org/photo"}
    with patch("uri_core.services.office_document_writer.find_openverse_images", return_value=[image]) as lookup:
        result, _, store = generate(tmp_path, variants=VARIANTS, include_images=True, image_queries=["students campus"])
    assert lookup.call_count == 1
    assert result["images_added"] == 1 and result["created_count"] == 3
    for record in store.list_for_session("mine"):
        deck = Presentation(tmp_path / record.stored_name)
        assert result["slide_count"] == len(deck.slides)
        assert any(shape.shape_type == 13 for slide in deck.slides for shape in slide.shapes)


def test_image_failure_is_honest_for_all_variants(tmp_path):
    with patch("uri_core.services.office_document_writer.find_openverse_images", return_value=[]):
        result, _, _ = generate(tmp_path, variants=VARIANTS, include_images=True, image_queries=["campus"])
    assert result["created_count"] == 3 and result["images_added"] == 0
    assert "without photos" in result["image_note"]


def test_retrieved_but_unembeddable_image_is_not_counted(tmp_path):
    with patch("uri_core.services.office_document_writer.find_openverse_images", return_value=[{}]):
        result, _, _ = generate(tmp_path, variants=VARIANTS, include_images=True, image_queries=["campus"])
    assert result["images_added"] == 0
    assert "without photos" in result["image_note"]


def test_canonical_gate_receives_variants_and_runtime_identity():
    calls = []
    brain = SimpleNamespace(approval_gate=SimpleNamespace(execute_tool=lambda tool, **kw:
        calls.append((tool, kw)) or {"status": "success", "data": {"files": []}}))
    principal = SimpleNamespace(user_id="real-user")
    contract = {"capability": "generate_document", "actions": [{"name": "generate_document", "inputs": {
        "variants": VARIANTS, "output_format": "pptx", "title": "Original subject",
        "session_id": "forged", "principal": "admin", "request_text": "forged"}}]}
    _execute_legacy_capability(contract, orchestrator=brain, session_id="real-session",
        user_text="real request", principal=principal)
    assert calls[0][1]["variants"] == VARIANTS
    assert calls[0][1]["session_id"] == "real-session"
    assert calls[0][1]["principal"] is principal
    assert calls[0][1]["request_text"] == "real request"


def test_metadata_projection_uses_only_this_user_session_and_is_bounded(tmp_path):
    store = FileStore(str(tmp_path / "user-one"))
    for n in range(12):
        store.save(filename=f"draft-{n}.txt", content=b"data", session_id="mine")
    store.save(filename="another-session.txt", content=b"private", session_id="other")
    other_user = FileStore(str(tmp_path / "user-two"))
    other_user.save(filename="other-user.txt", content=b"private", session_id="mine")
    state = {"turn": {"user_text": "Check the local drafts"}, "recent_conversation": [],
             "active_pointer": {}, "capability_summaries": [], "session_facts": [], "attempt_history": []}
    with patch("uri_core.core.decision_engine.CapabilityDirectory"), \
         patch("uri_core.core.decision_engine.assemble_turn_state", return_value=SimpleNamespace(data=state)):
        result, _ = build_turn_state_and_directory(orchestrator=SimpleNamespace(file_store=store, capability_feasibility=object()),
            session_id="mine", user_text="check", principal=SimpleNamespace(user_id="one"))
    metadata = result.data["session_files"]
    assert metadata["listed_count"] == 12 and metadata["truncated"]
    assert len(metadata["files"]) == 10
    assert all(set(ref) == {"file_id", "filename", "media_type", "size_bytes"} for ref in metadata["files"])
    assert not any("other" in ref["filename"] for ref in metadata["files"])
    request = build_decision_request(state)
    assert "session_files" in request


def test_native_brain_receives_dialogue_and_local_file_metadata():
    from uri_core.core.native_tool_loop import run_native_tool_loop
    data = {"recent_conversation": [{"user": "Original presentation topic", "uri": "One draft created"}],
            "session_files": {"listed_count": 1, "files": [{"file_id": "real", "filename": "draft.pptx"}], "truncated": False}}
    prompts = []
    def model(**kw):
        prompts.append(kw["user"])
        return SimpleNamespace(content="Only one local file is listed.", tool_calls=())
    with patch("uri_core.core.native_tool_loop.build_turn_state_and_directory", return_value=(SimpleNamespace(data=data), object())), \
         patch("uri_core.core.native_tool_loop.build_tool_schemas", return_value=[]):
        result = run_native_tool_loop(orchestrator=SimpleNamespace(multi_action_dispatch=None), session_id="mine",
            user_text="Check again", principal=None, model_callable=model)
    assert result["status"] == "success"
    assert "Original presentation topic" in prompts[0]
    assert '"listed_count": 1' in prompts[0]
    assert "not content or authority" in prompts[0]


def test_native_reply_preserves_only_runtime_file_references():
    from uri_core.core.native_tool_loop import _terminal_envelope
    refs = [{"file_id": str(n), "filename": f"draft-{n}.pptx"} for n in range(3)]
    result = _terminal_envelope(content="Three drafts created.", tier="tier1_continuation", session_id="mine",
        branch_results=[{"response": {"files": refs, "file": refs[0]}},
                        {"response": {"file": refs[1]}}, {"response": None}])
    assert result["response"]["files"] == refs
    assert result["response"]["file"] == refs[0]
    assert "files" not in _terminal_envelope(content="I invented fake.pptx", tier="tier0",
        session_id="mine", branch_results=[])['response']


def test_canonical_brain_discovers_actual_generation_content_schema():
    state = {"turn": {"user_text": "Render designs"}, "recent_conversation": [],
             "active_pointer": {}, "capability_summaries": [
                 {"capability_id": "generate_document", "source": "legacy"}],
             "session_facts": [], "attempt_history": []}
    request = build_decision_request(state, preselected_ids=["generate_document"], capability_directory=SimpleNamespace())
    assert list(request['capability_action_affordances']['generate_document']) == ['generate_document']
    fields = request['capability_input_schemas']['generate_document']['generate_document']['properties']
    assert fields['variants']['maxItems'] == 3
    assert fields['include_images']['type'] == 'boolean'
    assert not {'principal', 'session_id', 'request_text', 'file_id'} & fields.keys()
    assert 'capability_input_schemas' not in build_decision_request(state)


def test_real_composer_presentation_contract_does_not_inherit_note_template():
    from uri_core.services.document_composer import DocumentComposer
    composer = DocumentComposer()
    brief = {'output_format': 'pptx', 'suggested_sections': ['OFFICE HEADER'],
             'institutional_rules': {'header': 'OFFICE HEADER'}, 'policy_rules': '',
             'document_type': 'presentation'}
    slides = composer._build_system_prompt(brief, 'URI identity rules')
    note = composer._build_system_prompt({**brief, 'output_format': 'docx'}, 'URI identity rules')
    assert 'URI identity rules' in slides
    assert "'# Slide Title'" in slides and 'exact content-slide count' in slides
    assert 'OFFICE HEADER' not in slides
    assert 'Institutional conventions' in note and 'OFFICE HEADER' in note


def test_narrative_cannot_deny_actually_embedded_photos():
    from uri_core.core.response_validation import validate_drafted_response, ResponseValidationError
    outcome = {'execution': {'status': 'success'}, 'response': {'status': 'success', 'images_added': 2}}
    with pytest.raises(ResponseValidationError, match='actually embedded'):
        validate_drafted_response('The files were created without images.', outcome)
    assert validate_drafted_response('The files contain photos.', outcome)
    assert validate_drafted_response('The files were created without photos.',
        {**outcome, 'response': {'status': 'success', 'images_added': 0}})
