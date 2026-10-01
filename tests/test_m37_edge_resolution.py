"""M37 Edge: inspection, diagnosis, repair, resolution, local completion and boundaries."""
import ast
import hashlib
import json
from pathlib import Path
import subprocess
import textwrap

import pytest

from tests.m37_edge_support import SALES, Env, acceptance, make_workbook, script
from tests.test_m36_memory_support import bind, intake as m36_intake, opened, setup as m36_setup
from uri_v1.edge.compat import diagnose, normalize
from uri_v1.edge.contracts import DecisionKind, Reason
from uri_v1.edge.registry import CAPABILITIES, CapabilityRegistry, QualificationRecord
from uri_v1.edge.repair import RepairRefused, apply_remap
from uri_v1.edge.resolver import resolve
from uri_v1.edge.scriptinfo import inspect_script
from uri_v1.execution.officefile import read_workbook_structure
from uri_v1.memory.context_package import build_context
from uri_v1.memory.contracts import Authority, RecordKind, VerificationStatus

ROOT = Path(__file__).resolve().parents[1]


def test_script_inspection_patterns():
    info = inspect_script(textwrap.dedent('''
        import pandas as pd
        from openpyxl import load_workbook
        df = pd.read_excel("in.xlsx", sheet_name="Data")
        x = df["Price"]; y = df[["A", "B"]]; df["New"] = 1
        wb = load_workbook("in.xlsx"); ws = wb["Other"]; v = ws["A1"].value
        header = [c.value for c in ws[1]]
        i = header.index("Qty")
        df.to_excel("out.xlsx")
    ''').encode())
    d = info.to_dict()
    assert d["sheets"] == ["Data", "Other"] and d["columns"] == ["A", "B", "Price", "Qty"]
    assert d["workbook_paths"] == ["in.xlsx"] and d["output_paths"] == ["out.xlsx"] and d["dynamic"] == []
    # F-1: `.index(literal)` on a receiver that does NOT trace to a header-row derivation
    # is never a recognized column reference, however header-like the literal looks.
    not_header = inspect_script(b'modes = ["amount", "Amount"]\nmode = modes.index("amount")\n')
    assert not_header.values("column") == [] and "DYNAMIC_COLUMN_ACCESS" in not_header.dynamic
    dyn = inspect_script(b'from openpyxl import load_workbook\nwb = load_workbook("a.xlsx")\nname = "S" + "1"\nws = wb[name]\nh.index(f"{c}")\n')
    assert set(dyn.dynamic) == {"DYNAMIC_SHEET_ACCESS", "DYNAMIC_COLUMN_ACCESS"}
    assert inspect_script(b"def (:\n").dynamic == ("PARSE_FAILED",)


def test_normalized_unique_remap_only():
    assert normalize("  Unit_Price ") == normalize("unit-price") == normalize("UNIT  PRICE")
    assert normalize("Amount") != normalize("Amounts")


def _structure(tmp_path, sheets):
    return read_workbook_structure(make_workbook(tmp_path / "sales.xlsx", sheets))


def test_diagnosis_and_exact_literal_repair(tmp_path):
    structure = _structure(tmp_path, SALES)
    src = script(sheet="sales", column="unit_price").encode() + b'print("unit_price is not touched here")\n'
    diag = diagnose(inspect_script(src), structure, "sales.xlsx")
    assert diag["status"] == "REMAPPABLE"
    fixed = apply_remap(src, diag, structure, "sales.xlsx")
    assert b'wb["Sales"]' in fixed and b'header.index("Unit Price")' in fixed
    assert b'print("unit_price is not touched here")' in fixed
    # Every byte outside the two rewritten literals is preserved.
    assert fixed.replace(b'"Sales"', b'"sales"').replace(b'"Unit Price"', b'"unit_price"') == src


def test_ambiguous_absent_dynamic_and_path_mismatch_are_unresolved(tmp_path):
    amb = _structure(tmp_path, {"Sales": (("Amount", "amount "), ())})
    assert diagnose(inspect_script(script(column="AMOUNT").encode()), amb, "sales.xlsx")["missing_columns"][0]["status"] == "AMBIGUOUS"
    s = _structure(tmp_path, SALES)
    for src in (script(column="Amt"), script(workbook="other.xlsx"),
                script(extra='col2 = header.index(sheet_name)\n')):
        d = diagnose(inspect_script(src.encode()), s, "sales.xlsx")
        assert d["status"] == "UNRESOLVED"
        with pytest.raises(RepairRefused): apply_remap(src.encode(), d, s, "sales.xlsx")


def test_repair_refuses_non_simple_literals(tmp_path):
    s = _structure(tmp_path, SALES)
    src = script().replace('header.index("Amount")', "header.index('''amount''')").encode()
    d = diagnose(inspect_script(src), s, "sales.xlsx")
    assert d["status"] == "REMAPPABLE"
    with pytest.raises(RepairRefused, match="LITERAL_NOT_REWRITABLE"): apply_remap(src, d, s, "sales.xlsx")


def _setup(env, column="amount", sheets=SALES):
    make_workbook(env.root / "sales.xlsx", sheets); env.write("report.py", script(column=column))


def test_local_repair_is_claimed_then_verified_by_uri(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch); _setup(env)
    original = (env.root / "report.py").read_bytes()
    result = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert result.route == "LOCAL_VERIFIED" and result.decision.reason is Reason.QUALIFIED_LOCAL
    index = env.log.load()
    claim, verdict = index.records[result.local.claim_record_id], index.records[result.local.verifier_record_id]
    assert claim.authority is Authority.CLAIMED_OUTCOME and claim.payload["verification_status"] == VerificationStatus.CLAIMED_ONLY
    assert verdict.authority is Authority.VERIFIED_OUTCOME and verdict.supersedes == (claim.record_id,)
    assert claim.payload["attempt_id"] == result.local.run_id
    assert (env.root / "report.py").read_bytes() == original  # originals are never modified by M37
    run = env.store.get(result.local.run_id)
    assert [c["relpath"] for c in run["result"]["uri_computed_diff"]["changes"]] == ["report.py"]


def test_compatible_but_broken_script_fails_verification_and_escalates(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch)
    make_workbook(env.root / "sales.xlsx", SALES); env.write("report.py", script(extra="raise SystemExit(3)\n"))
    result = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert result.decision.reason is Reason.ALREADY_COMPATIBLE and result.route == "ESCALATE"
    assert result.local.verification_status == "FAILED" and not result.local.completed
    assert result.escalation.local_findings["local_attempt"]["verification_status"] == "FAILED"
    verdict = env.log.load().records[result.local.verifier_record_id]
    assert verdict.payload["verification_status"] == VerificationStatus.FAILED  # the failed attempt stays in Memory


def test_escalation_carries_only_unresolved_remainder_and_hashed_facts(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch); _setup(env, column="Amt")
    result = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    esc = result.escalation.to_dict()
    assert result.route == "ESCALATE" and result.decision.reason is Reason.NO_REMAP
    assert esc["unresolved"] == ["'Amt': none match in workbook"]
    assert esc["acceptance_sha256"] == acceptance().sha256
    assert all(len(i["content_sha256"]) == 64 for i in esc["inputs"])
    blob = json.dumps(esc)
    assert "North" not in blob and "print(" not in blob  # no data cells and no script text leave the Edge


def test_report_tasks_complete_locally_and_verify_by_recomputation(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch); _setup(env, column="Amt")
    for cls in ("WORKBOOK_INSPECT", "SCRIPT_WORKBOOK_DIAGNOSE"):
        r = env.edge.handle(env.task(cls))
        assert r.route == "LOCAL_VERIFIED", r
    run = env.store.get(r.local.run_id)
    ws_report = Path(run["started"]["workspace_dir"]) / "files" / "edge_report.json"
    ws_report.write_text(ws_report.read_text().replace("Amt", "Amount"))
    verdict = env.recorder.invoke_verifier(r.local.claim_record_id, "uri.verifier.execution.script_workbook.v1",
                                           trace_id="a" * 32, session_id="session")
    rec = env.log.load().records[verdict.record_ids[0]]
    assert rec.payload["verification_status"] == VerificationStatus.FAILED  # candidate changed after the claim


def test_unqualified_registry_never_completes_locally(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch, qualification_run=False); _setup(env)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.NO_QUALIFIED_CAPABILITY and r.local is None


def test_qualification_record_is_digest_pinned_and_version_exact(tmp_path):
    raw = {"schema_version": "m37.edge.qualification.v1", "capabilities": [
        {"capability_id": c.capability_id, "version": c.version, "mechanism": c.mechanism.value,
         "status": "QUALIFIED", "passed": True, "evidence": "x"} for c in CAPABILITIES]}
    p = tmp_path / "q.json"; p.write_text(json.dumps(raw))
    sha = hashlib.sha256(p.read_bytes()).hexdigest()
    reg = CapabilityRegistry(QualificationRecord.load(p, sha))
    assert [c.capability_id for c in reg.eligible("SCRIPT_WORKBOOK_REPAIR")] == ["uri.edge.script_workbook_literal_remap"]
    with pytest.raises(ValueError, match="DIGEST"): QualificationRecord.load(p, "0" * 64)
    raw["capabilities"][2]["version"] = "2"; p.write_text(json.dumps(raw))
    reg = CapabilityRegistry(QualificationRecord.load(p, hashlib.sha256(p.read_bytes()).hexdigest()))
    assert reg.eligible("SCRIPT_WORKBOOK_REPAIR") == []
    raw["capabilities"][2].update(version="1", mechanism="EDGE_MODEL"); p.write_text(json.dumps(raw))
    reg = CapabilityRegistry(QualificationRecord.load(p, hashlib.sha256(p.read_bytes()).hexdigest()))
    assert reg.eligible("SCRIPT_WORKBOOK_REPAIR") == []


# -- E-5: the Edge consumes a real M36 MemoryContextPackage without contract change --------------

def _m36_env(tmp_path, monkeypatch, column="amount"):
    log, s, rec, root = m36_setup(tmp_path, files=("sales.xlsx", "report.py"))
    make_workbook(root / "sales.xlsx", SALES); (root / "report.py").write_text(script(column=column))
    return Env(tmp_path, monkeypatch, existing=(log, s, rec, root))


def test_e5_current_binding_allows_local_completion(tmp_path, monkeypatch):
    env = _m36_env(tmp_path, monkeypatch)
    task, head = opened(env.recorder)
    b = bind(env.log, env.sources, "sales.xlsx")
    assert env.recorder.task_state(task, head.record_id, intake=m36_intake("sales.xlsx"), bindings=(b,)).persisted
    package = build_context(env.log.load(), env.sources, task)
    assert package.to_dict()["references"][0]["freshness"] == "CURRENT"
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), package)
    assert r.route == "LOCAL_VERIFIED"


def test_e5_changed_file_after_binding_clarifies(tmp_path, monkeypatch):
    env = _m36_env(tmp_path, monkeypatch)
    task, head = opened(env.recorder)
    b = bind(env.log, env.sources, "sales.xlsx")
    assert env.recorder.task_state(task, head.record_id, intake=m36_intake("sales.xlsx"), bindings=(b,)).persisted
    make_workbook(env.root / "sales.xlsx", {"Sales": (("Region", "Amount", "Unit Price", "Extra"), ())})
    package = build_context(env.log.load(), env.sources, task)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), package)
    assert r.route == "CLARIFY" and r.decision.reason is Reason.STALE_INPUT


def test_e5_user_corrected_reference_clarifies(tmp_path, monkeypatch):
    log, s, rec, root = m36_setup(tmp_path, files=("sales.xlsx", "sales_v2.xlsx", "report.py"))
    make_workbook(root / "sales.xlsx", SALES); make_workbook(root / "sales_v2.xlsx", SALES)
    (root / "report.py").write_text(script(column="amount"))
    env = Env(tmp_path, monkeypatch, existing=(log, s, rec, root))
    task, head = opened(rec)
    a, b = bind(log, s, "sales.xlsx"), bind(log, s, "sales_v2.xlsx")
    assert rec.task_state(task, head.record_id, intake=m36_intake("sales.xlsx"), bindings=(a,)).persisted
    assert rec.correction(task, "file", log.load().task(task).record_id, intake=m36_intake("use sales_v2.xlsx"), replacement=b).persisted
    package = build_context(log.load(), s, task)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), package)
    assert r.route == "CLARIFY" and r.decision.reason is Reason.CORRECTED_REFERENCE


def test_e5_memory_conflict_and_invalid_context_clarify(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch); _setup(env)
    task = env.task("SCRIPT_WORKBOOK_REPAIR")
    class Pkg:
        def __init__(self, d): self.d = d
        def to_dict(self): return self.d
    assert resolve(task, Pkg({"schema_version": "m36.context.v1", "degraded": "CONFLICT"}), env.registry, env.sources).reason is Reason.MEMORY_CONFLICT
    assert resolve(task, Pkg({"schema_version": "other"}), env.registry, env.sources).reason is Reason.INVALID_CONTEXT
    assert resolve(task, object(), env.registry, env.sources).kind is DecisionKind.CLARIFY


def test_telemetry_is_content_free(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch); _setup(env)
    env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    ev = env.telemetry.events()[-1]
    assert set(ev) == {"trace_id", "task_class", "decision", "reason", "capability_id", "route", "verification_status",
                       "escalation_avoided", "qualification_run", "latency_ms", "schema_version"}
    assert ev["escalation_avoided"] is True and "sales" not in json.dumps(ev).lower()


# -- boundaries ---------------------------------------------------------------------------------

BANNED = {"uri_core", "socket", "requests", "httpx", "urllib", "openai", "anthropic", "ollama"}


def _imports(path):
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in ast.walk(tree):
        if isinstance(node, ast.Import): yield from (a.name for a in node.names)
        elif isinstance(node, ast.ImportFrom) and node.level == 0: yield node.module or ""


def test_edge_and_execution_import_boundaries():
    for pkg in ("uri_v1/edge", "uri_v1/execution"):
        for path in (ROOT / pkg).rglob("*.py"):
            if "harness" in path.parts: continue
            for mod in _imports(path):
                assert mod.split(".")[0] not in BANNED, (path, mod)
                if mod.startswith("uri_v1."):
                    assert mod.split(".")[1] in ("memory", "results", "user_storage", "edge", "execution"), (path, mod)
                    assert not mod.startswith("uri_v1.memory.") or mod in ("uri_v1.memory.contracts",), (path, mod)


# Baseline = the reconciled M36 state this candidate carries (cloud M36 line 80b7359 plus
# the ported remainder of Antigravity-audited 70e6065). It was 405db1f, the pre-repair M36
# commit, which made this guard fail as soon as any accepted M36 repair was applied
# (Antigravity a87e158 F-5); M36 repairs are M36's, not a scope breach by M37/M38.
# Moved to 7f82400 for the M36 apostrophe port from the Codex line (379769f), then to
# 179fe10 for its NFKC / -n't parity port.
M36_BASELINE = "179fe10b7fe620cef60a979431e337c135f58215"

# M37-L3 fix (User, Chetan, 2026-09-29 13:57 UTC, "Fix in M37"): the FROZEN_M37_M38 prefixes
# only say where M37/M38 may differ from the M36 baseline, so on their own they let any edit
# under uri_v1/edge/ or uri_v1/execution/ pass. Since the freeze those paths are frozen as
# well and must match the freeze commit exactly (renames counted as delete + add). Only this
# guard file and its hash pin, tests/test_m37_frozen_guard_pin.py, are exempt; the pin holds
# this file to its authorized content.
FREEZE_COMMIT = "5bbf73856011fac7f8abad6970e664616a4c01ec"
FROZEN_M37_M38 = ("uri_v1/edge/", "uri_v1/execution/", "tests/test_m37_", "tests/test_m38_", "tests/m37_edge_support.py",
                  "tests/m38_harness_support.py", "scripts/m37_", "scripts/m38_", "fixtures/m37_edge/", "fixtures/m38_harness/")
FREEZE_EXEMPT = frozenset({"tests/test_m37_edge_resolution.py", "tests/test_m37_frozen_guard_pin.py"})


def frozen_scope_offenders():
    """Every path the guard below rejects: outside the allowed scope, or (M37-L3) changed since the freeze."""
    changed = subprocess.check_output(["git", "diff", "--name-only", M36_BASELINE, "--",
                                       "uri_v1", "uri_core", "uri_ui", "tests", "scripts", "fixtures"], cwd=ROOT, text=True).splitlines()
    untracked = subprocess.check_output(["git", "ls-files", "--others", "--exclude-standard", "--", "uri_v1", "uri_core", "uri_ui",
                                         "tests", "scripts", "fixtures"], cwd=ROOT, text=True).splitlines()
    allowed = FROZEN_M37_M38
    outside = [p for p in changed + untracked if not p.startswith(allowed)]
    since_freeze = subprocess.check_output(["git", "diff", "--name-only", "--no-renames", FREEZE_COMMIT, "--",
                                            "uri_v1", "uri_core", "uri_ui", "tests", "scripts", "fixtures"], cwd=ROOT, text=True).splitlines()
    unfrozen = [p for p in since_freeze + untracked if p.startswith(FROZEN_M37_M38) and p not in FREEZE_EXEMPT]
    return outside + unfrozen


def test_m36_and_frozen_files_unchanged():
    offenders = frozen_scope_offenders()
    assert not offenders, offenders
    from scripts.m33_3_r_anchors import require_anchors
    assert require_anchors()["ok"]
