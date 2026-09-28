"""M37 execution layer: acceptance, workbook reader, workspace, write-back, run lineage."""
import json
import os
import zipfile

import pytest

from tests.m37_edge_support import SALES, Env, acceptance, intake, make_workbook, script
from uri_v1.execution.acceptance import AcceptanceCriteria
from uri_v1.execution.officefile import WorkbookReadError, read_workbook_structure
from uri_v1.execution.run_record import RunRecordError, RunStore
from uri_v1.execution.verify import verify_candidate
from uri_v1.execution.workspace import (Workspace, WorkspaceError, WorkspaceInput, WorkspaceStrategy,
                                        WritebackApproval, new_run_id)


def test_acceptance_is_closed_and_hash_stable():
    a = acceptance(); b = AcceptanceCriteria.from_dict(json.loads(a.canonical()))
    assert a.sha256 == b.sha256
    for bad in ({"kind": "run_shell", "cmd": "rm"}, {"kind": "script_exit_zero", "script": "../x.py"},
                {"kind": "script_exit_zero", "script": "/abs.py"}, {"kind": "only_changed", "paths": ["a"], "extra": 1},
                {"kind": "workbook_has_sheets", "path": "a.xlsx", "sheets": []},
                {"kind": "script_exit_zero", "script": "a.py", "timeout_s": 10_000}):
        with pytest.raises(ValueError): AcceptanceCriteria(1, (bad,))
    with pytest.raises(ValueError): AcceptanceCriteria(1, ({"kind": "only_changed", "paths": []},) * 2)


def test_workbook_reader_real_file_and_errors(tmp_path):
    p = make_workbook(tmp_path / "w.xlsx", {"Sales": (("Region", "Amount"), (("N", 1),)), "Notes": (("Text",), ())},
                      formula_error=False)
    s = read_workbook_structure(p)
    assert s.sheet_names == ("Sales", "Notes") and s.sheet("Sales").header == ("Region", "Amount")
    (tmp_path / "bad.xlsx").write_bytes(b"not a zip")
    with pytest.raises(WorkbookReadError, match="XLSX_NOT_ZIP"): read_workbook_structure(tmp_path / "bad.xlsx")
    bomb = tmp_path / "bomb.xlsx"
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("xl/workbook.xml", b"<a>" + b" " * (10 * 1024 * 1024) + b"</a>")
    with pytest.raises(WorkbookReadError, match="XLSX_COMPRESSION_RATIO"): read_workbook_structure(bomb)
    dtd = tmp_path / "dtd.xlsx"
    with zipfile.ZipFile(dtd, "w") as z:
        z.writestr("xl/workbook.xml", b'<?xml version="1.0"?><!DOCTYPE a [<!ENTITY x "y">]><a>&x;</a>')
    with pytest.raises(WorkbookReadError, match="XLSX_DTD_REJECTED"): read_workbook_structure(dtd)


def test_formula_error_cells_are_counted(tmp_path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.append(["A"]); ws["A2"] = "#DIV/0!"; ws["A2"].data_type = "e"
    wb.save(tmp_path / "e.xlsx")
    assert read_workbook_structure(tmp_path / "e.xlsx").sheets[0].error_cells == 1


def _ws(env, *names):
    refs = [env.ref(n) for n in names]
    return Workspace.create(user_id=env.user, scratch_root=env.scratch, sources=env.sources,
                            inputs=[WorkspaceInput("x", r.source_id, r.content_sha256) for r in refs])


def test_copy_in_only_named_hash_bound_inputs(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch)
    make_workbook(env.root / "sales.xlsx", SALES); env.write("report.py", script()); env.write("secret.txt", "private")
    ws = _ws(env, "sales.xlsx", "report.py")
    assert sorted(p.name for p in ws.files.iterdir()) == ["report.py", "sales.xlsx"]
    assert ws.diff() == {"changes": [], "containment_violations": []}
    stale = WorkspaceInput("x", env.ref("report.py").source_id, "0" * 64)
    with pytest.raises(WorkspaceError, match="INPUT_NOT_CURRENT"):
        Workspace.create(user_id=env.user, scratch_root=env.scratch, sources=env.sources, inputs=[stale])
    with pytest.raises(WorkspaceError, match="DIRECT_ORIGINAL"):
        Workspace.create(user_id=env.user, scratch_root=env.scratch, sources=env.sources, inputs=[stale],
                         strategy=WorkspaceStrategy.DIRECT_ORIGINAL)
    with pytest.raises(WorkspaceError, match="OVERLAPS"):
        Workspace.create(user_id=env.user, scratch_root=env.root, sources=env.sources,
                         inputs=[WorkspaceInput("x", env.ref("report.py").source_id, env.ref("report.py").content_sha256)])


def test_uri_computes_changes_and_detects_containment_violations(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch)
    make_workbook(env.root / "sales.xlsx", SALES); env.write("report.py", script())
    ws = _ws(env, "sales.xlsx", "report.py")
    (ws.files / "report.py").write_text("print(1)\n"); (ws.files / "new.txt").write_text("n"); (ws.files / "sales.xlsx").unlink()
    ops = {c["relpath"]: c["op"] for c in ws.diff()["changes"]}
    assert ops == {"report.py": "MODIFIED", "new.txt": "ADDED", "sales.xlsx": "DELETED"}
    os.symlink(env.root / "report.py", ws.files / "escape.py")
    assert ws.diff()["containment_violations"] == ["escape.py"]
    report = verify_candidate(ws, acceptance())
    assert report["status"] == "FAILED" and report["items"][0]["detail"] == "CONTAINMENT_VIOLATION"


def test_single_writer_lock(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch)
    env.write("report.py", script())
    ws = _ws(env, "report.py"); ws.acquire("a")
    with pytest.raises(WorkspaceError, match="BUSY"): ws.acquire("b")
    with pytest.raises(WorkspaceError, match="NOT_OWNED"): ws.release("b")
    ws.release("a"); ws.acquire("b")


def test_writeback_requires_verified_run_explicit_trusted_approval_and_unchanged_original(tmp_path, monkeypatch):
    env = Env(tmp_path, monkeypatch)
    make_workbook(env.root / "sales.xlsx", SALES); original = env.write("report.py", script(column="amount"))
    ws = _ws(env, "sales.xlsx", "report.py")
    (ws.files / "report.py").write_bytes(script().encode())
    after = {c["relpath"]: c["after_sha256"] for c in ws.diff()["changes"]}
    for text in ("don't apply it", "maybe later"):
        with pytest.raises(WorkspaceError, match="EXPLICIT_APPROVAL_REQUIRED"):
            WritebackApproval.from_user(intake(text, user=env.user), ws.run_id, after)
    approval = WritebackApproval.from_user(intake("yes, apply the fix", user=env.user), ws.run_id, after)
    with pytest.raises(WorkspaceError, match="REQUIRES_VERIFIED"): ws.write_back(approval, sources=env.sources, verified_status="FAILED")
    forged = WritebackApproval(ws.run_id, ("report.py",), after, object())
    with pytest.raises(WorkspaceError, match="APPROVAL_REQUIRED"): ws.write_back(forged, sources=env.sources, verified_status="VERIFIED")
    other = WritebackApproval.from_user(intake("approve", user="00000000-0000-0000-0000-000000000009"), ws.run_id, after)
    with pytest.raises(WorkspaceError, match="WRONG_USER"): ws.write_back(other, sources=env.sources, verified_status="VERIFIED")
    original.write_text("# user edited meanwhile\n")
    with pytest.raises(WorkspaceError, match="ORIGINAL_CHANGED"): ws.write_back(approval, sources=env.sources, verified_status="VERIFIED")
    original.write_bytes(script(column="amount").encode())
    written = ws.write_back(approval, sources=env.sources, verified_status="VERIFIED")
    assert written[0]["relpath"] == "report.py" and original.read_bytes() == script().encode()


def test_run_lineage_is_append_only_and_tamper_evident(tmp_path):
    store = RunStore("00000000-0000-0000-0000-000000000007", tmp_path)
    rid = new_run_id()
    with pytest.raises(RunRecordError, match="RUN_UNKNOWN"): store.append(rid, "RUN_RESULT", {})
    store.append(rid, "RUN_STARTED", {"executor": "LOCAL_EDGE"})
    with pytest.raises(RunRecordError, match="RUN_EXISTS"): store.append(rid, "RUN_STARTED", {"executor": "LOCAL_EDGE"})
    store.append(rid, "RUN_RESULT", {"x": 1})
    assert store.get(rid)["events"] == ["RUN_STARTED", "RUN_RESULT"]
    raw = store.path.read_bytes().replace(b'"x":1', b'"x":2'); store.path.write_bytes(raw)
    with pytest.raises(RunRecordError, match="CHAIN_BROKEN"): store.get(rid)
