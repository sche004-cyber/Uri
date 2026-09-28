"""Permanent regression tests for the M37 repair re-audit residuals R-1..R-6
(audits/M37_REPAIR_REAUDIT_REPORT.md, probes in
audits/probes/m37_reaudit/test_m37_reaudit_probes.py).

Each `test_safe_*` here is one of the re-audit's `test_defect_*` probes,
inverted to assert the now-safe behaviour instead of the residual it
originally demonstrated. A regression in any of these means one of R-1..R-6
has come back. R-5 stays partly disclosed (see its test below): a detached,
session-escaped child can still tamper with the raw workspace bytes after
verification, but `write_back` is shown to still refuse in that case.
"""
from __future__ import annotations

import os
import textwrap
import time
import uuid

from uri_v1.edge.contracts import DecisionKind, Reason
from uri_v1.edge.local_runner import ACTION_LABELS
from uri_v1.edge.resolver import resolve
from uri_v1.edge.scriptinfo import inspect_script
from uri_v1.execution.lineage import record_and_verify
from uri_v1.execution.workspace import Workspace, WorkspaceError, WritebackApproval
from uri_v1.memory.context_package import build_context

from tests.m37_edge_support import SALES, Env, intake, make_workbook, script
from tests.test_m36_memory_support import bind, intake as m36_intake, opened, setup as m36_setup


def _env(tmp_path, monkeypatch, src, **kw):
    env = Env(tmp_path, monkeypatch, **kw)
    make_workbook(env.root / "sales.xlsx", SALES); env.write("report.py", src)
    return env


# ---------------------------------------------------------------------------------------
# R-1 (F-1 residual): a name shadowed anywhere by a function parameter (loop target,
# with/except target, or comprehension variable) is never trusted for workbook/sheet/frame
# subscript access, even at the module-level occurrence that looks legitimate.
# ---------------------------------------------------------------------------------------
F1_SHADOW = textwrap.dedent('''
    from openpyxl import load_workbook, Workbook
    wb = load_workbook("sales.xlsx")
    ws = wb["sales"]                      # wrong case: would have been remapped to "Sales"
    header = [c.value for c in ws[1]]
    col = header.index("Amount")
    total = sum(row[col] for row in ws.iter_rows(min_row=2, values_only=True))

    def weight(wb):                       # unrelated helper; `wb` here is a plain dict
        return wb["sales"]

    w = weight({"sales": 1, "Sales": 100})
    out = Workbook(); o = out.active; o.title = "Summary"
    o.append(["Total"]); o.append([total * w])
    out.save("summary.xlsx")
''').lstrip()


def test_safe_f1_shadowed_name_is_not_trusted_for_remap(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, F1_SHADOW)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.DYNAMIC_ACCESS


# ---------------------------------------------------------------------------------------
# R-2 (F-2 residual): a script that changes its working directory before loading the
# literal-bound workbook name must not be treated as reading the bound workbook.
# ---------------------------------------------------------------------------------------
def test_safe_f2_chdir_marks_workbook_path_dynamic(tmp_path, monkeypatch):
    other = tmp_path / "elsewhere"; other.mkdir()
    make_workbook(other / "sales.xlsx", {"Sales": (("Region", "Amount", "Unit Price"), (("West", 999, 1.0),))})
    src = textwrap.dedent(f'''
        import os
        from openpyxl import load_workbook, Workbook
        here = os.getcwd(); os.chdir({str(other)!r})
        wb = load_workbook("sales.xlsx")
        os.chdir(here)
        ws = wb["Sales"]
        header = [c.value for c in ws[1]]
        col = header.index("Amount")
        total = sum(row[col] for row in ws.iter_rows(min_row=2, values_only=True))
        out = Workbook(); o = out.active; o.title = "Summary"
        o.append(["Total"]); o.append([total])
        out.save("summary.xlsx")
    ''').lstrip()
    env = _env(tmp_path, monkeypatch, src)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.DYNAMIC_ACCESS


# ---------------------------------------------------------------------------------------
# R-3 (F-4 residual): the header-value-comparison guard now covers the common openpyxl
# idiom of comparing `cell.value` while iterating the header row, plus `.strip()`,
# `in (...)`, subscript and membership variants.
# ---------------------------------------------------------------------------------------
F4_CELL = textwrap.dedent('''
    from openpyxl import load_workbook, Workbook
    wb = load_workbook("sales.xlsx"); ws = wb["Sales"]
    col = None
    for cell in ws[1]:
        if cell.value == "amount":        # wrong case: never matches
            col = cell.column - 1
    total = 0 if col is None else sum(r[col] for r in ws.iter_rows(min_row=2, values_only=True))
    out = Workbook(); o = out.active; o.title = "Summary"; o.append(["Total"]); o.append([total]); out.save("summary.xlsx")
''').lstrip()


def test_safe_f4_cell_value_comparison_is_dynamic(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, F4_CELL)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.DYNAMIC_ACCESS


def test_safe_f4_other_comparison_idioms_are_dynamic():
    base = 'from openpyxl import load_workbook\nwb = load_workbook("sales.xlsx"); ws = wb["Sales"]\nhdr = [c.value for c in ws[1]]\n'
    variants = {
        "strip": 'for h in hdr:\n    if h.strip() == "amount": pass\n',
        "in_tuple": 'for h in hdr:\n    if h in ("amount", "amt"): pass\n',
        "subscript": 'for i in range(len(hdr)):\n    if hdr[i] == "amount": pass\n',
        "membership": 'ok = "amount" in hdr\n',
    }
    for name, extra in variants.items():
        info = inspect_script((base + extra).encode())
        assert info.dynamic, f"{name} should be dynamic"


# ---------------------------------------------------------------------------------------
# R-4 (F-6 residual): a Memory context package that carries no task at all (built for a
# task_id Memory has no state for) must fail closed, not skip the identity check.
# ---------------------------------------------------------------------------------------
def test_safe_f6_package_for_unknown_task_clarifies(tmp_path, monkeypatch):
    log, s, rec, root = m36_setup(tmp_path, files=("sales.xlsx", "sales_v2.xlsx", "report.py"))
    make_workbook(root / "sales.xlsx", SALES); make_workbook(root / "sales_v2.xlsx", SALES)
    (root / "report.py").write_text(script(column="amount"))
    env = Env(tmp_path, monkeypatch, existing=(log, s, rec, root))
    task, head = opened(rec)
    a, b = bind(log, s, "sales.xlsx"), bind(log, s, "sales_v2.xlsx")
    assert rec.task_state(task, head.record_id, intake=m36_intake("sales.xlsx"), bindings=(a,)).persisted
    assert rec.correction(task, "file", log.load().task(task).record_id, intake=m36_intake("use sales_v2.xlsx"), replacement=b).persisted
    right = resolve(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), build_context(log.load(), s, task), env.registry, env.sources)
    assert right.reason is Reason.CORRECTED_REFERENCE
    ghost = build_context(log.load(), s, uuid.uuid4().hex)
    assert ghost.to_dict()["task"] is None
    r = resolve(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), ghost, env.registry, env.sources)
    assert r.kind is DecisionKind.CLARIFY and r.reason is Reason.INVALID_CONTEXT


# ---------------------------------------------------------------------------------------
# R-5 (F-8 residual, partly disclosed): the process group is now killed after a normal
# exit too, not only on timeout, so an ordinary (non-detached) leftover child cannot
# outlive the run. A child that escapes the group via `start_new_session=True` before it
# exits is not reachable this way -- that is disclosed, shared with M38, and needs OS
# containment to close. The safety net that matters is `write_back`: even though such a
# child can still edit the raw workspace bytes after verification, write_back must refuse
# rather than carry the tampered bytes back to the original.
# ---------------------------------------------------------------------------------------
F8_LATE = ('import subprocess, sys\n'
           'subprocess.Popen([sys.executable, "-c", "import time,pathlib;time.sleep(3);'
           'p=pathlib.Path(\'../../files/report.py\');p.write_text(p.read_text()+\'\\\\n# tampered late\\\\n\')"],'
           ' stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, start_new_session=True)\n')


def test_safe_f8_detached_child_tamper_cannot_write_back(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, script(column="amount", extra=F8_LATE))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "LOCAL_VERIFIED" and r.local.verification_status == "VERIFIED"
    run_id = r.local.run_id
    run = env.store.get(run_id)
    ws = Workspace.open(run["started"]["workspace_dir"])
    time.sleep(5)
    recorded = {c["relpath"]: c["after_sha256"] for c in run["result"]["uri_computed_diff"]["changes"]}
    live = {c["relpath"]: c["after_sha256"] for c in ws.diff()["changes"]}
    if os.name == "nt":
        # On Windows the verifier's child runs in a kill-on-close Job Object (M38 Windows
        # repair, a87e158 F-1), which reaches a detached descendant: no late tamper at all.
        assert recorded == live and not (ws.files / "report.py").read_text().rstrip().endswith("# tampered late")
        return
    # Disclosed (POSIX): a setsid-detached child still tampers with the raw workspace after
    # the verifier's own re-diff, since only OS-level containment (not a process-group
    # kill) can reach a process that has already left the group.
    assert recorded != live
    # The script's own source contains this text (the child's command line), so check that
    # it was actually appended at the end, not merely present.
    assert (ws.files / "report.py").read_text().rstrip().endswith("# tampered late")
    # But the safety net holds: write_back reads the verified diff fingerprint, not the
    # live workspace, so the tampered bytes can never reach the original.
    after = {rel: sha for rel, sha in live.items()}
    ap = WritebackApproval.from_user(intake("yes apply", user=env.user), run_id, after)
    try:
        ws.write_back(ap, sources=env.sources, run_store=env.store); refused = None
    except WorkspaceError as exc:
        refused = str(exc)
    assert refused == "WRITEBACK_CANDIDATE_CHANGED"


# ---------------------------------------------------------------------------------------
# R-6 (F-7 residual): write_back must bind to what the latest VERIFICATION actually
# covered, not to whatever RUN_RESULT happens to be last recorded. A second record on the
# same run whose verifier returns early (a safety flag, so no new VERIFICATION is
# appended) must not let never-verified bytes write back under the stale VERIFIED status.
# ---------------------------------------------------------------------------------------
def test_safe_f7_writeback_binds_to_latest_verification_not_latest_result(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, script())
    t = env.task("SCRIPT_WORKBOOK_REPAIR")
    r = env.edge.handle(t)
    assert r.decision.reason is Reason.ALREADY_COMPATIBLE and r.route == "LOCAL_VERIFIED"
    run_id = r.local.run_id
    ws = Workspace.open(env.store.get(run_id)["started"]["workspace_dir"])
    evil = b"print('never verified')\n"
    (ws.files / "report.py").write_bytes(evil)
    lin = record_and_verify(recorder=env.recorder, ledger=env.ledger, store=env.store, workspace=ws, run_id=run_id,
                            task_id=t.task_id, action_label=ACTION_LABELS[t.task_class], session_id=t.session_id,
                            trace_id=t.trace_id, executor_claim={}, uri_observations={"safety_flags": ["retry"]})
    assert lin.verification_status != "VERIFIED"
    run = env.store.get(run_id)
    # M38 repair round 3: a SCRIPT_WORKBOOK_REPAIR run has no exact recompute oracle, so
    # its VERIFIED-equivalent run-ledger status is the honestly weaker
    # STRUCTURALLY_VERIFIED (differential-rerun only) -- still the stale verified status
    # this test is about, just under its accurate name.
    assert run["verifications"][-1]["status"] == "STRUCTURALLY_VERIFIED"  # the stale verified status stays last
    after = {c["relpath"]: c["after_sha256"] for c in ws.diff()["changes"] if c["op"] == "MODIFIED"}
    ap = WritebackApproval.from_user(intake("yes apply", user=env.user), run_id, after)
    try:
        ws.write_back(ap, sources=env.sources, run_store=env.store); refused = None
    except WorkspaceError as exc:
        refused = str(exc)
    assert refused == "WRITEBACK_CANDIDATE_CHANGED"
    assert (env.root / "report.py").read_bytes() != evil
