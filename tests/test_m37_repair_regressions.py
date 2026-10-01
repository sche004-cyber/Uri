"""Permanent regression tests for the M37 audit findings F-1..F-9
(audits/M37_INTERNAL_AUDIT_REPORT.md, probes in audits/probes/m37/test_m37_audit_probes.py).

Each test here is the audit's own probe (P1-P9, P11b), inverted to assert the
now-safe behaviour instead of the defect it originally demonstrated. A
regression in any of these means one of F-1..F-9 has come back.
"""
from __future__ import annotations

import hashlib
import json
import textwrap
from pathlib import Path

from uri_v1.edge.contracts import DecisionKind, Reason
from uri_v1.edge.resolver import resolve
from uri_v1.execution.acceptance import AcceptanceCriteria
from uri_v1.execution.run_record import RunStore
from uri_v1.execution.workspace import Workspace, WorkspaceError, WorkspaceInput, WritebackApproval
from uri_v1.memory.context_package import build_context

from tests.m37_edge_support import SALES, Env, intake, make_workbook, script
from tests.test_m36_memory_support import bind, intake as m36_intake, opened, setup as m36_setup


def _env(tmp_path, monkeypatch, src, **kw):
    env = Env(tmp_path, monkeypatch, **kw)
    make_workbook(env.root / "sales.xlsx", SALES); env.write("report.py", src)
    return env


# F-3: a REPAIR task whose acceptance omits script_exit_zero can no longer be LOCAL_VERIFIED
# without the script ever actually running.
def test_f3_weak_acceptance_no_longer_completes_locally(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, script(column="amount", extra="raise SystemExit(3)\n"))
    weak = AcceptanceCriteria(1, ({"kind": "only_changed", "paths": ["report.py"]},))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR", accept=weak))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.ACCEPTANCE_BELOW_MINIMUM


# F-4: a silent access-pattern mismatch (wrong-case string comparison against a header value)
# is now flagged dynamic and escalates instead of being recorded ALREADY_COMPATIBLE/LOCAL_VERIFIED.
def test_f4_silent_mismatch_no_longer_local_verified(tmp_path, monkeypatch):
    src = textwrap.dedent('''
        from openpyxl import load_workbook, Workbook
        wb = load_workbook("sales.xlsx"); ws = wb["Sales"]
        total = 0
        hdr = [c.value for c in ws[1]]
        for i, h in enumerate(hdr):
            if h == "amount":          # wrong case: never matches, total stays 0
                total = sum(r[i] for r in ws.iter_rows(min_row=2, values_only=True))
        out = Workbook(); o = out.active; o.title = "Summary"; o.append(["Total"]); o.append([total]); out.save("summary.xlsx")
    ''').lstrip()
    env = _env(tmp_path, monkeypatch, src)
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.DYNAMIC_ACCESS


# F-1: `.index("x")` on a receiver that is not a recognized header-row derivation is never
# treated as a column lookup, so the literal remap can no longer rewrite unrelated code.
def test_f1_remap_no_longer_rewrites_non_workbook_index_literal(tmp_path, monkeypatch):
    extra = 'modes = ["amount", "Amount"]\nmode = modes.index("amount")\nassert mode == 0\n'
    env = _env(tmp_path, monkeypatch, script(extra=extra))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.decision.reason is Reason.DYNAMIC_ACCESS


# F-5: a failed local attempt's escalation no longer carries raw script output (cell values).
def test_f5_failed_attempt_escalation_no_longer_leaks_cell_values(tmp_path, monkeypatch):
    extra = 'import sys\nsys.stderr.write(str(list(ws.iter_rows(min_row=2, values_only=True))))\nraise SystemExit(3)\n'
    env = _env(tmp_path, monkeypatch, script(extra=extra))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE"
    blob = json.dumps(r.escalation.to_dict())
    assert "North" not in blob and "South" not in blob
    detail = r.escalation.local_findings["local_attempt"]["items"]
    tail_details = [i["detail"] for i in detail if isinstance(i.get("detail"), dict) and "stderr_tail_sha256" in i["detail"]]
    assert tail_details and all("stderr_tail" not in d for d in tail_details)


# F-6: a clean Memory context package built for a different task_id can no longer bypass a
# real CONFLICT/correction on the actual task.
def test_f6_package_task_identity_is_checked(tmp_path, monkeypatch):
    log, s, rec, root = m36_setup(tmp_path, files=("sales.xlsx", "sales_v2.xlsx", "report.py"))
    make_workbook(root / "sales.xlsx", SALES); make_workbook(root / "sales_v2.xlsx", SALES)
    (root / "report.py").write_text(script(column="amount"))
    env = Env(tmp_path, monkeypatch, existing=(log, s, rec, root))
    task, head = opened(rec)
    a, b = bind(log, s, "sales.xlsx"), bind(log, s, "sales_v2.xlsx")
    assert rec.task_state(task, head.record_id, intake=m36_intake("sales.xlsx"), bindings=(a,)).persisted
    assert rec.correction(task, "file", log.load().task(task).record_id, intake=m36_intake("use sales_v2.xlsx"), replacement=b).persisted
    other, _ = opened(rec, "unrelated task")
    wrong_pkg = build_context(log.load(), s, other)
    r = resolve(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), wrong_pkg, env.registry, env.sources)
    assert r.kind is DecisionKind.CLARIFY and r.reason is Reason.INVALID_CONTEXT
    right = resolve(env.task("SCRIPT_WORKBOOK_REPAIR", task_id=task), build_context(log.load(), s, task), env.registry, env.sources)
    assert right.reason is Reason.CORRECTED_REFERENCE


# F-9: an input changing between resolve() and copy-in no longer raises out of handle();
# it CLARIFYs with STALE_INPUT and a telemetry event is written.
def test_f9_toctou_clarifies_instead_of_raising(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, script(column="amount"))
    t = env.task("SCRIPT_WORKBOOK_REPAIR")
    import uri_v1.edge.service as svc
    real = svc.resolve
    def racing(*a, **k):
        d = real(*a, **k); make_workbook(env.root / "sales.xlsx", {"Sales": (("Region", "Amount", "Unit Price", "X"), ())}); return d
    monkeypatch.setattr(svc, "resolve", racing)
    r = env.edge.handle(t)
    assert r.route == "CLARIFY" and r.decision.reason is Reason.STALE_INPUT
    assert env.telemetry.events()


# F-7: write_back no longer trusts a caller-supplied status string; with no run record at
# all, it refuses.
def test_f7_writeback_no_longer_trusts_caller_string(tmp_path, monkeypatch):
    env = _env(tmp_path, monkeypatch, script(column="amount"))
    ws = Workspace.create(user_id=env.user, scratch_root=env.scratch, sources=env.sources,
                          inputs=[WorkspaceInput("script", env.ref("report.py").source_id, env.ref("report.py").content_sha256)])
    (ws.files / "report.py").write_text("print('never verified')\n")
    after = {c["relpath"]: c["after_sha256"] for c in ws.diff()["changes"]}
    ap = WritebackApproval.from_user(intake("yes apply", user=env.user), ws.run_id, after)
    store = RunStore(env.user, env.tmp / "runs2")
    try:
        ws.write_back(ap, sources=env.sources, run_store=store)
        raised = False
    except WorkspaceError:
        raised = True
    assert raised
    assert (env.root / "report.py").read_bytes() != b"print('never verified')\n"


# F-8: a script that tampers with the candidate it is being verified against (escaping the
# verification copy) no longer results in LOCAL_VERIFIED.
def test_f8_verification_tamper_no_longer_verifies(tmp_path, monkeypatch):
    extra = 'from pathlib import Path\np = Path("../../files/report.py")\np.write_text(p.read_text() + "\\n# tampered\\n")\n'
    env = _env(tmp_path, monkeypatch, script(column="amount", extra=extra))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.route == "ESCALATE" and r.local.verification_status != "VERIFIED"


# F-2: a script whose workbook path has any directory component (even ending in the bound
# workbook's basename) is a path mismatch, not ALREADY_COMPATIBLE.
def test_f2_absolute_path_outside_workspace_is_mismatch(tmp_path, monkeypatch):
    other = tmp_path / "elsewhere"; other.mkdir()
    make_workbook(other / "sales.xlsx", {"Sales": (("Region", "Amount", "Unit Price"), (("West", 999, 1.0),))})
    env = _env(tmp_path, monkeypatch, script(workbook=(other / "sales.xlsx").as_posix()))
    r = env.edge.handle(env.task("SCRIPT_WORKBOOK_REPAIR"))
    assert r.decision.reason is Reason.WORKBOOK_PATH_MISMATCH and r.route == "ESCALATE"
