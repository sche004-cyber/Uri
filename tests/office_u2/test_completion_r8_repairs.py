"""Completion review round 8 (driver URI-OFFICE-COMPLETION-r2) and the save-publishing gap found while checking the
R1-01 repair: each test reproduces the original defect on real files or packages."""
import io
import os
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "office"))
from test_xlsx_struct import budget_bytes, parts_of  # noqa: E402

from uri_office import ops as O
from uri_office import xlsx_patch as X
from uri_office import xlsx_struct as S
from uri_office.authoring import docx_ops as W
from uri_office.errors import OfficeError


# ---- R8-01: Word find and replace in a merged table cell -----------------------------------------------------------
def test_replace_in_a_merged_cell_changes_it_once(tmp_path):
    d = Document()
    t = d.add_table(rows=1, cols=3)
    merged = t.cell(0, 0).merge(t.cell(0, 1))
    merged.text = "a"
    t.cell(0, 2).text = "a"
    r = W.replace_text(d, "a", "aa")
    path = tmp_path / "merged.docx"
    d.save(path)
    out = Document(path).tables[0]
    assert out.cell(0, 0).text == "aa" and out.cell(0, 2).text == "aa"
    assert r["replaced"] == 2


# ---- R8-02: an insertion whose new rows would pass row 1048576 ----------------------------------------------------
def test_ops_refuse_new_rows_past_the_last_row():
    with pytest.raises(OfficeError) as e:
        O.to_cell_edits({"op": "insert_rows", "sheet": "Budget", "at": 1048576, "count": 2}, engine="local",
                        sheets=["Budget"])
    assert e.value.code == "BAD_ROWS"
    ok = O.to_cell_edits({"op": "insert_rows", "sheet": "Budget", "at": 1048576, "count": 1}, engine="local",
                         sheets=["Budget"])
    assert ok and ok[0]["at"] == 1048576


def test_row_engine_refuses_new_rows_past_the_last_row_and_changes_nothing():
    parts = parts_of(budget_bytes())
    before = dict(parts)
    with pytest.raises(OfficeError) as e:
        S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 1048576, "count": 2})
    assert e.value.code == "STRUCTURE_NOT_SUPPORTED"
    assert parts == before


# ---- publishing without hard links (FAT/exFAT folders the owner may allow, OD-R1) ---------------------------------
@pytest.fixture
def no_hard_links(monkeypatch):
    def refuse(src, dst, *a, **k):
        raise OSError(1, "hard links are not supported on this drive")
    monkeypatch.setattr(X.os, "link", refuse)


@pytest.mark.skipif(os.name != "nt", reason="the rename fallback is Windows-only (rename never replaces there)")
def test_save_and_save_as_work_without_hard_links_and_never_replace(tmp_path, no_hard_links):
    import hashlib
    # Save as: a new name is created; an existing name is refused and left as it was
    new = tmp_path / "new.docx"
    X.atomic_write(new, lambda fh: fh.write(b"NEW"), create_only=True)
    assert new.read_bytes() == b"NEW"
    with pytest.raises(OfficeError) as e:
        X.atomic_write(new, lambda fh: fh.write(b"OTHER"), create_only=True)
    assert e.value.code == "DESTINATION_EXISTS" and new.read_bytes() == b"NEW"
    # Save: the checked original is backed up and replaced
    orig = tmp_path / "orig.docx"
    orig.write_bytes(b"V1")
    X.atomic_write(orig, lambda fh: fh.write(b"V2"), expected_sha256=hashlib.sha256(b"V1").hexdigest(),
                   backup=tmp_path / "backup.docx")
    assert orig.read_bytes() == b"V2" and (tmp_path / "backup.docx").read_bytes() == b"V1"
    # Save over a file changed meanwhile: refused, and the changed file stays under its own name
    with pytest.raises(OfficeError) as e:
        X.atomic_write(orig, lambda fh: fh.write(b"V3"), expected_sha256=hashlib.sha256(b"V1").hexdigest(),
                       backup=tmp_path / "backup2.docx")
    assert e.value.code == "TARGET_CHANGED" and orig.read_bytes() == b"V2"
    assert not list(tmp_path.glob("*.uri-office.hold")) and not list(tmp_path.glob("*.uri-office.tmp"))
