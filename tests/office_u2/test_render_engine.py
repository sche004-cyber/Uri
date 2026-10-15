"""Real LibreOffice engine tests (marked `engine`): AC-O2-3, AC-O2-4, AC-O2-5, AC-O2-18; F1, F2.

The pinned engine is verified once per test session (about 100 s). Canary checks use an HTTP canary server on
127.0.0.1 and a system-wide census of connections to the SMB canary host 127.0.0.2 (ports 445, 139), polled every
50 ms, plus a census of every engine process's own connections. A positive control (the unsanitized canary file with
a default, non-hardened profile) shows the canary server does detect LibreOffice fetches.
"""
from __future__ import annotations

import http.server
import os
import shutil
import statistics
import subprocess
import sys
import threading
import time
from pathlib import Path

import psutil
import pytest

import canaries as CN
from conftest import lo_present, needs_windows
from uri_office import engines as E
from uri_office import render as R
from uri_office import sanitize as S

pytestmark = [pytest.mark.engine, needs_windows,
              pytest.mark.skipif(not lo_present(), reason="pinned LibreOffice not present on this PC")]


@pytest.fixture(scope="session")
def integrity(tmp_path_factory):
    st = tmp_path_factory.mktemp("lo_integrity")
    ei = E.EngineIntegrity("libreoffice", st)
    t0 = time.perf_counter()
    ei.verify_full()
    ei.verify_seconds = round(time.perf_counter() - t0, 1)
    return ei


@pytest.fixture
def lo(tmp_path, integrity):
    eng = R.LibreOffice(tmp_path / "state", integrity, idle_seconds=600)
    yield eng
    eng.shutdown()


class Canary:
    def __init__(self):
        self.hits = []
        canary = self

        class H(http.server.BaseHTTPRequestHandler):
            def do_GET(self):
                canary.hits.append(self.path)
                self.send_response(404)
                self.end_headers()
            do_HEAD = do_POST = do_OPTIONS = do_PROPFIND = do_GET

            def log_message(self, *a):
                pass
        self.srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}/canary"
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def close(self):
        self.srv.shutdown()
        self.srv.server_close()


class Census:
    """Connections to the SMB canary host (system-wide; the SMB redirector is PID 4) and every connection of the given
    process ids, polled every 50 ms."""
    def __init__(self, pids_fn):
        self.pids_fn = pids_fn
        self.smb, self.proc = set(), set()
        self._stop = threading.Event()
        self.t = threading.Thread(target=self._run, daemon=True)
        self.t.start()

    def _run(self):
        while not self._stop.is_set():
            try:
                for c in psutil.net_connections(kind="inet"):
                    if c.raddr and c.raddr.ip == "127.0.0.2" and c.raddr.port in (445, 139):
                        self.smb.add((c.pid, c.raddr.port))
            except psutil.Error:
                pass
            for pid in self.pids_fn():
                try:
                    for p in [psutil.Process(pid)] + psutil.Process(pid).children(recursive=True):
                        for c in p.net_connections(kind="inet"):
                            if c.raddr:
                                self.proc.add((p.pid, c.raddr.ip, c.raddr.port))
                except psutil.Error:
                    pass
            self._stop.wait(0.05)

    def stop(self):
        self._stop.set()
        self.t.join(5)


def _convert(lo, src_file: Path, kind: str, *, sanitize=True):
    d = lo.work_root / "docs" / src_file.stem
    d.mkdir(parents=True, exist_ok=True)
    if sanitize:
        r = S.make_render_copy(src_file, d, kind)
        src = r.path
    else:
        src = d / f"raw.{kind}"
        shutil.copy(src_file, src)
    out = d / "out.pdf"
    lo.submit(src, out, owner=src_file.stem).result(240)
    return out


@pytest.mark.parametrize("kind,maker", [("docx", CN.docx_canary), ("pptx", CN.pptx_canary),
                                        ("xlsx", CN.xlsx_canary)])
def test_sanitized_canaries_make_no_connection(lo, tmp_path, kind, maker):
    can = Canary()
    census = Census(lo.pids)
    try:
        src = maker(tmp_path / f"c.{kind}", can.base)
        out = _convert(lo, src, kind)
        time.sleep(1.0)
    finally:
        census.stop()
        can.close()
    assert out.is_file() and out.read_bytes()[:5] == b"%PDF-"
    assert can.hits == [], can.hits
    assert census.smb == set(), census.smb
    assert census.proc == set(), census.proc           # the engine made no network connection at all


def test_hardened_profile_alone_blocks_unsanitized_canaries(lo, tmp_path):
    """Defence in depth: even the unsanitized file fetches nothing under the hardened profile."""
    can = Canary()
    census = Census(lo.pids)
    try:
        src = CN.docx_canary(tmp_path / "raw.docx", can.base)
        _convert(lo, src, "docx", sanitize=False)
        time.sleep(1.0)
    finally:
        census.stop()
        can.close()
    assert can.hits == [] and census.smb == set(), (can.hits, census.smb)


def test_positive_control_default_profile_does_fetch(tmp_path, integrity, monkeypatch):
    """Without the hardened profile and without the sanitizer, LibreOffice fetches the canary. This proves the canary
    server can see a fetch, so the zero results above are meaningful."""
    monkeypatch.setattr(R, "HARDENED_XCU", '<?xml version="1.0" encoding="UTF-8"?><oor:items '
                                           'xmlns:oor="http://openoffice.org/2001/registry"></oor:items>')
    eng = R.LibreOffice(tmp_path / "state", integrity)
    can = Canary()
    try:
        src = CN.docx_canary(tmp_path / "raw.docx", can.base)
        _convert(eng, src, "docx", sanitize=False)
        time.sleep(1.0)
    finally:
        eng.shutdown()
        can.close()
    assert can.hits, "positive control: the default profile was expected to fetch the canary"


def test_warm_conversion_timing_and_cold_start(lo, tmp_path):
    import docx
    p = tmp_path / "report20.docx"
    d = docx.Document()
    for sec in range(1, 21):
        d.add_heading(f"{sec}. Section", 1)
        for k in range(6):
            d.add_paragraph(f"Paragraph {k} of section {sec}. " + "The committee recorded its decisions. " * 12)
        d.add_page_break()
    d.save(p)
    dd = lo.work_root / "docs" / "t"
    dd.mkdir(parents=True)
    r = S.make_render_copy(p, dd, "docx")
    t0 = time.perf_counter()
    lo.submit(r.path, dd / "o0.pdf", owner="t").result(240)
    cold = time.perf_counter() - t0
    warm = []
    for i in range(5):
        t0 = time.perf_counter()
        lo.submit(r.path, dd / f"o{i + 1}.pdf", owner="t").result(240)
        warm.append(time.perf_counter() - t0)
    import pypdf
    pages = len(pypdf.PdfReader(str(dd / "o1.pdf")).pages)
    print(f"cold {cold:.2f}s (launch check incl.), last_start {lo.last_start_seconds}s, warm {warm}, pages {pages}")
    assert pages >= 20
    assert statistics.median(warm) <= 2.0                  # AC-O2-3 (provisional threshold)
    assert cold <= 30.0                                     # first start in a new state folder


def test_engine_crash_restarts_once_and_leaves_no_orphans(lo, tmp_path):
    import docx
    p = tmp_path / "a.docx"
    docx.Document().save(p)
    dd = lo.work_root / "docs" / "a"
    dd.mkdir(parents=True)
    r = S.make_render_copy(p, dd, "docx")
    lo.submit(r.path, dd / "o1.pdf", owner="a").result(240)
    before = lo.pids()
    bins = [c.pid for pid in before for c in psutil.Process(pid).children(recursive=True)]
    for pid in bins + before:
        try:
            psutil.Process(pid).kill()
        except psutil.Error:
            pass
    res = lo.submit(r.path, dd / "o2.pdf", owner="a").result(240)
    assert (dd / "o2.pdf").is_file()
    assert lo.starts == 2 or res["restarted"]
    time.sleep(0.5)
    assert not [p for p in bins + before if psutil.pid_exists(p) and psutil.Process(p).status() != "zombie"]


def test_idle_shutdown_leaves_no_process_and_empty_tmp(tmp_path, integrity):
    eng = R.LibreOffice(tmp_path / "state", integrity, idle_seconds=2)
    try:
        import docx
        p = tmp_path / "a.docx"
        docx.Document().save(p)
        dd = eng.work_root / "docs" / "a"
        dd.mkdir(parents=True)
        r = S.make_render_copy(p, dd, "docx")
        eng.submit(r.path, dd / "o.pdf", owner="a").result(240)
        pids = eng.pids() + [c.pid for pid in eng.pids() for c in psutil.Process(pid).children(recursive=True)]
        deadline = time.monotonic() + 40
        while eng.pids() and time.monotonic() < deadline:
            time.sleep(0.5)
        time.sleep(1)
        assert eng.pids() == []
        assert not [p for p in pids if psutil.pid_exists(p)]
        assert R._tree_summary(eng.tmp)["files"] == 0
        assert eng.status()["state"] == "idle_stopped"
        eng.submit(r.path, dd / "o2.pdf", owner="a").result(240)          # restarts on demand
        assert eng.starts == 2
    finally:
        eng.shutdown()


def test_killing_the_service_process_kills_the_engine(tmp_path, integrity):
    """AC-O2-5: a service that dies leaves no LibreOffice process (Job Object KILL_ON_JOB_CLOSE)."""
    root = Path(__file__).resolve().parents[2]
    script = tmp_path / "svc.py"
    script.write_text(
        "import sys, json, time\n"
        f"sys.path.insert(0, {str(root)!r})\n"
        "from pathlib import Path\n"
        "from uri_office import render as R, engines as E\n"
        "ei = E.EngineIntegrity('libreoffice', Path(sys.argv[1]))\n"
        "ei._entries = __import__('json').loads(Path(sys.argv[2]).read_text())['files']\n"
        "lo = R.LibreOffice(Path(sys.argv[1]), ei)\n"
        "import docx; p = lo.work_root / 'a.docx'; docx.Document().save(p)\n"
        "lo.submit(p, lo.work_root / 'a.pdf', owner='a').result(240)\n"
        "kids = []\n"
        "import psutil\n"
        "for pid in lo.pids(): kids += [pid] + [c.pid for c in psutil.Process(pid).children(recursive=True)]\n"
        "print(json.dumps(kids), flush=True)\n"
        "time.sleep(300)\n")
    p = subprocess.Popen([sys.executable, str(script), str(tmp_path / "svcstate"), str(integrity.store)],
                         stdout=subprocess.PIPE, text=True)
    import json
    kids = json.loads(p.stdout.readline())
    assert kids and all(psutil.pid_exists(k) for k in kids)
    p.kill()
    p.wait(10)
    deadline = time.monotonic() + 5
    while any(psutil.pid_exists(k) for k in kids) and time.monotonic() < deadline:
        time.sleep(0.2)
    assert not [k for k in kids if psutil.pid_exists(k)]


def _split_field_docx(path, base):
    import docx
    d = docx.Document()
    d.add_paragraph("Body")
    d.save(path)
    xml = ('<w:p><w:r><w:fldChar w:fldCharType="begin"/></w:r><w:r><w:instrText> INCLUDEPIC</w:instrText></w:r>'
           '<w:r><w:instrText>TURE "' + base + '/split.png" </w:instrText></w:r><w:r><w:fldChar '
           'w:fldCharType="separate"/></w:r><w:r><w:t>cached</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/>'
           '</w:r></w:p>')
    CN.rewrite(path, lambda parts: parts.__setitem__("word/document.xml", parts["word/document.xml"].replace(
        b"<w:sectPr", xml.encode() + b"<w:sectPr", 1)))
    return path


@pytest.mark.parametrize("sanitized", [True, False], ids=["sanitized", "hardened_profile_only"])
def test_split_field_canary_makes_no_connection(lo, tmp_path, sanitized):
    """Review round 1, finding 2: a field instruction split over runs. Sanitized it is removed; unsanitized, the
    hardened profile alone must still fetch nothing."""
    can = Canary()
    census = Census(lo.pids)
    try:
        src = _split_field_docx(tmp_path / "split.docx", can.base)
        _convert(lo, src, "docx", sanitize=sanitized)
        time.sleep(1.0)
    finally:
        census.stop()
        can.close()
    assert can.hits == [] and census.smb == set() and census.proc == set(), (can.hits, census.smb, census.proc)


def test_native_chart_is_drawn(lo, tmp_path):
    """Live run 2026-10-15 regression: a native PowerPoint chart must be drawn (axis labels and series in the PDF),
    not replaced by a broken-object icon (it was, while DisableActiveContent was set)."""
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    import pypdf
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.title.text = "Scanned per month"
    cd = CategoryChartData()
    cd.categories = ["Junexample", "Julyexample"]
    cd.add_series("Scannedseries", (320, 410))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    p = tmp_path / "chart.pptx"
    pr.save(p)
    out = _convert(lo, p, "pptx")
    text = pypdf.PdfReader(str(out)).pages[0].extract_text() or ""
    assert "Junexample" in text and "Julyexample" in text, text
