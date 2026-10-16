"""Second independent Claude review of the completion candidate (e8e000bb), findings C2-01..C2-07: each test
reproduces the original defect on real files or packages."""
import hashlib
import io
import os
import sys
import time
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from openpyxl import Workbook
from openpyxl.styles import GradientFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from uri_office import xlsx_patch as X
from uri_office import xlsx_struct as S
from uri_office import xlsx_styles as ST
from uri_office.authoring import docx_ops as W
from uri_office.authoring import pptx_ops as P
from uri_office.errors import OfficeError


def sha12(b):
    return hashlib.sha256(b).hexdigest()[:12]


def _hold(dst: Path, content: bytes, recorded: bytes | None = None, age: float = 0.0) -> Path:
    h = dst.parent / f".{dst.name}.{sha12(recorded if recorded is not None else content)}-0123abcd{X.HOLD_SUFFIX}"
    h.write_bytes(content)
    if age:
        t = time.time() - age
        os.utime(h, (t, t))
    return h


# ---- C2-01: a hold is restored only under the exact name it belongs to ------------------------------------------
def test_a_hold_of_a_longer_name_is_never_restored_under_a_shorter_one(tmp_path):
    plan, v2 = tmp_path / "Plan.pptx", tmp_path / "Plan.pptx.v2.pptx"
    _hold(v2, b"ORIGINAL-OF-V2")
    assert X.recover_held(plan) is None and not plan.exists()
    assert X.recover_held(v2) == v2 and v2.read_bytes() == b"ORIGINAL-OF-V2"


# ---- C2-04: wrong or empty holds are not restored; stale leftovers are removed -----------------------------------
def test_an_empty_or_wrong_hold_is_never_restored(tmp_path):
    deck = tmp_path / "Deck.pptx"
    _hold(deck, b"", recorded=b"ORIGINAL")                        # an empty file under a real original's hash
    assert X.recover_held(deck) is None and not deck.exists()
    real = _hold(deck, b"ORIGINAL", age=5)
    assert X.recover_held(deck) == deck and deck.read_bytes() == b"ORIGINAL" and not real.exists()


def test_stale_holds_and_temp_files_are_cleaned_when_the_file_exists(tmp_path):
    deck = tmp_path / "Deck.pptx"
    deck.write_bytes(b"CURRENT")
    stale = _hold(deck, b"OLD", age=3600)
    fresh = _hold(deck, b"OTHER", age=0)
    tmp = tmp_path / ".Deck.pptx.abcd1234.uri-office.tmp"
    tmp.write_bytes(b"NEW")
    t = time.time() - 3600
    os.utime(tmp, (t, t))
    other = tmp_path / ".Deck.pptx.v2.pptx.abcd1234.uri-office.tmp"            # another file's: not ours
    other.write_bytes(b"x")
    os.utime(other, (t, t))
    assert X.recover_held(deck) is None
    assert deck.read_bytes() == b"CURRENT"
    assert not stale.exists() and not tmp.exists()
    assert fresh.exists() and other.exists()                     # young leftovers and other files' are kept


def test_during_a_save_a_young_hold_is_not_taken(tmp_path):
    deck = tmp_path / "Deck.pptx"
    _hold(deck, b"ORIGINAL", age=0)
    assert X.recover_held(deck, restore_min_age=X.STALE_SECONDS) is None and not deck.exists()


# ---- C2-02: Excel 2010+ validations on another sheet ----------------------------------------------------------------
def test_a_row_change_is_refused_when_another_sheet_has_an_x14_validation_naming_it():
    wb = Workbook()
    ls = wb.active
    ls.title = "Lists"
    for i in range(1, 11):
        ls.cell(i, 1, f"item{i}")
    wb.create_sheet("Form")["B2"] = "item3"
    b = io.BytesIO()
    wb.save(b)
    z = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    parts = {n: z.read(n) for n in z.namelist()}
    ext = (b'<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
           b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
           b'<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
           b'<x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Lists!$A$2:$A$10</xm:f>'
           b'</x14:formula1><xm:sqref>B2</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
    parts["xl/worksheets/sheet2.xml"] = parts["xl/worksheets/sheet2.xml"].replace(b"</worksheet>", ext + b"</worksheet>")
    before = dict(parts)
    with pytest.raises(OfficeError) as e:
        S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Lists", "at": 1, "count": 3})
    assert e.value.code == "STRUCTURE_NOT_SUPPORTED" and "Form" in e.value.detail
    assert parts == before


# ---- C2-03 / C2-05: Word content that is not plain text ------------------------------------------------------------
NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
CASES = {
    "ruby": f'<w:p {NS}><w:r><w:ruby><w:rubyPr/><w:rt><w:r><w:t>kan</w:t></w:r></w:rt><w:rubyBase><w:r><w:t>KANJI'
            f'</w:t></w:r></w:rubyBase></w:ruby></w:r><w:r><w:t> tail</w:t></w:r></w:p>',
    "dir": f'<w:p {NS}><w:r><w:t>start </w:t></w:r><w:dir w:val="rtl"><w:r><w:t>RTLTEXT</w:t></w:r></w:dir></w:p>',
    "ptab": f'<w:p {NS}><w:r><w:t>Left</w:t><w:ptab w:relativeTo="margin" w:alignment="right" w:leader="none"/>'
            f'<w:t>Right</w:t></w:r></w:p>',
    "pgNum": f'<w:p {NS}><w:r><w:t>Page </w:t></w:r><w:r><w:pgNum/></w:r></w:p>',
}


def _doc(xml):
    d = Document()
    body = d.element.body
    for p in list(body.iterfind(W._W + "p")):
        body.remove(p)
    body.insert(0, parse_xml(xml))
    return d


@pytest.mark.parametrize("name", sorted(CASES))
def test_paragraphs_with_more_than_text_are_locked_by_name_and_kept(name):
    d = _doc(CASES[name])
    b = W.describe(d)["blocks"][0]
    assert b["locked"] and name in b["locked"]
    before = d.element.body.xml
    with pytest.raises(OfficeError):
        W.set_paragraph_text(d, 0, b["expect"], "NEW TEXT")
    assert d.element.body.xml == before


def test_replace_counts_only_what_it_changed():
    d = _doc(CASES["ptab"])
    d.add_paragraph("Left")
    r = W.replace_text(d, "Left", "X")
    assert r["replaced"] == 1 and r["skipped"] == 1
    assert d.paragraphs[1].text == "X"


# ---- C2-06: a chart linked to another file ---------------------------------------------------------------------------
def test_duplicating_a_slide_with_a_linked_chart_is_refused_by_name():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    cd = CategoryChartData()
    cd.categories = ["a", "b"]
    cd.add_series("S", (1, 2))
    cp = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, 0, 0, Inches(4), Inches(3), cd).chart.part
    ed = cp._element.find(qn("c:externalData"))
    rid = cp.relate_to(r"file:///C:/data/sales.xlsx",
                       "http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject", is_external=True)
    ed.set(qn("r:id"), rid)
    b = io.BytesIO()
    prs.save(b)
    prs = Presentation(io.BytesIO(b.getvalue()))
    with pytest.raises(OfficeError) as e:
        P.duplicate_slide(prs, prs.slides[0].slide_id)
    assert e.value.code == "AUTHORING_NOT_SUPPORTED" and "linked" in e.value.detail
    assert len(prs.slides) == 1


# ---- C2-07: a gradient fill ---------------------------------------------------------------------------------------
def test_a_gradient_fill_is_refused_not_a_crash(tmp_path):
    wb = Workbook()
    wb.active["A1"].fill = GradientFill(stop=("FF0000", "0000FF"))
    p = tmp_path / "g.xlsx"
    wb.save(p)
    with pytest.raises(OfficeError) as e:
        ST.check_visible_fills(p, [{"sheet": "Sheet", "cell": "A1", "before": {}, "after": {"bg": "#FFFF00"}}])
    assert e.value.code == "STYLE_NOT_SUPPORTED"
