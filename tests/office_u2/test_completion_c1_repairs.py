"""Independent Claude review of the completion candidate (57e42097), findings C1-01..C1-11: each test reproduces the
original defect on real files or packages and fails before the repair."""
import hashlib
import io
import os
import stat
import subprocess
import sys
import textwrap
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula
from pptx import Presentation
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from lxml import etree

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "office"))
from test_authoring_pptx import FakeDocuments, act_body, deck_bytes  # noqa: E402
from test_xlsx_struct import parts_of  # noqa: E402

from uri_office import xlsx_patch as X
from uri_office import xlsx_struct as S
from uri_office.approvals import Approvals
from uri_office.authoring import docx_ops as W
from uri_office.authoring import pptx_ops as P
from uri_office.authoring.session import Authoring
from uri_office.errors import OfficeError
from uri_office.grants import Grants

ROOT = Path(__file__).resolve().parents[2]


def sha(b):
    return hashlib.sha256(b).hexdigest()


# ---- C1-01: an interrupted Save never leaves the owner's file missing --------------------------------------------
def test_backup_is_durable_before_the_original_is_touched_and_a_held_original_is_recovered(tmp_path):
    f = tmp_path / "Report.docx"
    f.write_bytes(b"ORIGINAL")
    bk = tmp_path / "bk.docx"
    # The process dies right after the original name was moved aside (the narrowest window left).
    script = textwrap.dedent(f"""
        import os, sys
        sys.path.insert(0, {str(ROOT)!r})
        from pathlib import Path
        from uri_office import xlsx_patch as X
        real = X._publish_no_replace
        def die(src, dst):
            os._exit(9)
        X._publish_no_replace = die
        X.atomic_write(Path({str(f)!r}), lambda fh: fh.write(b"NEW"),
                       expected_sha256={sha(b"ORIGINAL")!r}, backup=Path({str(bk)!r}))
    """)
    r = subprocess.run([sys.executable, "-B", "-c", script], capture_output=True)
    assert r.returncode == 9
    assert not f.exists()                                       # the moment of the crash
    assert bk.read_bytes() == b"ORIGINAL"                       # the backup was written and synced first
    assert X.recover_held(f) == f                               # the next start or Save puts it back
    assert f.read_bytes() == b"ORIGINAL"
    assert not list(tmp_path.glob("*.uri-office.hold"))


def test_a_changed_original_is_refused_before_anything_moves(tmp_path):
    f = tmp_path / "a.docx"
    f.write_bytes(b"CHANGED")
    with pytest.raises(OfficeError) as e:
        X.atomic_write(f, lambda fh: fh.write(b"NEW"), expected_sha256=sha(b"V1"), backup=tmp_path / "bk")
    assert e.value.code == "TARGET_CHANGED" and f.read_bytes() == b"CHANGED" and not (tmp_path / "bk").exists()
    assert sorted(p.name for p in tmp_path.iterdir()) == ["a.docx"]


# ---- C1-02: a read-only original is refused, not replaced and then reported as failed ----------------------------
def test_a_read_only_original_is_refused_and_left_alone(tmp_path):
    f = tmp_path / "a.pptx"
    f.write_bytes(b"V1")
    os.chmod(f, stat.S_IREAD)
    try:
        with pytest.raises(OfficeError) as e:
            X.atomic_write(f, lambda fh: fh.write(b"V2"), expected_sha256=sha(b"V1"), backup=tmp_path / "bk")
        assert e.value.code == "FILE_READ_ONLY"
        assert f.read_bytes() == b"V1"
        assert sorted(p.name for p in tmp_path.iterdir()) == ["a.pptx"]     # no .tmp or .hold left beside it
    finally:
        os.chmod(f, stat.S_IWRITE | stat.S_IREAD)


# ---- C1-03 / C1-04 / C1-07: Excel rows ------------------------------------------------------------------------------
def _lists_and_form():
    wb = Workbook()
    ls = wb.active
    ls.title = "Lists"
    for i in range(1, 11):
        ls.cell(i, 1, i)
    ls.conditional_formatting.add("A2:A10", FormulaRule(formula=["A2>5"]))
    form = wb.create_sheet("Form")
    dv = DataValidation(type="list", formula1="=Lists!$A$2:$A$10")
    form.add_data_validation(dv)
    dv.add("B2")
    form.conditional_formatting.add("C2", FormulaRule(formula=["Lists!$A$5>3"]))
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def _xml(parts, name):
    return parts[name].decode("utf-8")


def test_validation_and_formats_on_other_sheets_move_with_the_rows():
    parts = parts_of(_lists_and_form())
    S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Lists", "at": 4, "count": 2})
    form = _xml(parts, "xl/worksheets/sheet2.xml")
    assert "=Lists!$A$2:$A$12" in form or "Lists!$A$2:$A$12" in form
    assert "Lists!$A$7&gt;3" in form


def test_deleting_the_first_row_of_a_conditional_format_is_refused_not_broken():
    parts = parts_of(_lists_and_form())
    before = dict(parts)
    with pytest.raises(OfficeError) as e:
        S.apply_row_op(parts, {"op": "delete_rows", "sheet": "Lists", "at": 2, "count": 1})
    assert e.value.code == "STRUCTURE_NOT_SUPPORTED" and "first row" in e.value.detail
    assert parts == before
    S.apply_row_op(parts, {"op": "delete_rows", "sheet": "Lists", "at": 3, "count": 1})   # not the first row: fine
    assert "A2&gt;5" in _xml(parts, "xl/worksheets/sheet1.xml")


def test_rows_cannot_be_inserted_inside_an_array_formula(tmp_path):
    wb = Workbook()
    ws = wb.active
    for i in range(1, 6):
        ws.cell(i, 1, i)
    ws["B2"] = ArrayFormula("B2:B4", "=A2:A4*2")
    src, dst = tmp_path / "arr.xlsx", tmp_path / "out.xlsx"
    wb.save(src)
    with pytest.raises(OfficeError) as e:
        X.patch(src, dst, [], [{"op": "insert_rows", "sheet": "Sheet", "at": 3, "count": 1,
                                "values": [[{"v": 1}, {"v": 99}]]}])
    assert "array" in e.value.detail and not dst.exists()


# ---- C1-05: save-as grants live only while their card is open -------------------------------------------------------
def test_save_as_grants_are_released_when_superseded_and_when_the_edit_closes(tmp_path):
    home, extra = tmp_path / "home", tmp_path / "extra"
    home.mkdir()
    extra.mkdir()
    (tmp_path / "state").mkdir()
    grants = Grants([home], folders_file=tmp_path / "state" / "folders.json")
    grants.add_folder(str(extra))
    docs = FakeDocuments(grants)
    appr = Approvals(tmp_path / "state", "owner")
    a = Authoring(tmp_path / "state", docs, grants, appr)
    aid = a.new_presentation()["authoring_id"]
    a.propose_save(aid, str(extra / "Deck.pptx"))
    a.propose_save(aid, str(home / "Deck.pptx"))                 # supersedes the first card
    alive = [g.path for g in grants._by_id.values() if g.kind == "save_as"]
    assert alive == [home / "Deck.pptx"]
    assert grants.remove_folder(str(extra))["removed"] is True   # the folder can be revoked
    a.close(aid)
    assert not [g for g in grants._by_id.values() if g.kind == "save_as"]


# ---- C1-06 / C1-09: duplicating a slide -----------------------------------------------------------------------------
def test_duplicate_copies_chart_relationships_background_and_timing(tmp_path):
    prs = Presentation(io.BytesIO(deck_bytes()))
    s3 = prs.slides[2]
    cp = [sh for sh in s3.shapes if sh.has_chart][0].chart.part
    us = Part(PackURI("/ppt/drawings/drawing1.xml"),
              "application/vnd.openxmlformats-officedocument.drawingml.chartshapes+xml", prs.part.package,
              b'<?xml version="1.0"?><c:userShapes xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
    rid = cp.relate_to(us, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartUserShapes")
    etree.SubElement(cp._element, qn("c:userShapes")).set(qn("r:id"), rid)
    s3._element.cSld.insert(0, etree.fromstring(
        '<p:bg xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><p:bgPr><a:solidFill><a:srgbClr val="FF0000"/>'
        '</a:solidFill><a:effectLst/></p:bgPr></p:bg>'))
    s3._element.append(etree.fromstring(
        '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="slow"><p:fade/>'
        '</p:transition>'))
    b = io.BytesIO()
    prs.save(b)
    prs = Presentation(io.BytesIO(b.getvalue()))
    new = P.duplicate_slide(prs, prs.slides[2].slide_id)
    out = io.BytesIO()
    prs.save(out)
    z = zipfile.ZipFile(io.BytesIO(out.getvalue()))
    import re
    for n in sorted(x for x in z.namelist() if re.fullmatch(r"ppt/charts/chart\d+\.xml", x)):
        used = set(re.findall(rb'r:id="(rId\d+)"', z.read(n)))
        have = set(re.findall(rb'Id="(rId\d+)"', z.read(n.replace("charts/", "charts/_rels/") + ".rels")))
        assert used <= have, n
    drawings = [x for x in z.namelist() if x.startswith("ppt/drawings/drawing")]
    assert len(drawings) == 2                                     # the copy has its own user shapes part
    dup = next(s for s in Presentation(io.BytesIO(out.getvalue())).slides if s.slide_id == new)
    assert dup._element.cSld.find(qn("p:bg")) is not None
    assert dup._element.find(qn("p:transition")) is not None


# ---- C1-08 / C1-11: Word ------------------------------------------------------------------------------------------
def _para_with(xml: str):
    d = Document()
    d.add_paragraph("Before")
    p = d.add_paragraph("text ")
    p._p.append(etree.fromstring(xml))
    d.add_paragraph("After")
    return d


@pytest.mark.parametrize("xml,why", [
    ('<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:r><m:t>x</m:t></m:r></m:oMath>',
     "an equation"),
    ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:rPr><w:b/><w:rPrChange w:id="1" '
     'w:author="A"><w:rPr/></w:rPrChange></w:rPr><w:t>bold</w:t></w:r>', "a tracked change"),
    ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:sym w:font="Wingdings" '
     'w:char="F0FC"/></w:r>', "a symbol"),
])
def test_unsupported_word_content_is_named_and_kept(xml, why):
    d = _para_with(xml)
    v = W.describe(d)
    assert v["blocks"][1]["locked"] == why
    with pytest.raises(OfficeError):
        W.delete_block(d, 1, v["blocks"][1]["expect"])
    with pytest.raises(OfficeError):
        W.set_paragraph_text(d, 1, v["blocks"][1]["expect"], "x")


def test_replace_does_not_rewrite_a_run_that_holds_a_symbol():
    d = Document()
    p = d.add_paragraph()
    r = p.add_run("box")
    r._r.append(etree.fromstring('<w:sym xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
                                 'w:font="Wingdings" w:char="F0FC"/>'))
    d.add_paragraph("box")
    W.replace_text(d, "box", "tick")
    assert d.paragraphs[0]._p.find(".//" + qn("w:sym")) is not None
    assert d.paragraphs[1].text == "tick"


def test_inserts_after_a_block_check_its_address():
    d = Document()
    for t in ("a", "b", "c"):
        d.add_paragraph(t)
    v = W.describe(d)
    with pytest.raises(OfficeError) as e:
        W.OPS["insert_paragraph"](d, {"op": "insert_paragraph", "after": 1, "after_expect": "000000000000",
                                      "text": "x"})
    assert e.value.code == "STALE_BLOCK"
    with pytest.raises(KeyError):                                # the address is required by the op map
        W.OPS["insert_page_break"](d, {"op": "insert_page_break", "after": 1})
    W.OPS["insert_paragraph"](d, {"op": "insert_paragraph", "after": 1, "after_expect": v["blocks"][1]["expect"],
                                  "text": "x"})
    assert [p.text for p in d.paragraphs] == ["a", "b", "x", "c"]
