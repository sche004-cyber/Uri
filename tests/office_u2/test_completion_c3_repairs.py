"""Third independent Claude review of the completion candidate (3c2bcab6), findings C3-01..C3-05: each test
reproduces the original defect on real files or packages."""
import hashlib
import io
import os
import time
import zipfile

import pytest
from docx import Document
from docx.oxml import parse_xml
from openpyxl import Workbook

from uri_office import xlsx_patch as X
from uri_office import xlsx_struct as S
from uri_office.authoring import docx_ops as W
from uri_office.errors import OfficeError

NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _doc(*xmls):
    d = Document()
    body = d.element.body
    for p in list(body.iterfind(W._W + "p")):
        body.remove(p)
    for i, x in enumerate(xmls):
        body.insert(i, parse_xml(x))
    return d


# ---- C3-01: breaks and special hyphens are not plain text -----------------------------------------------------------
@pytest.mark.parametrize("xml,why", [
    (f'<w:p {NS}><w:r><w:t>Before</w:t><w:br w:type="page"/><w:t>After</w:t></w:r></w:p>', "a page break"),
    (f'<w:p {NS}><w:r><w:t>Before</w:t><w:br w:type="column"/><w:t>After</w:t></w:r></w:p>', "a column break"),
    (f'<w:p {NS}><w:r><w:t>co</w:t><w:noBreakHyphen/><w:t>operate</w:t></w:r></w:p>', "a non-breaking hyphen"),
    (f'<w:p {NS}><w:r><w:t>hy</w:t><w:softHyphen/><w:t>phen</w:t></w:r></w:p>', "an optional hyphen"),
])
def test_breaks_and_special_hyphens_are_locked_and_kept(xml, why):
    d = _doc(xml)
    b = W.describe(d)["blocks"][0]
    assert b["locked"] == why
    before = d.element.body.xml
    with pytest.raises(OfficeError):
        W.set_paragraph_text(d, 0, b["expect"], b["text"])       # even the unchanged text
    assert d.element.body.xml == before


def test_a_line_break_still_round_trips_and_a_page_break_paragraph_can_be_deleted():
    d = _doc(f'<w:p {NS}><w:r><w:t>one</w:t><w:br/><w:t>two</w:t></w:r></w:p>', f'<w:p {NS}><w:r><w:t>x</w:t></w:r></w:p>')
    b = W.describe(d)["blocks"][0]
    assert b["locked"] is None
    W.set_paragraph_text(d, 0, b["expect"], "one\ntwo!")
    assert d.paragraphs[0].text == "one\ntwo!"
    i = W.insert_page_break(d, None)
    v = W.describe(d)
    assert v["blocks"][i]["locked"] == "a page break"
    W.delete_block(d, i, v["blocks"][i]["expect"])               # the owner removes the break
    assert not any(br.get(W._W + "type") == "page" for br in d.element.body.iter(W._W + "br"))


def test_deleting_a_break_paragraph_refuses_other_unsupported_content():
    d = _doc(f'<w:p {NS}><w:r><w:br w:type="page"/><w:softHyphen/></w:r></w:p>',
             f'<w:p {NS}><w:r><w:t>Keep</w:t></w:r></w:p>')
    b = W.describe(d)["blocks"][0]
    assert b["locked"] == "a page break"
    before = d.element.body.xml
    with pytest.raises(OfficeError) as exc:
        W.delete_block(d, 0, b["expect"])
    assert exc.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "optional hyphen" in exc.value.detail
    assert d.element.body.xml == before


def test_replace_skips_a_run_with_a_page_break():
    d = _doc(f'<w:p {NS}><w:r><w:t>Intro</w:t><w:br w:type="page"/><w:t>Chapter Two</w:t></w:r></w:p>')
    r = W.replace_text(d, "Two", "X")
    assert r == {"replaced": 0, "split_across_formatting": 0, "skipped": 1}
    assert any(br.get(W._W + "type") == "page" for br in d.element.body.iter(W._W + "br"))
    assert "Two" in W.block_text(W.blocks(d)[0])


# ---- C3-02: long file names can still be saved --------------------------------------------------------------------
def test_save_works_for_a_long_file_name(tmp_path):
    name = "L" * (225 - len(".pptx")) + ".pptx"
    f = tmp_path / name
    try:
        f.write_bytes(b"V1")
    except OSError:
        pytest.skip("this volume does not allow such long names")
    X.atomic_write(f, lambda fh: fh.write(b"V2"), expected_sha256=hashlib.sha256(b"V1").hexdigest(),
                   backup=tmp_path / "bk")
    assert f.read_bytes() == b"V2"


# ---- C3-03: a fresh hold is fresh even when the original was old ---------------------------------------------------
def test_a_hold_is_aged_by_its_creation_not_by_the_original_mtime(tmp_path):
    deck = tmp_path / "Deck.pptx"
    deck.write_bytes(b"ORIGINAL")
    old = time.time() - 3600
    os.utime(deck, (old, old))
    h = X._hold_name(deck, hashlib.sha256(b"ORIGINAL").hexdigest()[:12], int(time.time()), "0123abcd")
    os.replace(deck, h)                                          # as Save does: the rename keeps the old mtime
    assert X.recover_held(deck, restore_min_age=X.STALE_SECONDS) is None and not deck.exists()
    assert X.recover_held(deck) == deck                          # service start: nothing in progress


# ---- C3-04: bytes that are not the checked original are kept, never removed ------------------------------------------
def test_another_programs_bytes_left_aside_are_kept_and_restorable(tmp_path, monkeypatch):
    deck = tmp_path / "Deck.pptx"
    deck.write_bytes(b"ORIGINAL")
    real_replace = os.replace

    def other_program_writes_then_move(src, dst):
        if str(src) == str(deck) and str(dst).endswith(X.HOLD_SUFFIX):
            deck.write_bytes(b"OTHER-PROGRAM")               # written just before Save moves the file aside
        return real_replace(src, dst)
    monkeypatch.setattr(X.os, "replace", other_program_writes_then_move)
    real_publish = X._publish_no_replace

    def put_back_fails(src, dst):
        if str(src).endswith(X.HOLD_SUFFIX):
            raise PermissionError("put back failed")
        return real_publish(src, dst)
    monkeypatch.setattr(X, "_publish_no_replace", put_back_fails)
    with pytest.raises(OfficeError) as e:
        X.atomic_write(deck, lambda fh: fh.write(b"NEW"), expected_sha256=hashlib.sha256(b"ORIGINAL").hexdigest(),
                       backup=tmp_path / "bk")
    assert e.value.code == "TARGET_CHANGED"
    monkeypatch.undo()
    keeps = list(tmp_path.glob("*" + X.KEEP_SUFFIX))
    assert len(keeps) == 1 and keeps[0].read_bytes() == b"OTHER-PROGRAM"
    deck.write_bytes(b"RECREATED")
    later = time.time() + 3600
    os.utime(keeps[0], (later, later))
    assert X.recover_held(deck, cleanup_min_age=0) is None and keeps[0].exists()   # never removed
    deck.unlink()
    assert X.recover_held(deck) == deck and deck.read_bytes() == b"OTHER-PROGRAM"


# ---- C3-05: a sheet whose name ends another sheet's name -------------------------------------------------------------
def test_x14_check_matches_the_sheet_name_not_a_longer_one():
    wb = Workbook()
    wb.active.title = "Sheet1"
    wb.create_sheet("MySheet1")
    wb.create_sheet("Other")
    b = io.BytesIO()
    wb.save(b)
    z = zipfile.ZipFile(io.BytesIO(b.getvalue()))
    parts = {n: z.read(n) for n in z.namelist()}

    def ext(ref):
        return (b'<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" '
                b'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
                b'<x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
                b'<x14:dataValidation type="list"><x14:formula1><xm:f>' + ref + b'</xm:f></x14:formula1>'
                b'<xm:sqref>B2</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>')
    other = "xl/worksheets/sheet3.xml"
    base = parts[other]
    parts[other] = base.replace(b"</worksheet>", ext(b"MySheet1!$A$1:$A$5") + b"</worksheet>")
    S.apply_row_op(dict(parts), {"op": "insert_rows", "sheet": "Sheet1", "at": 1, "count": 1})   # not refused
    parts[other] = base.replace(b"</worksheet>", ext(b"Sheet1!$A$1:$A$5") + b"</worksheet>")
    with pytest.raises(OfficeError):
        S.apply_row_op(dict(parts), {"op": "insert_rows", "sheet": "Sheet1", "at": 1, "count": 1})
