"""Plan 16 increment 1, OD-A: Notesheet and Circular layouts (OD-1), Brain drafting into Word (OD-2), PDF text out and
OCR on request (OD-3). The Brain is a fake client (no Runtime is started); the engine-marked tests skip cleanly when
the pinned LibreOffice is absent; the OCR test runs the installed Tesseract when it is present and skips otherwise."""
from __future__ import annotations

import io
import json
import threading
from pathlib import Path
from types import SimpleNamespace

import jsonschema
import pytest
from docx import Document

from conftest import pdf_with_text
from test_authoring_docx import FakeDocs
from test_authoring_pptx import act_body
from test_service_o2 import _seed, _wait, call, engine, verified_manifest  # noqa: F401 (fixture)
from uri_office import brain as B
from uri_office import capabilities as CAP
from uri_office import contracts as K
from uri_office import intent as I
from uri_office import ocr as OCR
from uri_office.authoring import docx_templates as T
from uri_office.authoring import drafting as DR
from uri_office.brain import OfficeBrain
from uri_office.docs import DocSession, Documents
from uri_office.errors import OfficeError
from uri_office.service import Office, serve

REAL_OCR_PNG = OCR.ocr_png


def load(data: bytes):
    return Document(io.BytesIO(data))


# ---- OD-1: layouts ---------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("lid,title", [("notesheet", "Note Sheet"), ("circular", "Circular")])
def test_notesheet_and_circular_layouts_open_with_placeholders_only(lid, title):
    assert lid in [x["id"] for x in T.layout_list()]
    d = load(T.from_layout(lid))
    paras = [(p.style.name, p.text) for p in d.paragraphs]
    assert paras[0] == ("Title", title)
    assert round(d.sections[0].page_width.cm, 1) == 21.0 and round(d.sections[0].page_height.cm, 1) == 29.7
    text = "\n".join(t for _, t in paras)
    assert not any(c.isdigit() for c in text)                       # no pre-filled numbers or dates
    assert "[date]" in text and "Subject: [subject]" in text
    if lid == "notesheet":
        assert paras[1] == ("Normal", "File No. [file number]")
        assert [s for s, _ in paras].count("List Number") >= 2          # numbered paragraphs of the note
        assert {"[Signature]", "[Designation]"} <= {t for _, t in paras}
    else:
        assert any(t.startswith("No. [number]") and "Date: [date]" in t for _, t in paras)
        assert ("Heading 2", "Distribution") in paras and ("List Bullet", "[Recipient]") in paras
        assert "[Designation of the issuing authority]" in text


@pytest.fixture
def env(tmp_path):
    from uri_office.approvals import Approvals
    from uri_office.authoring.session import Authoring
    from uri_office.grants import Grants
    (tmp_path / "state").mkdir()
    grants = Grants([tmp_path])
    appr = Approvals(tmp_path / "state", "owner")
    return Authoring(tmp_path / "state", FakeDocs(grants), grants, appr), appr, tmp_path


@pytest.mark.parametrize("lid", ["notesheet", "circular"])
def test_new_notesheet_and_circular_save_as_a_new_file(env, lid):
    a, appr, tmp = env
    v = a.new_document(lid)
    assert v["new"] and v["name"] == T.LAYOUTS[lid][0] + ".docx"
    dest = tmp / f"{lid}.docx"
    a.apply_save(act_body(appr, a.propose_save(v["authoring_id"], str(dest))))
    assert load(dest.read_bytes()).paragraphs[0].text == T.LAYOUTS[lid][1][0][1]


@engine
def test_notesheet_and_circular_render_in_libreoffice(svc_real, verified_manifest):  # noqa: F811
    office, port, token, _ = svc_real
    _seed(office, verified_manifest)
    for lid in ("notesheet", "circular"):
        st, v = call(port, token, "POST", "/authoring/new", {"kind": "docx", "layout": lid})
        assert st == 200, v
        pv = _wait(lambda: (lambda x: x if x["preview"]["state"] != "rendering" else None)(
            call(port, token, "GET", f"/document/{v['preview_document_id']}")[1]))
        assert pv["preview"]["state"] == "ready" and pv["pages"] >= 1, (lid, pv)
        call(port, token, "POST", f"/authoring/{v['authoring_id']}/close", {})


# ---- OD-2: deterministic parts ---------------------------------------------------------------------------------------
@pytest.mark.parametrize("text,lid", [
    ("Write an office order transferring the section officer", "office_order"),
    ("draft a note sheet proposing purchase of laptops", "notesheet"),
    ("Put up a notesheet for approval", "notesheet"),
    ("prepare a circular to all staff about the holiday", "circular"),
    ("write a memo to the accounts section", "memo"),
    ("Draft an office memorandum on leave", "memo"),
    ("write the minutes of the review meeting", "minutes"),
    ("issue a notice for the staff meeting", "notice"),
    ("write a letter to the district collector", "letter"),
    ("prepare a report on the audit", "report"),
    ("write an office order asking staff to report at 10 am", "office_order"),     # earliest keyword wins
    ("please write something nice", None),
    ("draft an order", None),                                                        # "order" alone is not a kind
])
def test_layout_is_chosen_from_keywords(text, lid):
    assert DR.choose_layout(text) == lid


INSTR = ("Write an office order No. 12/3/2026-Admin dated 23 October 2026 transferring Shri Ramesh Kumar, Section "
         "Officer, to the Accounts Section with effect from 01.11.2026. Pay Rs. 25,000 advance.")


def test_facts_are_extracted_deterministically():
    facts = DR.extract_facts(INSTR)
    assert facts == [
        {"kind": "reference", "text": "12/3/2026-Admin"}, {"kind": "date", "text": "23 October 2026"},
        {"kind": "name", "text": "Shri Ramesh Kumar"}, {"kind": "name", "text": "Section Officer"},
        {"kind": "name", "text": "Accounts Section"}, {"kind": "date", "text": "01.11.2026"},
        {"kind": "amount", "text": "25,000"}]
    assert DR.extract_facts("Draft a note sheet for File No. A-12011/3/2026 to buy 5 laptops for the Ministry of "
                            "Home Affairs") == [{"kind": "reference", "text": "A-12011/3/2026"},
                                                {"kind": "number", "text": "5"},
                                                {"kind": "name", "text": "Ministry of Home Affairs"}]
    assert DR.extract_facts("please write something nice") == []


def test_missing_facts_are_named_never_silent():
    draft = ("Office Order No. 12/3/2026-Admin Date: 23 October 2026. Shri Ramesh Kumar, Section Officer, is "
             "transferred to the accounts section with effect from 01.11.2026.")
    w = DR.missing_facts(INSTR, draft)
    assert [(x["code"], x["kind"], x["text"]) for x in w] == [("FACT_MISSING", "amount", "25,000")]
    assert "25,000" in w[0]["message"]
    assert DR.missing_facts(INSTR, draft.replace("23 October 2026", "23 Oct 2026") + " 25,000")[0]["text"] == \
        "23 October 2026"                                       # literal: a changed date is reported


# ---- OD-2: the Brain call -------------------------------------------------------------------------------------------
class FakeClient:
    def __init__(self, *answers, status="COMPLETE"):
        self.answers, self.status, self.bodies = list(answers), status, []

    def stream_chat(self, body, *, cancel=None, deadlines=None):
        self.bodies.append(body)
        a = self.answers.pop(0) if self.answers else {}
        return SimpleNamespace(status=self.status, text=a if isinstance(a, str) else json.dumps(a), ttft_s=0.1,
                               total_s=0.2, timings={})


def answer(status="propose", blocks=None, message="", missing=None):
    return {"status": status, "message": message, "blocks": blocks if blocks is not None else [], "missing":
            missing or []}


ORDER_BLOCKS = [
    {"role": "title", "text": "Office Order"},
    {"role": "paragraph", "text": "No. 12/3/2026-Admin                Date: 23 October 2026"},
    {"role": "heading", "text": "Subject: Transfer of Shri Ramesh Kumar, Section Officer"},
    {"role": "numbered", "text": "Shri Ramesh Kumar, Section Officer, is transferred to the Accounts Section with "
                                 "effect from 01.11.2026."},
    {"role": "numbered", "text": "An advance of Rs. 25,000 is sanctioned."},
    {"role": "paragraph", "text": "[Name]"},
    {"role": "bullet", "text": "[Recipient]"},
]


def test_draft_request_is_fast_tier_strict_json_and_bounded():
    c = FakeClient(answer(blocks=ORDER_BLOCKS))
    r = OfficeBrain(c).draft_document("Write an office order", layout_title="Office order",
                                      layout_paras=T.LAYOUTS["office_order"][1])
    body = c.bodies[0]
    fmt = body["response_format"]["json_schema"]
    assert fmt["strict"] is True and fmt["schema"] is B.DRAFT_SCHEMA
    assert body["max_tokens"] == B.DRAFT_MAX_TOKENS
    assert body["temperature"] == 0
    assert "Never invent" in body["messages"][0]["content"]
    assert "UNTRUSTED DATA" in body["messages"][1]["content"]            # the layout goes in as data
    jsonschema.validate(answer(blocks=ORDER_BLOCKS), B.DRAFT_SCHEMA)
    assert r["status"] == "propose" and len(r["blocks"]) == len(ORDER_BLOCKS)


@pytest.mark.parametrize("bad", [
    answer(blocks=[{"role": "title", "text": "x"}] * (B.DRAFT_MAX_BLOCKS + 1)),
    answer(blocks=[{"role": "paragraph", "text": "x" * (B.DRAFT_MAX_BLOCK_CHARS + 1)}]),
    answer(blocks=[{"role": "paragraph", "text": "x" * 1400}] * 15),                   # total cap
    answer(blocks=[{"role": "table", "text": "x"}]),
    answer(status="maybe", blocks=[{"role": "title", "text": "x"}]),
    answer(blocks=[]),                                                                 # propose with nothing
    "not json",
])
def test_over_cap_or_malformed_drafts_are_refused_not_truncated(bad):
    with pytest.raises(OfficeError) as e:
        OfficeBrain(FakeClient(bad)).draft_document("Write a letter", layout_title="Letter", layout_paras=[])
    assert e.value.code == "BRAIN_OUTPUT_INVALID"


def test_unreachable_brain_is_a_named_refusal():
    with pytest.raises(OfficeError) as e:
        OfficeBrain(FakeClient(status="UNREACHABLE")).draft_document("Write a letter", layout_title="Letter",
                                                                     layout_paras=[])
    assert e.value.code == "BRAIN_UNAVAILABLE" and e.value.status == 503


# ---- OD-2: through the service ---------------------------------------------------------------------------------------
class FakeHub:
    def states(self):
        return {}

    def shutdown(self):
        pass


def office_with(tmp_path, client):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path], brain=OfficeBrain(client) if client else None)
    o._hub = FakeHub()
    o._documents = FakeDocs(o.grants)
    return o


def test_propose_opens_a_new_word_working_copy_in_the_chosen_layout(tmp_path):
    c = FakeClient(answer(blocks=ORDER_BLOCKS, message="Drafted."))
    o = office_with(tmp_path, c)
    before = sorted(p.name for p in tmp_path.iterdir())
    v = o.draft_document({"instruction": INSTR})
    assert v["status"] == "propose" and v["layout"] == "office_order" and v["new"] is True
    assert v["name"] == "Office order (draft).docx" and v["kind"] == "docx" and v["changes"] == []
    assert v["warnings"] == []                                         # every fact is in the draft
    d = load(o.authoring.get(v["authoring_id"]).data)
    assert [(p.style.name, p.text) for p in d.paragraphs][:4] == [
        ("Title", "Office Order"), ("Normal", ORDER_BLOCKS[1]["text"]), ("Heading 1", ORDER_BLOCKS[2]["text"]),
        ("List Number", ORDER_BLOCKS[3]["text"])]
    assert d.paragraphs[-1].style.name == "List Bullet"
    assert not any("[Text of the order.]" in p.text for p in d.paragraphs)    # the placeholder body is replaced
    assert sorted(p.name for p in tmp_path.iterdir()) == before                # nothing written beside the files
    assert "Office order" in c.bodies[0]["messages"][1]["content"]
    with pytest.raises(OfficeError) as e:
        o.authoring.propose_save(v["authoring_id"])
    assert e.value.code == "SAVE_AS_REQUIRED"                            # saved only through the Save as card


def test_fact_missing_warning_names_the_fact(tmp_path):
    blocks = [b for b in ORDER_BLOCKS if "25,000" not in b["text"]]
    v = office_with(tmp_path, FakeClient(answer(blocks=blocks))).draft_document({"instruction": INSTR})
    assert v["status"] == "propose"
    assert [(w["code"], w["text"]) for w in v["warnings"]] == [("FACT_MISSING", "25,000")]


@pytest.mark.parametrize("status,msg", [("clarify", "Who is transferred?"), ("cannot", "That is not a document.")])
def test_clarify_and_cannot_open_nothing(tmp_path, status, msg):
    o = office_with(tmp_path, FakeClient(answer(status=status, message=msg, blocks=ORDER_BLOCKS,
                                                missing=["name of the officer"])))
    v = o.draft_document({"instruction": "write an office order"})
    assert v["status"] == status and v["message"] == msg and "authoring_id" not in v
    assert v["missing"] == ["name of the officer"] and o.authoring.sessions == {}


def test_no_layout_keyword_asks_which_kind_without_calling_the_brain(tmp_path):
    c = FakeClient()
    v = office_with(tmp_path, c).draft_document({"instruction": "please write something about the holiday"})
    assert v["status"] == "clarify" and c.bodies == []
    assert {"notesheet", "circular", "office_order"} <= {x["id"] for x in v["layouts"]}
    assert "blank" not in {x["id"] for x in v["layouts"]}


def test_explicit_layout_overrides_keywords_and_unknown_layout_is_refused(tmp_path):
    c = FakeClient(answer(blocks=[{"role": "title", "text": "Note Sheet"}, {"role": "numbered", "text": "Para."}]))
    o = office_with(tmp_path, c)
    v = o.draft_document({"instruction": "write about the laptops", "layout": "notesheet"})
    assert v["layout"] == "notesheet" and v["name"] == "Note sheet (draft).docx"
    for bad in ("nope", "blank", 3):
        with pytest.raises(OfficeError) as e:
            o.draft_document({"instruction": "write a letter", "layout": bad})
        assert e.value.code == "UNKNOWN_LAYOUT"


def test_no_brain_is_a_named_refusal(tmp_path):
    with pytest.raises(OfficeError) as e:
        office_with(tmp_path, None).draft_document({"instruction": "write a letter"})
    assert e.value.code == "BRAIN_UNAVAILABLE" and e.value.status == 503
    for bad in ({}, {"instruction": "  "}, {"instruction": "x" * 2001}):
        with pytest.raises(OfficeError):
            office_with(tmp_path, FakeClient()).draft_document(bad)


def test_write_document_is_available_and_routes_to_the_draft_route():
    cap = CAP.get("write_document")
    assert cap["status"] == "available" and cap["route"] == "/authoring/draft"
    assert CAP.get("edit_document")["status"] == "planned:O3"
    r = I.route("write an office order transferring the section officer", None, [])
    assert r["result"] == "matched" and r["intent"] == "write_document" and r["route"] == "/authoring/draft"
    assert I.route("change the wording of this document", "docx", [])["result"] != "matched" or \
        I.route("change the wording of this document", "docx", [])["intent"] != "edit_document"


# ---- HTTP: route security and shapes ---------------------------------------------------------------------------------
@pytest.fixture
def svc_draft(tmp_path):
    office = office_with(tmp_path, FakeClient(answer(blocks=ORDER_BLOCKS)))
    srv, token = serve(office, "127.0.0.1", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield office, srv.server_address[1], token, tmp_path
    srv.shutdown()
    srv.server_close()


@pytest.fixture
def svc_real(tmp_path):
    office = Office(tmp_path / "state", allowed_roots=[tmp_path])
    srv, token = serve(office, "127.0.0.1", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield office, srv.server_address[1], token, tmp_path
    srv.shutdown()
    srv.server_close()
    if office._hub is not None:
        office._hub.shutdown()


DID = "0123456789abcdef"


@pytest.mark.parametrize("path", ["/authoring/draft", f"/document/{DID}/text", f"/document/{DID}/to_word"])
def test_new_routes_refuse_bad_token_host_and_origin(svc_draft, path):
    _, port, token, _ = svc_draft
    assert call(port, " wrong", "POST", path, {})[0] == 401
    assert call(port, token, "POST", path, {}, {"Host": "evil.example"})[0] == 421
    assert call(port, token, "POST", path, {}, {"Origin": "http://x"})[0] == 403


def test_draft_route_end_to_end_with_a_fake_brain(svc_draft):
    office, port, token, _ = svc_draft
    st, v = call(port, token, "POST", "/authoring/draft", {"instruction": INSTR})
    assert st == 200 and v["status"] == "propose" and v["doc"]["blocks"][0]["text"] == "Office Order"
    assert st == 200 and v["warnings"] == [] and v["preview_document_id"]
    st, err = call(port, token, "POST", "/authoring/draft", {"instruction": ""})
    assert st == 400 and err["error"] == "EMPTY_INSTRUCTION" and err["message"]


def test_draft_route_without_brain_names_the_refusal(svc_real):
    _, port, token, _ = svc_real
    st, err = call(port, token, "POST", "/authoring/draft", {"instruction": "write a letter"})
    assert st == 503 and err["error"] == "BRAIN_UNAVAILABLE" and "Brain" in err["message"]
    for verb in ("text", "to_word"):
        st, err = call(port, token, "POST", f"/document/{DID}/{verb}", {})
        assert st == 404 and err["error"] == "UNKNOWN_DOCUMENT"


def test_contract_lists_the_new_document_routes_and_error_texts():
    assert ("POST", "/document/<id>/text") in K.ROUTES and ("POST", "/document/<id>/to_word") in K.ROUTES
    for code in ("BRAIN_UNAVAILABLE", "OCR_UNAVAILABLE", "NOT_A_PDF", "EMPTY_INSTRUCTION"):
        assert K.ERRORS[code].message.endswith(".")


# ---- OD-3: PDF text --------------------------------------------------------------------------------------------------
def test_text_worker_returns_page_text_with_a_total_cap(tmp_path, work):
    from test_workers import text_worker
    pdf = pdf_with_text(work / "a.pdf", ["Office order 41", "Annexure page two", "Notice of meeting"])
    w = text_worker(tmp_path / "state", work)
    try:
        r = w.request("text", pdf=str(pdf), max_chars=2_000_000)
        assert r["ok"] and [p["page"] for p in r["pages"]] == [1, 2, 3] and r["truncated"] is False
        assert "Annexure page two" in r["pages"][1]["text"] and r["chars"] == sum(len(p["text"]) for p in r["pages"])
        r = w.request("text", pdf=str(pdf), max_chars=20)
        assert r["truncated"] is True and r["chars"] == 20 and len(r["pages"]) == 2
        r = w.request("text", pdf=str(pdf), max_chars=10**9)                 # the worker's own cap still holds
        assert r["truncated"] is False
        assert w.request("text", pdf=str(tmp_path / "x.pdf"))["error"] == "PATH_REFUSED"
    finally:
        w.stop()


class TextHub:
    """The worker answers and page images, without engines."""

    def __init__(self, pages, png=None):
        self.pages, self.png, self.rendered = pages, png, []

    def text_request(self, op, **kw):
        assert op == "text"
        cap = kw["max_chars"]
        out, total, trunc = [], 0, False
        for i, t in enumerate(self.pages):
            if total + len(t) > cap:
                t, trunc = t[:cap - total], True
            out.append({"page": i + 1, "text": t, "failed": False})
            total += len(t)
            if trunc:
                break
        return {"ok": True, "pages": out, "chars": total, "truncated": trunc}

    def page_png(self, pdf, cache, n, width):
        self.rendered.append((n, width))
        return self.png

    def cache_dir(self, *a):
        return Path(".")


def docs_with(tmp_path, pages, kind="pdf", png=None):
    hub = TextHub(pages, png)
    d = Documents.__new__(Documents)
    d.hub, d.audit, d.sessions = hub, (lambda *a: None), {}
    s = DocSession(DID, kind, "Scan.pdf", None, tmp_path / "Scan.pdf", "0" * 64, 1, tmp_path)
    s.state, s.preview_pdf, s.pages = "ready", tmp_path / "Scan.pdf", len(pages)
    d.sessions[DID] = s
    return d, hub


def test_document_text_is_page_labelled_and_names_pages_without_text(tmp_path):
    d, _ = docs_with(tmp_path, ["First page text", "   ", "Third"])
    r = d.text(DID)
    assert r["text"] == "--- Page 1 ---\nFirst page text\n\n--- Page 2 ---\n\n\n--- Page 3 ---\nThird"
    assert r["pages_without_text"] == [2] and r["ocr_used"] is False and r["truncated"] is False
    assert [p["source"] for p in r["pages"]] == ["text", "none", "text"]


def test_document_text_refuses_non_pdf_and_unready(tmp_path):
    d, _ = docs_with(tmp_path, ["x"], kind="docx")
    with pytest.raises(OfficeError) as e:
        d.text(DID)
    assert e.value.code == "NOT_A_PDF"
    d, _ = docs_with(tmp_path, ["x"])
    d.sessions[DID].state = "rendering"
    with pytest.raises(OfficeError) as e:
        d.text(DID)
    assert e.value.code == "PAGE_NOT_READY"


def test_ocr_runs_only_on_request_and_is_refused_by_name_without_tesseract(tmp_path, monkeypatch):
    d, hub = docs_with(tmp_path, ["text", ""], png=tmp_path / "p.png")
    calls = []
    monkeypatch.setattr(OCR, "tesseract_exe", lambda: Path("tess.exe"))
    monkeypatch.setattr(OCR, "ocr_png", lambda png, exe, **kw: calls.append(png) or "Scanned words")
    assert d.text(DID)["pages_without_text"] == [2] and calls == [] and hub.rendered == []
    r = d.text(DID, ocr=True)
    assert r["pages"][1] == {"page": 2, "text": "Scanned words", "source": "ocr", "failed": False}
    assert hub.rendered == [(2, OCR.OCR_WIDTH)] and r["pages_without_text"] == [] and r["ocr_used"]
    monkeypatch.setattr(OCR, "tesseract_exe", lambda: None)
    with pytest.raises(OfficeError) as e:
        d.text(DID, ocr=True)
    assert e.value.code == "OCR_UNAVAILABLE"
    with pytest.raises(OfficeError) as e:
        REAL_OCR_PNG(tmp_path / "p.png")
    assert e.value.code == "OCR_UNAVAILABLE"


def test_ocr_page_cap_lists_skipped_pages(tmp_path, monkeypatch):
    from uri_office import docs as DOCS
    d, _ = docs_with(tmp_path, [""] * (DOCS.MAX_OCR_PAGES + 3), png=tmp_path / "p.png")
    monkeypatch.setattr(OCR, "tesseract_exe", lambda: Path("tess.exe"))
    monkeypatch.setattr(OCR, "ocr_png", lambda png, exe, **kw: "w")
    r = d.text(DID, ocr=True)
    assert r["ocr_skipped"] == [DOCS.MAX_OCR_PAGES + 1, DOCS.MAX_OCR_PAGES + 2, DOCS.MAX_OCR_PAGES + 3]


@pytest.mark.skipif(OCR.tesseract_exe() is None, reason="Tesseract is not installed on this PC")
def test_real_tesseract_reads_a_page_image_in_a_job_object(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("L", (1240, 300), 255)
    try:
        font = ImageFont.load_default(size=64)
    except TypeError:                                  # Pillow without FreeType sizes
        pytest.skip("Pillow cannot draw a large enough font here")
    ImageDraw.Draw(img).text((40, 100), "OFFICE ORDER 41", fill=0, font=font)
    png = tmp_path / "p.png"
    img.save(png)
    assert "OFFICE ORDER 41" in OCR.ocr_png(png)


def test_to_word_makes_a_new_working_copy_with_page_headings(tmp_path):
    o = office_with(tmp_path, None)
    o._documents, _ = docs_with(tmp_path, ["Line one\nLine two", "", "Last"])
    o._authoring = None
    from uri_office.authoring.session import Authoring
    o._authoring = Authoring(o.state, FakeDocs(o.grants), o.grants, o.approvals)
    v = o.document_to_word(DID, {})
    assert v["new"] and v["name"] == "Scan (text).docx" and v["pages_without_text"] == [2]
    d = load(o.authoring.get(v["authoring_id"]).data)
    assert [(p.style.name, p.text) for p in d.paragraphs] == [
        ("Title", "Scan"), ("Heading 1", "Page 1"), ("Normal", "Line one"), ("Normal", "Line two"),
        ("Heading 1", "Page 2"), ("Normal", "[No text on this page]"), ("Heading 1", "Page 3"), ("Normal", "Last")]


def test_xml_unsafe_characters_from_pdf_text_are_dropped():
    d = load(T.from_paragraphs([("Normal", "a\x01b\x0bc")]))
    assert d.paragraphs[0].text == "abc"


def test_capabilities_say_drafting_needs_the_brain(tmp_path):
    caps = {c["id"]: c for c in office_with(tmp_path, None).capabilities()["capabilities"]}
    assert caps["write_document"]["ready"] is False and caps["write_document"]["unready_reason"] == "BRAIN_UNAVAILABLE"
    assert caps["new_workbook_from_files"]["ready"] is True                    # unchanged: no draft route
    caps = {c["id"]: c for c in office_with(tmp_path, FakeClient()).capabilities()["capabilities"]}
    assert caps["write_document"]["ready"] is True
