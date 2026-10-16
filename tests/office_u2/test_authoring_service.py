"""Milestones A/B through the real HTTP service: route security for the authoring routes, and end to end with the real
engine (marked `engine`): open a presentation, edit it, the preview renders the working copy, the save card replaces
the file only after approval (AB-1, AB-5, AB-6, AB-8)."""
from __future__ import annotations

import hashlib
import io
import threading

import pytest
from pptx import Presentation

from test_authoring_pptx import deck_bytes
from test_service_o2 import _run_card, _seed, _wait, call, engine, verified_manifest  # noqa: F401 (fixture)
from uri_office.service import Office, serve


@pytest.fixture
def svc(tmp_path):
    office = Office(tmp_path / "state", allowed_roots=[tmp_path])
    srv, token = serve(office, "127.0.0.1", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield office, srv.server_address[1], token, tmp_path
    srv.shutdown()
    srv.server_close()
    if office._hub is not None:
        office._hub.shutdown()


AID = "0123456789abcdef"
ROUTES = [("GET", f"/authoring/{AID}"), ("POST", "/authoring/edit"), ("POST", "/authoring/new"),
          ("POST", f"/authoring/{AID}/op"), ("POST", f"/authoring/{AID}/undo"), ("POST", f"/authoring/{AID}/redo"),
          ("POST", f"/authoring/{AID}/save"), ("POST", f"/authoring/{AID}/close")]


@pytest.mark.parametrize("method,path", ROUTES)
def test_authoring_routes_refuse_bad_token_host_and_origin(svc, method, path):
    _, port, token, _ = svc
    body = {} if method == "POST" else None
    assert call(port, token=" wrong", method=method, path=path, body=body)[0] == 401
    assert call(port, token, method, path, body, {"Host": "evil.example"})[0] == 421
    assert call(port, token, method, path, body, {"Origin": "http://x"})[0] == 403


def test_unknown_edits_and_files_outside_the_folders(svc, tmp_path):
    _, port, token, root = svc
    assert call(port, token, "GET", f"/authoring/{AID}")[1]["error"] == "UNKNOWN_EDIT"
    assert call(port, token, "GET", "/authoring/zz")[0] == 404
    st, out = call(port, token, "POST", "/authoring/new", {"template_path": str(root.anchor) + "Windows\\x.potx"})
    assert st in (400, 403) and out["error"] == "PATH_REFUSED"


@engine
def test_edit_a_presentation_preview_and_save_end_to_end(svc, verified_manifest):
    office, port, token, root = svc
    _seed(office, verified_manifest)
    src = root / "Quarterly Review.pptx"
    src.write_bytes(deck_bytes())
    before = src.read_bytes()
    st, doc = call(port, token, "POST", "/open/document", {"path": str(src)})
    assert st == 200 and doc["kind"] == "pptx"
    st, v = call(port, token, "POST", "/authoring/edit", {"document_id": doc["document_id"]})
    assert st == 200 and len(v["deck"]["slides"]) == 4
    aid, first, chart = v["authoring_id"], v["deck"]["slides"][0]["id"], v["deck"]["slides"][2]["id"]
    st, v = call(port, token, "POST", f"/authoring/{aid}/op",
                 {"op": {"op": "set_title", "slide": first, "text": "Quarterly Review 2026"}})
    assert st == 200
    st, v = call(port, token, "POST", f"/authoring/{aid}/op", {"op": {"op": "duplicate_slide", "slide": chart}})
    assert st == 200 and len(v["deck"]["slides"]) == 5
    pid = v["preview_document_id"]
    pv = _wait(lambda: (lambda x: x if x["preview"]["state"] != "rendering" else None)(
        call(port, token, "GET", f"/document/{pid}")[1]))
    assert pv["preview"]["state"] == "ready" and pv["pages"] == 5            # the engine renders the working copy
    st, png, hdr = call(port, token, "GET", f"/document/{pid}/page/1?w=440", raw=True)
    assert st == 200 and png[:4] == b"\x89PNG"
    assert src.read_bytes() == before                                        # nothing written yet
    st, card = call(port, token, "POST", f"/authoring/{aid}/save", {})
    assert st == 200 and "Quarterly Review 2026" in str(card["messages"])
    st, r = _run_card(port, token, card["messages"], "apply")
    assert st == 200 and r["status"] == "applied", r
    saved = Presentation(io.BytesIO(src.read_bytes()))
    assert saved.slides[0].shapes.title.text == "Quarterly Review 2026" and len(saved.slides) == 5
    assert sum(1 for s in saved.slides for sh in s.shapes if sh.has_chart) == 2
    assert _run_card(port, token, card["messages"], "apply")[0] != 200       # a used card is refused
    backups = list((office.state / "authoring" / "backups").rglob("*.pptx"))
    assert [hashlib.sha256(b.read_bytes()).hexdigest() for b in backups] == [hashlib.sha256(before).hexdigest()]
    st, c = call(port, token, "POST", f"/authoring/{aid}/close", {})
    assert st == 200 and c["unsaved_changes"] == 0
