"""Milestone C (plan 12_AUTHORING_C_WORD_PLAN.md): Word ops and the authoring session for .docx (C-1..C-7)."""
import copy
import hashlib
import io
import json
import zipfile

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm

from test_authoring_pptx import FakeDocuments, act_body, card_lines
from uri_office.approvals import Approvals
from uri_office.authoring import docx_ops as W
from uri_office.authoring import docx_templates as T
from uri_office.authoring.session import Authoring
from uri_office.errors import OfficeError
from uri_office.grants import Grants


def report_bytes(png_path=None) -> bytes:
    d = Document()
    d.add_heading("Annual Report", 0)                              # 0
    d.add_heading("Summary", 1)                                     # 1
    p = d.add_paragraph()                                           # 2
    r = p.add_run("Revenue grew by ten percent.")
    r.bold = True
    d.add_paragraph("Costs were flat across the year.")             # 3
    t = d.add_table(rows=2, cols=2)                                 # 4
    t.style = d.styles["Table Grid"]
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Cost", "Paper", "100"
    commented = d.add_paragraph("Reviewed by the committee.")       # 5
    d.add_comment(commented.runs, text="Check the date", author="Owner")
    tracked = d.add_paragraph("Draft line ")                        # 6
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "91")
    ins.set(qn("w:author"), "Owner")
    rr = OxmlElement("w:r")
    tt = OxmlElement("w:t")
    tt.text = "inserted"
    rr.append(tt)
    ins.append(rr)
    tracked._p.append(ins)
    if png_path is not None:
        d.add_paragraph().add_run().add_picture(str(png_path), width=Cm(2))   # 7
    d.add_paragraph("Closing remarks.")
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "Department of Records"
    sec.footer.paragraphs[0].text = "Internal"
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def load(b: bytes):
    return Document(io.BytesIO(b))


def save(d) -> bytes:
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def parts(b: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(b)) as z:
        return {n: z.read(n) for n in z.namelist()}


def view(d):
    return W.describe(d)


def ex(d, i):
    return view(d)["blocks"][i]["expect"]


@pytest.fixture
def png(tmp_path):
    from PIL import Image
    p = tmp_path / "logo.png"
    Image.new("RGB", (40, 20), (200, 30, 30)).save(p)
    return p


# ---- ops (C-2, C-3, C-4, C-7) -------------------------------------------------------------------------------------
def test_describe_lists_blocks_styles_page_and_header(png):
    v = view(load(report_bytes(png)))
    kinds = [b["kind"] for b in v["blocks"]]
    assert kinds[:5] == ["paragraph", "paragraph", "paragraph", "paragraph", "table"]
    assert v["blocks"][0]["style"] == "Title" and v["blocks"][0]["level"] == 0
    assert v["blocks"][1]["level"] == 1
    assert v["blocks"][4]["cells"] == [["Item", "Cost"], ["Paper", "100"]]
    assert v["blocks"][5]["locked"] == "a comment" and v["blocks"][6]["locked"] == "a tracked change"
    assert v["blocks"][7]["picture"] is True
    assert v["blocks"][2]["bold"] is True and v["blocks"][3]["bold"] is False      # the Bold button can turn it off
    assert v["header"] == "Department of Records" and v["footer"] == "Internal"
    assert {"Normal", "Heading 1", "Title"} <= set(v["styles"])
    assert v["page"]["sections"] == 1


def test_paragraph_ops_keep_formatting_and_reopen():
    d = load(report_bytes())
    W.set_paragraph_text(d, 2, ex(d, 2), "Revenue grew by twelve percent.")
    i = W.insert_paragraph(d, 3, "A new finding.", "Heading 2")
    assert i == 4
    W.set_style(d, 3, ex(d, 3), "Quote")
    W.format_text(d, 3, ex(d, 3), {"i": True, "size": 13, "color": "#1F4E79", "align": "justify"})
    W.move_block(d, 4, ex(d, 4), 1)
    W.insert_page_break(d, 0)
    out = load(save(d))
    texts = [p.text for p in out.paragraphs]
    assert texts[:5] == ["Annual Report", "", "A new finding.", "Summary", "Revenue grew by twelve percent."]
    p = out.paragraphs[4]
    assert p.runs[0].bold                                           # first run's format kept
    q = next(p for p in out.paragraphs if p.text.startswith("Costs"))
    assert q.style.name == "Quote" and q.runs[0].italic and q.runs[0].font.size.pt == 13


def test_a_stale_or_unknown_address_is_refused_and_nothing_changes():
    d = load(report_bytes())
    stale = ex(d, 3)
    W.insert_paragraph(d, None, "Front matter")                    # block 3 is now block 4
    before = save(d)
    with pytest.raises(OfficeError) as e:
        W.set_paragraph_text(d, 3, stale, "x")
    assert e.value.code == "STALE_BLOCK"
    with pytest.raises(OfficeError) as e:
        W.delete_block(d, 99, "0" * 12)
    assert e.value.code == "UNKNOWN_BLOCK"
    assert save(d) == before


def test_moving_equal_text_with_different_styles_invalidates_the_old_address():
    d = Document()
    d.add_paragraph("Same", style="Heading 1")
    d.add_paragraph("Same", style="Normal")
    stale = ex(d, 0)
    W.move_block(d, 0, stale, 1)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.set_paragraph_text(d, 0, stale, "Changed")
    assert error.value.code == "STALE_BLOCK"
    assert d.element.body.xml == before


def test_moving_byte_identical_paragraphs_invalidates_the_old_address():
    d = Document()
    first = d.add_paragraph("same")._p
    second = d.add_paragraph("same")._p
    assert first.xml == second.xml
    stale = ex(d, 0)
    original_xml = d.element.body.xml

    W.move_block(d, 1, ex(d, 1), 0)
    assert d.element.body.xml == original_xml
    assert W.blocks(d)[0] is second
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.set_paragraph_text(d, 0, stale, "Changed")
    assert error.value.code == "STALE_BLOCK"
    assert d.element.body.xml == before
    assert [p.text for p in d.paragraphs] == ["same", "same"]


def test_content_uri_does_not_own_is_refused_by_name(png):
    d = load(report_bytes(png))
    for i, why in ((5, "a comment"), (6, "a tracked change")):
        with pytest.raises(OfficeError) as e:
            W.set_paragraph_text(d, i, ex(d, i), "x")
        assert e.value.code == "AUTHORING_NOT_SUPPORTED" and why in e.value.detail
        with pytest.raises(OfficeError):
            W.delete_block(d, i, ex(d, i))
    with pytest.raises(OfficeError) as e:
        W.set_paragraph_text(d, 7, ex(d, 7), "x")                   # a picture paragraph
    assert "picture" in e.value.detail
    with pytest.raises(OfficeError):
        W.set_paragraph_text(d, 4, ex(d, 4), "x")                   # a table is not a paragraph
    with pytest.raises(OfficeError):
        W.set_style(d, 2, ex(d, 2), "No Such Style")


def test_delete_block_refuses_a_field_and_preserves_the_paragraph():
    d = Document()
    d.add_paragraph("A field follows")
    d.add_paragraph("A second paragraph")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    d.paragraphs[0]._p.append(field)
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.delete_block(d, 0, expect)
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "field" in error.value.detail
    assert d.element.body.xml == before


def test_delete_table_block_refuses_a_field_and_preserves_the_table():
    d = Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    cell.paragraphs[0]._p.append(field)
    d.add_paragraph("Keep a second block")
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.delete_block(d, 0, expect)
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "field" in error.value.detail
    assert d.element.body.xml == before


def test_delete_table_row_refuses_a_field_and_preserves_the_row():
    d = Document()
    t = d.add_table(rows=2, cols=1)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    t.cell(0, 0).paragraphs[0]._p.append(field)
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.delete_table_row(d, 0, expect, 0)
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "field" in error.value.detail
    assert d.element.body.xml == before


def test_table_ops():
    d = load(report_bytes())
    W.set_table_cell(d, 4, ex(d, 4), 1, 1, "150")
    W.add_table_row(d, 4, ex(d, 4), 1)
    W.set_table_cell(d, 4, ex(d, 4), 2, 0, "Ink")
    i = W.insert_table(d, 3, [["A", "B", "C"], ["1"]])
    out = load(save(d))
    t0, t1 = out.tables[0], out.tables[1]
    assert [[c.text for c in r.cells] for r in t0.rows] == [["A", "B", "C"], ["1", "", ""]]
    assert [[c.text for c in r.cells] for r in t1.rows] == [["Item", "Cost"], ["Paper", "150"], ["Ink", ""]]
    d = load(save(d))
    W.delete_table_row(d, i + 1, ex(d, i + 1), 0)
    assert len(load(save(d)).tables[1].rows) == 2
    for bad in ([], [[]], [["x"] * 21]):
        with pytest.raises(OfficeError):
            W.insert_table(d, 0, bad)


def test_table_cell_with_field_in_later_paragraph_is_preserved_on_refusal():
    d = Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("First")
    second = cell.add_paragraph("Page: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    second._p.append(field)
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.set_table_cell(d, 0, expect, 0, 0, "Replacement")
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "field" in error.value.detail
    assert d.element.body.xml == before


def test_add_table_row_refuses_source_field_without_duplication():
    d = Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("Date: ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "DATE")
    cell.paragraphs[0]._p.append(field)
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.add_table_row(d, 0, expect, 0)
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "field" in error.value.detail
    assert d.element.body.xml == before
    assert len(d.tables[0].rows) == 1


def test_table_cell_with_nested_table_is_preserved_on_refusal():
    d = Document()
    cell = d.add_table(rows=1, cols=1).cell(0, 0)
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested"
    expect = ex(d, 0)
    before = d.element.body.xml
    with pytest.raises(OfficeError) as error:
        W.set_table_cell(d, 0, expect, 0, 0, "Replacement")
    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert d.element.body.xml == before


def test_picture_page_header_footer_and_setup(png):
    d = load(report_bytes())
    W.insert_image(d, 1, png, 0.3)
    W.set_header_text(d, "Records Division")
    W.set_footer_text(d, "Confidential")
    W.insert_page_number(d, "footer", "right")
    W.set_page_setup(d, "landscape", {"left": 2, "right": 2})
    out = load(save(d))
    assert out.paragraphs[2]._p.find(".//" + qn("w:drawing")) is not None
    sec = out.sections[0]
    assert sec.header.paragraphs[0].text == "Records Division"
    assert sec.footer.paragraphs[0].text == "Confidential"
    assert any(p._p.find(qn("w:fldSimple")) is not None for p in sec.footer.paragraphs)
    assert sec.page_width > sec.page_height and round(sec.left_margin.cm, 1) == 2.0
    with pytest.raises(OfficeError):
        W.insert_page_number(load(save(out)), "footer")              # already has a field
    with pytest.raises(OfficeError):
        W.insert_image(d, 0, png, 2)


def test_replace_text_inside_runs_counts_split_matches():
    d = load(report_bytes())
    p = d.paragraphs[3]
    p.add_run(" per ")
    p.add_run("cen")
    p.add_run("t")                                                   # "percent" split across runs
    r = W.replace_text(d, "flat", "level")
    assert r["replaced"] == 1 and r["skipped"] == 0
    assert "level" in load(save(d)).paragraphs[3].text
    r = W.replace_text(d, "cent", "CENT")
    assert r["split_across_formatting"] == 1                        # " per " + "cen" + "t" not changed
    with pytest.raises(OfficeError) as e:
        W.replace_text(d, "zebra", "x")
    assert e.value.code == "NOT_FOUND"
    r = W.replace_text(d, "inserted", "x")
    assert r["skipped"] == 1 and r["replaced"] == 0                  # a tracked change is not rewritten


def test_replace_text_skips_text_in_run_with_inline_drawing():
    from PIL import Image

    d = Document()
    run = d.add_paragraph().add_run("Revenue ")
    picture = io.BytesIO()
    Image.new("RGB", (2, 2), (200, 30, 30)).save(picture, format="PNG")
    picture.seek(0)
    run.add_picture(picture, width=Cm(2))
    run.add_text(" rose")
    before = d.element.body.xml
    result = W.replace_text(d, "Revenue", "Sales")
    assert result == {"replaced": 0, "split_across_formatting": 0, "skipped": 1}
    assert d.element.body.xml == before
    assert len(d.element.body.findall(".//" + qn("w:drawing"))) == 1


def test_untouched_parts_stay_byte_identical(png):
    src = report_bytes(png)
    d = load(src)
    W.set_paragraph_text(d, 3, ex(d, 3), "Costs fell.")
    W.insert_table(d, 0, [["x"]])
    after, before = parts(save(d)), parts(src)
    keep = [n for n in before if n.startswith(("word/media/", "word/theme/")) or n in
            ("word/comments.xml", "word/numbering.xml", "word/fontTable.xml", "word/header1.xml", "word/footer1.xml")]
    assert "word/comments.xml" in keep and any(n.startswith("word/media/") for n in keep)
    for n in keep:
        assert after[n] == before[n], n


def test_layouts_and_owner_templates_make_documents(tmp_path):
    for lay in T.layout_list():
        d = load(T.from_layout(lay["id"]))
        assert round(d.sections[0].page_width.cm, 1) == 21.0
        if lay["id"] != "blank":
            assert any("[" in p.text for p in d.paragraphs)          # placeholders, not invented content
    data = parts(report_bytes())
    data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(
        b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, b in data.items():
            z.writestr(n, b)
    out = load(T.from_docx_template(buf.getvalue()))
    assert out.sections[0].header.paragraphs[0].text == "Department of Records"
    with pytest.raises(OfficeError):
        T.from_layout("nope")


# ---- session (C-1, C-5) ---------------------------------------------------------------------------------------------
class FakeDocs(FakeDocuments):
    preview_state = "ready"

    def open_bytes(self, name, kind, data, *, owner):
        d = super().open_bytes(name, kind, data, owner=owner)
        d.state = self.preview_state
        return d

    def add_opened(self, path, grants, kind="docx"):
        g = grants.grant_open(str(path), kind=kind)
        data = path.read_bytes()
        did = f"d{len(self.sessions)}"
        self.sessions[did] = type("D", (), {"document_id": did, "kind": kind, "name": path.name,
                                            "grant_id": g.grant_id, "sha256": hashlib.sha256(data).hexdigest()})()
        return did


@pytest.fixture
def env(tmp_path):
    p = tmp_path / "Annual Report.docx"
    p.write_bytes(report_bytes())
    (tmp_path / "state").mkdir()
    grants = Grants([tmp_path])
    docs = FakeDocs(grants)
    appr = Approvals(tmp_path / "state", "owner")
    return Authoring(tmp_path / "state", docs, grants, appr), docs, grants, appr, p, tmp_path


def test_word_edit_change_list_undo_and_save(env):
    a, docs, grants, appr, p, tmp = env
    original = p.read_bytes()
    v = a.edit_document(docs.add_opened(p, grants))
    aid = v["authoring_id"]
    assert a.get(aid).data == original and v["kind"] == "docx" and "doc" in v       # C-1
    b3 = v["doc"]["blocks"][3]
    v = a.apply(aid, {"op": "set_paragraph_text", "block": 3, "expect": b3["expect"], "text": "Costs fell."})
    v = a.apply(aid, {"op": "insert_paragraph", "after": None, "text": "Draft", "style": "Heading 1"})
    assert [c["text"] for c in v["changes"]] == [
        "Change the text of paragraph 4 ('Costs were flat across the year.') to 'Costs fell.'",
        "Add a paragraph at the start: 'Draft'"]
    a.undo(aid)
    assert len(a.view(aid)["changes"]) == 1
    prop = a.propose_save(aid)
    assert "Save this document?" in card_lines(prop) and "Costs fell." in card_lines(prop)
    assert p.read_bytes() == original
    a.apply_save(act_body(appr, prop))
    assert load(p.read_bytes()).paragraphs[3].text == "Costs fell."


def test_identical_paragraph_move_invalidates_session_addresses_across_undo_redo(env):
    a, docs, grants, _, p, _ = env
    d = Document()
    d.add_paragraph("same")
    d.add_paragraph("same")
    p.write_bytes(save(d))
    v = a.edit_document(docs.add_opened(p, grants))
    aid = v["authoring_id"]
    stale = v["doc"]["blocks"][0]["expect"]

    moved = a.apply(aid, {"op": "move_block", "block": 1,
                          "expect": v["doc"]["blocks"][1]["expect"], "to": 0})
    assert moved["doc"]["blocks"][0]["expect"] != stale
    before = a.get(aid).data
    with pytest.raises(OfficeError) as error:
        a.apply(aid, {"op": "set_paragraph_text", "block": 0, "expect": stale, "text": "Changed"})
    assert error.value.code == "STALE_BLOCK"
    assert a.get(aid).data == before

    undone = a.undo(aid)
    assert undone["doc"]["blocks"][0]["expect"] != stale
    assert undone["doc"]["blocks"][0]["expect"] != moved["doc"]["blocks"][0]["expect"]
    redone = a.undo(aid, redo=True)
    assert redone["doc"]["blocks"][0]["expect"] != moved["doc"]["blocks"][0]["expect"]
    assert [p.text for p in load(a.get(aid).data).paragraphs] == ["same", "same"]


def test_a_working_copy_that_did_not_render_is_not_saved(env):
    a, docs, grants, appr, p, tmp = env
    v = a.edit_document(docs.add_opened(p, grants))
    aid = v["authoring_id"]
    docs.preview_state = "unavailable"
    a.apply(aid, {"op": "insert_page_break", "after": 0})
    with pytest.raises(OfficeError) as e:
        a.propose_save(aid)
    assert e.value.code == "PREVIEW_FAILED"
    docs.preview_state = "rendering"
    a.apply(aid, {"op": "insert_page_break", "after": 0})
    with pytest.raises(OfficeError) as e:
        a.propose_save(aid)
    assert e.value.code == "PREVIEW_NOT_READY"


def test_new_word_document_from_a_layout_needs_save_as(env):
    a, docs, grants, appr, p, tmp = env
    v = a.new_document("letter")
    assert v["name"] == "Letter.docx" and v["new"] and v["doc"]["blocks"][0]["text"] == "[Your name and address]"
    with pytest.raises(OfficeError) as e:
        a.propose_save(v["authoring_id"])
    assert e.value.code == "SAVE_AS_REQUIRED"
    dest = tmp / "Letter to the bank.docx"
    a.apply_save(act_body(appr, a.propose_save(v["authoring_id"], str(dest))))
    assert load(dest.read_bytes()).paragraphs[0].text == "[Your name and address]"
