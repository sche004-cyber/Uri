"""Milestones A/B (plan 11_AUTHORING_A_B_PLAN.md): PowerPoint ops and the authoring session (AB-1..AB-7)."""
import hashlib
import io
import json
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from uri_office.approvals import Approvals
from uri_office.authoring import pptx_ops as P
from uri_office.authoring.session import Authoring
from uri_office.errors import OfficeError
from uri_office.grants import Grants


def deck_bytes() -> bytes:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Quarterly Review"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Highlights"
    body = s2.placeholders[1].text_frame
    body.text = "Revenue up"
    body.paragraphs[0].runs[0].font.bold = True
    body.paragraphs[0].runs[0].font.size = Pt(28)
    body.add_paragraph().text = "Costs flat"
    s2.notes_slide.notes_text_frame.text = "Say thanks"
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Chart"
    cd = CategoryChartData()
    cd.categories = ["Q1", "Q2"]
    cd.add_series("Sales", (10, 12))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(6), Inches(4), cd)
    s4 = prs.slides.add_slide(prs.slide_layouts[5])
    s4.shapes.title.text = "Table"
    t = s4.shapes.add_table(2, 2, Inches(1), Inches(1.5), Inches(6), Inches(2)).table
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Item", "Cost", "Paper", "100"
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def load(b: bytes):
    return Presentation(io.BytesIO(b))


def save(prs) -> bytes:
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def ids(prs):
    return [s.slide_id for s in prs.slides]


def parts(b: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(b)) as z:
        return {n: z.read(n) for n in z.namelist()}


# ---- ops (AB-2, AB-3, AB-7) ------------------------------------------------------------------------------------
def test_slide_structure_ops():
    prs = load(deck_bytes())
    a, b, c, d = ids(prs)
    new = P.add_slide(prs, "Title Only", after=a)
    assert ids(prs) == [a, new, b, c, d]
    P.move_slide(prs, new, 4)
    assert ids(prs) == [a, b, c, d, new]
    P.delete_slide(prs, new)
    assert ids(prs) == [a, b, c, d]
    with pytest.raises(OfficeError):
        P.add_slide(prs, "No Such Layout")
    prs2 = load(save(prs))                              # still a valid package
    assert ids(prs2) == [a, b, c, d]


def test_duplicate_copies_chart_independently_and_notes():
    prs = load(deck_bytes())
    a, b, c, d = ids(prs)
    dup = P.duplicate_slide(prs, c)
    assert ids(prs) == [a, b, c, dup, d]
    dup2 = P.duplicate_slide(prs, b)
    out = load(save(prs))
    charts = [sh.chart for s in out.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 2 and charts[0].part.partname != charts[1].part.partname
    wb = [ch.part.chart_workbook.xlsx_part.partname for ch in charts]
    assert wb[0] != wb[1]
    copy = next(s for s in out.slides if s.slide_id == dup2)
    assert copy.notes_slide.notes_text_frame.text == "Say thanks"
    assert copy.placeholders[1].text_frame.text == "Revenue up\nCosts flat"


def test_duplicate_refuses_content_it_cannot_copy_and_changes_nothing():
    prs = load(deck_bytes())
    s = prs.slides[0]
    s.part.relate_to("https://example.invalid/v.mp4",
                     "http://schemas.openxmlformats.org/officeDocument/2006/relationships/video", is_external=True)
    before = ids(prs)
    with pytest.raises(OfficeError) as e:
        P.duplicate_slide(prs, s.slide_id)
    assert e.value.code == "AUTHORING_NOT_SUPPORTED" and "video" in e.value.detail
    assert ids(prs) == before


def test_duplicate_refuses_linked_speaker_notes_before_changing_deck():
    prs = load(deck_bytes())
    source = prs.slides[1]
    linked = source.notes_slide.notes_text_frame.paragraphs[0].add_run()
    linked.text = " Review"
    linked.hyperlink.address = "https://example.org/review"
    before = save(prs)

    with pytest.raises(OfficeError) as error:
        P.duplicate_slide(prs, source.slide_id)

    assert error.value.code == "AUTHORING_NOT_SUPPORTED"
    assert "notes" in error.value.detail and "hyperlink" in error.value.detail
    assert save(prs) == before
    assert [s.slide_id for s in prs.slides] == ids(load(before))


def test_text_ops_keep_formatting():
    prs = load(deck_bytes())
    b = ids(prs)[1]
    sl = P.slide(prs, b)
    body = sl.placeholders[1]
    P.set_text(prs, b, body.shape_id, ["Revenue up 12%", "Costs flat", "Hiring paused"])
    P.set_title(prs, b, "Highlights 2026")
    P.set_notes(prs, b, "Thank the team")
    out = load(save(prs))
    s = out.slides[1]
    paras = s.placeholders[1].text_frame.paragraphs
    assert [p.text for p in paras] == ["Revenue up 12%", "Costs flat", "Hiring paused"]
    assert paras[0].runs[0].font.bold and paras[0].runs[0].font.size == Pt(28)       # first run's format kept
    assert s.shapes.title.text == "Highlights 2026"
    assert s.notes_slide.notes_text_frame.text == "Thank the team"


def test_set_text_refuses_link_in_later_run_without_changing_presentation():
    prs = load(deck_bytes())
    s = prs.slides[1]
    body = s.placeholders[1]
    linked = body.text_frame.paragraphs[0].add_run()
    linked.text = " details"
    linked.hyperlink.address = "https://example.com"
    before = save(prs)

    with pytest.raises(OfficeError) as e:
        P.set_text(prs, s.slide_id, body.shape_id, ["Replacement"])

    assert e.value.code == "AUTHORING_NOT_SUPPORTED" and "hyperlink" in e.value.detail
    assert save(prs) == before
    reopened = load(save(prs)).slides[1].placeholders[1]
    assert reopened.text_frame.paragraphs[0].runs[1].hyperlink.address == "https://example.com"


def test_set_text_refuses_inline_field_without_changing_presentation():
    prs = load(deck_bytes())
    s = prs.slides[1]
    body = s.placeholders[1]
    body.text_frame.paragraphs[0]._p.append(etree.fromstring(
        '<a:fld xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'id="{00000000-0000-0000-0000-000000000001}" type="slidenum"><a:t>2</a:t></a:fld>'))
    before = save(prs)

    with pytest.raises(OfficeError) as e:
        P.set_text(prs, s.slide_id, body.shape_id, ["Replacement"])

    assert e.value.code == "AUTHORING_NOT_SUPPORTED" and "inline text" in e.value.detail
    assert save(prs) == before


def test_format_box_table_and_shape_ops():
    prs = load(deck_bytes())
    a, b, c, d = ids(prs)
    tb = P.insert_text_box(prs, a, {"x": 0.1, "y": 0.8, "w": 0.5, "h": 0.1}, ["Prepared by Admin"])
    P.format_text(prs, a, tb, {"b": True, "size": 18, "color": "#C00000", "align": "center"})
    tbl = next(sh for sh in P.slide(prs, d).shapes if sh.has_table).shape_id
    P.set_table_cell(prs, d, tbl, 1, 1, "150")
    new_tbl = P.insert_table(prs, c, [["A", "B"], ["1"]], {"x": 0.05, "y": 0.05, "w": 0.3, "h": 0.2})
    P.move_shape(prs, a, tb, {"x": 0.2, "y": 0.7, "w": 0.6, "h": 0.1})
    title = P.slide(prs, a).shapes.title.shape_id
    P.delete_shape(prs, a, title)                     # a placeholder is emptied, not removed
    out = load(save(prs))
    box = next(sh for sh in out.slides[0].shapes if sh.shape_id == tb)
    r = box.text_frame.paragraphs[0].runs[0]
    assert r.font.bold and r.font.size == Pt(18) and str(r.font.color.rgb) == "C00000"
    assert out.slides[0].shapes.title.text == ""
    t = next(sh for sh in out.slides[3].shapes if sh.has_table).table
    assert t.cell(1, 1).text == "150"
    nt = next(sh for sh in out.slides[2].shapes if sh.shape_id == new_tbl).table
    assert nt.cell(1, 1).text == ""
    for bad in ({"x": 2, "y": 0, "w": 1, "h": 1}, {"x": 0, "y": 0, "w": 0, "h": 1}, {"x": 0}):
        with pytest.raises(OfficeError):
            P.insert_text_box(prs, a, bad, ["x"])


def test_untouched_parts_stay_byte_identical():
    src = deck_bytes()
    prs = load(src)
    P.set_title(prs, ids(prs)[0], "New title")
    after = parts(save(prs))
    before = parts(src)
    for name in before:
        if name.startswith(("ppt/charts/", "ppt/embeddings/", "ppt/slideLayouts/", "ppt/slideMasters/",
                            "ppt/theme/")) or name in ("ppt/slides/slide3.xml", "ppt/slides/slide4.xml"):
            assert after[name] == before[name], name


def test_describe_lists_slides_items_and_layouts():
    d = P.describe(load(deck_bytes()))
    assert [s["title"] for s in d["slides"]] == ["Quarterly Review", "Highlights", "Chart", "Table"]
    kinds = {it["kind"] for s in d["slides"] for it in s["items"]}
    assert {"text", "chart", "table"} <= kinds
    assert d["slides"][1]["notes"] == "Say thanks"
    title = next(it for it in d["slides"][1]["items"] if it["id"] == d["slides"][1]["title_id"])
    assert title["paragraphs"] == ["Highlights"]
    body = next(it for it in d["slides"][1]["items"] if it.get("paragraphs") == ["Revenue up", "Costs flat"])
    assert body["bold"] is False                                   # only the first paragraph is bold
    assert "Title and Content" in d["layouts"]


# ---- session (AB-1, AB-4, AB-5, AB-6) ------------------------------------------------------------------------------
class FakeDocuments:
    def __init__(self, grants):
        self.grants, self.sessions, self.n = grants, {}, 0
        self.opened = []

    def get(self, did):
        return self.sessions[did]

    def open_bytes(self, name, kind, data, *, owner):
        self.n += 1
        did = f"p{self.n}"
        self.opened.append(hashlib.sha256(data).hexdigest())
        self.sessions[did] = type("D", (), {"document_id": did})()
        return self.sessions[did]

    def close(self, did):
        self.sessions.pop(did, None)

    def add_opened(self, path, grants):
        g = grants.grant_open(str(path), kind="pptx")
        data = path.read_bytes()
        did = f"d{len(self.sessions)}"
        self.sessions[did] = type("D", (), {"document_id": did, "kind": "pptx", "name": path.name,
                                            "grant_id": g.grant_id, "sha256": hashlib.sha256(data).hexdigest()})()
        return did


@pytest.fixture
def env(tmp_path):
    p = tmp_path / "Quarterly Review.pptx"
    p.write_bytes(deck_bytes())
    (tmp_path / "state").mkdir()
    grants = Grants([tmp_path])
    docs = FakeDocuments(grants)
    appr = Approvals(tmp_path / "state", "owner")
    a = Authoring(tmp_path / "state", docs, grants, appr)
    return a, docs, grants, appr, p, tmp_path


def card_lines(prop):
    return json.dumps(prop["messages"], ensure_ascii=False)


def test_failed_initial_preview_does_not_consume_authoring_capacity(env, monkeypatch):
    a, docs, grants, appr, p, tmp = env
    open_bytes = docs.open_bytes

    def broken(*args, **kwargs):
        raise OfficeError("ENGINE_UNAVAILABLE", "")

    monkeypatch.setattr(docs, "open_bytes", broken)
    for _ in range(4):
        with pytest.raises(OfficeError) as error:
            a.new_presentation()
        assert error.value.code == "ENGINE_UNAVAILABLE"
        assert not a.sessions

    monkeypatch.setattr(docs, "open_bytes", open_bytes)
    view = a.new_presentation()
    assert view["authoring_id"] in a.sessions


def test_failed_edit_preview_releases_its_grant_and_session(env, monkeypatch):
    a, docs, grants, appr, p, tmp = env
    did = docs.add_opened(p, grants)
    prior_grants = set(grants._by_id)
    open_bytes = docs.open_bytes

    def broken(*args, **kwargs):
        raise OfficeError("ENGINE_UNAVAILABLE", "")

    monkeypatch.setattr(docs, "open_bytes", broken)
    with pytest.raises(OfficeError) as error:
        a.edit_document(did)
    assert error.value.code == "ENGINE_UNAVAILABLE"
    assert not a.sessions
    assert set(grants._by_id) == prior_grants

    monkeypatch.setattr(docs, "open_bytes", open_bytes)
    view = a.edit_document(did)
    assert view["authoring_id"] in a.sessions
    assert grants.get(a.get(view["authoring_id"]).grant_id, "open").path == p


def act_body(appr, prop):
    sid = prop["messages"][0]["createSurface"]["surfaceId"]
    row = appr.db.execute("SELECT body_json FROM surface WHERE surface_id=?", (sid,)).fetchone()
    return json.loads(row["body_json"])


def test_working_copy_is_the_original_and_save_replaces_only_after_approval(env):
    a, docs, grants, appr, p, tmp = env
    original = p.read_bytes()
    did = docs.add_opened(p, grants)
    v = a.edit_document(did)
    s = a.get(v["authoring_id"])
    assert s.data == original                                    # AB-1
    first = v["deck"]["slides"][0]["id"]
    v = a.apply(s.aid, {"op": "set_title", "slide": first, "text": "Quarterly Review 2026"})
    v = a.apply(s.aid, {"op": "add_slide", "layout": "Title Only"})
    assert [c["text"] for c in v["changes"]] == ["Slide 1: title 'Quarterly Review 2026'", "Add a 'Title Only' slide"]
    assert p.read_bytes() == original                            # nothing written yet
    prop = a.propose_save(s.aid)
    text = card_lines(prop)
    assert "Slide 1: title 'Quarterly Review 2026'" in text and "replaces the file you opened" in text
    a.apply_save(act_body(appr, prop))
    saved = load(p.read_bytes())
    assert saved.slides[0].shapes.title.text == "Quarterly Review 2026" and len(saved.slides) == 5
    backups = list((tmp / "state" / "authoring" / "backups").rglob("*.pptx"))
    assert backups and backups[0].read_bytes() == original       # AB-6 backup first
    assert a.view(s.aid)["changes"] == []


def test_undo_redo_restore_exact_versions(env):
    a, docs, grants, appr, p, tmp = env
    did = docs.add_opened(p, grants)
    v = a.edit_document(did)
    aid = v["authoring_id"]
    h0 = a.get(aid).data
    a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    h1 = a.get(aid).data
    a.apply(aid, {"op": "delete_slide", "slide": v["deck"]["slides"][0]["id"]})
    a.undo(aid)
    assert a.get(aid).data == h1
    a.undo(aid)
    assert a.get(aid).data == h0 and a.view(aid)["changes"] == []          # AB-4
    a.undo(aid, redo=True)
    assert a.get(aid).data == h1 and len(a.view(aid)["changes"]) == 1
    a.undo(aid, redo=True)
    with pytest.raises(OfficeError):                                    # nothing left to redo
        a.undo(aid, redo=True)


def test_a_later_change_makes_the_card_stale_and_a_changed_file_is_not_replaced(env):
    a, docs, grants, appr, p, tmp = env
    did = docs.add_opened(p, grants)
    v = a.edit_document(did)
    aid = v["authoring_id"]
    a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    prop = a.propose_save(aid)
    body = act_body(appr, prop)
    a.apply(aid, {"op": "add_slide", "layout": "Blank"})          # changed after the card was shown
    with pytest.raises(OfficeError) as e:
        a.apply_save(body)
    assert e.value.code == "STALE_CHANGE"
    prop = a.propose_save(aid)
    p.write_bytes(p.read_bytes() + b"x")                          # someone else changed the file
    with pytest.raises(OfficeError) as e:
        a.apply_save(act_body(appr, prop))
    assert e.value.code == "TARGET_CHANGED"


@pytest.mark.parametrize("save_as", [False, True])
def test_save_does_not_replace_a_file_changed_during_publication(env, monkeypatch, save_as):
    from uri_office import xlsx_patch

    a, docs, grants, appr, source, tmp = env
    did = docs.add_opened(source, grants)
    aid = a.edit_document(did)["authoring_id"]
    a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    dest = tmp / "new.pptx" if save_as else source
    body = act_body(appr, a.propose_save(aid, str(dest) if save_as else None))
    actual_mkstemp = xlsx_patch.tempfile.mkstemp
    raced = False

    def create_then_race(*args, **kwargs):
        nonlocal raced
        result = actual_mkstemp(*args, **kwargs)
        if kwargs.get("suffix") == ".uri-office.tmp" and not raced:
            raced = True
            dest.write_bytes(b"external edit")
        return result

    monkeypatch.setattr(xlsx_patch.tempfile, "mkstemp", create_then_race)
    with pytest.raises(OfficeError) as error:
        a.apply_save(body)
    assert error.value.code == ("DESTINATION_EXISTS" if save_as else "TARGET_CHANGED")
    assert raced and dest.read_bytes() == b"external edit"
    assert a.get(aid).changes  # a failed Save keeps the working copy and its pending change


def test_save_does_not_replace_a_name_created_after_the_old_file_is_claimed(env, monkeypatch):
    from uri_office import xlsx_patch

    a, docs, grants, appr, dest, tmp = env
    original = dest.read_bytes()
    aid = a.edit_document(docs.add_opened(dest, grants))["authoring_id"]
    a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    body = act_body(appr, a.propose_save(aid))
    actual_link = xlsx_patch.os.link
    raced = False

    def create_before_publish(source, target, *args, **kwargs):
        nonlocal raced
        if not raced and str(source).endswith(".uri-office.tmp") and target == dest:
            raced = True
            dest.write_bytes(b"new external file")
        return actual_link(source, target, *args, **kwargs)

    monkeypatch.setattr(xlsx_patch.os, "link", create_before_publish)
    with pytest.raises(OfficeError) as error:
        a.apply_save(body)
    assert error.value.code == "TARGET_CHANGED"
    assert raced and dest.read_bytes() == b"new external file"
    backups = list((tmp / "state" / "authoring" / "backups").rglob("*.pptx"))
    assert len(backups) == 1 and backups[0].read_bytes() == original


def test_new_presentation_needs_save_as_and_never_replaces(env):
    a, docs, grants, appr, p, tmp = env
    v = a.new_presentation()
    aid = v["authoring_id"]
    assert v["new"] and len(v["deck"]["slides"]) == 1 and v["deck"]["aspect"] == pytest.approx(16 / 9, 0.01)
    with pytest.raises(OfficeError) as e:
        a.propose_save(aid)
    assert e.value.code == "SAVE_AS_REQUIRED"
    dest = tmp / "Plan.pptx"
    prop = a.propose_save(aid, str(dest))
    assert "An existing file is never replaced" in card_lines(prop)
    a.apply_save(act_body(appr, prop))
    assert len(load(dest.read_bytes()).slides) == 1
    # the edit continues on the new file: the next Save replaces it, with a backup
    v = a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    assert v["new"] is False and v["name"] == "Plan.pptx"
    prop = a.propose_save(aid)
    assert "replaces the file you opened" in card_lines(prop)
    a.apply_save(act_body(appr, prop))
    assert len(load(dest.read_bytes()).slides) == 2
    prop = a.propose_save(aid, str(p))                            # an existing file as Save as
    with pytest.raises(OfficeError) as e:
        a.apply_save(act_body(appr, prop))
    assert e.value.code == "DESTINATION_EXISTS"


def test_new_from_the_owners_template_keeps_its_layouts(env, tmp_path):
    a, docs, grants, appr, p, tmp = env
    tpl = tmp / "Corporate.potx"
    data = parts(deck_bytes())
    data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(
        b"presentationml.presentation.main+xml", b"presentationml.template.main+xml")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in data.items():
            z.writestr(n, d)
    tpl.write_bytes(buf.getvalue())
    v = a.new_presentation(str(tpl))
    assert v["name"] == "Corporate (new).pptx" and "Title and Content" in v["deck"]["layouts"]
    assert len(v["deck"]["slides"]) == 4


def test_word_and_unknown_ops_are_refused(env):
    a, docs, grants, appr, p, tmp = env
    v = a.new_presentation()
    with pytest.raises(OfficeError):
        a.apply(v["authoring_id"], {"op": "run_macro"})
    with pytest.raises(OfficeError) as e:
        a.apply(v["authoring_id"], {"op": "set_title"})
    assert e.value.code == "BAD_OP"


def test_the_edit_keeps_its_own_grant_and_bad_values_are_refused(env):
    a, docs, grants, appr, p, tmp = env
    did = docs.add_opened(p, grants)
    v = a.edit_document(did)
    aid = v["authoring_id"]
    grants.revoke(docs.sessions[did].grant_id)                    # the viewer that opened the file was closed
    first = v["deck"]["slides"][0]["id"]
    a.apply(aid, {"op": "set_title", "slide": first, "text": "Still saves"})
    a.apply_save(act_body(appr, a.propose_save(aid)))
    assert load(p.read_bytes()).slides[0].shapes.title.text == "Still saves"
    with pytest.raises(OfficeError) as e:
        a.apply(aid, {"op": "move_slide", "slide": first, "to": "last"})
    assert e.value.code == "BAD_POSITION"

    def wrong_shape(prs, op):
        raise TypeError("x")
    P.OPS["set_title"], real = wrong_shape, P.OPS["set_title"]
    try:
        with pytest.raises(OfficeError) as e:
            a.apply(aid, {"op": "set_title", "slide": first, "text": 3})
    finally:
        P.OPS["set_title"] = real
    assert e.value.code == "BAD_OP"
    gid = a.get(aid).grant_id
    a.close(aid)
    with pytest.raises(OfficeError):
        grants.get(gid, "open")                                   # closing the edit revokes its grant


def test_undo_history_has_a_memory_cap(env, monkeypatch):
    from uri_office.authoring import session as S
    a, docs, grants, appr, p, tmp = env
    monkeypatch.setattr(S, "MAX_HISTORY_BYTES", 3 * len(p.read_bytes()))
    did = docs.add_opened(p, grants)
    aid = a.edit_document(did)["authoring_id"]
    for _ in range(6):
        a.apply(aid, {"op": "add_slide", "layout": "Blank"})
    s = a.get(aid)
    assert len(s.versions) <= 3 and s.dropped >= 4
    assert len(a.view(aid)["changes"]) == 6                     # the change list is complete
