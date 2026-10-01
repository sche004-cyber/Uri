"""Shared fixtures for M37/M38 tests: real .xlsx files, real scripts, real M36 stores."""
from __future__ import annotations

from pathlib import Path
import textwrap
import uuid

from openpyxl import Workbook

import uri_v1.memory.recorder as memory_recorder
from uri_v1.edge.contracts import EdgeTask, TaskClass, TaskInput
from uri_v1.edge.local_runner import LocalRunner, verifier_entries
from uri_v1.edge.registry import CapabilityRegistry
from uri_v1.edge.service import EdgeService, EdgeTelemetry
from uri_v1.execution.acceptance import AcceptanceCriteria
from uri_v1.execution.run_record import RunStore
from uri_v1.memory.log import MemoryLog
from uri_v1.memory.recorder import Recorder, _bootstrap_verifiers
from uri_v1.memory.sources import SourceRegistry
from uri_v1.memory.uri_adapter.envelope import URIIntake
from uri_v1.results.version_ledger import ResultVersionLedger

USER = str(uuid.UUID(int=7))


def intake(raw, *, user=USER, session="session"):
    return URIIntake().capture(user, session, uuid.uuid4().hex, uuid.uuid4().hex, raw, spans={"file": (0, len(raw))})


def make_workbook(path: Path, sheets: dict, *, formula_error=False) -> Path:
    wb = Workbook(); first = True
    for name, (header, rows) in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name; first = False
        ws.append(list(header))
        for r in rows: ws.append(list(r))
        if formula_error: ws["Z1"] = "=1/0"
    wb.save(path)
    return path


SALES = {"Sales": (("Region", "Amount", "Unit Price"), (("North", 10, 2.5), ("South", 20, 3.0)))}

SCRIPT_OK = textwrap.dedent('''
    from openpyxl import load_workbook, Workbook
    wb = load_workbook("sales.xlsx")
    ws = wb["Sales"]
    header = [c.value for c in ws[1]]
    col = header.index("Amount")
    total = sum(row[col] for row in ws.iter_rows(min_row=2, values_only=True))
    out = Workbook(); o = out.active; o.title = "Summary"
    o.append(["Total"]); o.append([total])
    out.save("summary.xlsx")
''').lstrip()


def script(sheet="Sales", column="Amount", *, workbook="sales.xlsx", extra=""):
    return (SCRIPT_OK.replace('"Sales"', repr(sheet).replace("'", '"')).replace('"Amount"', repr(column).replace("'", '"'))
            .replace('"sales.xlsx"', '"%s"' % workbook) + extra)


def acceptance(script_name="report.py", *, output="summary.xlsx", allowed=("report.py",)):
    return AcceptanceCriteria(1, ({"kind": "only_changed", "paths": list(allowed)},
                                  {"kind": "script_exit_zero", "script": script_name, "timeout_s": 60},
                                  {"kind": "output_workbook_exists", "path": output},
                                  {"kind": "workbook_has_sheets", "path": output, "sheets": ["Summary"]},
                                  {"kind": "no_formula_errors", "path": output}))


def report_acceptance(task_class):
    return AcceptanceCriteria(1, ({"kind": "only_changed", "paths": ["edge_report.json"]},
                                  {"kind": "report_recomputes", "path": "edge_report.json", "task_class": task_class}))


class Env:
    """Real M36 log/sources/recorder, S11 ledger, run store and Edge service in temporary roots."""

    def __init__(self, tmp: Path, monkeypatch, *, qualification_run=True, record=None, python=None, existing=None):
        self.tmp = tmp
        self.scratch = tmp / "scratch"; self.runs = tmp / "runs"
        if existing:  # (log, sources, recorder, root) from the real M36 test helpers
            self.log, self.sources, self.recorder, self.root = existing
        else:
            self.root = tmp / "authorized"; self.root.mkdir(exist_ok=True)
            self.log = MemoryLog(USER, tmp / "private")
            self.sources = SourceRegistry(self.log); self.sources.register(self.root)
            self.recorder = Recorder(self.log, self.sources)
        self.user = self.log.user_id
        self.ledger = ResultVersionLedger(self.user, root=str(tmp / "results"))
        self.store = RunStore(self.user, self.runs)
        monkeypatch.setattr(memory_recorder, "_VERIFIERS", memory_recorder.MappingProxyType({}))
        _bootstrap_verifiers(verifier_entries(self.runs, python=python))
        self.registry = CapabilityRegistry(record, qualification_run=qualification_run)
        self.runner = LocalRunner(user_id=self.user, scratch_root=self.scratch, run_store=self.store, recorder=self.recorder,
                                  ledger=self.ledger, sources=self.sources)
        self.telemetry = EdgeTelemetry(self.user, tmp / "telemetry")
        self.edge = EdgeService(registry=self.registry, sources=self.sources, runner=self.runner, telemetry=self.telemetry)

    def write(self, name, data) -> Path:
        p = self.root / name; p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(data if isinstance(data, bytes) else data.encode())
        return p

    def ref(self, name):
        scan = self.sources.scan()
        loc = next(l for l in scan.locators if l.relpath == name)
        return self.sources.fingerprint(loc)

    def open_task(self, label="make the script work with the spreadsheet"):
        task_id, result = self.recorder.open_task(intake(label, user=self.user), label)
        assert result.persisted, result
        return task_id

    def task(self, task_class, *, script_name="report.py", workbook_name="sales.xlsx", accept=None, task_id=None):
        task_class = TaskClass(task_class)
        inputs = [TaskInput("workbook", self.ref(workbook_name).source_id, self.ref(workbook_name).content_sha256)]
        if task_class is not TaskClass.WORKBOOK_INSPECT:
            inputs.append(TaskInput("script", self.ref(script_name).source_id, self.ref(script_name).content_sha256))
        if accept is None:
            accept = acceptance(script_name) if task_class is TaskClass.SCRIPT_WORKBOOK_REPAIR else report_acceptance(task_class.value)
        return EdgeTask(task_id or self.open_task(), task_class, tuple(inputs), accept, uuid.uuid4().hex, "session")
