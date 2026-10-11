"""AC3, AC5, AC6, AC7 through the real HTTP service on an ephemeral loopback port."""
import http.client
import json
import os
import shutil
import subprocess
import threading

import pytest
from openpyxl import load_workbook

from uri_office.errors import OfficeError
from uri_office.service import Office, serve


@pytest.fixture
def svc(tmp_path):
    office = Office(tmp_path / "state")
    srv, token = serve(office, "127.0.0.1", 0)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    yield office, srv.server_address[1], token
    srv.shutdown()
    srv.server_close()


def call(port, token, method, path, body=None, headers=None):
    c = http.client.HTTPConnection("127.0.0.1", port, timeout=10)
    h = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
    h.update(headers or {})
    data = json.dumps(body).encode() if body is not None else None
    c.request(method, path, body=data, headers=h)
    r = c.getresponse()
    out = (r.status, json.loads(r.read() or b"{}"), dict(r.getheaders()))
    c.close()
    return out


def action_for(messages, button, name):
    sid = messages[0]["createSurface"]["surfaceId"]
    model = messages[2]["updateDataModel"]["value"]
    return {"name": name, "surfaceId": sid, "sourceComponentId": button,
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


# ---- AC6 transport -------------------------------------------------------------------------------------------------
def test_non_loopback_bind_refused(tmp_path):
    for host in ("0.0.0.0", "192.168.1.5", "example.com"):
        with pytest.raises(OfficeError) as e:
            serve(Office(tmp_path / "s"), host, 0)
        assert e.value.code == "BIND_REFUSED"


def test_token_required_on_every_request(svc):
    _, port, token = svc
    for method, path in (("GET", "/status"), ("GET", "/a2ui/surfaces"), ("POST", "/open/local"),
                         ("POST", "/a2ui/action"), ("POST", "/workbook/x/edit")):
        assert call(port, "wrong", method, path, {} if method == "POST" else None)[0] == 401
        assert call(port, "", method, path, {} if method == "POST" else None)[0] == 401
    assert call(port, token, "GET", "/status")[0] == 200


def test_token_file_is_per_launch_and_current_user_only(tmp_path):
    office = Office(tmp_path / "state")
    tokens = []
    for _ in range(2):
        srv, token = serve(office, "127.0.0.1", 0)
        srv.server_close()
        f = tmp_path / "state" / "api-token"
        assert f.read_text(encoding="utf-8") == token
        tokens.append(token)
    assert tokens[0] != tokens[1] and len(tokens[0]) >= 40
    if os.name == "nt":
        acl = subprocess.run(["icacls", str(f)], capture_output=True, text=True).stdout
        entries = [ln.strip() for ln in acl.splitlines()[0:-2] if ln.strip()]
        entries[0] = entries[0][len(str(f)):].strip()
        user = os.environ["USERNAME"]
        allowed = {os.environ.get("USERDOMAIN", "") + "\\" + user, os.environ.get("COMPUTERNAME", "") + "\\" + user}
        assert len(entries) == 1, entries                       # exactly one ACE: no Users, Everyone or inherited
        account, _, perm = entries[0].rpartition(":")
        assert account in allowed and perm == "(F)", entries
        assert "(I)" not in acl


def test_get_never_changes_state(svc, rich):
    office, port, token = svc
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})
    chain = office.approvals.verify_chain()["events"]
    snap = json.dumps(office.view(wid)["draft"], sort_keys=True)
    for path in ("/status", f"/workbook/{wid}", "/a2ui/surfaces", f"/workbook/{wid}/undo", f"/workbook/{wid}/edit",
                 f"/workbook/{wid}/propose_save", f"/workbook/{wid}/close", "/open/local", "/a2ui/action"):
        call(port, token, "GET", path)
    assert office.approvals.verify_chain()["events"] == chain
    assert json.dumps(office.view(wid)["draft"], sort_keys=True) == snap
    assert wid in office.sessions


def test_dns_rebinding_host_refused(svc):
    _, port, token = svc
    assert call(port, token, "GET", "/status", headers={"Host": f"evil.example:{port}"})[0] == 421
    assert call(port, token, "GET", "/status", headers={"Host": "127.0.0.1:1"})[0] == 421


def test_browser_origin_refused_and_no_cors(svc):
    _, port, token = svc
    st, _, hdrs = call(port, token, "GET", "/status", headers={"Origin": "https://office.uri.local"})
    assert st == 403
    assert not any(k.lower().startswith("access-control-") for k in hdrs)
    st, _, hdrs = call(port, token, "OPTIONS", "/status")
    assert st == 405 and not any(k.lower().startswith("access-control-") for k in hdrs)
    st, _, hdrs = call(port, token, "GET", "/status")
    assert not any(k.lower().startswith("access-control-") for k in hdrs)


def test_unknown_routes_and_principal(svc):
    _, port, token = svc
    assert call(port, token, "POST", "/run", {})[0] == 404
    assert call(port, token, "POST", "/open/local", {"path": "x", "principal": "someone"})[0] == 403


# ---- AC7 file access -------------------------------------------------------------------------------------------------
@pytest.mark.parametrize("bad", ["relative.xlsx", r"\\server\share\a.xlsx", "C:/a/../b.xlsx", "C:/a/b.xlsx:ads",
                                 "C:/a/b.csv", "", None])
def test_open_refuses_bad_paths(svc, bad):
    _, port, token = svc
    st, out, _ = call(port, token, "POST", "/open/local", {"path": bad})
    assert st == 400 and out["error"] == "PATH_REFUSED"


def test_edit_and_save_only_after_hash_bound_approval(svc, rich):
    office, port, token = svc
    st, wb, _ = call(port, token, "POST", "/open/local", {"path": str(rich)})
    assert st == 200 and wb["engine"] == "local"
    wid = wb["workbook_id"]
    before = rich.read_bytes()
    assert call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})[0] == 200
    assert rich.read_bytes() == before                                   # the draft never touches the file
    st, prop, _ = call(port, token, "POST", f"/workbook/{wid}/propose_save", {})
    assert st == 200
    act = action_for(prop["messages"], "apply", "office.apply_change")
    bad = json.loads(json.dumps(act))
    bad["context"]["change_sha256"] = "0" * 64
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": bad})
    assert st == 400 and out["error"] == "BOUND_VALUE_MISMATCH"
    stale = json.loads(json.dumps(act))
    stale["context"]["surface_digest"] = "f" * 64
    assert call(port, token, "POST", "/a2ui/action", {"action": stale})[1]["error"] == "STALE_SURFACE"
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": act})
    assert st == 200 and out["status"] == "applied"
    assert load_workbook(rich)["Sales"]["B2"].value == 12
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": act})          # replay
    assert st == 400 and out["error"] == "UNKNOWN_SURFACE"
    assert office.approvals.verify_chain()["ok"]


def test_reject_is_hash_bound_and_closes(svc, rich):
    _, port, token = svc
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})
    prop = call(port, token, "POST", f"/workbook/{wid}/propose_save", {})[1]
    act = action_for(prop["messages"], "reject", "office.reject_change")
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": act})
    assert st == 200 and out["status"] == "declined"
    assert load_workbook(rich)["Sales"]["B2"].value == 10
    assert call(port, token, "POST", "/a2ui/action", {"action": act})[1]["error"] == "UNKNOWN_SURFACE"


def test_wrong_button_refused(svc, rich):
    _, port, token = svc
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})
    prop = call(port, token, "POST", f"/workbook/{wid}/propose_save", {})[1]
    act = action_for(prop["messages"], "reject", "office.apply_change")
    assert call(port, token, "POST", "/a2ui/action", {"action": act})[1]["error"] == "NOT_FROM_THIS_BUTTON"


# ---- AC5 conflicts -----------------------------------------------------------------------------------------------
def _prepared(port, token, rich):
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})
    prop = call(port, token, "POST", f"/workbook/{wid}/propose_save", {})[1]
    return wid, action_for(prop["messages"], "apply", "office.apply_change")


def test_file_changed_on_disk_is_conflict(svc, rich, tmp_path):
    _, port, token = svc
    _, act = _prepared(port, token, rich)
    other = tmp_path / "other.xlsx"
    shutil.copy(rich, other)
    wb = load_workbook(other)
    wb["Sales"]["A2"] = "changed elsewhere"
    wb.save(rich)
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": act})
    assert st == 409 and out["error"] == "TARGET_CHANGED"
    assert load_workbook(rich)["Sales"]["A2"].value == "changed elsewhere"


def test_file_open_in_excel_is_conflict(svc, rich):
    _, port, token = svc
    _, act = _prepared(port, token, rich)
    (rich.parent / ("~$" + rich.name)).write_bytes(b"lock")
    st, out, _ = call(port, token, "POST", "/a2ui/action", {"action": act})
    assert st == 409 and out["error"] == "FILE_OPEN_ELSEWHERE"
    assert load_workbook(rich)["Sales"]["B2"].value == 10


def test_save_as_goes_only_to_the_granted_new_file(svc, rich, tmp_path):
    _, port, token = svc
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    call(port, token, "POST", f"/workbook/{wid}/edit", {"sheet": "Sales", "cell": "B2", "value": 12})
    dest = tmp_path / "copy.xlsx"
    prop = call(port, token, "POST", f"/workbook/{wid}/propose_save", {"save_as_path": str(dest)})[1]
    out = call(port, token, "POST", "/a2ui/action", {"action": action_for(prop["messages"], "apply",
                                                                        "office.apply_change")})[1]
    assert out["status"] == "applied" and out["saved_to"] == "copy.xlsx"
    assert load_workbook(dest)["Sales"]["B2"].value == 12 and load_workbook(rich)["Sales"]["B2"].value == 10
    st, out, _ = call(port, token, "POST", f"/workbook/{wid}/propose_save", {"save_as_path": str(rich)})
    assert st == 409 and out["error"] == "DESTINATION_EXISTS"


def test_closed_workbook_grant_is_revoked(svc, rich):
    office, port, token = svc
    wid = call(port, token, "POST", "/open/local", {"path": str(rich)})[1]["workbook_id"]
    gid = office.sessions[wid].grant_id
    call(port, token, "POST", f"/workbook/{wid}/close", {})
    with pytest.raises(OfficeError) as e:
        office.grants.get(gid, "open")
    assert e.value.code == "NO_GRANT"
