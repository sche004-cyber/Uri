"""Independent review S3 of the Office LOW follow-up (0d9b20dd, UI V2 session's reviewer): each test failed before its
repair (owner 2026-10-17: "Repair S3, re-review, then ask")."""
import hashlib
import os
import stat
import sys
import time

import pytest
from openpyxl import Workbook

from uri_office import xlsx_patch as X
from uri_office.brain import OfficeBrain
from uri_office.ops import describe_op
from uri_office.service import Office

from tests.office.test_local_rows_service import FakeClient, op


# ---- S3-04: the delete card always says how many rows are not shown ----------------------------------------------------
def test_a_long_delete_whose_captured_rows_are_empty_still_counts_the_rest():
    text = describe_op({"op": "delete_rows", "sheet": "S", "at": 1, "count": 60, "shown": [[] for _ in range(50)]})
    assert text == "Delete rows 1-60 on 'S' (the first 50 rows are empty; and 10 more rows not shown)"


# ---- S3-05 / S3-06: per-op refusals keep the rest of the answer ----------------------------------------------------------
def _two_sheets(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    for r in range(1, 8):
        ws.cell(r, 1, f"b{r}")
    wb.create_sheet("Other")["A1"] = "o1"
    wb.save(path)
    return path


def _office(tmp_path, *answers):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path], brain=OfficeBrain(FakeClient(*answers)))
    p = _two_sheets(tmp_path / "two.xlsx")
    return o, o.open_local(str(p))["workbook_id"]


def test_a_pending_row_refusal_keeps_a_cell_edit_on_another_sheet(tmp_path):
    o, wid = _office(tmp_path,
                     {"message": "i", "status": "propose", "ops": [op(op="insert_rows", sheet="Budget", at=2, count=1)]},
                     {"message": "x", "status": "propose", "ops": [op(op="delete_rows", sheet="Budget", at=5, count=1),
                                                                   op(op="set_values", sheet="Other", range="A1",
                                                                      values=[["changed"]])]})
    r = o.instruct(wid, {"text": "add a row", "active_sheet": "Budget"})
    o.decide_all(wid, {"decision": "accept", "draft_version": r["draft"]["version"],
                       "set_sha256": o._session(wid).draft.proposed_set_sha256()})
    r = o.instruct(wid, {"text": "delete row 5 and change Other!A1", "active_sheet": "Budget"})
    assert [x["code"] for x in r["refused"]] == ["ROW_CHANGE_PENDING"]
    assert [(c["kind"], c.get("sheet")) for c in r["draft"]["changes"] if c["state"] == "proposed"] == [("cell", "Other")]


def test_an_invalid_row_op_does_not_refuse_cell_edits_on_its_sheet(tmp_path):
    o, wid = _office(tmp_path, {"message": "x", "status": "propose", "ops": [
        op(op="delete_rows", sheet="Budget", at=3, count=0), op(op="set_values", sheet="Budget", range="B2", values=[[5]])]})
    r = o.instruct(wid, {"text": "fix it", "active_sheet": "Budget"})
    assert [x["code"] for x in r["refused"]] == ["BAD_ROWS"]                  # not ROW_AND_CELL_CHANGES
    assert [(c["kind"], c.get("cell")) for c in r["draft"]["changes"]] == [("cell", "B2")]


# ---- S3-03: on Windows the saved file keeps the read-only bit of the original ------------------------------------------
@pytest.mark.skipif(sys.platform != "win32", reason="PLATFORM-GATED: Windows read-only attribute; NOT VERIFIED here")
def test_save_as_from_a_read_only_original_keeps_the_saved_file_read_only(tmp_path):
    src, dst = tmp_path / "a.xlsx", tmp_path / "new.xlsx"
    wb = Workbook()
    wb.active["A1"] = 1
    wb.save(src)
    os.chmod(src, stat.S_IREAD)
    try:
        X.patch(src, dst, [{"sheet": "Sheet", "cell": "A1", "value": 2}], create_only=True)
        assert not os.stat(dst).st_mode & stat.S_IWRITE                    # was cleared by the leftover cleanup
        assert sorted(p.name for p in tmp_path.iterdir()) == ["a.xlsx", "new.xlsx"]      # no leftover
    finally:
        for p in tmp_path.iterdir():
            os.chmod(p, stat.S_IWRITE | stat.S_IREAD)


@pytest.mark.skipif(sys.platform != "win32", reason="PLATFORM-GATED: Windows read-only attribute; NOT VERIFIED here")
def test_a_restored_read_only_hold_stays_read_only(tmp_path):
    deck = tmp_path / "Deck.pptx"
    data = b"ORIGINAL"
    h = X._hold_name(deck, hashlib.sha256(data).hexdigest()[:12], int(time.time()) - 5, "0123abcd")
    h.write_bytes(data)
    os.chmod(h, stat.S_IREAD)
    try:
        assert X.recover_held(deck) == deck and deck.read_bytes() == data and not h.exists()
        assert not os.stat(deck).st_mode & stat.S_IWRITE
    finally:
        for p in tmp_path.iterdir():
            os.chmod(p, stat.S_IWRITE | stat.S_IREAD)
