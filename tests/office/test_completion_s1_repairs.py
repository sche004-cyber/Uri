"""Independent reviews S1 of the completion candidate (3e202efc): each test reproduces the original defect on real files.
Own subagent review: S1-01 defined names that look like cells past column XFD, S1-02 deleting an animated shape.
UI V2 session's reviewer: S1-01 a second row change on one sheet before a save, S1-02 the Excel save path."""
import hashlib
import io
import time
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.util import Inches

from uri_office import xlsx_patch as X
from uri_office import xlsx_struct as S
from uri_office.authoring import pptx_ops as P
from uri_office.errors import OfficeError

from tests.office.test_local_rows_service import FakeClient, act, budget, op  # noqa: F401  (fixture helpers)
from uri_office.brain import OfficeBrain
from uri_office.service import Office


# ---- own S1-01: a defined name is never shifted as a cell -------------------------------------------------------------
def test_row_insert_keeps_a_defined_name_that_looks_like_a_cell_past_column_xfd(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["C1"] = 0.1
    for r in range(2, 9):
        ws.cell(r, 2, r * 100)
    ws["D8"] = "=B8*YTD2024"
    ws["E8"] = "=XFE10+B8"
    wb.defined_names["YTD2024"] = DefinedName("YTD2024", attr_text="Budget!$C$1")
    wb.defined_names["XFE10"] = DefinedName("XFE10", attr_text="Budget!$C$1")
    src = tmp_path / "b.xlsx"
    wb.save(src)
    X.patch(src, src, [], [{"op": "insert_rows", "sheet": "Budget", "at": 6, "count": 1}])
    ws = load_workbook(src)["Budget"]
    assert ws["D9"].value == "=B9*YTD2024" and ws["E9"].value == "=XFE10+B9"
    sh = S._Shift("delete_rows", 2024, 1)
    assert S.shift_formula("YTD2024*2", sh, "Budget", "Budget") == "YTD2024*2"     # never #REF!
    assert S.shift_formula("XFD2025", sh, "Budget", "Budget") == "XFD2024"          # the last column is a cell


# ---- UI V2 reviewer S1-01: one row change per sheet until it is saved -------------------------------------------------
@pytest.fixture
def office(tmp_path):
    def make(*answers):
        o = Office(tmp_path / "state", allowed_roots=[tmp_path], brain=OfficeBrain(FakeClient(*answers)))
        p = budget(tmp_path / "budget.xlsx")
        return o, o.open_local(str(p))["workbook_id"], p
    return make


def test_a_second_row_change_on_the_sheet_before_a_save_is_refused(office):
    o, wid, p = office(
        {"message": "Insert.", "status": "propose", "ops": [op(op="insert_rows", sheet="Budget", at=4, count=1)]},
        {"message": "Delete.", "status": "propose", "ops": [op(op="delete_rows", sheet="Budget", at=6, count=1)]})
    r = o.instruct(wid, {"text": "Add a row above row 4", "active_sheet": "Budget"})
    o.decide_all(wid, {"decision": "accept", "draft_version": r["draft"]["version"],
                       "set_sha256": o._session(wid).draft.proposed_set_sha256()})
    before = [dict(c) for c in o.view(wid)["draft"]["changes"]]
    r = o.instruct(wid, {"text": "Delete row 6", "active_sheet": "Budget"})   # S2-02: refused per op, not an error
    assert [x["code"] for x in r["refused"]] == ["ROW_CHANGE_PENDING"]
    assert o.view(wid)["draft"]["changes"] == before                    # nothing else changed
    with pytest.raises(OfficeError) as e:                               # the draft still refuses it as a last check
        o._session(wid).draft.propose([{"op": "delete_rows", "sheet": "Budget", "at": 6, "count": 1}])
    assert e.value.code == "ROW_CHANGE_PENDING"


def test_two_row_changes_on_one_sheet_in_one_answer_keep_only_the_first(office):
    o, wid, p = office({"message": "x", "status": "propose", "ops": [
        op(op="insert_rows", sheet="Budget", at=4, count=1), op(op="delete_rows", sheet="Budget", at=6, count=1)]})
    r = o.instruct(wid, {"text": "Add a row and delete another", "active_sheet": "Budget"})
    assert [x["code"] for x in r["refused"]] == ["ROW_CHANGE_PENDING"]
    assert [c["op"]["op"] for c in r["draft"]["changes"]] == ["insert_rows"]


def test_after_the_save_a_new_row_change_is_allowed(office):
    o, wid, p = office(
        {"message": "Insert.", "status": "propose", "ops": [op(op="insert_rows", sheet="Budget", at=4, count=1)]},
        {"message": "Delete.", "status": "propose", "ops": [op(op="delete_rows", sheet="Budget", at=7, count=1)]})
    r = o.instruct(wid, {"text": "Add a row above row 4", "active_sheet": "Budget"})
    o.decide_all(wid, {"decision": "accept", "draft_version": r["draft"]["version"],
                       "set_sha256": o._session(wid).draft.proposed_set_sha256()})
    assert o.action(act(o.propose_save(wid, {})))["status"] == "applied"
    r = o.instruct(wid, {"text": "Delete row 7", "active_sheet": "Budget"})
    assert [c["op"]["op"] for c in r["draft"]["changes"]] == ["delete_rows"]


# ---- UI V2 reviewer S1-02: the Excel save path publishes like authoring ------------------------------------------------
def _book(path, value=1):
    wb = Workbook()
    wb.active["A1"] = value
    wb.save(path)
    return path


def test_excel_save_as_never_replaces_a_file_created_meanwhile(tmp_path):
    src, dst = _book(tmp_path / "a.xlsx"), tmp_path / "new.xlsx"
    dst.write_bytes(b"ANOTHER PROGRAM")                                 # appears after the service checked the name
    with pytest.raises(OfficeError) as e:
        X.patch(src, dst, [{"sheet": "Sheet", "cell": "A1", "value": 2}], create_only=True)
    assert e.value.code == "DESTINATION_EXISTS" and dst.read_bytes() == b"ANOTHER PROGRAM"


def test_excel_save_rechecks_the_original_at_publish_and_keeps_every_backup(tmp_path):
    src, bk = _book(tmp_path / "a.xlsx"), tmp_path / "bk"
    bk.mkdir()
    sha = hashlib.sha256(src.read_bytes()).hexdigest()
    other = _book(tmp_path / "other.xlsx", value=9).read_bytes()
    src.write_bytes(other)                                              # changed after the service's check
    with pytest.raises(OfficeError) as e:
        X.patch(src, src, [{"sheet": "Sheet", "cell": "A1", "value": 2}], expected_sha256=sha,
                backup=bk / f"{time.time_ns()}-a.xlsx")
    assert e.value.code == "TARGET_CHANGED" and src.read_bytes() == other
    for v in (3, 4):                                                    # two saves in the same second: two backups
        sha = hashlib.sha256(src.read_bytes()).hexdigest()
        X.patch(src, src, [{"sheet": "Sheet", "cell": "A1", "value": v}], expected_sha256=sha,
                backup=bk / f"{time.time_ns()}-a.xlsx")
    assert len(list(bk.iterdir())) == 2 and load_workbook(src).active["A1"].value == 4


# ---- own S1-02: an animated shape is not deleted ---------------------------------------------------------------------
def _animated_deck():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    tb = s.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    tb.text_frame.text = "Fly in"
    plain = s.shapes.add_textbox(Inches(1), Inches(3), Inches(3), Inches(1))
    sid = tb.shape_id
    s._element.append(parse_xml(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:anim><p:cBhvr>'
        '<p:cTn id="2" dur="500"/><p:tgtEl><p:spTgt spid="%d"/></p:tgtEl></p:cBhvr></p:anim></p:childTnLst>'
        '</p:cTn></p:par></p:tnLst></p:timing>' % sid))
    b = io.BytesIO()
    prs.save(b)
    prs = Presentation(io.BytesIO(b.getvalue()))
    return prs, prs.slides[0].slide_id, sid, plain.shape_id


def test_deleting_an_animated_shape_is_refused_and_the_slide_is_unchanged():
    prs, slide_id, animated, plain = _animated_deck()
    before = prs.slides[0]._element.xml
    with pytest.raises(OfficeError) as e:
        P.delete_shape(prs, slide_id, animated)
    assert e.value.code == "AUTHORING_NOT_SUPPORTED" and "animated" in e.value.detail
    assert prs.slides[0]._element.xml == before
    P.delete_shape(prs, slide_id, plain)                                # a shape without an animation still goes
    assert all(sh.shape_id != plain for sh in prs.slides[0].shapes)
