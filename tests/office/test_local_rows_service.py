"""OD-R3 / XB-2 end to end in the service: the Brain proposes a real row insert or delete for a local workbook, the
owner accepts it, the save card says it in plain words, and the saved file holds the moved rows."""
import json
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from uri_office import ops as O
from uri_office.brain import OfficeBrain
from uri_office.errors import OfficeError
from uri_office.service import Office


class FakeClient:
    def __init__(self, *answers):
        self.answers, self.bodies = list(answers), []

    def stream_chat(self, body, *, cancel=None, deadlines=None):
        self.bodies.append(body)
        return SimpleNamespace(status="COMPLETE", text=json.dumps(self.answers.pop(0)), ttft_s=0.1, total_s=0.2,
                               timings={})


def op(**kw):
    return {**{k: None for k in O.BRAIN_SCHEMA["properties"]["ops"]["items"]["required"]}, **kw}


def budget(path, table=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    for r in (("Head", "Allocation", "Spent"), ("Salaries", 1250000, 610500), ("Digitization", 480000, 212750),
              ("Stationery", 35000, 18420), ("Travel", 60000, 22100), ("Maintenance", 90000, 41380)):
        ws.append(list(r))
    ws["A8"], ws["B8"] = "Total", "=SUM(B2:B6)"
    if table:
        ws.add_table(Table(displayName="T", ref="A1:C6"))
    wb.save(path)
    return path


def act(prop):
    m = prop["messages"]
    model = m[2]["updateDataModel"]["value"]
    return {"name": "office.apply_change", "surfaceId": m[0]["createSurface"]["surfaceId"], "sourceComponentId": "apply",
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


def card_text(prop) -> str:
    return json.dumps(prop["messages"], ensure_ascii=False)


@pytest.fixture
def session(tmp_path):
    def make(*answers, table=False):
        o = Office(tmp_path / "state", allowed_roots=[tmp_path], brain=OfficeBrain(FakeClient(*answers)))
        p = budget(tmp_path / "budget.xlsx", table=table)
        return o, o.open_local(str(p))["workbook_id"], p
    return make


def test_insert_below_digitization_with_values_is_a_real_insert(session):
    o, wid, p = session({"message": "Adds Travel 2 below Digitization.", "status": "propose", "ops": [
        op(op="insert_rows", sheet="Budget", at=4, count=1, values=[["Travel 2", 25000]])]})
    r = o.instruct(wid, {"text": "Add a row for Travel 2 with 25000 below Digitization", "active_sheet": "Budget"})
    [ch] = r["draft"]["changes"]
    assert ch["kind"] == "struct" and ch["state"] == "proposed"
    assert ch["text"] == "Insert 1 row above row 4 on 'Budget': Travel 2  |  25000"
    o.decide_all(wid, {"decision": "accept", "draft_version": r["draft"]["version"],
                       "set_sha256": o._session(wid).draft.proposed_set_sha256()})
    prop = o.propose_save(wid, {})
    assert "Insert 1 row above row 4 on 'Budget': Travel 2" in card_text(prop)
    res = o.action(act(prop))
    assert res["status"] == "applied" and res["layout_reloaded"] is True and res["layout_version"] == 1
    import zipfile
    with zipfile.ZipFile(p) as z:
        assert b'fullCalcOnLoad="1"' in z.read("xl/workbook.xml")      # totals recalculate when the file is opened
    ws = load_workbook(p)["Budget"]
    assert [ws.cell(r, 1).value for r in range(1, 10)] == ["Head", "Salaries", "Digitization", "Travel 2", "Stationery",
                                                            "Travel", "Maintenance", None, "Total"]
    assert ws["B4"].value == 25000 and ws["B9"].value == "=SUM(B2:B7)"
    v = o.view(wid)
    assert v["layout_version"] == 1 and v["draft"]["changes"] == []
    assert v["workbook"]["sheets"][0]["cells"]["A4"]["v"] == "Travel 2"


def test_delete_card_shows_what_the_row_holds_now(session):
    o, wid, p = session({"message": "Deletes Travel.", "status": "propose", "ops": [
        op(op="delete_rows", sheet="Budget", at=5, count=1)]})
    r = o.instruct(wid, {"text": "Remove the Travel row", "active_sheet": "Budget"})
    [ch] = r["draft"]["changes"]
    assert ch["text"] == "Delete row 5 on 'Budget' (now: Travel  |  60000  |  22100)"


def test_row_change_and_cell_change_on_the_same_sheet_are_not_mixed(session):
    o, wid, p = session({"message": "x", "status": "propose", "ops": [
        op(op="insert_rows", sheet="Budget", at=4, count=1, values=None),
        op(op="set_formula", sheet="Budget", cell="B8", formula="=SUM(B2:B7)")]})
    r = o.instruct(wid, {"text": "Add a row and fix the total", "active_sheet": "Budget"})
    assert [x["code"] for x in r["refused"]] == ["ROW_AND_CELL_CHANGES"]
    assert [c["kind"] for c in r["draft"]["changes"]] == ["struct"]


def test_an_unsupported_workbook_is_refused_at_save_and_the_file_is_unchanged(session):
    o, wid, p = session({"message": "x", "status": "propose", "ops": [
        op(op="insert_rows", sheet="Budget", at=3, count=1, values=None)]}, table=True)
    before = p.read_bytes()
    r = o.instruct(wid, {"text": "Insert a row above row 3", "active_sheet": "Budget"})
    o.decide_all(wid, {"decision": "accept", "draft_version": r["draft"]["version"],
                       "set_sha256": o._session(wid).draft.proposed_set_sha256()})
    with pytest.raises(OfficeError) as e:
        o.action(act(o.propose_save(wid, {})))
    assert e.value.code == "STRUCTURE_NOT_SUPPORTED" and "table" in e.value.detail
    assert p.read_bytes() == before


def test_values_on_a_delete_or_on_google_are_refused():
    with pytest.raises(OfficeError):
        O.to_cell_edits(op(op="delete_rows", sheet="S", at=2, count=1, values=[[1]]), engine="local", sheets=["S"])
    with pytest.raises(OfficeError):
        O.to_cell_edits(op(op="insert_rows", sheet="S", at=2, count=1, values=[[1]]), engine="google", sheets=["S"])
    with pytest.raises(OfficeError):
        O.to_cell_edits(op(op="insert_rows", sheet="S", at=2, count=1, values=[[1], [2]]), engine="local", sheets=["S"])
    got = O.to_cell_edits(op(op="insert_rows", sheet="S", at=2, count=2, values=[["a", "=SUM(B1:B3)"]]),
                          engine="local", sheets=["S"])
    assert got[0]["values"] == [[{"v": "a"}, {"f": "=SUM(B1:B3)"}]]
    with pytest.raises(OfficeError):        # network formulas stay refused inside inserted rows too
        O.to_cell_edits(op(op="insert_rows", sheet="S", at=2, count=1, values=[["=WEBSERVICE(\"x\")"]]),
                        engine="local", sheets=["S"])


def test_the_local_prompt_explains_real_row_operations(session):
    o, wid, p = session({"message": "ok", "status": "propose", "ops": []})
    o.instruct(wid, {"text": "x", "active_sheet": "Budget"})
    system = o.brain.client.bodies[0]["messages"][0]["content"]
    assert "insert_rows" in system and "ABOVE row `at`" in system and "never imitate an insert" in system
