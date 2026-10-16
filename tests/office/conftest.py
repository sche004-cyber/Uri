"""Fixtures for URI Office tests: a rich workbook (chart, image, styles, merges, dates, shared/inline strings, shared and
array formulas, optional calcChain) built locally; no network, no Brain, no Google."""
from __future__ import annotations

import datetime as dt
import io
import struct
import sys
import zipfile
import zlib
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import Workbook  # noqa: E402
from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.drawing.image import Image  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402


def _png() -> bytes:
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * 4 for _ in range(4))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 4, 4, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def rewrite(path: Path, fn):
    """Rewrite package parts in place: fn(parts: dict) -> None."""
    with zipfile.ZipFile(path) as z:
        infos = z.infolist()
        parts = {i.filename: z.read(i.filename) for i in infos}
    fn(parts)
    names = [i.filename for i in infos]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for i in infos:
            if i.filename in parts:
                z.writestr(i, parts[i.filename])
        for n, data in parts.items():
            if n not in names:
                z.writestr(n, data)


def build_rich(path: Path, *, excel_like=False, calc_chain=False, special_formulas=True):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Amount", "Note"])
    for r, (reg, amt) in enumerate([("North", 10), ("South", 20), ("East", 30)], start=2):
        ws.append([reg, amt, f"row {r}"])
    ws["B5"] = "=SUM(B2:B4)"
    ws["A1"].font = Font(bold=True, color="FF0000")
    ws["B1"].fill = PatternFill("solid", fgColor="FFFF00")
    ws["B3"].font = Font(italic=True)
    ws.merge_cells("D1:E1")
    ws["D1"] = "Merged header"
    ws.column_dimensions["A"].width = 22
    ws["E2"] = dt.datetime(2026, 1, 15)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "G2")
    ws.add_image(Image(io.BytesIO(_png())), "G20")
    other = wb.create_sheet("Other")
    other["A1"] = "untouched sheet"
    other["A2"] = "=Sales!B5*2"
    wb.save(path)

    def edit(parts):
        x = parts["xl/worksheets/sheet1.xml"].decode()
        # openpyxl writes inline strings; Excel writes a shared-string table. Convert C2..C4 to shared strings.
        for i, r in enumerate((2, 3, 4)):
            x = x.replace(f'<c r="C{r}" t="inlineStr"><is><t>row {r}</t></is></c>', f'<c r="C{r}" t="s"><v>{i}</v></c>')
        parts["xl/sharedStrings.xml"] = (
            b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="3" uniqueCount="3">'
            b'<si><t>row 2</t></si><si><t>row 3</t></si><si><t>row 4</t></si></sst>')
        parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
            b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
            b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
        parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
            b"</Relationships>", b'<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.org/'
            b'officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
        extra = '<row r="6"><c r="C6" t="inlineStr"><is><t>inline text</t></is></c></row>'
        if special_formulas:
            extra += ('<row r="7"><c r="B7"><f t="shared" ref="B7:B8" si="0">B2*2</f><v>20</v></c></row>'
                      '<row r="8"><c r="B8"><f t="shared" si="0"/><v>40</v></c>'
                      '<c r="C8"><f t="array" ref="C8">SUM(B2:B4*1)</f><v>60</v></c></row>')
        parts["xl/worksheets/sheet1.xml"] = x.replace("</sheetData>", extra + "</sheetData>").encode()
        if excel_like:
            parts["xl/workbook.xml"] = parts["xl/workbook.xml"].replace(b' fullCalcOnLoad="1"', b"")
        if calc_chain:
            parts["xl/calcChain.xml"] = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                                         b'<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                                         b'<c r="B5" i="1"/><c r="A2" i="2"/></calcChain>')
            parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
                b"</Types>", b'<Override PartName="/xl/calcChain.xml" ContentType="application/'
                b'vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
            parts["xl/_rels/workbook.xml.rels"] = parts["xl/_rels/workbook.xml.rels"].replace(
                b"</Relationships>", b'<Relationship Id="rIdCalc" Type="http://schemas.openxmlformats.org/'
                b'officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
    rewrite(path, edit)
    return path


@pytest.fixture
def rich(tmp_path) -> Path:
    return build_rich(tmp_path / "rich.xlsx")
