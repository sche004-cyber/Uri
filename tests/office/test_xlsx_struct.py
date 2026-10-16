"""OD-R3 / XB-2: whole-row insert and delete at the XML level (uri_office/xlsx_struct.py)."""
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation

from uri_office import xlsx_struct as S
from uri_office.errors import OfficeError


def budget_bytes(extra=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    rows = [("Head", "Allocation", "Spent"), ("Salaries", 1250000, 610500), ("Digitization", 480000, 212750),
            ("Stationery", 35000, 18420), ("Travel", 60000, 22100), ("Maintenance", 90000, 41380)]
    for r in rows:
        ws.append(list(r))
    ws["A8"] = "Total"
    ws["B8"] = "=SUM(B2:B6)"
    ws["C8"] = "=SUM($C$2:$C$6)"
    ws["D3"] = "=B3-C3"
    other = wb.create_sheet("Notes & 'Plan'")
    other["A1"] = "=Budget!B8"
    other["A2"] = "='Budget'!B5*2"
    other["A3"] = "=SUM(Budget!B2:B6)"
    other["A4"] = "=B5"                       # unqualified: this sheet's B5, must not move
    if extra:
        extra(wb, ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parts_of(data: bytes) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def to_wb(parts: dict):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    return load_workbook(io.BytesIO(buf.getvalue()))


def test_insert_shifts_cells_and_every_reference():
    parts = parts_of(budget_bytes())
    S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 4, "count": 1})
    wb = to_wb(parts)
    ws, other = wb["Budget"], wb["Notes & 'Plan'"]
    assert [ws.cell(r, 1).value for r in range(1, 10)] == ["Head", "Salaries", "Digitization", None, "Stationery",
                                                            "Travel", "Maintenance", None, "Total"]
    assert ws["B9"].value == "=SUM(B2:B7)"          # range around the new row expands
    assert ws["C9"].value == "=SUM($C$2:$C$7)"
    assert ws["D3"].value == "=B3-C3"               # above the insert: unchanged
    assert other["A1"].value == "=Budget!B9"
    assert other["A2"].value == "='Budget'!B6*2"
    assert other["A3"].value == "=SUM(Budget!B2:B7)"
    assert other["A4"].value == "=B5"               # another sheet's own cell is not moved


def test_insert_refuses_to_push_a_cell_past_excel_last_row_without_changes():
    parts = parts_of(budget_bytes(lambda wb, ws: setattr(ws["A1048576"], "value", "Last row")))
    # A stale or absent dimension must not be the only thing guarding real row elements.
    from uri_office.xlsx_patch import M, _ser, _xml
    sheet = "xl/worksheets/sheet1.xml"
    root = _xml(parts[sheet])
    root.remove(root.find(M + "dimension"))
    parts[sheet] = _ser(root)
    before = dict(parts)
    with pytest.raises(OfficeError) as error:
        S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 1, "count": 1})
    assert error.value.code == "STRUCTURE_NOT_SUPPORTED"
    assert parts == before


def test_delete_shrinks_ranges_and_breaks_references_to_deleted_cells():
    parts = parts_of(budget_bytes())
    S.apply_row_op(parts, {"op": "delete_rows", "sheet": "Budget", "at": 5, "count": 1})   # Travel
    wb = to_wb(parts)
    ws, other = wb["Budget"], wb["Notes & 'Plan'"]
    assert [ws.cell(r, 1).value for r in range(1, 8)] == ["Head", "Salaries", "Digitization", "Stationery",
                                                           "Maintenance", None, "Total"]
    assert ws["B7"].value == "=SUM(B2:B5)"
    assert other["A1"].value == "=Budget!B7"
    assert other["A2"].value == "='Budget'!#REF!*2"  # pointed at the deleted Travel row
    assert other["A3"].value == "=SUM(Budget!B2:B5)"


def test_merges_conditional_formats_validation_hyperlinks_and_names_move():
    def extra(wb, ws):
        ws.merge_cells("A10:C11")
        ws.conditional_formatting.add("B2:B6", CellIsRule(operator="greaterThan", formula=["$B$5"]))
        dv = DataValidation(type="whole", operator="greaterThan", formula1="0")
        dv.add("B2:B6")
        ws.add_data_validation(dv)
        ws["A6"].hyperlink = "https://example.invalid/m"
        from openpyxl.workbook.defined_name import DefinedName
        wb.defined_names["Amounts"] = DefinedName("Amounts", attr_text="Budget!$B$2:$B$6")
    parts = parts_of(budget_bytes(extra))
    S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 3, "count": 2})
    wb = to_wb(parts)
    ws = wb["Budget"]
    assert [str(m) for m in ws.merged_cells.ranges] == ["A12:C13"]
    cf = list(ws.conditional_formatting)
    assert str(cf[0].sqref) == "B2:B8" and cf[0].rules[0].formula == ["$B$7"]
    assert str(ws.data_validations.dataValidation[0].sqref) == "B2:B8"
    assert ws["A8"].hyperlink is not None
    assert wb.defined_names["Amounts"].attr_text == "Budget!$B$2:$B$8"


def test_shared_formulas_are_written_out_before_moving():
    parts = parts_of(budget_bytes())
    sheet = "xl/worksheets/sheet1.xml"
    xml = parts[sheet].decode()
    # D2:D6 as one shared group (master D2 = B2-C2), the way Excel saves a fill-down
    xml = xml.replace('<c r="D3"><f>B3-C3</f><v></v></c>', "")
    cells = ['<c r="D2"><f t="shared" ref="D2:D6" si="0">B2-C2</f></c>'] + \
            [f'<c r="D{r}"><f t="shared" si="0"/></c>' for r in range(3, 7)]
    for r, c in zip(range(2, 7), cells):
        xml = xml.replace(f'</c></row><row r="{r + 1}"', f'</c>{c}</row><row r="{r + 1}"', 1) if r < 6 else \
            xml.replace('</c></row><row r="8"', f'</c>{c}</row><row r="8"', 1)
    parts[sheet] = xml.encode()
    assert xml.count('t="shared"') == 5 and xml.count('r="D3"') == 1, "fixture: the shared group must be in place"
    S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 4, "count": 1})
    ws = to_wb(parts)["Budget"]
    got = {r: ws[f"D{r}"].value for r in range(2, 8)}
    assert got[2] == "=B2-C2" and got[3] == "=B3-C3" and got[4] is None
    assert got[5] == "=B5-C5" and got[7] == "=B7-C7"
    assert b't="shared"' not in parts[sheet]


@pytest.mark.parametrize("extra,why", [
    (lambda wb, ws: ws.protection.__setattr__("sheet", True), "protected"),
])
def test_unsupported_structures_are_refused_and_nothing_changes(extra, why):
    parts = parts_of(budget_bytes(extra))
    before = dict(parts)
    with pytest.raises(OfficeError) as e:
        S.apply_row_op(parts, {"op": "insert_rows", "sheet": "Budget", "at": 2, "count": 1})
    assert e.value.code == "STRUCTURE_NOT_SUPPORTED" and why in e.value.detail
    assert parts == before


def test_tables_and_charts_are_refused():
    from openpyxl.chart import BarChart, Reference
    from openpyxl.worksheet.table import Table

    def with_table(wb, ws):
        ws.add_table(Table(displayName="T1", ref="A1:C6"))

    def with_chart(wb, ws):
        ch = BarChart()
        ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=6), titles_from_data=True)
        wb["Notes & 'Plan'"].add_chart(ch, "C3")        # the chart lives on the other sheet but reads Budget
    for extra in (with_table, with_chart):
        parts = parts_of(budget_bytes(extra))
        before = dict(parts)
        with pytest.raises(OfficeError) as e:
            S.apply_row_op(parts, {"op": "delete_rows", "sheet": "Budget", "at": 3, "count": 1})
        assert e.value.code == "STRUCTURE_NOT_SUPPORTED"
        assert parts == before


def test_a_deletion_through_merged_cells_is_refused():
    parts = parts_of(budget_bytes(lambda wb, ws: ws.merge_cells("A4:B5")))
    with pytest.raises(OfficeError) as e:
        S.apply_row_op(parts, {"op": "delete_rows", "sheet": "Budget", "at": 5, "count": 1})
    assert "merged" in e.value.detail


def test_shift_formula_unit_cases():
    ins = S._Shift("insert_rows", 4, 2)
    assert S.shift_formula("=A3+A4+$A$10", ins, "S", "S") == "=A3+A6+$A$12"
    assert S.shift_formula("=SUM(A:A)+SUM(3:5)", ins, "S", "S") == "=SUM(A:A)+SUM(3:7)"
    assert S.shift_formula("=Other!A5+A5", ins, "S", "Other") == "=Other!A5+A5"
    assert S.shift_formula("=s!A5", ins, "S", "X") == "=s!A7"              # sheet names are case-insensitive
    assert S.shift_formula('="A5"&A5', ins, "S", "S") == '="A5"&A7'          # text is not a reference
    assert S.shift_formula("=MyName+A1", ins, "S", "S") == "=MyName+A1"
    dele = S._Shift("delete_rows", 4, 2)                                      # rows 4..5
    assert S.shift_formula("=SUM(A2:A10)", dele, "S", "S") == "=SUM(A2:A8)"
    assert S.shift_formula("=SUM(A4:A5)", dele, "S", "S") == "=SUM(#REF!)"
    assert S.shift_formula("=SUM(A5:A9)", dele, "S", "S") == "=SUM(A4:A7)"
    with pytest.raises(OfficeError):
        S.shift_formula("=SUM(Sheet1:Sheet3!A1)", dele, "S", "S")


def test_repeated_same_sheet_range_shifts_both_qualified_endpoints():
    shift = S._Shift("insert_rows", 3, 1)
    assert S.shift_formula("SUM(Sheet1!A1:Sheet1!A5)", shift, "Sheet1", "Other") == \
        "SUM(Sheet1!A1:Sheet1!A6)"
    assert S.shift_formula("SUM(Sheet1!A1:Sheet1!A5)", shift, "Sheet2", "Other") == \
        "SUM(Sheet1!A1:Sheet1!A5)"
    with pytest.raises(OfficeError) as error:
        S.shift_formula("SUM(Sheet1!A1:Sheet2!A5)", shift, "Sheet1", "Other")
    assert error.value.code == "STRUCTURE_NOT_SUPPORTED" and "3-D" in error.value.detail


@pytest.mark.parametrize("op", [{"op": "insert_rows", "sheet": "Budget", "at": 0, "count": 1},
                                {"op": "insert_rows", "sheet": "Budget", "at": 2, "count": 0},
                                {"op": "insert_rows", "sheet": "Budget", "at": 2, "count": 5000},
                                {"op": "delete_rows", "sheet": "Nope", "at": 2, "count": 1},
                                {"op": "move_rows", "sheet": "Budget", "at": 2, "count": 1}])
def test_bad_ops_are_refused(op):
    parts = parts_of(budget_bytes())
    with pytest.raises(OfficeError):
        S.apply_row_op(parts, op)
