"""AC4 on a workbook saved by real Microsoft Excel (tests/office/fixtures/README.md has its provenance)."""
import hashlib
import re
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from uri_office import xlsx_model, xlsx_patch as XP
from uri_office.errors import OfficeError

FIX = Path(__file__).parent / "fixtures" / "excel_financial_sample.xlsx"
SHA = "c3f17156ab7c192571ecfc742e88c16a4b7243a4d4b4a420fcb68dec44e10196"


@pytest.fixture
def excel(tmp_path):
    assert hashlib.sha256(FIX.read_bytes()).hexdigest() == SHA, "fixture changed"
    dst = tmp_path / "financial.xlsx"
    shutil.copy(FIX, dst)
    return dst


def parts(p):
    with zipfile.ZipFile(p) as z:
        return {n: z.read(n) for n in z.namelist()}


def test_fixture_really_is_excel_saved(excel):
    app = parts(excel)["docProps/app.xml"].decode()
    assert "<Application>Microsoft Excel</Application>" in app
    assert b'count="4216" uniqueCount="50"' in parts(excel)["xl/sharedStrings.xml"]


def test_patch_keeps_every_other_part_and_the_shared_string_table(excel, tmp_path):
    out = tmp_path / "out.xlsx"
    XP.patch(excel, out, [{"sheet": "Sheet1", "cell": "A2", "value": "Changed segment"},   # a shared-string cell
                          {"sheet": "Sheet1", "cell": "F2", "value": 1234},
                          {"sheet": "Sheet1", "cell": "R5", "value": "note outside the table"}])
    a, b = parts(excel), parts(out)
    assert set(a) == set(b)
    assert {n for n in a if a[n] != b[n]} == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    assert a["xl/sharedStrings.xml"] == b["xl/sharedStrings.xml"]
    x = b["xl/worksheets/sheet1.xml"].decode()
    for ref in ("A2", "R5"):
        c = re.search(r'<c r="%s"[^>]*>.*?</c>' % ref, x, re.S).group(0)
        assert 't="inlineStr"' in c and "<v>" not in c
    # Excel's markup survives: mc:Ignorable prefixes declared, row attributes kept, other cells untouched.
    root = x[:x.index(">", x.index("<worksheet")) + 1]
    assert 'mc:Ignorable="x14ac' in root and 'xmlns:x14ac=' in root
    assert '<row r="1" spans="1:16" x14ac:dyDescent="0.25">' in x
    assert '<c r="B2" t="s"><v>16</v></c>' in x
    ws = load_workbook(out)["Sheet1"]
    assert (ws["A2"].value, ws["F2"].value, ws["R5"].value) == ("Changed segment", 1234, "note outside the table")
    assert ws["A3"].value == load_workbook(excel)["Sheet1"]["A3"].value


def test_table_header_cells_are_locked(excel, tmp_path):
    assert "A1" in XP.locked_cells(excel)["Sheet1"] and "P1" in XP.locked_cells(excel)["Sheet1"]
    with pytest.raises(OfficeError) as e:
        XP.patch(excel, tmp_path / "out.xlsx", [{"sheet": "Sheet1", "cell": "A1", "value": "Renamed"}])
    assert e.value.code == "CELL_NOT_EDITABLE"


def test_model_reads_excel_file(excel):
    wb = xlsx_model.read(excel)
    s = wb["sheets"][0]
    assert s["name"] == "Sheet1" and "A1" in s["locked"] and s["cells"]["A1"]["v"] == "Segment"
