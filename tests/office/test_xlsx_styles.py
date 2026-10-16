"""OD-R3 / XF-1: cell formatting is saved (uri_office/xlsx_styles.py) and only what changed is written."""
import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from uri_office import xlsx_model
from uri_office import xlsx_styles as X
from uri_office.errors import OfficeError


def book_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "Head"
    ws["A1"].font = Font(name="Georgia", sz=14, color="FF336699")          # size, family, colour must survive
    ws["A1"].border = Border(bottom=Side(style="thin"))
    ws["B2"] = 1234.5
    ws["C3"] = "theme"
    ws["C3"].font = Font(color=None)
    ws["C3"].font = Font(name="Calibri", sz=11)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parts_of(data):
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def reopen(parts):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    return load_workbook(io.BytesIO(buf.getvalue()))


def test_bold_keeps_font_size_family_colour_and_border():
    parts = parts_of(book_bytes())
    X.apply_formats(parts, [{"sheet": "S", "cell": "A1", "before": {"fc": "#336699"},
                             "after": {"fc": "#336699", "b": True}}])
    c = reopen(parts)["S"]["A1"]
    assert c.font.b is True and c.font.name == "Georgia" and c.font.sz == 14
    assert c.font.color.rgb == "FF336699"
    assert c.border.bottom.style == "thin"


def test_fill_number_format_alignment_wrap_and_new_cell():
    parts = parts_of(book_bytes())
    X.apply_formats(parts, [
        {"sheet": "S", "cell": "B2", "before": {}, "after": {"bg": "#ffff00", "nf": "#,##0.00", "ha": "right"}},
        {"sheet": "S", "cell": "D9", "before": {}, "after": {"i": True, "wrap": True, "nf": "0.0%"}},
        {"sheet": "S", "cell": "B3", "before": {}, "after": {"nf": "dd-mmm-yyyy hh:mm \"custom\""}},
    ])
    ws = reopen(parts)["S"]
    assert ws["B2"].fill.fill_type == "solid" and ws["B2"].fill.fgColor.rgb == "FFFFFF00"
    assert ws["B2"].number_format == "#,##0.00" and ws["B2"].alignment.horizontal == "right"
    assert ws["B2"].value == 1234.5
    assert ws["D9"].font.i is True and ws["D9"].alignment.wrap_text and ws["D9"].number_format == "0.0%"
    assert ws["B3"].number_format == 'dd-mmm-yyyy hh:mm "custom"'


def test_removing_bold_and_fill():
    parts = parts_of(book_bytes())
    X.apply_formats(parts, [{"sheet": "S", "cell": "B2", "before": {}, "after": {"b": True, "bg": "#00FF00"}}])
    X.apply_formats(parts, [{"sheet": "S", "cell": "B2", "before": {"b": True, "bg": "#00FF00"}, "after": {}}])
    c = reopen(parts)["S"]["B2"]
    assert not c.font.b and c.fill.fill_type in (None, "none")


def test_round_trip_through_urIs_reader():
    """What URI reads back after the save is what the owner set (the grid shows the saved formatting)."""
    parts = parts_of(book_bytes())
    X.apply_formats(parts, [{"sheet": "S", "cell": "B2", "before": {}, "after": {"b": True, "u": True, "fc": "#C00000",
                                                                               "bg": "#DDEBF7", "ha": "center"}}])
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in parts.items():
            z.writestr(n, d)
    import tempfile
    import pathlib
    p = pathlib.Path(tempfile.mkdtemp()) / "b.xlsx"
    p.write_bytes(buf.getvalue())
    wb = xlsx_model.read(p)
    st = wb["styles"][wb["sheets"][0]["cells"]["B2"]["s"]]
    assert st == {"b": True, "u": True, "fc": "#C00000", "bg": "#DDEBF7", "ha": "center"}


def test_identical_formats_reuse_one_record():
    parts = parts_of(book_bytes())
    X.apply_formats(parts, [{"sheet": "S", "cell": c, "before": {}, "after": {"b": True}} for c in ("E1", "E2", "E3")])
    xml = parts["xl/styles.xml"].decode()
    ws = reopen(parts)["S"]
    assert ws["E1"].style_id == ws["E2"].style_id == ws["E3"].style_id
    assert xml.count("<b/>") == 1


@pytest.mark.parametrize("st", [{"bold": True}, {"b": 1}, {"fc": "red"}, {"ha": "justify"}, {"nf": "x" * 300},
                                {"nf": "\x01"}, "b"])
def test_bad_styles_are_refused(st):
    with pytest.raises(OfficeError):
        X.check_style(st)
