"""AC4: only the allowed parts change; everything else stays byte-identical (unpacked part hashes)."""
import datetime as dt
import hashlib
import re
import zipfile

import pytest
from openpyxl import load_workbook

from tests.office.conftest import build_rich
from uri_office import xlsx_patch as XP
from uri_office.errors import OfficeError


def hashes(p):
    with zipfile.ZipFile(p) as z:
        return {n: hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()}


def part(p, name):
    with zipfile.ZipFile(p) as z:
        return z.read(name).decode()


def cell_xml(xml, ref):
    m = re.search(r'<c r="%s"[^>]*?(?:/>|>.*?</c>)' % ref, xml, re.S)
    return m.group(0) if m else None


def changed(a, b):
    ha, hb = hashes(a), hashes(b)
    return {n for n in ha if hb.get(n) != ha[n]} | (set(hb) - set(ha))


def test_values_only_change_one_sheet_part(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    XP.patch(rich, dst, [{"sheet": "Sales", "cell": "B3", "value": 25}, {"sheet": "Sales", "cell": "A4", "value": "West"},
                         {"sheet": "Sales", "cell": "C6", "value": "inline edited"},
                         {"sheet": "Sales", "cell": "F40", "value": True}])
    assert changed(rich, dst) == {"xl/worksheets/sheet1.xml"}
    ws = load_workbook(dst)["Sales"]
    assert (ws["B3"].value, ws["A4"].value, ws["C6"].value, ws["F40"].value) == (25, "West", "inline edited", True)
    assert ws["B5"].value == "=SUM(B2:B4)"
    assert load_workbook(dst)["Other"]["A2"].value == "=Sales!B5*2"


def test_chart_image_drawing_and_shared_strings_untouched(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    XP.patch(rich, dst, [{"sheet": "Sales", "cell": "A2", "value": "Norte"}])
    before, after = hashes(rich), hashes(dst)
    for n in before:
        if n.startswith(("xl/charts/", "xl/media/", "xl/drawings/")) or n == "xl/sharedStrings.xml":
            assert after[n] == before[n], n
    ss = part(dst, "xl/sharedStrings.xml")
    assert ss == part(rich, "xl/sharedStrings.xml")       # count/uniqueCount untouched: strings are written inline
    assert 'count="3" uniqueCount="3"' in ss


def test_new_and_changed_text_is_inline_and_shared_strings_byte_identical(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    XP.patch(rich, dst, [{"sheet": "Sales", "cell": "C3", "value": "edited"},      # was a shared string
                         {"sheet": "Sales", "cell": "H9", "value": "brand new"}])  # new text
    assert changed(rich, dst) == {"xl/worksheets/sheet1.xml"}
    with zipfile.ZipFile(rich) as a, zipfile.ZipFile(dst) as b:
        assert a.read("xl/sharedStrings.xml") == b.read("xl/sharedStrings.xml")
    x = part(dst, "xl/worksheets/sheet1.xml")
    for ref in ("C3", "H9"):
        c = cell_xml(x, ref)
        assert 't="inlineStr"' in c and "<v>" not in c, c
    assert 't="s"' in cell_xml(x, "C2") and 't="s"' in cell_xml(x, "C4")
    ws = load_workbook(dst)["Sales"]
    assert (ws["C2"].value, ws["C3"].value, ws["C4"].value) == ("row 2", "edited", "row 4")


def test_cell_style_kept_and_cached_value_stripped_on_formula(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    s_before = re.search(r's="(\d+)"', cell_xml(part(rich, "xl/worksheets/sheet1.xml"), "B3")).group(1)
    XP.patch(rich, dst, [{"sheet": "Sales", "cell": "B3", "formula": "=B2+1"},
                         {"sheet": "Sales", "cell": "B5", "formula": "=SUM(B2:B3)"}])
    x = part(dst, "xl/worksheets/sheet1.xml")
    b3 = cell_xml(x, "B3")
    assert f's="{s_before}"' in b3
    assert "<v>" not in b3 and "<v>" not in cell_xml(x, "B5")
    assert load_workbook(dst)["Sales"]["B3"].font.i is True


def test_date_serial_cells_unchanged(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    e2 = cell_xml(part(rich, "xl/worksheets/sheet1.xml"), "E2")
    XP.patch(rich, dst, [{"sheet": "Sales", "cell": "A3", "value": "Sur"}])
    assert cell_xml(part(dst, "xl/worksheets/sheet1.xml"), "E2") == e2
    assert load_workbook(dst)["Sales"]["E2"].value == dt.datetime(2026, 1, 15)


def test_excel_like_calcpr_gets_full_calc_on_load(tmp_path):
    src = build_rich(tmp_path / "x.xlsx", excel_like=True)
    dst = tmp_path / "out.xlsx"
    XP.patch(src, dst, [{"sheet": "Sales", "cell": "B2", "value": 11}])
    assert changed(src, dst) == {"xl/worksheets/sheet1.xml", "xl/workbook.xml"}
    assert 'fullCalcOnLoad="1"' in part(dst, "xl/workbook.xml")


def test_formula_change_drops_calc_chain_and_its_references(tmp_path):
    src = build_rich(tmp_path / "x.xlsx", calc_chain=True)
    dst = tmp_path / "out.xlsx"
    XP.patch(src, dst, [{"sheet": "Sales", "cell": "B6", "formula": "=B5*2"}])
    assert changed(src, dst) <= {"xl/worksheets/sheet1.xml", "xl/workbook.xml", "xl/calcChain.xml",
                                 "[Content_Types].xml", "xl/_rels/workbook.xml.rels"}
    with zipfile.ZipFile(dst) as z:
        assert "xl/calcChain.xml" not in z.namelist()
    assert "calcChain" not in part(dst, "[Content_Types].xml")
    assert "calcChain" not in part(dst, "xl/_rels/workbook.xml.rels")


def test_value_edit_keeps_calc_chain(tmp_path):
    src = build_rich(tmp_path / "x.xlsx", calc_chain=True)
    dst = tmp_path / "out.xlsx"
    XP.patch(src, dst, [{"sheet": "Sales", "cell": "B2", "value": 1}])
    with zipfile.ZipFile(dst) as z:
        assert "xl/calcChain.xml" in z.namelist()


@pytest.mark.parametrize("edit,code", [
    ({"sheet": "Sales", "cell": "E1", "value": "x"}, "CELL_NOT_EDITABLE"),      # merged, not the anchor
    ({"sheet": "Sales", "cell": "B7", "value": 1}, "CELL_NOT_EDITABLE"),        # shared formula master
    ({"sheet": "Sales", "cell": "B8", "value": 1}, "CELL_NOT_EDITABLE"),        # shared formula child
    ({"sheet": "Sales", "cell": "C8", "value": 1}, "CELL_NOT_EDITABLE"),        # array formula
    ({"sheet": "Nope", "cell": "A1", "value": 1}, "UNKNOWN_SHEET"),
    ({"sheet": "Sales", "cell": "1A", "value": 1}, "BAD_CELL_REF"),
    ({"sheet": "Sales", "cell": "A1", "formula": "SUM(1)"}, "BAD_FORMULA"),
    ({"sheet": "Sales", "cell": "A1", "value": float("nan")}, "BAD_VALUE"),
])
def test_refusals_leave_no_output(rich, tmp_path, edit, code):
    dst = tmp_path / "out.xlsx"
    with pytest.raises(OfficeError) as e:
        XP.patch(rich, dst, [edit])
    assert e.value.code == code
    assert not dst.exists() and not list(tmp_path.glob("*.tmp"))


def test_locked_cells_reported(rich):
    locked = XP.locked_cells(rich)["Sales"]
    assert {"B7", "B8", "C8", "E1"} <= set(locked) and "D1" not in locked


def test_in_place_patch_is_atomic(rich):
    XP.patch(rich, rich, [{"sheet": "Sales", "cell": "B2", "value": 99}])
    assert load_workbook(rich)["Sales"]["B2"].value == 99
    assert not list(rich.parent.glob("*.uri-office.tmp"))


def test_not_an_xlsx(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    with pytest.raises(OfficeError) as e:
        XP.patch(bad, tmp_path / "o.xlsx", [{"sheet": "S", "cell": "A1", "value": 1}])
    assert e.value.code == "NOT_AN_XLSX"
