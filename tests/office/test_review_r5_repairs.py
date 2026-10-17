"""Regression tests for the URI-OFFICE-M1 review round 5 findings (each failed before its repair):

R5-01  an approved Save As committed the draft although the session keeps editing the original file, so a later save
       of the original (or a later Save As) silently left out the earlier edits;
R5-02  the service's workbook snapshot never followed a save (the panel reloaded the as-opened values) and the Brain
       planned from the as-opened values instead of the draft.
"""
import pytest
from openpyxl import Workbook, load_workbook

from uri_office.service import Office


def _act(prop):
    m = prop["messages"]
    model = m[2]["updateDataModel"]["value"]
    return {"name": "office.apply_change", "surfaceId": m[0]["createSurface"]["surfaceId"], "sourceComponentId": "apply",
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


@pytest.fixture
def office(tmp_path):
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 1
    p = tmp_path / "w.xlsx"
    wb.save(p)
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    return o, o.open_local(str(p))["workbook_id"], p, tmp_path


def _cells(view_or_result):
    return view_or_result["workbook"]["sheets"][0]["cells"]


# ---- R5-01 -----------------------------------------------------------------------------------------------------------
def test_save_as_then_save_keeps_earlier_edits(office):
    o, wid, p, tmp = office
    o.edit(wid, {"sheet": "S", "cell": "A1", "value": 99})
    r = o.action(_act(o.propose_save(wid, {"save_as_path": str(tmp / "copy.xlsx")})))
    assert r["status"] == "applied"
    o.edit(wid, {"sheet": "S", "cell": "B1", "value": 5})
    prop = o.propose_save(wid, {})
    shown = [(e["cell"], e["after"]) for e in prop["messages"][2]["updateDataModel"]["value"]["bound"]["edits"]] \
        if "edits" in prop["messages"][2]["updateDataModel"]["value"]["bound"] else None
    r = o.action(_act(prop))
    assert r["status"] == "applied"
    ws = load_workbook(p)["S"]
    assert (ws["A1"].value, ws["B1"].value) == (99, 5)              # was (1, 5): the A1 edit was lost
    c = load_workbook(tmp / "copy.xlsx")["S"]
    assert (c["A1"].value, c["B1"].value) == (99, None)
    if shown is not None:
        assert ("A1", {"v": 99}) in shown


def test_second_save_as_carries_earlier_edits(office):
    o, wid, p, tmp = office
    o.edit(wid, {"sheet": "S", "cell": "A1", "value": 99})
    o.action(_act(o.propose_save(wid, {"save_as_path": str(tmp / "c1.xlsx")})))
    o.edit(wid, {"sheet": "S", "cell": "B1", "value": 5})
    o.action(_act(o.propose_save(wid, {"save_as_path": str(tmp / "c2.xlsx")})))
    c2 = load_workbook(tmp / "c2.xlsx")["S"]
    assert (c2["A1"].value, c2["B1"].value) == (99, 5)              # was (1, 5)
    assert load_workbook(p)["S"]["A1"].value == 1                   # the original is untouched by Save As
    assert {e["cell"] for e in o.sessions[wid].draft.pending()["edits"]} == {"A1", "B1"}


# ---- R5-02 -----------------------------------------------------------------------------------------------------------
def test_save_refreshes_workbook_snapshot_and_action_result(office):
    o, wid, p, _ = office
    o.edit(wid, {"sheet": "S", "cell": "A1", "value": 99})
    o.edit(wid, {"sheet": "S", "cell": "C4", "value": "new"})
    r = o.action(_act(o.propose_save(wid, {})))
    assert r["status"] == "applied" and r["draft"]["changes"] == []
    assert load_workbook(p)["S"]["A1"].value == 99
    assert _cells(r)["A1"]["v"] == 99                                # the panel reloads from this
    assert _cells(o.view(wid))["A1"]["v"] == 99                      # was {'v': 1}
    sh = o.view(wid)["workbook"]["sheets"][0]
    assert sh["cells"]["C4"]["v"] == "new" and sh["max_row"] >= 4 and sh["max_col"] >= 3


def test_save_as_leaves_snapshot_on_the_original(office):
    o, wid, _, tmp = office
    o.edit(wid, {"sheet": "S", "cell": "A1", "value": 99})
    r = o.action(_act(o.propose_save(wid, {"save_as_path": str(tmp / "copy.xlsx")})))
    assert _cells(r)["A1"]["v"] == 1                                 # the session's file still holds 1
    assert [c["cell"] for c in r["draft"]["changes"]] == ["A1"]      # and the edit is still pending for it


class _Brain:
    def __init__(self):
        self.sheets = None

    def plan(self, text, *, sheets, active, engine, locked, history_rows, **_):
        self.sheets = sheets
        return {"message": "ok", "edits": [], "refused": [], "metrics": {}}


def test_brain_plans_from_draft_values(office):
    o, wid, _, _ = office
    o.brain = _Brain()
    o.edit(wid, {"sheet": "S", "cell": "A1", "value": 99})
    o.edit(wid, {"sheet": "S", "cell": "B2", "formula": "=A1*2"})
    o.instruct(wid, {"text": "double it", "active_sheet": "S"})
    cells = o.brain.sheets[0]["cells"]
    assert cells["A1"]["v"] == 99                                    # was 1: the as-opened value
    assert cells["B2"]["f"] == "=A1*2"
    assert _cells(o.view(wid))["A1"]["v"] == 1                       # the snapshot itself is not changed by a draft edit
