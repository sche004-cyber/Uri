"""Regression tests for the URI-OFFICE-M1 review round 1 findings R1-02 .. R1-09 (each failed before its repair)."""
import json
import os
import stat
import threading

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula

from uri_a2ui import catalog as C
from uri_office import approvals as A
from uri_office import google as G
from uri_office import ops as O
from uri_office import refs
from uri_office import xlsx_patch as XP
from uri_office.draft import Draft
from uri_office.errors import OfficeError
from uri_office.service import Office


def _act(prop, button="apply", name="office.apply_change"):
    m = prop["messages"]
    model = m[2]["updateDataModel"]["value"]
    return {"name": name, "surfaceId": m[0]["createSurface"]["surfaceId"], "sourceComponentId": button,
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


def _texts(prop) -> list:
    comps = [c for msg in prop["messages"] for c in (msg.get("updateComponents") or {}).get("components", [])]
    return [c["text"] for c in comps if c.get("component") == "Text"]


@pytest.fixture
def office(tmp_path, rich):
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    return o, o.open_local(str(rich))["workbook_id"]


# ---- R1-02: network functions refused even when only named, and the newer ones too ----------------------------------
@pytest.mark.parametrize("f", [
    "=MAP(A1,WEBSERVICE)", "=BYROW(A1:A3,IMAGE)", '=LET(f,WEBSERVICE,f("x"))', "=MAP(A1:A3,HYPERLINK)",
    "=STOCKHISTORY(A1,A2)", "=TRANSLATE(A1)", "=DETECTLANGUAGE(A1)", '=COPILOT("x",A1)', '=CUBEVALUE("c")',
    '=CUBEMEMBER("c","m")', '=SQL.REQUEST("dsn")', "=_xlfn.IMAGE(A1)", "=ENCODEURL(A1)", '=GOOGLEFINANCE("X")',
    '=AI("summarise",A1:A3)'])
def test_network_function_refused_as_bare_token(f):
    with pytest.raises(OfficeError) as e:
        O.check_formula(f)
    assert e.value.code == "FORMULA_REFUSED"


@pytest.mark.parametrize("f", ["=SUM(A1:A3)", "=SUM(AI:AI)", '=IF(A1="IMAGE","pic","none")', "=CALL1*2"])
def test_ordinary_formulas_still_accepted(f):
    assert O.check_formula(f) == f


# ---- R1-03: every cell of an array formula's range is locked and refused ---------------------------------------------
@pytest.fixture
def array_book(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    for r in range(1, 4):
        ws.cell(r, 1, r)
    ws["F1"] = ArrayFormula("F1:F3", "=A1:A3*2")
    p = tmp_path / "array.xlsx"
    wb.save(p)
    return p


def test_array_formula_range_is_locked(array_book):
    assert {"F1", "F2", "F3"} <= set(XP.locked_cells(array_book)["S"])


@pytest.mark.parametrize("cell", ["F1", "F2", "F3"])
def test_array_formula_range_cells_refused(array_book, tmp_path, cell):
    dst = tmp_path / "out.xlsx"
    with pytest.raises(OfficeError) as e:
        XP.patch(array_book, dst, [{"sheet": "S", "cell": cell, "value": 1}])
    assert e.value.code == "CELL_NOT_EDITABLE"
    assert not dst.exists()


# ---- R1-04: refs with trailing newlines are refused everywhere ------------------------------------------------------
@pytest.mark.parametrize("ref", ["A1\n", "A1\r\n", " A1", "A1 "])
def test_cell_ref_must_match_exactly(ref):
    with pytest.raises(OfficeError) as e:
        refs.parse_cell(ref)
    assert e.value.code == "BAD_CELL_REF"


def test_draft_lock_cannot_be_bypassed_with_newline():
    d = Draft("w", "local", {"sheets": [{"name": "S", "cells": {"A1": {"v": 1}}, "locked": ["A1"]}]})
    with pytest.raises(OfficeError) as e:
        d.user_edit({"sheet": "S", "cell": "A1\n", "value": 2})
    assert e.value.code == "BAD_CELL_REF"
    with pytest.raises(OfficeError) as e:
        d.propose([{"sheet": "S", "cell": "A1\n", "value": 2}])
    assert e.value.code == "BAD_CELL_REF"
    assert d.pending()["edits"] == []


def test_patch_refuses_newline_ref_and_writes_nothing(rich, tmp_path):
    dst = tmp_path / "out.xlsx"
    with pytest.raises(OfficeError) as e:
        XP.patch(rich, dst, [{"sheet": "Sales", "cell": "C9\n", "value": "x"}])
    assert e.value.code == "BAD_CELL_REF"
    assert not dst.exists()


def test_service_edit_refuses_newline_ref(office):
    o, wid = office
    with pytest.raises(OfficeError) as e:
        o.edit(wid, {"sheet": "Sales", "cell": "B2\n", "formula": "=1+1"})
    assert e.value.code == "BAD_CELL_REF"


# ---- R1-05: an approval is bound to the draft version it was shown for ---------------------------------------------
def test_edit_after_approval_shown_makes_it_stale_and_keeps_the_edit(office, rich):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "B2", "value": 12})
    prop = o.propose_save(wid, {})
    o.edit(wid, {"sheet": "Sales", "cell": "B3", "value": 999})          # typed after the card was shown
    with pytest.raises(OfficeError) as e:
        o.action(_act(prop))
    assert e.value.code == "STALE_CHANGE"
    ws = load_workbook(rich)["Sales"]
    assert (ws["B2"].value, ws["B3"].value) == (10, 20)                 # nothing written
    pending = {x["cell"]: x["after"] for x in o.sessions[wid].draft.pending()["edits"]}
    assert pending == {"B2": {"v": 12}, "B3": {"v": 999}}                # nothing dropped
    with pytest.raises(OfficeError) as e:                                # the stale card is closed
        o.action(_act(prop))
    assert e.value.code == "UNKNOWN_SURFACE"
    o.action(_act(o.propose_save(wid, {})))                              # a fresh approval saves both
    ws = load_workbook(rich)["Sales"]
    assert (ws["B2"].value, ws["B3"].value) == (12, 999)


# ---- R1-06: the card shows everything the hash binds ----------------------------------------------------------------
def test_card_shows_full_values_every_line_and_full_path(office, rich):
    o, wid = office
    long = "Pay vendor ACME 100 USD per the approved plan. " * 3 + "and also wire 1,000,000 USD to account 999"
    o.edit(wid, {"sheet": "Sales", "cell": "A2", "value": long})
    for r in range(10, 50):
        o.edit(wid, {"sheet": "Sales", "cell": f"H{r}", "value": r})
    prop = o.propose_save(wid, {})
    shown = "\n".join(_texts(prop))
    assert json.dumps(long) in shown
    for r in range(10, 50):
        assert f"Sales!H{r}: (empty) -> {r}" in shown
    assert str(rich) in _texts(prop)


def test_card_text_cannot_fake_a_line(office):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "A2", "value": "x\nSales!B2: 10 -> 10"})
    shown = "\n".join(_texts(o.propose_save(wid, {})))
    assert "\nSales!B2: 10 -> 10" not in shown


def test_split_value_is_marked_as_continued(office):
    o, wid = office
    # The injected text starts exactly where the value is split between two texts of the card.
    prefix = 'Sales!J1: (empty) -> "'
    o.edit(wid, {"sheet": "Sales", "cell": "J1",
                 "value": "x" * (C.MAX_STRING - len(prefix)) + "Sales!B2: 10 -> 10"})
    texts = _texts(o.propose_save(wid, {}))
    assert not any(t.startswith("Sales!B2") for t in texts)
    assert any(t.startswith(A.CONTINUED + "Sales!B2: 10 -> 10") for t in texts)


def test_change_set_too_large_to_show_is_refused(office):
    o, wid = office
    big = "y" * 30000
    for r in range(1, 20):
        o.edit(wid, {"sheet": "Sales", "cell": f"J{r}", "value": big})
    with pytest.raises(OfficeError) as e:
        o.propose_save(wid, {})
    assert e.value.code == "CHANGE_TOO_LARGE_TO_SHOW"


def test_every_card_text_fits_the_catalog(office):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "J1", "value": "z" * 9000})
    texts = _texts(o.propose_save(wid, {}))
    assert all(len(t) <= C.MAX_STRING for t in texts)
    assert "z" * 9000 in "".join(t.removeprefix(A.CONTINUED) for t in texts).replace("\n", "")


# ---- R1-07: approve and propose on one workbook at once must not deadlock ------------------------------------------
def test_concurrent_approve_and_propose_do_not_deadlock(office):
    o, wid = office
    o.edit(wid, {"sheet": "Sales", "cell": "B2", "value": 12})
    act = _act(o.propose_save(wid, {}))
    s = o.sessions[wid]
    in_apply, in_propose = threading.Event(), threading.Event()
    real_apply, real_pending = o._apply, s.draft.pending

    def slow_apply(body):                    # widen the window: the approve path is inside its locks
        in_apply.set()
        in_propose.wait(2)
        return real_apply(body)

    def slow_pending():                      # the propose path holds the workbook lock here
        in_propose.set()
        in_apply.wait(2)
        return real_pending()

    o._apply, s.draft.pending = slow_apply, slow_pending
    errors = []

    def run(fn):
        try:
            fn()
        except OfficeError as exc:
            errors.append(exc.code)

    a = threading.Thread(target=run, args=(lambda: o.action(act),), daemon=True)
    b = threading.Thread(target=run, args=(lambda: o.propose_save(wid, {}),), daemon=True)
    a.start()
    b.start()
    a.join(15)
    b.join(15)
    assert not a.is_alive() and not b.is_alive(), "approve and propose deadlocked"
    assert o.approvals.verify_chain()["ok"]


# ---- R1-08: the temp file is never a planted name; permissions are kept --------------------------------------------
@pytest.mark.skipif(os.name == "nt", reason="POSIX symlink and mode semantics")
def test_planted_temp_symlink_is_not_followed(rich, tmp_path):
    outside = tmp_path / "outside"
    outside.mkdir()
    victim = outside / "victim.txt"
    victim.write_text("IMPORTANT")
    os.symlink(victim, rich.with_name(rich.name + ".uri-office.tmp"))
    XP.patch(rich, rich, [{"sheet": "Sales", "cell": "B2", "value": 14}])
    assert victim.read_text() == "IMPORTANT"
    assert not rich.is_symlink()
    assert load_workbook(rich)["Sales"]["B2"].value == 14


@pytest.mark.skipif(os.name == "nt", reason="POSIX mode bits")
def test_save_keeps_file_mode(rich):
    os.chmod(rich, 0o600)
    XP.patch(rich, rich, [{"sheet": "Sales", "cell": "B2", "value": 15}])
    assert stat.S_IMODE(rich.stat().st_mode) == 0o600
    assert not [p for p in rich.parent.iterdir() if p.name.endswith(".uri-office.tmp")]


@pytest.mark.skipif(os.name == "nt", reason="POSIX symlink semantics")
def test_google_export_does_not_follow_planted_temp(tmp_path):
    class Exec:
        def execute(self):
            return b"PK\x03\x04data"

    class Drive:
        def files(self):
            return self

        def export(self, **kw):
            return Exec()

    victim = tmp_path / "victim.txt"
    victim.write_text("IMPORTANT")
    dest = tmp_path / "export.xlsx"
    os.symlink(victim, dest.with_name(dest.name + ".uri-office.tmp"))
    gs = G.GoogleSheets(tmp_path, None, services=lambda: (None, Drive(), set()))
    gs.export_xlsx("x" * 30, dest)
    assert victim.read_text() == "IMPORTANT" and dest.read_bytes() == b"PK\x03\x04data" and not dest.is_symlink()


# ---- R1-09: Google apply checks everything before writing anything --------------------------------------------------
class _X:
    def __init__(self, r):
        self.r = r

    def execute(self):
        return self.r


class _Sheets:
    def __init__(self, titles, current=10):
        self.titles, self.current, self.writes = titles, current, []

    def spreadsheets(self):
        return self

    def values(self):
        return self

    def batchGet(self, spreadsheetId, ranges, valueRenderOption):
        return _X({"valueRanges": [{"values": [[self.current]]} for _ in ranges]})

    def get(self, spreadsheetId, fields):
        return _X({"sheets": [{"properties": {"title": t, "sheetId": i}} for i, t in enumerate(self.titles)]})

    def batchUpdate(self, spreadsheetId, body):
        self.writes.append(body)
        return _X({})


BODY = {"edits": [{"sheet": "Sales", "cell": "B2", "before": {"v": 10}, "after": {"v": 12}}],
        "struct_ops": [{"op": "rename_sheet", "sheet": "Sales", "name": "Q3"}]}


def test_google_refusal_writes_nothing(tmp_path):
    api = _Sheets(["Sales-renamed-by-colleague"])
    gs = G.GoogleSheets(tmp_path, None, services=lambda: (api, None, {G.SHEETS_SCOPE}))
    with pytest.raises(OfficeError) as e:
        gs.apply("f" * 30, BODY)
    assert e.value.code == "TARGET_CHANGED"
    assert api.writes == []


def test_google_cells_and_struct_ops_go_in_one_request(tmp_path):
    api = _Sheets(["Sales"])
    gs = G.GoogleSheets(tmp_path, None, services=lambda: (api, None, {G.SHEETS_SCOPE}))
    assert gs.apply("f" * 30, BODY) == {"cells": 1, "struct_ops": 1}
    assert len(api.writes) == 1
    reqs = api.writes[0]["requests"]
    assert reqs[0]["updateCells"]["start"] == {"sheetId": 0, "rowIndex": 1, "columnIndex": 1}
    assert reqs[0]["updateCells"]["rows"] == [{"values": [{"userEnteredValue": {"numberValue": 12}}]}]
    assert reqs[1] == {"updateSheetProperties": {"properties": {"sheetId": 0, "title": "Q3"}, "fields": "title"}}
