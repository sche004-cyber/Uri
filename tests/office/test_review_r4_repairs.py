"""Regression tests for the URI-OFFICE-M1 review round 4 finding R4-01 (each failed before its repair): text an .xlsx
cannot store (XML-1.0-invalid characters) is refused when proposed, and an unexpected failure while applying an
approved change is audited and closes the approval instead of escaping as a bare exception."""
import pytest
from openpyxl import Workbook, load_workbook

from uri_office import ops as O
from uri_office.errors import OfficeError
from uri_office.service import Office

BAD_CHARS = ["\x0b", "\x01", "\x00", "\x1f", "￾", "￿", "\ud800"]


def _act(prop):
    m = prop["messages"]
    model = m[2]["updateDataModel"]["value"]
    return {"name": "office.apply_change", "surfaceId": m[0]["createSurface"]["surfaceId"], "sourceComponentId": "apply",
            "context": {"surface_digest": model["bound"]["surface"]["digest"],
                        "change_sha256": model["bound"]["change_sha256"]}}


@pytest.fixture
def office(tmp_path):
    wb = Workbook()
    wb.active.title = "S"
    wb.active["A1"] = 1
    p = tmp_path / "b.xlsx"
    wb.save(p)
    o = Office(tmp_path / "state", allowed_roots=[tmp_path])
    return o, o.open_local(str(p))["workbook_id"], p


def _events(o):
    return [r[0] for r in o.approvals.db.execute("SELECT kind FROM event")]


@pytest.mark.parametrize("ch", BAD_CHARS)
def test_xml_invalid_text_refused(ch):
    with pytest.raises(OfficeError) as e:
        O.check_value(f"line1{ch}line2")
    assert e.value.code == "BAD_VALUE"
    with pytest.raises(OfficeError) as e:
        O.check_formula(f'="a{ch}b"')
    assert e.value.code == "BAD_FORMULA"


@pytest.mark.parametrize("v", ["tab\there", "lf\nhere", "cr\rhere", "é中\U0001F600", "�"])
def test_xml_valid_text_accepted(v):
    assert O.check_value(v) == v


@pytest.mark.parametrize("ch", BAD_CHARS)
def test_edit_with_invalid_char_refused_and_save_still_works(office, ch):
    o, wid, p = office
    with pytest.raises(OfficeError) as e:
        o.edit(wid, {"sheet": "S", "cell": "B1", "value": f"line1{ch}line2"})
    assert e.value.code == "BAD_VALUE"
    o.edit(wid, {"sheet": "S", "cell": "C1", "value": "good"})
    o.action(_act(o.propose_save(wid, {})))
    ws = load_workbook(p)["S"]
    assert ws["B1"].value is None and ws["C1"].value == "good"


def test_unexpected_apply_failure_is_audited_and_closes_approval(office):
    o, wid, p = office
    # Bypass validation to get text the writer cannot store into the draft (as any unforeseen apply failure would).
    o.sessions[wid].draft.user_edit({"sheet": "S", "cell": "B1", "value": "x\x0by"})
    o.edit(wid, {"sheet": "S", "cell": "C1", "value": "good"})
    act = _act(o.propose_save(wid, {}))
    with pytest.raises(OfficeError) as e:
        o.action(act)
    assert e.value.code == "APPLY_FAILED" and e.value.status == 500
    states = [r[0] for r in o.approvals.db.execute("SELECT state FROM surface")]
    assert "OPEN" not in states
    assert _events(o)[-1] == "office.action_refused"
    assert o.approvals.verify_chain()["ok"]
    assert load_workbook(p)["S"]["C1"].value is None                     # nothing was written
    with pytest.raises(OfficeError) as e:                                # the closed approval cannot be replayed
        o.action(act)
    assert e.value.code == "UNKNOWN_SURFACE"
