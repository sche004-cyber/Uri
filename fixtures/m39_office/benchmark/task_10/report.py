from openpyxl import load_workbook, Workbook
wb = load_workbook("sales.xlsx")
ws = wb["Sales"]
header = [c.value for c in ws[1]]
col = header.index("Date")
count = sum(1 for row in ws.iter_rows(min_row=2, values_only=True) if row[col] == "2026-01-15")
out = Workbook(); o = out.active; o.title = "Summary"
o.append(["MatchCount"]); o.append([count])
out.save("summary.xlsx")
