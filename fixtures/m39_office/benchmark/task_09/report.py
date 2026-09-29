from openpyxl import load_workbook, Workbook
wb = load_workbook("sales_old.xlsx")
ws = wb["Sales"]
header = [c.value for c in ws[1]]
col = header.index("Amt")
total = sum(row[col] for row in ws.iter_rows(min_row=2, values_only=True) if row[col] is not None)
out = Workbook(); o = out.active; o.title = "Summary"
o.append(["Total"]); o.append([total])
out.save("summary.xlsx")
