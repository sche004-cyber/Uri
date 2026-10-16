"""PowerPoint ops on a python-pptx Presentation (milestone B; plan 11_AUTHORING_A_B_PLAN.md section 3).

Every op names its target by stable ids: slide id (`p:sldId/@id`, unchanged by moves) and shape id (`p:cNvPr/@id`).
Ops only touch the parts they are about; anything else in the package (charts, media, other slides) is left as
python-pptx loaded it. An op that would need content it does not handle is refused with AUTHORING_NOT_SUPPORTED and a
named reason, before anything changes.
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart
from pptx.util import Emu, Pt

from ..errors import OfficeError

MAX_TEXT = 20_000
MAX_TABLE = 50
_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
SLIDE_W_FALLBACK = 9144000


def _refuse(why: str):
    raise OfficeError("AUTHORING_NOT_SUPPORTED", why)


def _text(v, name="text") -> str:
    if not isinstance(v, str) or len(v) > MAX_TEXT:
        raise OfficeError("BAD_TEXT", name)
    return v


def slide(prs, sid):
    for s in prs.slides:
        if s.slide_id == sid:
            return s
    raise OfficeError("UNKNOWN_SLIDE", str(sid)[:16])


def shape(sl, shid):
    for sh in sl.shapes:
        if sh.shape_id == shid:
            return sh
    raise OfficeError("UNKNOWN_SHAPE", str(shid)[:16])


def layout(prs, name):
    for master in prs.slide_masters:
        for lay in master.slide_layouts:
            if lay.name == name:
                return lay
    raise OfficeError("UNKNOWN_LAYOUT", str(name)[:60])


def layouts(prs) -> list:
    return [lay.name for m in prs.slide_masters for lay in m.slide_layouts]


def _box(prs, box):
    """{x, y, w, h} as fractions of the slide (0..1), so URI and the UI never deal in EMUs."""
    if not isinstance(box, dict) or set(box) != {"x", "y", "w", "h"} or \
            not all(isinstance(box[k], (int, float)) and 0 <= box[k] <= 1 for k in box) or box["w"] <= 0 or box["h"] <= 0:
        raise OfficeError("BAD_BOX", "x, y, w, h as fractions of the slide")
    W, H = prs.slide_width or SLIDE_W_FALLBACK, prs.slide_height or 6858000
    return Emu(int(box["x"] * W)), Emu(int(box["y"] * H)), Emu(int(box["w"] * W)), Emu(int(box["h"] * H))


# ---- slides --------------------------------------------------------------------------------------------------------
def add_slide(prs, layout_name: str, after=None) -> int:
    new = prs.slides.add_slide(layout(prs, layout_name))
    if after is not None:
        move_slide(prs, new.slide_id, _index(prs, after) + 1)
    return new.slide_id


def _index(prs, sid) -> int:
    for i, s in enumerate(prs.slides):
        if s.slide_id == sid:
            return i
    raise OfficeError("UNKNOWN_SLIDE", str(sid)[:16])


def move_slide(prs, sid, to: int):
    lst = prs.slides._sldIdLst
    items = list(lst)
    i = _index(prs, sid)
    if not isinstance(to, int) or not 0 <= to < len(items):
        raise OfficeError("BAD_POSITION", str(to)[:8])
    el = items[i]
    lst.remove(el)
    lst.insert(to, el)


def delete_slide(prs, sid):
    if len(prs.slides) <= 1:
        _refuse("a presentation keeps at least one slide")
    lst = prs.slides._sldIdLst
    el = list(lst)[_index(prs, sid)]
    prs.part.drop_rel(el.rId)
    lst.remove(el)


_CHART_PART_TEMPLATES = {
    "application/vnd.ms-office.chartstyle+xml": "/ppt/charts/style%d.xml",
    "application/vnd.ms-office.chartcolorstyle+xml": "/ppt/charts/colors%d.xml",
    "application/vnd.openxmlformats-officedocument.drawingml.chartshapes+xml": "/ppt/drawings/drawing%d.xml",
    "application/vnd.openxmlformats-officedocument.themeOverride+xml": "/ppt/theme/themeOverride%d.xml",
}


def _chart_copy_plan(old) -> list:
    """What copying a chart part needs, or a named refusal (checked before any change)."""
    plan = []
    for rId, rel in old.rels.items():
        if rel.reltype in (RT.PACKAGE, RT.OLE_OBJECT) and rel.is_external:
            _refuse("the chart's data is linked to another file")      # independent review C2-06
        if rel.reltype == RT.PACKAGE:
            continue                                  # the embedded workbook: copied through chart_workbook below
        if rel.is_external:
            plan.append((rId, rel, None))
            continue
        part = rel.target_part
        if len(part.rels):
            _refuse("the chart has " + rel.reltype.rsplit("/", 1)[-1] + " content with its own links")
        template = _CHART_PART_TEMPLATES.get(part.content_type)
        if template is None:
            _refuse("the chart has content URI cannot copy yet (" + rel.reltype.rsplit("/", 1)[-1] + ")")
        plan.append((rId, rel, (template, part)))
    return plan


def _copy_chart(prs, old):
    """A chart part copy with its own copy of the embedded workbook and of every other part the chart uses (style,
    colours, user shapes, theme override); external links kept. Independent review C1-06: only the workbook was copied
    before, so a chart with user shapes kept an r:id with no relationship. A related part that has relationships of
    its own (for example user shapes with pictures) is refused by name instead of half-copied."""
    from pptx.opc.package import Part, XmlPart
    package = prs.part.package
    plan = _chart_copy_plan(old)
    pn = package.next_partname(ChartPart.partname_template)
    chart = ChartPart.load(pn, old.content_type, package, old.blob)
    for ed in chart._element.findall(qn("c:externalData")):
        chart._element.remove(ed)                    # named the old chart's workbook; a copy is attached below
    ids = {}
    for rId, rel, copy_of in plan:
        if copy_of is None:
            ids[rId] = chart.relate_to(rel.target_ref, rel.reltype, is_external=True)
            continue
        template, part = copy_of
        cls = XmlPart if part.content_type.endswith("+xml") else Part
        new_part = cls.load(package.next_partname(template), part.content_type, package, part.blob)
        ids[rId] = chart.relate_to(new_part, rel.reltype)
    for e in chart._element.iter():                   # one pass, old id -> new id
        for a, v in list(e.attrib.items()):
            if a.startswith(_R) and v in ids:
                e.set(a, ids[v])
    xl = old.chart_workbook.xlsx_part
    if xl is not None:
        chart.chart_workbook.update_from_xlsx_blob(xl.blob)
    return chart


def duplicate_slide(prs, sid) -> int:
    """A copy right after the original: shapes, tables, pictures (image parts shared, they are never edited in place),
    hyperlinks, notes text, and charts with their own copy of the chart part and its embedded workbook."""
    src = slide(prs, sid)
    allowed = (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE, RT.IMAGE, RT.CHART, RT.HYPERLINK)
    for rel in src.part.rels.values():
        if rel.reltype not in allowed:
            _refuse("the slide has content URI cannot copy yet (" + rel.reltype.rsplit("/", 1)[-1] + ")")
    for rel in src.part.rels.values():               # every refusal before anything changes (all or nothing)
        if rel.reltype == RT.CHART:
            _chart_copy_plan(rel.target_part)
    new = prs.slides.add_slide(src.slide_layout)
    tree = new.shapes._spTree
    for el in list(tree)[2:]:
        tree.remove(el)
    remap = {}
    for rId, rel in src.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            remap[rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.CHART:
            remap[rId] = new.part.relate_to(_copy_chart(prs, rel.target_part), RT.CHART)
        else:
            remap[rId] = new.part.relate_to(rel.target_part, rel.reltype)
    def remapped(el):
        c = copy.deepcopy(el)
        for e in c.iter():
            for a, v in list(e.attrib.items()):
                if a.startswith(_R) and v in remap:
                    e.set(a, remap[v])
        return c

    for el in list(src.shapes._spTree)[2:]:
        tree.append(remapped(el))
    # Independent review C1-09: the slide's own background, colour map override, transition and animation timing.
    # Timing refers to shapes by id, and the copied shapes keep their ids.
    src_sld, new_sld = src._element, new._element
    bg = src_sld.cSld.find(qn("p:bg"))
    if bg is not None:
        old_bg = new_sld.cSld.find(qn("p:bg"))
        if old_bg is not None:
            new_sld.cSld.remove(old_bg)
        new_sld.cSld.insert(0, remapped(bg))
    for tag in ("p:clrMapOvr", "p:transition", "p:timing"):
        el = src_sld.find(qn(tag))
        if el is None:
            continue
        mine = new_sld.find(qn(tag))
        if mine is not None:
            new_sld.remove(mine)
        ext = new_sld.find(qn("p:extLst"))
        if ext is not None:
            ext.addprevious(remapped(el))
        else:
            new_sld.append(remapped(el))
    if src.has_notes_slide and src.notes_slide.notes_text_frame is not None:
        new.notes_slide.notes_text_frame.text = src.notes_slide.notes_text_frame.text
    move_slide(prs, new.slide_id, _index(prs, sid) + 1)
    return new.slide_id


# ---- text ----------------------------------------------------------------------------------------------------------
def _set_paragraphs(tf, paragraphs: list):
    """Replace the text, keeping each paragraph's first-run formatting (and the last one's for added paragraphs)."""
    if not isinstance(paragraphs, list) or not paragraphs or len(paragraphs) > 200:
        raise OfficeError("BAD_TEXT", "paragraphs")
    paras = list(tf.paragraphs)
    templates = []
    for p in paras:
        if any(el.tag in (qn("a:hlinkClick"), qn("a:hlinkMouseOver")) for el in p._p.iter()):
            _refuse("a text hyperlink cannot be preserved by replacement")
        for el in p._p:
            if el.tag not in (qn("a:pPr"), qn("a:r"), qn("a:endParaRPr")):
                _refuse("an inline text item cannot be preserved by replacement")
            if el.tag == qn("a:r") and any(child.tag not in (qn("a:rPr"), qn("a:t")) for child in el):
                _refuse("an inline text item cannot be preserved by replacement")
        r = p.runs[0]._r if p.runs else None
        templates.append((copy.deepcopy(p._p.pPr) if p._p.pPr is not None else None,
                          copy.deepcopy(r.rPr) if r is not None and r.rPr is not None else None))
    body = tf._txBody
    for p in list(body.findall(qn("a:p"))):
        body.remove(p)
    for i, text in enumerate(paragraphs):
        _text(text)
        ppr, rpr = templates[min(i, len(templates) - 1)] if templates else (None, None)
        p = body.add_p()
        if ppr is not None:
            p.insert(0, copy.deepcopy(ppr))
        if text:
            r = p.add_r()
            if rpr is not None:
                r.insert(0, copy.deepcopy(rpr))
            r.text = text


def set_text(prs, sid, shid, paragraphs: list):
    sh = shape(slide(prs, sid), shid)
    if not sh.has_text_frame:
        _refuse("that item has no text")
    _set_paragraphs(sh.text_frame, paragraphs)


def set_title(prs, sid, text: str):
    sl = slide(prs, sid)
    if sl.shapes.title is None:
        _refuse("this slide's layout has no title")
    _set_paragraphs(sl.shapes.title.text_frame, [_text(text)])


def set_notes(prs, sid, text: str):
    slide(prs, sid).notes_slide.notes_text_frame.text = _text(text)


def format_text(prs, sid, shid, fmt: dict, paragraph=None):
    sh = shape(slide(prs, sid), shid)
    if not sh.has_text_frame:
        _refuse("that item has no text")
    if not isinstance(fmt, dict) or set(fmt) - {"b", "i", "u", "size", "color", "align"}:
        raise OfficeError("BAD_FORMAT", "unknown formatting")
    paras = list(sh.text_frame.paragraphs)
    if paragraph is not None:
        if not isinstance(paragraph, int) or not 0 <= paragraph < len(paras):
            raise OfficeError("BAD_POSITION", "paragraph")
        paras = [paras[paragraph]]
    for p in paras:
        if "align" in fmt:
            if fmt["align"] not in _ALIGN:
                raise OfficeError("BAD_FORMAT", "align")
            p.alignment = _ALIGN[fmt["align"]]
        for r in p.runs:
            f = r.font
            for k, attr in (("b", "bold"), ("i", "italic"), ("u", "underline")):
                if k in fmt:
                    if not isinstance(fmt[k], bool):
                        raise OfficeError("BAD_FORMAT", k)
                    setattr(f, attr, fmt[k])
            if "size" in fmt:
                if not isinstance(fmt["size"], (int, float)) or not 6 <= fmt["size"] <= 200:
                    raise OfficeError("BAD_FORMAT", "size")
                f.size = Pt(fmt["size"])
            if "color" in fmt:
                c = fmt["color"]
                if not (isinstance(c, str) and len(c) == 7 and c[0] == "#"):
                    raise OfficeError("BAD_FORMAT", "color")
                f.color.rgb = RGBColor.from_string(c[1:].upper())


# ---- shapes --------------------------------------------------------------------------------------------------------
def insert_text_box(prs, sid, box: dict, paragraphs: list) -> int:
    sl = slide(prs, sid)
    x, y, w, h = _box(prs, box)
    tb = sl.shapes.add_textbox(x, y, w, h)
    tb.text_frame.word_wrap = True
    _set_paragraphs(tb.text_frame, paragraphs)
    return tb.shape_id


def insert_image(prs, sid, image_path: Path, box: dict) -> int:
    sl = slide(prs, sid)
    x, y, w, h = _box(prs, box)
    pic = sl.shapes.add_picture(str(image_path), x, y, width=w)
    if pic.height > h:                     # keep the aspect ratio inside the box
        ratio = h / pic.height
        pic.height, pic.width = int(h), int(pic.width * ratio)
    return pic.shape_id


def insert_table(prs, sid, cells: list, box: dict) -> int:
    if not isinstance(cells, list) or not cells or len(cells) > MAX_TABLE or \
            any(not isinstance(r, list) or not r or len(r) > MAX_TABLE for r in cells):
        raise OfficeError("BAD_TABLE", "rows of cell texts")
    rows, cols = len(cells), max(len(r) for r in cells)
    sl = slide(prs, sid)
    x, y, w, h = _box(prs, box)
    gf = sl.shapes.add_table(rows, cols, x, y, w, h)
    for i, r in enumerate(cells):
        for j in range(cols):
            gf.table.cell(i, j).text = _text(r[j] if j < len(r) else "")
    return gf.shape_id


def set_table_cell(prs, sid, shid, row: int, col: int, text: str):
    sh = shape(slide(prs, sid), shid)
    if not sh.has_table:
        _refuse("that item is not a table")
    t = sh.table
    if not (isinstance(row, int) and isinstance(col, int) and 0 <= row < len(t.rows) and 0 <= col < len(t.columns)):
        raise OfficeError("BAD_POSITION", "cell")
    tf = t.cell(row, col).text_frame
    _set_paragraphs(tf, _text(text).split("\n"))


def move_shape(prs, sid, shid, box: dict):
    sh = shape(slide(prs, sid), shid)
    sh.left, sh.top, sh.width, sh.height = _box(prs, box)


def delete_shape(prs, sid, shid):
    sl = slide(prs, sid)
    sh = shape(sl, shid)
    if sh.is_placeholder:                      # removing it would bring the layout's empty prompt back: empty it
        if sh.has_text_frame:
            _set_paragraphs(sh.text_frame, [""])
            return
        _refuse("that placeholder cannot be removed")
    sh._element.getparent().remove(sh._element)


# ---- reading ----------------------------------------------------------------------------------------------------
def describe(prs) -> dict:
    """What the edit panes show: slides in order with their editable items (no file paths, no binary data)."""
    W, H = prs.slide_width or SLIDE_W_FALLBACK, prs.slide_height or 6858000
    out = []
    for i, s in enumerate(prs.slides):
        items = []
        for sh in s.shapes:
            it = {"id": sh.shape_id, "name": sh.name, "kind": "text" if sh.has_text_frame else
                  "table" if sh.has_table else "chart" if sh.has_chart else
                  "picture" if sh.shape_type is not None and "PICTURE" in str(sh.shape_type) else "other",
                  "placeholder": bool(sh.is_placeholder),
                  "box": None if sh.width is None else {"x": round((sh.left or 0) / W, 4), "y": round((sh.top or 0) / H, 4),
                                                         "w": round(sh.width / W, 4), "h": round(sh.height / H, 4)}}
            if sh.has_text_frame:
                it["paragraphs"] = [p.text for p in sh.text_frame.paragraphs]
                runs = [r for p in sh.text_frame.paragraphs for r in p.runs if r.text.strip()]
                # every run bold / italic (explicitly); the buttons turn it off again
                it["bold"] = bool(runs) and all(r.font.bold is True for r in runs)
                it["italic"] = bool(runs) and all(r.font.italic is True for r in runs)
            if sh.has_table:
                it["cells"] = [[c.text for c in r.cells] for r in sh.table.rows][:MAX_TABLE]
            items.append(it)
        title = s.shapes.title.text if s.shapes.title is not None else ""
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and \
            s.notes_slide.notes_text_frame is not None else ""
        out.append({"id": s.slide_id, "index": i, "layout": s.slide_layout.name, "title": title,
                    "title_id": s.shapes.title.shape_id if s.shapes.title is not None else None, "items": items,
                    "notes": notes})
    return {"slides": out, "layouts": layouts(prs), "aspect": round(W / H, 4)}


OPS = {
    "add_slide": lambda prs, op: add_slide(prs, op["layout"], op.get("after")),
    "delete_slide": lambda prs, op: delete_slide(prs, op["slide"]),
    "duplicate_slide": lambda prs, op: duplicate_slide(prs, op["slide"]),
    "move_slide": lambda prs, op: move_slide(prs, op["slide"], op["to"]),
    "set_title": lambda prs, op: set_title(prs, op["slide"], op["text"]),
    "set_text": lambda prs, op: set_text(prs, op["slide"], op["shape"], op["paragraphs"]),
    "set_notes": lambda prs, op: set_notes(prs, op["slide"], op["text"]),
    "format_text": lambda prs, op: format_text(prs, op["slide"], op["shape"], op["format"], op.get("paragraph")),
    "insert_text_box": lambda prs, op: insert_text_box(prs, op["slide"], op["box"], op["paragraphs"]),
    "insert_table": lambda prs, op: insert_table(prs, op["slide"], op["cells"], op["box"]),
    "set_table_cell": lambda prs, op: set_table_cell(prs, op["slide"], op["shape"], op["row"], op["col"], op["text"]),
    "move_shape": lambda prs, op: move_shape(prs, op["slide"], op["shape"], op["box"]),
    "delete_shape": lambda prs, op: delete_shape(prs, op["slide"], op["shape"]),
    # insert_image needs a granted file path; the session resolves the grant and calls insert_image itself.
}
