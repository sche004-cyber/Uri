"""New Word documents (milestone C, OD-A5; plan 12 section 4): built-in layouts made by code, and the owner's own
.docx/.dotx as a template. Layouts are headings and placeholder paragraphs in the document's own styles; nothing is
pre-filled with names, numbers or dates."""
from __future__ import annotations

import io
import re
import zipfile

from docx import Document
from docx.shared import Cm

LAYOUTS = {
    "blank": ("Blank document", []),
    "letter": ("Letter", [
        ("Normal", "[Your name and address]"), ("Normal", "[Date]"), ("Normal", ""),
        ("Normal", "[Recipient name and address]"), ("Normal", ""), ("Normal", "Subject: [subject]"),
        ("Normal", ""), ("Normal", "Dear [name],"), ("Normal", "[Write your letter here.]"), ("Normal", ""),
        ("Normal", "Yours sincerely,"), ("Normal", ""), ("Normal", "[Your name]"),
    ]),
    "office_order": ("Office order", [
        ("Title", "Office Order"), ("Normal", "No. [number]                                        Date: [date]"),
        ("Heading 1", "Subject: [subject]"), ("Normal", "[Text of the order.]"), ("Normal", ""),
        ("Normal", "[Name]"), ("Normal", "[Designation]"), ("Heading 2", "Copy to"), ("List Number", "[Recipient]"),
    ]),
    "notice": ("Notice", [
        ("Title", "Notice"), ("Normal", "Date: [date]"), ("Heading 1", "[Subject]"),
        ("Normal", "[Text of the notice.]"), ("Normal", ""), ("Normal", "[Name and designation]"),
    ]),
    "memo": ("Memo", [
        ("Title", "Memorandum"), ("Normal", "To: [name]"), ("Normal", "From: [name]"), ("Normal", "Date: [date]"),
        ("Normal", "Subject: [subject]"), ("Heading 1", "[Heading]"), ("Normal", "[Text of the memo.]"),
    ]),
    "minutes": ("Minutes of a meeting", [
        ("Title", "Minutes of the meeting"), ("Normal", "Date and place: [date, place]"),
        ("Heading 1", "Present"), ("List Bullet", "[Name]"), ("Heading 1", "Agenda and decisions"),
        ("Heading 2", "1. [Item]"), ("Normal", "[Discussion and decision.]"), ("Heading 1", "Action points"),
        ("List Bullet", "[Action, by whom, by when]"),
    ]),
    "report": ("Report", [
        ("Title", "[Report title]"), ("Subtitle", "[Prepared by, date]"), ("Heading 1", "Summary"),
        ("Normal", "[Summary.]"), ("Heading 1", "Background"), ("Normal", "[Background.]"),
        ("Heading 1", "Findings"), ("Normal", "[Findings.]"), ("Heading 1", "Recommendations"),
        ("List Number", "[Recommendation]"),
    ]),
    # OD-1 (plan 16): Government of India style noting sheet and circular. Placeholders only.
    "notesheet": ("Note sheet", [
        ("Title", "Note Sheet"), ("Normal", "File No. [file number]"), ("Heading 1", "Subject: [subject]"),
        ("List Number", "[First paragraph of the note.]"), ("List Number", "[Next paragraph of the note.]"),
        ("Normal", ""), ("Normal", "[Signature]"), ("Normal", "[Name]"), ("Normal", "[Designation]"),
        ("Normal", "Date: [date]"),
    ]),
    "circular": ("Circular", [
        ("Title", "Circular"), ("Normal", "No. [number]                                        Date: [date]"),
        ("Heading 1", "Subject: [subject]"), ("Normal", "[Text of the circular.]"), ("Normal", ""),
        ("Normal", "[Name]"), ("Normal", "[Designation of the issuing authority]"), ("Heading 2", "Distribution"),
        ("List Bullet", "[Recipient]"),
    ]),
}

# OD-2: the paragraph style each drafted block role is written in (every style exists in python-docx's default
# template, so a drafted document never names a style it does not have).
ROLE_STYLES = {"title": "Title", "heading": "Heading 1", "paragraph": "Normal", "bullet": "List Bullet",
               "numbered": "List Number"}


def layout_list() -> list:
    return [{"id": k, "title": v[0]} for k, v in LAYOUTS.items()]


def from_layout(layout_id: str) -> bytes:
    from ..errors import OfficeError
    if layout_id not in LAYOUTS:
        raise OfficeError("UNKNOWN_LAYOUT", str(layout_id)[:40])
    return _build(LAYOUTS[layout_id][1])


def from_blocks(layout_id: str, blocks: list) -> bytes:
    """OD-2: a new document in a layout's page setup whose body is the drafted blocks ({role, text}) instead of the
    layout's placeholder paragraphs. Pure: returns the bytes; nothing is written to disk."""
    from ..errors import OfficeError
    if layout_id not in LAYOUTS:
        raise OfficeError("UNKNOWN_LAYOUT", str(layout_id)[:40])
    paras = []
    for b in blocks:
        style = ROLE_STYLES.get(b.get("role")) if isinstance(b, dict) else None
        if style is None or not isinstance(b.get("text"), str):
            raise OfficeError("BAD_BLOCK", "a drafted block has an unknown role or no text")
        paras.append((style, b["text"]))
    if not paras:
        raise OfficeError("BAD_BLOCK", "nothing was drafted")
    return _build(paras)


def from_paragraphs(paras: list) -> bytes:
    """OD-3: a new document from (style, text) pairs (PDF text to Word). Pure."""
    return _build(paras)


_NOT_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff￾￿]")


def _xml_safe(text: str) -> str:
    """Characters XML cannot hold (PDF text and model output can carry them) are dropped, never written."""
    return _NOT_XML.sub("", text or "")


def _build(paras: list) -> bytes:
    doc = Document()
    for sec in doc.sections:                      # A4, 2.54 cm margins
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        for side in ("top", "bottom", "left", "right"):
            setattr(sec, f"{side}_margin", Cm(2.54))
    body = doc.paragraphs
    for i, (style, text) in enumerate(paras):
        p = body[0] if i == 0 and body else doc.add_paragraph()
        p.style = doc.styles[style]
        text = _xml_safe(text)
        if text:
            p.add_run(text)
    b = io.BytesIO()
    doc.save(b)
    return b.getvalue()


def from_docx_template(data: bytes) -> bytes:
    """A new document from the owner's .docx/.dotx: styles, headers, footers and page setup kept; a .dotx's main part
    type is changed to a document's."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                blob = blob.replace(b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
                                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml")
            z.writestr(info, blob)
    return out.getvalue()
