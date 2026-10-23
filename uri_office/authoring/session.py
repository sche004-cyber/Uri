"""Authoring sessions (milestone A; Word added in milestone C): a working copy, an ordered change list, undo/redo, a live preview and a
hash-bound save card. Plan: docs/office/unified/11_AUTHORING_A_B_PLAN.md section 2.

The working copy starts as the original's exact bytes (re-read from the owner's open grant and checked against the
SHA-256 recorded when the file was opened), never the sanitized render copy, so a save never drops what the sanitizer
removed for previewing. Each change is applied to the working copy at once (the owner's own edit); nothing reaches the
original or any destination until the owner approves a save card that lists every change in plain words.
"""
from __future__ import annotations

import hashlib
import io
import threading
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path

from ..errors import OfficeError
from ..xlsx_patch import atomic_write
from . import docx_ops as D
from . import pptx_ops as P

MAX_VERSIONS = 40
MAX_HISTORY_BYTES = 512 * 1024 * 1024      # undo history held in memory per edit; the oldest steps go first
MAX_SESSIONS = 4
MAX_BYTES = 200 * 1024 * 1024


def _sha(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def blank_pptx() -> bytes:
    from pptx import Presentation
    prs = Presentation()
    from pptx.util import Emu
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)          # 16:9
    prs.slides.add_slide(prs.slide_layouts[0])
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def from_template(data: bytes) -> bytes:
    """A new presentation from the owner's .pptx/.potx: same masters, layouts and theme; a .potx's main part type is
    changed to a presentation's, and its slides are kept (a template's sample slides can be deleted like any slide)."""
    import zipfile
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in zin.infolist():
            blob = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                blob = blob.replace(b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                                    b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml")
            z.writestr(info, blob)
    return out.getvalue()


KIND_WORDS = {"pptx": "presentation", "docx": "document"}


def describe_docx_change(op: dict, view: dict | None = None) -> str:
    """Plain words for one Word change: the block by its position and the start of its text."""
    blocks = (view or {}).get("blocks", [])

    def where(i):
        b = blocks[i] if isinstance(i, int) and 0 <= i < len(blocks) else None
        if b is None:
            return "a paragraph"
        if b["kind"] == "table":
            return f"the table at block {i + 1}"
        t = (b.get("text") or "").strip()
        return f"paragraph {i + 1}" + (f" ('{t[:40]}{'...' if len(t) > 40 else ''}')" if t else " (empty)")

    def after(op):
        return "at the start" if op.get("after") is None else f"after {where(op.get('after'))}"

    k = op.get("op")
    b = where(op.get("block"))
    texts = {
        "insert_paragraph": f"Add a paragraph {after(op)}: '{str(op.get('text'))[:60]}'",
        "set_paragraph_text": f"Change the text of {b} to '{str(op.get('text'))[:60]}'",
        "delete_block": f"Delete {b}",
        "move_block": f"Move {b} to position {(op.get('to') or 0) + 1}",
        "set_style": f"{b[0].upper() + b[1:]}: style '{op.get('style')}'",
        "format_text": f"{b[0].upper() + b[1:]}: format text",
        "insert_page_break": f"Add a page break {after(op)}",
        "replace_text": f"Replace '{str(op.get('find'))[:40]}' with '{str(op.get('replace'))[:40]}'",
        "insert_table": f"Add a table {after(op)}",
        "set_table_cell": f"Change a cell in {b}",
        "add_table_row": f"Add a row to {b}",
        "delete_table_row": f"Delete a row from {b}",
        "insert_image": f"Add a picture {after(op)}",
        "set_header_text": f"Header: '{str(op.get('text'))[:60]}'",
        "set_footer_text": f"Footer: '{str(op.get('text'))[:60]}'",
        "insert_page_number": f"Add page numbers in the {op.get('where', 'footer')}",
        "set_page_setup": "Change the page setup" + (f" ({op.get('orientation')})" if op.get("orientation") else ""),
    }
    return texts.get(k, str(k))


def describe_change(op: dict, deck: dict | None = None) -> str:
    """Plain words for one change (the change list and the save card)."""
    if deck is not None and "blocks" in deck:
        return describe_docx_change(op, deck)
    names = {s["id"]: f"slide {s['index'] + 1}" for s in (deck or {}).get("slides", [])}
    sl = names.get(op.get("slide"), "a slide")
    k = op.get("op")
    texts = {
        "add_slide": f"Add a '{op.get('layout')}' slide",
        "delete_slide": f"Delete {sl}",
        "duplicate_slide": f"Duplicate {sl}",
        "move_slide": f"Move {sl} to position {(op.get('to') or 0) + 1}",
        "set_title": f"{sl.capitalize()}: title '{str(op.get('text'))[:60]}'",
        "set_text": f"{sl.capitalize()}: change text",
        "set_notes": f"{sl.capitalize()}: change notes",
        "format_text": f"{sl.capitalize()}: format text",
        "insert_text_box": f"{sl.capitalize()}: add a text box",
        "insert_image": f"{sl.capitalize()}: add a picture",
        "insert_table": f"{sl.capitalize()}: add a table",
        "set_table_cell": f"{sl.capitalize()}: change a table cell",
        "move_shape": f"{sl.capitalize()}: move or resize an item",
        "delete_shape": f"{sl.capitalize()}: remove an item",
    }
    return texts.get(k, str(k))


@dataclass
class AuthSession:
    aid: str
    kind: str
    name: str
    origin: Path | None                 # the opened file (None: a new file)
    origin_sha: str | None
    grant_id: str | None
    versions: list = field(default_factory=list)       # bytes, versions[pos] is current
    pos: int = 0
    revision: int = 0                              # monotonic address epoch, including undo and redo
    changes: list = field(default_factory=list)        # [{"op", "text", "sha256"}]: changes[i] made version i+1
    dropped: int = 0                                   # versions dropped from the front (undo depth cap)
    preview_doc: str | None = None
    preview_sha256: str | None = None                  # the working copy the preview shows (driver review R2-02)
    lock: threading.RLock = field(default_factory=threading.RLock)

    @property
    def data(self) -> bytes:
        return self.versions[self.pos]


class Authoring:
    def __init__(self, state: Path, documents, grants, approvals, audit=None):
        self.root = Path(state) / "authoring"
        self.root.mkdir(parents=True, exist_ok=True)
        self.documents, self.grants, self.approvals = documents, grants, approvals
        self.audit = audit or (lambda kind, payload: None)
        self.sessions: dict[str, AuthSession] = {}
        self._lock = threading.Lock()

    # ---- start -------------------------------------------------------------------------------------------------
    def _new(self, kind, name, data, origin=None, origin_sha=None, grant_id=None) -> AuthSession:
        if kind not in KIND_WORDS:
            raise OfficeError("AUTHORING_NOT_SUPPORTED", "URI edits presentations and Word documents")
        if len(data) > MAX_BYTES:
            raise OfficeError("FILE_TOO_LARGE", "")
        self._load(data, kind)                            # refuses what python-pptx / python-docx cannot read
        with self._lock:
            if len(self.sessions) >= MAX_SESSIONS:
                raise OfficeError("TOO_MANY_OPEN", "finish or close an edit first")
            s = AuthSession(uuid.uuid4().hex[:16], kind, name, origin, origin_sha, grant_id, [data])
            self._preview(s)
            self.sessions[s.aid] = s
        self.audit("office.authoring_started", {"authoring_id": s.aid, "kind": kind, "new": origin is None})
        return s

    def edit_document(self, document_id: str) -> dict:
        d = self.documents.get(document_id)
        if d.kind not in KIND_WORDS:
            raise OfficeError("AUTHORING_NOT_SUPPORTED", "URI edits presentations and Word documents")
        path = self.grants.get(d.grant_id, "open").path
        # The edit holds its own grant on the same file (revoked at close), so closing the viewer that opened it does
        # not take Save away from the edit.
        g = self.grants.grant_open(str(path), kind=d.kind)
        try:
            data = g.path.read_bytes()
            if _sha(data) != d.sha256:
                raise OfficeError("TARGET_CHANGED", "the file changed on disk since it was opened; open it again", 409)
            s = self._new(d.kind, d.name, data, origin=g.path, origin_sha=d.sha256, grant_id=g.grant_id)
        except BaseException:
            self.grants.revoke(g.grant_id)
            raise
        return self.view(s.aid)

    def new_presentation(self, template_path=None) -> dict:
        if template_path:
            g = self.grants.grant_open(template_path, kind="pptx_template")
            try:
                data = from_template(g.path.read_bytes())
            finally:
                self.grants.revoke(g.grant_id)
            name = Path(g.path).stem + " (new).pptx"
        else:
            data, name = blank_pptx(), "New presentation.pptx"
        return self.view(self._new("pptx", name, data).aid)

    def new_document(self, layout=None, template_path=None) -> dict:
        """Milestone C: a new Word document from a built-in layout (default blank) or the owner's .docx/.dotx."""
        from . import docx_templates as T
        if template_path:
            g = self.grants.grant_open(template_path, kind="docx_template")
            try:
                data = T.from_docx_template(g.path.read_bytes())
            finally:
                self.grants.revoke(g.grant_id)
            name = Path(g.path).stem + " (new).docx"
        else:
            lid = layout or "blank"
            data = T.from_layout(lid)
            name = ("New document" if lid == "blank" else T.LAYOUTS[lid][0]) + ".docx"
        return self.view(self._new("docx", name, data).aid)

    def new_document_from_bytes(self, name: str, data: bytes) -> dict:
        """OD-2 / OD-3: a new Word document whose bytes URI made (a Brain draft, a PDF's text). It is a new file like
        one from a layout: nothing is on disk until the owner approves a Save as card."""
        return self.view(self._new("docx", name, data).aid)

    # ---- editing -----------------------------------------------------------------------------------------------
    def get(self, aid) -> AuthSession:
        s = self.sessions.get(aid) if isinstance(aid, str) else None
        if s is None:
            raise OfficeError("UNKNOWN_EDIT", "", 404)
        return s

    @staticmethod
    def _load(data: bytes, kind: str = "pptx"):
        try:
            if kind == "docx":
                from docx import Document
                return Document(io.BytesIO(data))
            from pptx import Presentation
            return Presentation(io.BytesIO(data))
        except Exception as e:  # noqa: BLE001 - any parse failure is one refusal
            raise OfficeError("NOT_A_DOCUMENT" if kind == "docx" else "NOT_A_PRESENTATION", type(e).__name__) from None

    @staticmethod
    def _ops(kind):
        return D if kind == "docx" else P

    def apply(self, aid, op: dict, image_grant: str | None = None) -> dict:
        s = self.get(aid)
        with s.lock:
            M = self._ops(s.kind)
            if not isinstance(op, dict) or (op.get("op") not in M.OPS and op.get("op") != "insert_image"):
                raise OfficeError("UNKNOWN_OP", str(op.get("op") if isinstance(op, dict) else op)[:32])
            prs = self._load(s.data, s.kind)
            if s.kind == "docx":
                prs._uri_edit_revision = s.revision
            before = M.describe(prs)
            try:
                if op["op"] == "insert_image":
                    path = self.grants.get(image_grant, "open").path
                    if s.kind == "docx":
                        result = D.insert_image(prs, D.check_after(prs, op), path, op.get("width", 0.5))
                    else:
                        result = P.insert_image(prs, op["slide"], path, op["box"])
                else:
                    result = M.OPS[op["op"]](prs, op)
            except KeyError as e:
                raise OfficeError("BAD_OP", f"missing {e.args[0]}") from None
            except (TypeError, ValueError, IndexError, AttributeError) as e:     # a value of the wrong shape
                raise OfficeError("BAD_OP", type(e).__name__) from None
            out = io.BytesIO()
            prs.save(out)
            data = out.getvalue()
            del s.versions[s.pos + 1:]
            del s.changes[s.dropped + s.pos:]
            s.versions.append(data)
            s.changes.append({"op": op, "text": describe_change(op, before), "sha256": _sha(data)})
            s.pos += 1
            if s.kind == "docx":
                s.revision += 1
            # oldest undo steps go (count and memory caps; the current and one earlier step stay); the change list
            # is kept
            while len(s.versions) > 2 and (len(s.versions) > MAX_VERSIONS
                                           or sum(map(len, s.versions)) > MAX_HISTORY_BYTES):
                s.versions.pop(0)
                s.pos -= 1
                s.dropped += 1
            self._preview(s)
            return {**self.view(aid), "result": result}

    def undo(self, aid, redo=False) -> dict:
        s = self.get(aid)
        with s.lock:
            if redo:
                if s.pos >= len(s.versions) - 1:
                    raise OfficeError("NOTHING_TO_REDO")
                s.pos += 1
            else:
                if s.pos == 0:
                    raise OfficeError("NOTHING_TO_UNDO")
                s.pos -= 1
            if s.kind == "docx":
                s.revision += 1
            self._preview(s)
            return self.view(aid)

    @staticmethod
    def _live_changes(s) -> list:
        """The changes that made the current working copy (since it was opened or last saved)."""
        return s.changes[: s.dropped + s.pos]

    def view(self, aid) -> dict:
        s = self.get(aid)
        with s.lock:
            loaded = self._load(s.data, s.kind)
            if s.kind == "docx":
                loaded._uri_edit_revision = s.revision
            content = self._ops(s.kind).describe(loaded)
            live = self._live_changes(s)
            # path: the file this edit saves to (picked or approved by the owner), so the app can open it again
            return {"authoring_id": s.aid, "kind": s.kind, "name": s.name, "new": s.origin is None,
                    "path": str(s.origin) if s.origin is not None else None,
                    ("doc" if s.kind == "docx" else "deck"): content,
                    "changes": [{"text": c["text"]} for c in live],
                    "can_undo": s.pos > 0, "can_redo": s.pos < len(s.versions) - 1,
                    "preview_document_id": s.preview_doc, "sha256_prefix": _sha(s.data)[:12]}

    def _preview(self, s: AuthSession):
        """The preview of exactly the current working copy. Driver review R2-02: when a new preview cannot be made,
        the previous one must not keep authorizing Save for different bytes, so it is dropped first."""
        old = s.preview_doc
        s.preview_doc, s.preview_sha256 = None, None
        data = s.data
        try:
            preview = self.documents.open_bytes(s.name, s.kind, data, owner=s.aid)
        finally:
            if old:
                try:
                    self.documents.close(old)
                except OfficeError:
                    pass
        s.preview_doc, s.preview_sha256 = preview.document_id, _sha(data)

    # ---- save --------------------------------------------------------------------------------------------------
    def propose_save(self, aid, save_as_path=None) -> dict:
        s = self.get(aid)
        with s.lock:
            live = self._live_changes(s)
            if not live and s.origin is not None and save_as_path is None:
                raise OfficeError("NOTHING_TO_SAVE", "there are no changes")
            if save_as_path is None and s.origin is None:
                raise OfficeError("SAVE_AS_REQUIRED", f"a new {KIND_WORDS[s.kind]} is saved as a new file")
            self._check_rendered(s)
            if save_as_path is not None:
                g = self.grants.grant_save_as(save_as_path, kind=s.kind)
                dest, replaces, grant_id = g.path, False, g.grant_id
            else:
                dest, replaces, grant_id = s.origin, True, s.grant_id
            body = {"subject_kind": "authoring_save", "authoring_id": aid, "kind": s.kind,
                    "target": {"kind": "authoring_save", "destination": str(dest), "grant_id": grant_id,
                               "replaces": replaces, "origin_sha256": s.origin_sha},
                    "changes": [c["text"] for c in live], "working_sha256": _sha(s.data)}
            try:
                surface = self.approvals.propose_authoring_save(aid, body)
            except BaseException:
                if not replaces:
                    self.grants.revoke(grant_id)
                raise
            return {"messages": surface.messages(), "surface_id": surface.surface_id}

    def apply_save(self, body) -> dict:
        """The approved card: write exactly the working copy the card named."""
        s = self.get(body["authoring_id"])
        with s.lock:
            if _sha(s.data) != body["working_sha256"]:
                raise OfficeError("STALE_CHANGE", f"the {KIND_WORDS[s.kind]} changed after this card was shown; review "
                                                  "the save again", 409)
            t = body["target"]
            kind = "open" if t["replaces"] else "save_as"
            dest = self.grants.get(t["grant_id"], kind).path
            self._load(s.data, s.kind)                      # still readable (it was loaded at every change)
            if t["replaces"]:
                backup = self.root / "backups" / s.aid
                backup.mkdir(parents=True, exist_ok=True)
            elif dest.exists():
                raise OfficeError("DESTINATION_EXISTS", "choose a new file name", 409)
            data = s.data
            atomic_write(dest, lambda fh: fh.write(data), create_only=not t["replaces"],
                         expected_sha256=t["origin_sha256"] if t["replaces"] else None,
                         backup=backup / f"{time.time_ns()}-{dest.name}" if t["replaces"] else None)
            if t["replaces"]:
                s.origin_sha = _sha(data)
            else:
                # after Save as the edit continues on the new file: the next Save replaces it (with a backup)
                self.grants.revoke(t["grant_id"])
                g = self.grants.grant_open(str(dest), kind=s.kind)
                if s.grant_id is not None:
                    self.grants.revoke(s.grant_id)
                s.origin, s.origin_sha, s.grant_id, s.name = g.path, _sha(data), g.grant_id, g.path.name
            # the saved state is the new starting point: change list and undo history restart there
            s.versions = [data]
            s.pos = 0
            s.changes = []
            s.dropped = 0
            self.audit("office.authoring_saved", {"authoring_id": s.aid, "replaced": t["replaces"]})
            return {"saved_to": dest.name, "authoring_id": s.aid}

    def _check_rendered(self, s: AuthSession):
        """Plan 11 section 2: a working copy LibreOffice cannot open is never saved. The save card is offered only
        once the preview of exactly this working copy has rendered."""
        if s.preview_doc is None or s.preview_sha256 != _sha(s.data):
            raise OfficeError("PREVIEW_NOT_READY", "the preview is not ready yet", 409)
        try:
            d = self.documents.get(s.preview_doc)
        except OfficeError:
            raise OfficeError("PREVIEW_NOT_READY", "the preview is not ready yet", 409) from None
        state = getattr(d, "state", "ready")
        if state == "rendering":
            raise OfficeError("PREVIEW_NOT_READY", "wait until the preview shows your changes, then save", 409)
        if state != "ready":
            raise OfficeError("PREVIEW_FAILED", "LibreOffice could not show the changed file, so it is not saved; "
                                                "undo the last change or try again", 409)

    def close(self, aid) -> dict:
        with self._lock:
            s = self.sessions.pop(aid, None)
        if s is None:
            raise OfficeError("UNKNOWN_EDIT", "", 404)
        self.approvals.close_workbook(aid)
        if s.preview_doc:
            try:
                self.documents.close(s.preview_doc)
            except OfficeError:
                pass
        if s.grant_id is not None:
            self.grants.revoke(s.grant_id)
        return {"closed": True, "unsaved_changes": len(self._live_changes(s))}
