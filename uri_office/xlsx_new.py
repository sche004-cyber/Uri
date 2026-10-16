"""Write a new workbook file (M2-AC9, AC10; decisions M2-D4, M2-D5).

Values and formulas only (no styles), sheets in the draft's order, sheets without cells left out, and
`calcPr fullCalcOnLoad` so Excel calculates the formulas when the file is opened. The file is written to a temporary
name in the destination folder and then linked to the destination name, which fails if a file of that name exists:
an existing file is never replaced, also when one appears between the approval and the write.
"""
from __future__ import annotations

import os
import uuid
from pathlib import Path

from openpyxl import Workbook

from . import grants, refs
from .errors import OfficeError


def content(sheet_order: list, edits: list) -> list:
    """[(sheet name, {cell: {"v"|"f"}})] for the non-empty sheets, in the draft's sheet order."""
    by_sheet: dict = {}
    for e in edits:
        a = e["after"]
        if a.get("f") or a.get("v") not in (None, ""):
            by_sheet.setdefault(e["sheet"], {})[e["cell"]] = a
    return [(name, by_sheet[name]) for name in sheet_order if by_sheet.get(name)]


def write(dest: Path, sheets: list) -> dict:
    if not sheets:
        raise OfficeError("NOTHING_TO_SAVE", "the new workbook has no cells")
    wb = Workbook()
    wb.remove(wb.active)
    cells = formulas = 0
    for name, data in sheets:
        ws = wb.create_sheet(title=name)
        if ws.title != name:
            raise OfficeError("NEW_FILE_SHEET_NAME", f"the written sheet name differs from {name[:40]}")
        for ref, a in sorted(data.items(), key=lambda kv: (refs.parse_cell(kv[0])[1], refs.parse_cell(kv[0])[0])):
            c = ws[ref]
            if a.get("f"):
                c.value = a["f"]
                formulas += 1
            else:
                c.value = a.get("v")
                if isinstance(c.value, str):
                    c.data_type = "s"           # text is stored as text, never read back as a formula
            cells += 1
    wb.calculation.fullCalcOnLoad = True
    tmp = dest.parent / f".~uri-{uuid.uuid4().hex[:12]}.xlsx"
    try:
        wb.save(tmp)
        if tmp.stat().st_size > grants.MAX_BYTES:
            raise OfficeError("NEW_FILE_TOO_LARGE", "the new workbook exceeds the 50 MB file limit", 413)
        _link_no_replace(tmp, dest)
    except BaseException:
        # A failed cleanup must not replace the refusal or write error.
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    result = {"sheets": len(sheets), "cells": cells, "formulas": formulas}
    try:
        tmp.unlink(missing_ok=True)
    except OSError:
        result["temp_left"] = True
    return result


def _link_no_replace(tmp: Path, dest: Path):
    """Create `dest` with tmp's content; refuse if `dest` exists (never replaces it)."""
    try:
        os.link(tmp, dest)
        return
    except FileExistsError:
        raise OfficeError("DESTINATION_EXISTS", "choose a new file name", 409) from None
    except OSError:
        pass                                       # file system without hard links: exclusive create instead
    try:
        fd = os.open(dest, os.O_WRONLY | os.O_CREAT | os.O_EXCL | getattr(os, "O_BINARY", 0))
    except FileExistsError:
        raise OfficeError("DESTINATION_EXISTS", "choose a new file name", 409) from None
    try:
        with os.fdopen(fd, "wb") as out, open(tmp, "rb") as src:
            out.write(src.read())
    except Exception:
        dest.unlink(missing_ok=True)                 # O_EXCL created this file; do not leave a partial workbook
        raise
