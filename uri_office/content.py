"""Document content model and outline readers (plan §5 `content.py`).

`Block` and `SourceRef` are the shared shapes later phases extend (O3 Word edits, O4 page-cited answers). In O2 only
the outline readers are used. They read the **sanitized render copy** (never the original) with python-docx and
python-pptx; PDF outlines come from the text worker. Titles are shown to the owner in the UI and are never sent to the
Brain (plan §9.5).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

MAX_OUTLINE = 500
MAX_TITLE = 200


@dataclass(frozen=True)
class SourceRef:
    file: int                    # index of the source in the workflow's inputs
    page: int | None = None      # 1-based page or slide
    block: str | None = None     # block id within the source (O3+)


@dataclass(frozen=True)
class Block:
    id: str
    kind: str                    # "heading" | "paragraph" | "table" | "slide_title" | ...
    level: int
    text: str
    ref: SourceRef


def docx_outline(path: Path) -> list:
    """Headings as [{title, level, page: None}] (Word has no fixed pages before rendering)."""
    import docx
    out = []
    d = docx.Document(str(path))
    for p in d.paragraphs:
        name = (p.style.name if p.style is not None else "") or ""
        if name.lower().startswith("heading") or name.lower() == "title":
            lvl = 0 if name.lower() == "title" else int("".join(ch for ch in name if ch.isdigit()) or 1)
            t = " ".join(p.text.split())[:MAX_TITLE]
            if t:
                out.append({"title": t, "level": lvl, "page": None})
        if len(out) >= MAX_OUTLINE:
            break
    return out


def pptx_outline(path: Path) -> tuple:
    """(slide count, [{title, level: 1, page: slide number}])."""
    from pptx import Presentation
    p = Presentation(str(path))
    out = []
    for i, slide in enumerate(p.slides, 1):
        title = ""
        try:
            if slide.shapes.title is not None:
                title = " ".join(slide.shapes.title.text_frame.text.split())
        except (AttributeError, ValueError):
            title = ""
        out.append({"title": (title or f"Slide {i}")[:MAX_TITLE], "level": 1, "page": i})
    return len(p.slides), out[:MAX_OUTLINE]
