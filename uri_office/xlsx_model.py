"""Read an .xlsx into the workbook JSON the panel shows (AC1). Read only; uses openpyxl (already a URI dependency).

Workbook JSON
  {"sheets": [{"name", "index", "max_row", "max_col", "truncated",
               "cells": {"A1": {"v": value, "f": "=SUM(..)"?, "t": "date"?, "s": style_key?}},
               "merges": ["D1:E1"], "col_widths": {"A": 22.0}, "row_heights": {"1": 15.0},
               "locked": ["A2", ...]}],   # cells the patcher refuses (shared/array formulas, merged non-anchors)
   "styles": {style_key: {"b": bool?, "i": bool?, "u": bool?, "fc": "#RRGGBB"?, "bg": "#RRGGBB"?, "nf": str?,
                          "ha": "left|center|right"?, "wrap": bool?}},
   "kept_not_shown": {"charts": n, "images": n, "pivot_tables": n}}

For a formula cell, `v` is the value Excel cached when the file was last saved (None if never calculated) and `f` is the
formula. Charts, images and pivot tables stay in the file untouched and are only counted.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from . import refs, xlsx_patch
from .errors import OfficeError

MAX_CELLS_PER_SHEET = 50_000


def _color(c):
    rgb = getattr(c, "rgb", None) if c is not None else None
    if isinstance(rgb, str) and len(rgb) == 8 and rgb[:2] in ("FF", "00") and rgb != "00000000":
        return "#" + rgb[2:]
    return None


def _style(cell) -> dict:
    s = {}
    f = cell.font
    if f is not None:
        if f.b:
            s["b"] = True
        if f.i:
            s["i"] = True
        if f.u:
            s["u"] = True
        fc = _color(f.color)
        if fc and fc != "#000000":
            s["fc"] = fc
    fill = cell.fill
    if fill is not None and getattr(fill, "fill_type", None) == "solid":
        bg = _color(fill.fgColor)
        if bg:
            s["bg"] = bg
    if cell.number_format and cell.number_format != "General":
        s["nf"] = cell.number_format
    al = cell.alignment
    if al is not None:
        if al.horizontal in ("left", "center", "right"):
            s["ha"] = al.horizontal
        if al.wrap_text:
            s["wrap"] = True
    return s


def _value(v):
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat(), "date"
    if isinstance(v, (bool, int, float, str)) or v is None:
        return v, None
    return str(v), None


def kept_not_shown(path: Path) -> dict:
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    return {"charts": sum(n.startswith("xl/charts/chart") for n in names),
            "images": sum(n.startswith("xl/media/") for n in names),
            "pivot_tables": sum(n.startswith("xl/pivotTables/pivotTable") for n in names)}


def read(path: Path) -> dict:
    try:
        wf = load_workbook(path, data_only=False, read_only=False)
        wv = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001 - any parse failure is one refusal for the owner
        raise OfficeError("NOT_AN_XLSX", type(exc).__name__) from None
    try:
        locked = xlsx_patch.locked_cells(path)
    except Exception:  # noqa: BLE001
        raise OfficeError("NOT_AN_XLSX", "unreadable package") from None
    styles: dict[str, dict] = {}
    sheets = []
    for idx, ws in enumerate(wf.worksheets):
        wsv = wv[ws.title]
        cells = {}
        truncated = False
        for row in ws.iter_rows():
            for c in row:
                if c.value is None and not c.has_style:
                    continue
                if len(cells) >= MAX_CELLS_PER_SHEET:
                    truncated = True
                    break
                entry = {}
                raw = c.value
                if isinstance(raw, str) and raw.startswith("="):
                    entry["f"] = raw
                    entry["v"], t = _value(wsv[c.coordinate].value)
                elif raw is not None and type(raw).__name__ in ("ArrayFormula", "DataTableFormula"):
                    entry["f"] = "=" + str(getattr(raw, "text", "")).lstrip("=")
                    entry["v"], t = _value(wsv[c.coordinate].value)
                    entry["locked"] = True
                else:
                    entry["v"], t = _value(raw)
                if t:
                    entry["t"] = t
                st = _style(c)
                if st:
                    key = hashlib.sha256(json.dumps(st, sort_keys=True).encode()).hexdigest()[:12]
                    styles[key] = st
                    entry["s"] = key
                if entry.get("v") is None and "f" not in entry and "s" not in entry:
                    continue
                cells[c.coordinate] = entry
            if truncated:
                break
        widths = {k: round(float(d.width), 2) for k, d in ws.column_dimensions.items()
                  if d.width and refs._CELL.fullmatch(k + "1")}
        heights = {str(k): round(float(d.height), 2) for k, d in ws.row_dimensions.items() if d.height}
        sheets.append({"name": ws.title, "index": idx, "max_row": ws.max_row, "max_col": ws.max_column,
                       "truncated": truncated, "cells": cells, "locked": locked.get(ws.title, []),
                       "merges": [str(r) for r in ws.merged_cells.ranges],
                       "col_widths": widths, "row_heights": heights})
    wf.close()
    wv.close()
    return {"sheets": sheets, "styles": styles, "kept_not_shown": kept_not_shown(path)}


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()
