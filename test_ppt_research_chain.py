"""Research-first PPT creation without live network or model calls."""

import base64
import io
from types import SimpleNamespace
from unittest.mock import patch

from pptx import Presentation
from PIL import Image

from uri_core.core.canonical_execution import _execute_legacy_capability, _ready_read_prefix, run_canonical_for_ask
from uri_core.core.decision_engine import DecisionOutcome
from uri_core.core.tool_schema import LEGACY_TOOL_PARAMETERS
from uri_core.core.response_validation import validate_drafted_response
from uri_core.services.office_document_writer import find_openverse_images, render_pptx_bytes
from uri_core.tools.generate_document import GenerateDocumentTool


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL/nwAAAABJRU5ErkJggg=="
)
valid_png_buffer = io.BytesIO()
Image.new("RGB", (1, 1), "red").save(valid_png_buffer, format="PNG")
VALID_PNG = valid_png_buffer.getvalue()


class Directory:
    def describe(self, capability_id):
        return {"action_schemas": {}}

    def summaries(self):
        return [
            {"capability_id": "generate_document", "available": True,
             "summary": "Creates local PPTX presentations with themes and licensed photos.",
             "approval_required": False},
            {"capability_id": "unavailable_tool", "available": False,
             "summary": "Unavailable", "approval_required": False},
        ]


def _contract(first="web_search"):
    return {
        "goal": "Research student welfare and make a PPT", "mode": "multi_action",
        "capability": first,
        "actions": [
            {"name": first, "inputs": {"request_text": "student welfare"}},
            {"name": "generate_document", "capability": "generate_document", "inputs": {}},
        ],
    }


def test_web_research_prefix_runs_and_defers_document():
    contract = _contract()
    brain = SimpleNamespace(
        approval_gate=SimpleNamespace(execute_tool=lambda capability, **kwargs: {
            "status": "success", "data": {"status": "success", "results": [{"title": "Welfare guidance"}]}
        }),
        multi_action_dispatch=None,
        _draft_narrative_safely=lambda **kwargs: None,
    )
    gate = SimpleNamespace(outcome="READY", reasons=[], sub_results=[])
    with patch("uri_core.core.canonical_execution.build_turn_state_and_directory", return_value=(SimpleNamespace(data={}), Directory())), \
         patch("uri_core.core.canonical_execution.propose_decision", return_value=DecisionOutcome(status="ok", contract=contract)), \
         patch("uri_core.core.decision_gates.evaluate_gates", return_value=gate), \
         patch("uri_core.core.canonical_execution.is_allowlisted", return_value=True), \
         patch("uri_core.core.canonical_execution.record_canonical_telemetry"):
        result = run_canonical_for_ask(orchestrator=brain, session_id="s", user_text="research then PPT", principal=None)
    assert result["execution"]["partial_read_only"] is True
    assert result["response"]["results"][0]["title"] == "Welfare guidance"
    assert result["plan"]["deferred_actions"] == [{"capability": "generate_document", "name": "generate_document"}]
    assert result["execution"]["status"] == "success"
    assert result["response"]["available_capabilities"] == [{
        "id": "generate_document", "summary": "Creates local PPTX presentations with themes and licensed photos.",
        "requires_approval": False,
    }]
    options = (
        "Research found welfare guidance. Option 1: academic overview with a campus services map. "
        "Option 2: student journey with warm photos and counselling support. "
        "Option 3: outcomes and access routes with a modern visual theme. "
        "I can build the chosen PPTX with photos after you pick. Which option would you like?"
    )
    assert validate_drafted_response(options, result) == options


def test_non_read_prefix_never_auto_runs():
    gate = SimpleNamespace(outcome="READY")
    with patch("uri_core.core.decision_gates.evaluate_gates", return_value=gate), \
         patch("uri_core.core.canonical_execution.is_allowlisted", return_value=True):
        assert _ready_read_prefix(_contract("generate_document"), Directory(), None, {}) == []


def test_ppt_schema_exposes_user_chosen_design_controls():
    properties = LEGACY_TOOL_PARAMETERS["generate_document"]["properties"]
    assert {"output_format", "title", "theme", "include_images", "image_queries"} <= properties.keys()


def test_choice_turn_forwards_research_and_design_to_generator():
    captured = {}

    def execute(capability, **kwargs):
        captured.update(kwargs)
        return {"status": "success", "data": {"status": "success"}}

    result = _execute_legacy_capability(
        {"capability": "generate_document", "actions": [{
            "name": "generate_document", "inputs": {
                "request_text": "model replacement is not authoritative",
                "requested_output": "PPTX on student welfare",
                "output_format": "pptx", "title": "Student Welfare at NIT Sikkim",
                "theme": "academic", "include_images": True,
                "image_queries": ["campus support"],
            },
        }]},
        orchestrator=SimpleNamespace(approval_gate=SimpleNamespace(execute_tool=execute)),
        session_id="s", user_text="Option 2", principal=None,
        recent_conversation=[{"user": "Research", "uri": "Research summary"}],
    )
    assert result["status"] == "success"
    assert captured["request_text"] == "Option 2"
    assert captured["theme"] == "academic"
    assert captured["output_format"] == "pptx"
    assert captured["title"] == "Student Welfare at NIT Sikkim"
    assert captured["decision_context"]["recent_conversation"][0]["uri"] == "Research summary"


def test_theme_picture_and_credits():
    data = render_pptx_bytes(
        "# Student welfare\n- Overview\n# Campus support\n- Counselling and peer support",
        theme="academic", images=[{"data": PNG, "credit": "Test creator; CC0; example.test"}],
    )
    deck = Presentation(io.BytesIO(data))
    assert len(deck.slides) == 3
    assert deck.slides[0].background.fill.fore_color.rgb != deck.slides[1].background.fill.fore_color.rgb
    assert any(shape.shape_type == 13 for shape in deck.slides[1].shapes)
    assert "Image credits" in " ".join(shape.text for shape in deck.slides[2].shapes if shape.has_text_frame)


class Composer:
    def compose(self, **kwargs):
        assert kwargs["evidence"]["recent_conversation"][0]["uri"] == "Research summary"
        return {"body": "# Student welfare\n- Support\n# Programmes\n- Counselling", "status": "success", "composed_by": "model"}


def test_followup_uses_brain_format_title_and_derived_image_queries():
    saved = {}

    class Store:
        def save(self, **kwargs):
            saved.update(kwargs)
            return SimpleNamespace(to_reference=lambda: {"file_id": "saved-pptx"})

    class DeckComposer:
        def compose(self, **kwargs):
            assert "Student Welfare at NIT Sikkim" in kwargs["request_text"]
            assert kwargs["evidence"]["recent_conversation"][0]["user"].startswith("Create a PPT")
            return {"body": "# Student Welfare at NIT Sikkim\n- Overview\n# Campus support\n- Counselling",
                    "status": "success", "composed_by": "model"}

    tool = GenerateDocumentTool(composer=DeckComposer(), file_store=Store())
    queries = []

    def images(found_queries, *, diagnostics=None):
        queries.extend(found_queries)
        return [{"data": PNG, "credit": "Test creator; CC0; example.test"}]

    with patch("uri_core.services.office_document_writer.find_openverse_images", side_effect=images):
        result = tool.generate(
            request_text="Option 2, please build it with photos",
            output_format="pptx", title="Student Welfare at NIT Sikkim",
            theme="warm", include_images=True, image_queries=[],
            decision_context={"recent_conversation": [{"user": "Create a PPT on Student Welfare at NIT Sikkim", "uri": "Three options"}]},
            session_id="s",
        )
    assert result["output_format"] == "pptx"
    assert result["images_added"] == 1
    assert saved["filename"] == "Student_Welfare_at_NIT_Sikkim.pptx"
    assert queries == ["Student Welfare at NIT Sikkim Campus support"]
    deck = Presentation(io.BytesIO(saved["content"]))
    assert any(shape.shape_type == 13 for shape in deck.slides[1].shapes)


def test_no_format_fallback_and_explicit_docx_are_preserved():
    class Store:
        def save(self, **kwargs):
            return SimpleNamespace(to_reference=lambda: {"file_id": kwargs["filename"]})

    class PlainComposer:
        def compose(self, **kwargs):
            return {"body": "# Summary\n- Detail", "status": "success", "composed_by": "model"}

    tool = GenerateDocumentTool(composer=PlainComposer(), file_store=Store())
    fallback = tool.generate(request_text="Create a PPT about campus support", session_id="s")
    assert fallback["output_format"] == "pptx"
    assert fallback["file"]["file_id"].endswith(".pptx")
    contextual = tool.generate(
        request_text="Option 2, build it", session_id="s",
        decision_context={"recent_conversation": [{"user": "Create a PPT about campus support", "uri": "Choose an option"}]},
    )
    assert contextual["output_format"] == "pptx"
    document = tool.generate(request_text="Option 2", output_format="docx", title="Student Welfare", session_id="s")
    assert document["output_format"] == "docx"
    assert document["file"]["file_id"] == "Student_Welfare.docx"
    explicit_docx = tool.generate(
        request_text="Create a DOCX report", session_id="s",
        decision_context={"recent_conversation": [{"user": "Create a PPT", "uri": "Choose an option"}]},
    )
    assert explicit_docx["output_format"] == "docx"


def test_image_failure_still_saves_pptx_with_research_context():
    class Store:
        def save(self, **kwargs):
            assert Presentation(io.BytesIO(kwargs["content"]))
            return SimpleNamespace(to_reference=lambda: {"file_id": "saved-file"})

    tool = GenerateDocumentTool(composer=Composer(), file_store=Store())
    with patch("uri_core.services.office_document_writer.requests.get", side_effect=OSError("offline")):
        result = tool.generate(
            request_text="Option 2", requested_output="PPTX student welfare",
            theme="modern", include_images=True, image_queries=["campus"],
            decision_context={"recent_conversation": [{"user": "Research", "uri": "Research summary"}]},
            session_id="s",
        )
    assert result["status"] == "success"
    assert result["images_added"] == 0
    assert "search was unavailable" in result["image_note"]
    assert result["file"]["file_id"]


class ImageResponse:
    def __init__(self, *, results=None, data=VALID_PNG, url="https://api.openverse.org/image/thumb/",
                 history=(), content_type="image/png"):
        self.results = results
        self.data = data
        self.url = url
        self.history = list(history)
        self.headers = {"Content-Type": content_type}

    def json(self):
        return {"results": self.results}

    def raise_for_status(self):
        pass

    def iter_content(self, size):
        yield self.data

    def __enter__(self):
        return self

    def __exit__(self, *args):
        pass


def _item(number):
    return {
        "license": "by", "title": f"Photo {number}", "creator": f"Creator {number}",
        "thumbnail": f"https://api.openverse.org/v1/images/{number}/thumb/",
        "url": f"https://photos.example/{number}.png",
        "foreign_landing_url": f"https://photos.example/{number}",
    }


def test_thumbnail_preferred_and_multiple_images_have_credits():
    calls = []

    def get(url, **kwargs):
        calls.append(url)
        if url.endswith("/images/"):
            assert kwargs["params"]["page_size"] == 5
            return ImageResponse(results=[_item(len(calls)), _item(len(calls) + 1)])
        return ImageResponse(url=url)

    with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
        images = find_openverse_images(["campus", "students"])
    assert len(images) == 2
    assert len([url for url in calls if "/thumb/" in url]) == 2
    assert not any("photos.example" in url for url in calls)
    deck = Presentation(io.BytesIO(render_pptx_bytes(
        "# Title\n- Intro\n# Campus\n- One\n# Students\n- Two", images=images,
    )))
    assert sum(any(shape.shape_type == 13 for shape in slide.shapes) for slide in deck.slides) == 2
    credits = " ".join(shape.text for shape in deck.slides[-1].shapes if shape.has_text_frame)
    assert "Creator 1" in credits and "Creator 3" in credits


def test_http_redirect_rejected_and_later_result_tried():
    calls = []

    def get(url, **kwargs):
        calls.append(url)
        if url.endswith("/images/"):
            return ImageResponse(results=[_item(1), _item(2)])
        if "/1/" in url or url.endswith("/1.png"):
            return ImageResponse(url="https://photos.example/final.png",
                                 history=[SimpleNamespace(url="http://photos.example/redirect")])
        return ImageResponse(url=url)

    with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
        images = find_openverse_images(["campus"])
    assert len(images) == 1
    assert "Creator 2" in images[0]["credit"]
    assert len(calls) == 4


def test_https_redirects_accepted_up_to_three():
    for hops, expected in ((3, 1), (4, 0)):
        def get(url, **kwargs):
            if url.endswith("/images/"):
                return ImageResponse(results=[{**_item(1), "url": None}])
            return ImageResponse(url="https://photos.example/final.png", history=[
                SimpleNamespace(url=f"https://photos.example/hop-{index}") for index in range(hops)
            ])

        with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
            assert len(find_openverse_images(["campus"])) == expected


def test_original_url_fallback_accepts_https_redirect():
    calls = []

    def get(url, **kwargs):
        calls.append(url)
        if url.endswith("/images/"):
            return ImageResponse(results=[_item(1)])
        if "/thumb/" in url:
            return ImageResponse(url=url, content_type="text/html")
        return ImageResponse(url="https://photos.example/final.png", history=[SimpleNamespace(url=url)])

    with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
        images = find_openverse_images(["campus"])
    assert len(images) == 1
    assert calls[-1] == "https://photos.example/1.png"


def test_oversize_rejected_and_download_failure_reason():
    def get(url, **kwargs):
        if url.endswith("/images/"):
            return ImageResponse(results=[_item(1)])
        return ImageResponse(url=url, data=b"x" * 5_000_001)

    diagnostics = {}
    with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
        assert find_openverse_images(["campus"], diagnostics=diagnostics) == []
    assert diagnostics["reason"] == "downloads_failed"


def test_no_results_reason():
    diagnostics = {}
    with patch("uri_core.services.office_document_writer.requests.get", return_value=ImageResponse(results=[])):
        assert find_openverse_images(["campus"], diagnostics=diagnostics) == []
    assert diagnostics["reason"] == "no_results"


def test_tool_image_note_distinguishes_no_results_and_failed_downloads():
    class Store:
        def save(self, **kwargs):
            return SimpleNamespace(to_reference=lambda: {"file_id": "saved-file"})

    class PlainComposer:
        def compose(self, **kwargs):
            return {"body": "# Topic\n# Details", "status": "success", "composed_by": "model"}

    tool = GenerateDocumentTool(composer=PlainComposer(), file_store=Store())
    for results, note in (([], "found no usable licensed images"),
                          ([_item(1)], "downloads failed")):
        def get(url, **kwargs):
            if url.endswith("/images/"):
                return ImageResponse(results=results)
            return ImageResponse(url=url, content_type="text/html")

        with patch("uri_core.services.office_document_writer.requests.get", side_effect=get):
            outcome = tool.generate(request_text="Make a PPT", include_images=True,
                                    image_queries=["campus"], session_id="s")
        assert outcome["images_added"] == 0
        assert note in outcome["image_note"]
