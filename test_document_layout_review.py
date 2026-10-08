"""Actual rendered geometry and bounded Brain corrections; no network/model."""
import io
import json
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from PIL import Image
from pptx import Presentation

from uri_core.core.file_store import FileStore
from uri_core.services.office_document_writer import render_pptx_bytes, inspect_pptx_layout
from uri_core.services.document_composer import DocumentComposer
from uri_core.tools.generate_document import GenerateDocumentTool

BODY = '# Student welfare\n- Suggested wellbeing activities\n# Campus activities\n- Suggested student clubs'


def test_fenced_json_review_is_formatting_only_and_extra_authority_is_rejected():
    provider = SimpleNamespace(complete=lambda **_: SimpleNamespace(content=
        '```json\n' + json.dumps({'verdict':'accept','issues':[], 'revised_body':None}) + '\n```'))
    composer = DocumentComposer(provider=provider)
    assert composer.review_document(request_text='Fixture', body=BODY, render_evidence={})['status'] == 'accepted'
    provider.complete = lambda **_: SimpleNamespace(content='```json\n' + json.dumps(
        {'verdict':'accept','issues':[], 'revised_body':None, 'approve':True}) + '\n```')
    assert composer.review_document(request_text='Fixture', body=BODY, render_evidence={})['status'] == 'unavailable'


def portrait():
    image = io.BytesIO()
    Image.new('RGB', (80, 500), 'green').save(image, format='PNG')
    return {'data': image.getvalue(), 'credit': 'Fixture, CC0, https://example.org/photo'}


def test_long_wrapped_bullets_and_portrait_stay_in_separate_regions():
    bullets = ['Establish confidential support services and provide practical guidance for student wellbeing. ' * 2] * 5
    data = render_pptx_bytes('# Student welfare\n- Introduction\n# Suggestions\n' + '\n'.join('- ' + b for b in bullets), images=[portrait()])
    report = inspect_pptx_layout(data)
    assert report['status'] == 'estimated_pass', report
    deck = Presentation(io.BytesIO(data))
    pictures = [s for slide in deck.slides for s in slide.shapes if s.shape_type == 13]
    assert pictures and all(s.top + s.height < int(6.6 * 914400) for s in pictures)
    assert report['visual_review'] == 'not_performed'


def test_all_nine_bullets_are_preserved_and_excess_density_is_detected():
    data = render_pptx_bytes('# Topics\n' + '\n'.join(f'- Distinct item {n}' for n in range(9)))
    report = inspect_pptx_layout(data)
    assert report['status'] == 'estimated_pass'
    assert all(f'Distinct item {n}' in report['extracted_text'] for n in range(9))
    dense = render_pptx_bytes('# Dense\n' + '\n'.join('- ' + ('Long content ' * 30) for _ in range(14)))
    assert inspect_pptx_layout(dense)['status'] == 'needs_repair'


class ReviewingComposer:
    def __init__(self, store, reviews):
        self.store, self.reviews, self.calls = store, iter(reviews), []
    def compose(self, **kwargs):
        return {'body': BODY, 'status': 'success', 'composed_by': 'mock'}
    def review_document(self, **kwargs):
        assert self.store.list_for_session('mine') == []
        self.calls.append(kwargs)
        return next(self.reviews)


def test_one_correction_is_rendered_then_rechecked_before_any_storage(tmp_path):
    store = FileStore(str(tmp_path))
    corrected = BODY.replace('student clubs', 'student-led campus clubs')
    composer = ReviewingComposer(store, [
        {'status': 'revised', 'body': corrected, 'issues': ['Clarify suggested activities']},
        {'status': 'accepted', 'issues': []}])
    result = GenerateDocumentTool(composer, store).generate(request_text='Render suggestions',
        output_format='pptx', session_id='mine', variants=[{'label':'Formal','theme':'academic'}, {'label':'Warm','theme':'warm'}])
    assert len(composer.calls) == 2
    assert 'student-led campus clubs' in composer.calls[1]['render_evidence']['extracted_text']
    assert result['quality_review']['correction_passes'] == 1
    assert result['quality_review']['brain_review'] == 'accepted'
    assert result['created_count'] == len(store.list_for_session('mine')) == 2


@pytest.mark.parametrize('reviews', [
    [{'status':'revised','body':'short'}],
    [{'status':'revised','body':BODY}, {'status':'revised','body':BODY}],
])
def test_invalid_or_unresolved_review_never_publishes_files(tmp_path, reviews):
    store = FileStore(str(tmp_path)); composer = ReviewingComposer(store, reviews)
    result = GenerateDocumentTool(composer, store).generate(request_text='Render', output_format='pptx', session_id='mine')
    assert result['status'] == 'error' and result['created_count'] == 0
    assert store.list_for_session('mine') == []


def test_unavailable_brain_review_is_explicitly_provisional(tmp_path):
    store = FileStore(str(tmp_path)); composer = ReviewingComposer(store, [{'status':'unavailable'}])
    result = GenerateDocumentTool(composer, store).generate(request_text='Render', output_format='pptx', session_id='mine')
    assert result['file'] and result['quality_review']['status'] == 'provisional'
    assert result['quality_review']['visual_review'] == 'not_performed'


def test_review_proposals_cannot_supply_authority_or_skip_json_validation():
    class Provider:
        def complete(self, **kwargs):
            return SimpleNamespace(content=json.dumps({'verdict':'accept', 'issues':[], 'revised_body':None, 'principal':'admin'}))
    result = DocumentComposer(provider=Provider()).review_document(request_text='Review', body=BODY, render_evidence={})
    assert result['status'] == 'unavailable'


def test_unrepairable_density_blocks_save_even_if_brain_accepts(tmp_path):
    store = FileStore(str(tmp_path))
    class Dense(ReviewingComposer):
        def compose(self, **kwargs):
            return {'body':'# Dense\n' + '\n'.join('- ' + ('Many words ' * 30) for _ in range(14)),
                    'status':'success','composed_by':'mock'}
    composer = Dense(store, [{'status':'accepted','issues':[]}])
    result = GenerateDocumentTool(composer, store).generate(request_text='Render', output_format='pptx', session_id='mine')
    assert result['status'] == 'error' and not store.list_for_session('mine')
