"""Amendment probe (research evidence, not product code), 2026-10-15.

Tests no-install alternatives to PyMuPDF for URI Office, and read-only Windows printer discovery:
  1. LibreOffice listener (hardened profile, Job Object): DOCX (20 pages), PPTX (11 slides), XLSX -> PDF, and
     per-page PNG export through UNO (replaces PyMuPDF rasterizing for Office documents).
  2. pypdf (BSD-3, already installed): page count, outline, per-page text, fonts, image extraction from a scanned
     (image-only) PDF, then Tesseract OCR on the extracted image (replaces PyMuPDF for text and scans).
  3. Printers via ctypes winspool (read-only): EnumPrintersW level 2, GetDefaultPrinterW, DeviceCapabilitiesW.
     Nothing is printed and no setting is changed.
No PyMuPDF import anywhere in this probe. Usage: python -I pdf_print_probe.py <out_dir>
"""
import ctypes, ctypes.wintypes as W, hashlib, json, os, shutil, subprocess, sys, tempfile, threading, time
from pathlib import Path
import psutil, docx, pptx, openpyxl, pypdf, win32api, win32con, win32job
from PIL import Image, ImageDraw, ImageFont
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

PROG = Path(os.environ["LOCALAPPDATA"]) / "URI" / "engines" / "libreoffice-26.8.1" / "program"
SOFFICE, LOPY = PROG / "soffice.com", PROG / "python.exe"
TESS = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
HERE = Path(__file__).resolve().parent
out = Path(sys.argv[1]).resolve(); shutil.rmtree(out, ignore_errors=True); out.mkdir(parents=True)
fx = out / "fixtures"; fx.mkdir(); lo_out = out / "lo"; lo_out.mkdir()
R = {"started": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "pymupdf_imported": False}

# ---------- fixtures (synthetic) ----------
report = fx / "report20.docx"
d = docx.Document()
for sec in range(1, 11):
    d.add_heading(f"{sec}. Section {sec} of the annual administrative report", 1)
    for k in range(6):
        d.add_paragraph(f"Paragraph {k + 1} of section {sec}. The department reviewed the files, the staff position and "
                        "the pending matters for the quarter, and recorded the decisions taken at each meeting. " * 2)
    t = d.add_table(rows=1, cols=3); t.style = "Table Grid"
    for c, v in zip(t.rows[0].cells, ("Item", "Target", "Achieved")): c.text = v
    for i in range(8):
        for c, v in zip(t.add_row().cells, (f"Item {i + 1}", str(10 + i), str(9 + i))): c.text = v
d.save(report)

deck = fx / "deck.pptx"
pr = pptx.Presentation()
for i in range(10):
    s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = f"Section {i + 1}"
    s.placeholders[1].text = "Point one\nPoint two\nPoint three"
s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Revenue by quarter"
cd = CategoryChartData(); cd.categories = ["Q1", "Q2", "Q3", "Q4"]; cd.add_series("Revenue", (12.4, 13.1, 13.9, 14.2))
s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
pr.save(deck)

book = fx / "budget.xlsx"
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Budget"
ws.append(["Head", "Sanctioned", "Spent", "Balance"])
for i in range(1, 41):
    ws.append([f"Head {i}", 1000 * i, 900 * i, f"=B{i + 1}-C{i + 1}"])
wb.save(book)

# scanned (image-only) PDF, 2 pages, made with Pillow
scan = fx / "scan.pdf"
pages = []
for n in range(2):
    im = Image.new("RGB", (1654, 2339), "white"); dr = ImageDraw.Draw(im)
    try:
        font = ImageFont.truetype("arial.ttf", 42)
    except OSError:
        font = ImageFont.load_default()
    for k, line in enumerate(["NOTICE No. 55/2026", "Corrigendum to tender No. 55/2026",
                              f"Page {n + 1}: the last date for bids is extended to 30 October 2026."]):
        dr.text((150, 200 + 90 * k), line, fill="black", font=font)
    pages.append(im)
pages[0].save(scan, save_all=True, append_images=pages[1:], resolution=200)
R["fixtures"] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (report, deck, book, scan)}

# ---------- 1. LibreOffice listener: PDF + per-page PNG ----------
XCU = (HERE.parent / "lo_qualify_20261015" / "lo_qualify.py").read_text(encoding="utf-8")
XCU = XCU[XCU.index('HARDENED_XCU = """') + len('HARDENED_XCU = """'):]
XCU = XCU[:XCU.index('"""')]
prof = Path(tempfile.mkdtemp(prefix="uri_lo_prof_")); (prof / "user").mkdir()
(prof / "user" / "registrymodifications.xcu").write_text(XCU, encoding="utf-8")
job = win32job.CreateJobObject(None, "")
info = win32job.QueryInformationJobObject(job, win32job.JobObjectExtendedLimitInformation)
info["BasicLimitInformation"]["LimitFlags"] = win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
win32job.SetInformationJobObject(job, win32job.JobObjectExtendedLimitInformation, info)
pipe = f"uri_lo_{os.getpid()}_pp"
lo = subprocess.Popen([str(SOFFICE), f"-env:UserInstallation={prof.as_uri()}", "--headless", "--invisible", "--norestore",
                       "--nolockcheck", "--nodefault", "--nologo", f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext"],
                      stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
h = win32api.OpenProcess(win32con.PROCESS_ALL_ACCESS, False, lo.pid); win32job.AssignProcessToJobObject(job, h)
t0 = time.perf_counter()
cl = subprocess.run([str(LOPY), str(HERE / "uno_page_png.py"), pipe, str(lo_out), "900", str(report), str(deck), str(book)],
                    capture_output=True, text=True, timeout=900)
R["lo_listener"] = {"wall_s_incl_startup": round(time.perf_counter() - t0, 2), "client_rc": cl.returncode,
                    "rows": [json.loads(l) for l in cl.stdout.splitlines() if l.startswith("{")],
                    "stderr_tail": cl.stderr[-300:]}
cl2 = subprocess.run([str(LOPY), str(HERE / "uno_page_png.py"), pipe, str(lo_out), "900", str(report)],
                     capture_output=True, text=True, timeout=900)
R["lo_listener_warm_repeat"] = [json.loads(l) for l in cl2.stdout.splitlines() if l.startswith("{")]
win32api.CloseHandle(h); job.Close(); time.sleep(2); shutil.rmtree(prof, ignore_errors=True)
R["lo_pngs"] = sorted(p.name for p in lo_out.glob("*.png"))

# ---------- 2. pypdf ----------
def pypdf_profile(path):
    t0 = time.perf_counter(); r = pypdf.PdfReader(str(path)); n = len(r.pages)
    t1 = time.perf_counter(); texts = [pg.extract_text() or "" for pg in r.pages]; t2 = time.perf_counter()
    fonts = set()
    for pg in r.pages:
        res = pg.get("/Resources") or {}
        for _, f in ((res.get("/Font") or {}).items() if res.get("/Font") else []):
            fonts.add(str(f.get_object().get("/BaseFont")))
    try:
        outline = len(r.outline)
    except Exception:
        outline = None
    hit = [i + 1 for i, t in enumerate(texts) if "Section 7" in t]
    return {"pages": n, "open_ms": round((t1 - t0) * 1000, 1), "extract_all_ms": round((t2 - t1) * 1000, 1),
            "chars": sum(len(t) for t in texts), "fonts": sorted(fonts)[:12], "outline_items": outline,
            "search_Section7_pages": hit[:5]}

R["pypdf"] = {}
for name in ("report20.pdf", "deck.pdf", "budget.pdf"):
    if (lo_out / name).exists():
        R["pypdf"][name] = pypdf_profile(lo_out / name)
r = pypdf.PdfReader(str(scan)); ocr = []
t0 = time.perf_counter()
for i, pg in enumerate(r.pages):
    txt = pg.extract_text() or ""
    imgs = list(pg.images)
    row = {"page": i + 1, "text_layer_chars": len(txt), "images": len(imgs)}
    if imgs:
        ip = out / f"scan_p{i + 1}.{imgs[0].name.rsplit('.', 1)[-1]}"; ip.write_bytes(imgs[0].data)
        t1 = time.perf_counter()
        o = subprocess.run([str(TESS), str(ip), "stdout"], capture_output=True, text=True, timeout=120)
        row["ocr_ms"] = round((time.perf_counter() - t1) * 1000, 1); row["ocr_text_head"] = o.stdout.strip()[:120]
    ocr.append(row)
R["pypdf_scan_ocr"] = {"rows": ocr, "total_ms": round((time.perf_counter() - t0) * 1000, 1)}

# ---------- 3. printers (read-only) ----------
winspool = ctypes.WinDLL("winspool.drv", use_last_error=True)


class PRINTER_INFO_2(ctypes.Structure):
    _fields_ = [("pServerName", W.LPWSTR), ("pPrinterName", W.LPWSTR), ("pShareName", W.LPWSTR),
                ("pPortName", W.LPWSTR), ("pDriverName", W.LPWSTR), ("pComment", W.LPWSTR), ("pLocation", W.LPWSTR),
                ("pDevMode", ctypes.c_void_p), ("pSepFile", W.LPWSTR), ("pPrintProcessor", W.LPWSTR),
                ("pDatatype", W.LPWSTR), ("pParameters", W.LPWSTR), ("pSecurityDescriptor", ctypes.c_void_p),
                ("Attributes", W.DWORD), ("Priority", W.DWORD), ("DefaultPriority", W.DWORD), ("StartTime", W.DWORD),
                ("UntilTime", W.DWORD), ("Status", W.DWORD), ("cJobs", W.DWORD), ("AveragePPM", W.DWORD)]


PRINTER_ENUM_LOCAL, PRINTER_ENUM_CONNECTIONS = 0x2, 0x4
need, count = W.DWORD(), W.DWORD()
winspool.EnumPrintersW(PRINTER_ENUM_LOCAL | PRINTER_ENUM_CONNECTIONS, None, 2, None, 0, ctypes.byref(need), ctypes.byref(count))
buf = ctypes.create_string_buffer(need.value)
ok = winspool.EnumPrintersW(PRINTER_ENUM_LOCAL | PRINTER_ENUM_CONNECTIONS, None, 2, buf, need, ctypes.byref(need), ctypes.byref(count))
infos = ctypes.cast(buf, ctypes.POINTER(PRINTER_INFO_2))
size = W.DWORD(0); winspool.GetDefaultPrinterW(None, ctypes.byref(size))
dbuf = ctypes.create_unicode_buffer(size.value); winspool.GetDefaultPrinterW(dbuf, ctypes.byref(size))
DC = {"DC_DUPLEX": 7, "DC_COPIES": 18, "DC_COLORDEVICE": 32, "DC_PAPERS": 2, "DC_ORIENTATION": 17, "DC_COLLATE": 22}
dcap = winspool.DeviceCapabilitiesW; dcap.restype = ctypes.c_int
printers = []
t0 = time.perf_counter()
for i in range(count.value if ok else 0):
    p = infos[i]
    caps = {k: dcap(p.pPrinterName, p.pPortName, v, None, None) for k, v in DC.items()}
    printers.append({"name": p.pPrinterName, "port": p.pPortName, "driver": p.pDriverName,
                     "attributes_hex": hex(p.Attributes), "status_hex": hex(p.Status), "jobs": p.cJobs,
                     "work_offline": bool(p.Attributes & 0x400), "network": bool(p.Attributes & 0x10),
                     "caps": caps})
R["printers"] = {"enum_ok": bool(ok), "count": count.value, "default": dbuf.value, "list": printers,
                 "enum_and_caps_ms": round((time.perf_counter() - t0) * 1000, 1)}
R["finished"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
R["pymupdf_imported"] = "fitz" in sys.modules or "pymupdf" in sys.modules
(out / "pdf_print_probe_result.json").write_text(json.dumps(R, indent=1))
print(json.dumps(R, indent=1))
