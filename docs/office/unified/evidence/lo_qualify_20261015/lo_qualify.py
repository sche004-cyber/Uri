"""O2 pre-freeze LibreOffice qualification probe (research evidence, not product code).

Checks on this PC, with no admin rights and no installer:
  1. executable, version, file hash;
  2. per-call headless conversion (cold + warm) for DOCX and PPTX;
  3. persistent listener (soffice --accept pipe + UNO client in LibreOffice's own python): per-file time, RSS;
  4. external-link canary: does a DOCX with an external linked image make LibreOffice open a connection,
     under (a) a default private profile and (b) a hardened private profile;
  5. Windows Job Object: closing the job kills soffice.com and its soffice.bin child.
Usage: python lo_qualify.py <out_dir>
"""
import hashlib, http.server, json, os, shutil, subprocess, sys, tempfile, threading, time
from pathlib import Path
import psutil, pymupdf, docx, pptx, win32api, win32con, win32job
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

PROG = Path(os.environ["LOCALAPPDATA"]) / "URI" / "engines" / "libreoffice-26.8.1" / "program"
SOFFICE, LOPY = PROG / "soffice.com", PROG / "python.exe"
HERE = Path(__file__).resolve().parent
out = Path(sys.argv[1]).resolve(); shutil.rmtree(out, ignore_errors=True); out.mkdir(parents=True)
fx = out / "fixtures"; fx.mkdir()
R = {"started": time.strftime("%Y-%m-%dT%H:%M:%S%z")}


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


# ---------- fixtures (synthetic, no personal data) ----------
def make_order(p):
    d = docx.Document(); s = d.sections[0]
    s.header.paragraphs[0].text = "GOVERNMENT OF EXAMPLE STATE - DEPARTMENT OF FINANCE"
    r = s.footer.paragraphs[0].add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin"); r._r.append(b)
    t = OxmlElement("w:instrText"); t.text = " PAGE "; r._r.append(t)
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)
    h = d.add_heading("OFFICE ORDER No. F/12/2026-Estt", 1); h.alignment = 1
    d.add_paragraph("The following transfers are ordered with immediate effect, in the public interest.")
    tb = d.add_table(rows=1, cols=4); tb.style = "Table Grid"
    for c, v in zip(tb.rows[0].cells, ("Sl.", "Name", "From", "To")):
        c.text = v
    for i in range(1, 41):
        for c, v in zip(tb.add_row().cells, (str(i), f"Officer {i:02d}", f"Division {i % 7}", f"Division {(i + 3) % 7}")):
            c.text = v
    d.add_paragraph("By order, Under Secretary").runs[0].font.size = Pt(11)
    d.save(p)


def make_deck(p):
    pr = pptx.Presentation()
    for i in range(10):
        s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = f"Section {i + 1}"
        s.placeholders[1].text = "Point one\nPoint two\nPoint three"
        s.notes_slide.notes_text_frame.text = "Notes"
    s = pr.slides.add_slide(pr.slide_layouts[5]); s.shapes.title.text = "Revenue by quarter"
    cd = CategoryChartData(); cd.categories = ["Q1", "Q2", "Q3", "Q4"]; cd.add_series("Revenue", (12.4, 13.1, 13.9, 14.2))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    pr.save(p)


def make_canary(p, url):
    d = docx.Document(); d.add_paragraph("Canary: this document links an external image.")
    rid = d.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image",
                           is_external=True)
    xml = ('<w:r xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:drawing>'
           '<wp:inline xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing">'
           '<wp:extent cx="914400" cy="914400"/><wp:docPr id="1" name="canary"/>'
           '<a:graphic xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/picture">'
           '<pic:pic xmlns:pic="http://schemas.openxmlformats.org/drawingml/2006/picture">'
           '<pic:nvPicPr><pic:cNvPr id="1" name="canary.png"/><pic:cNvPicPr/></pic:nvPicPr>'
           '<pic:blipFill><a:blip xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
           f' r:link="{rid}"/><a:stretch><a:fillRect/></a:stretch></pic:blipFill>'
           '<pic:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="914400"/></a:xfrm>'
           '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></pic:spPr></pic:pic></a:graphicData></a:graphic>'
           '</wp:inline></w:drawing></w:r>')
    d.add_paragraph()._p.append(parse_xml(xml)); d.save(p)


# canary listener on a free loopback port; records every request
hits = []


class H(http.server.BaseHTTPRequestHandler):
    def do_GET(self):
        hits.append(self.path); self.send_response(404); self.end_headers()

    def log_message(self, *a):
        pass


srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H); port = srv.server_address[1]
threading.Thread(target=srv.serve_forever, daemon=True).start()

order, deck, canary = fx / "order.docx", fx / "deck.pptx", fx / "canary.docx"
make_order(order); make_deck(deck); make_canary(canary, f"http://127.0.0.1:{port}/canary.png")
R["fixtures"] = {p.name: sha(p) for p in (order, deck, canary)}
R["soffice"] = {"path": str(SOFFICE), "sha256_soffice_bin": sha(PROG / "soffice.bin"),
                "version": subprocess.run([str(SOFFICE), "--version"], capture_output=True, text=True,
                                          timeout=180).stdout.strip()}

HARDENED_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="DisableMacrosExecution" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="BlockUntrustedRefererLinks" oor:op="fuse"><value>true</value></prop></item>
<item oor:path="/org.openoffice.Inet/Settings"><prop oor:name="ooInetProxyType" oor:op="fuse"><value>2</value></prop></item>
<item oor:path="/org.openoffice.Inet/Settings"><prop oor:name="ooInetHTTPProxyName" oor:op="fuse"><value>127.0.0.1</value></prop></item>
<item oor:path="/org.openoffice.Inet/Settings"><prop oor:name="ooInetHTTPProxyPort" oor:op="fuse"><value>9</value></prop></item>
<item oor:path="/org.openoffice.Inet/Settings"><prop oor:name="ooInetNoProxy" oor:op="fuse"><value></value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="UseOpenCL" oor:op="fuse"><value>false</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="CollectUsageInformation" oor:op="fuse"><value>false</value></prop></item>
</oor:items>
"""


def profile(hardened):
    p = Path(tempfile.mkdtemp(prefix="uri_lo_prof_"))
    if hardened:
        (p / "user").mkdir()
        (p / "user" / "registrymodifications.xcu").write_text(HARDENED_XCU, encoding="utf-8")
    return p


def tree_rss(pid):
    try:
        ps = [psutil.Process(pid)] + psutil.Process(pid).children(recursive=True)
        return sum(x.memory_info().rss for x in ps if x.is_running())
    except psutil.Error:
        return 0


def per_call(src, prof, outdir):
    cmd = [str(SOFFICE), f"-env:UserInstallation={prof.as_uri()}", "--headless", "--norestore", "--nolockcheck",
           "--convert-to", "pdf", "--outdir", str(outdir), str(src)]
    t0 = time.perf_counter(); p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE); peak = 0
    while p.poll() is None:
        peak = max(peak, tree_rss(p.pid)); time.sleep(0.05)
    pdf = outdir / (src.stem + ".pdf")
    return {"rc": p.returncode, "seconds": round(time.perf_counter() - t0, 2), "peak_rss_mb": round(peak / 2**20, 1),
            "pages": pymupdf.open(pdf).page_count if pdf.exists() else None}


# ---------- 2. per-call conversion (hardened private profile, reused for the warm call) ----------
pc = out / "per_call"; pc.mkdir(); prof = profile(True)
R["per_call_hardened_profile"] = {f"{s.name}#{i}": per_call(s, prof, pc) for s in (order, deck) for i in (1, 2)}
shutil.rmtree(prof, ignore_errors=True)


# ---------- 3 + 5. persistent listener inside a Job Object ----------
def listener_run(hardened, files, label):
    prof = profile(hardened); od = out / label; od.mkdir()
    job = win32job.CreateJobObject(None, "")
    info = win32job.QueryInformationJobObject(job, win32job.JobObjectExtendedLimitInformation)
    info["BasicLimitInformation"]["LimitFlags"] = win32job.JOB_OBJECT_LIMIT_KILL_ON_JOB_CLOSE
    win32job.SetInformationJobObject(job, win32job.JobObjectExtendedLimitInformation, info)
    pipe = f"uri_lo_{os.getpid()}_{label}"
    t0 = time.perf_counter()
    lo = subprocess.Popen([str(SOFFICE), f"-env:UserInstallation={prof.as_uri()}", "--headless", "--invisible",
                           "--norestore", "--nolockcheck", "--nodefault", "--nologo",
                           f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext"],
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    h = win32api.OpenProcess(win32con.PROCESS_ALL_ACCESS, False, lo.pid)
    win32job.AssignProcessToJobObject(job, h)
    peak = [0]; stop = threading.Event()

    def sample():
        while not stop.is_set():
            peak[0] = max(peak[0], tree_rss(lo.pid)); time.sleep(0.05)
    threading.Thread(target=sample, daemon=True).start()
    cl = subprocess.run([str(LOPY), str(HERE / "uno_convert.py"), pipe, str(od)] + [str(f) for f in files],
                        capture_output=True, text=True, timeout=600)
    total = round(time.perf_counter() - t0, 2)
    rows = [json.loads(line) for line in cl.stdout.splitlines() if line.startswith("{")]
    time.sleep(1.0); idle = tree_rss(lo.pid)
    kids = [c.pid for c in psutil.Process(lo.pid).children(recursive=True)] if psutil.pid_exists(lo.pid) else []
    stop.set(); win32api.CloseHandle(h); job.Close()
    time.sleep(2.0)
    alive = [p for p in [lo.pid] + kids if psutil.pid_exists(p)]
    pages = {f.name: pymupdf.open(f).page_count for f in od.glob("*.pdf")}
    shutil.rmtree(prof, ignore_errors=True)
    return {"files": rows, "client_rc": cl.returncode, "client_stderr_tail": cl.stderr[-300:],
            "wall_total_s_incl_startup": total, "peak_tree_rss_mb": round(peak[0] / 2**20, 1),
            "idle_tree_rss_mb_after": round(idle / 2**20, 1), "child_pids_before_job_close": len(kids),
            "pids_alive_after_job_close": alive, "pdf_pages": pages}


R["listener_hardened"] = listener_run(True, [order, order, order, deck, deck], "listener_hardened")

# ---------- 4. external-link canary: default vs hardened profile, per-call and listener ----------
for hardened in (False, True):
    tag = "hardened" if hardened else "default"
    hits.clear(); d = out / f"canary_percall_{tag}"; d.mkdir(); prof = profile(hardened)
    res = per_call(canary, prof, d); shutil.rmtree(prof, ignore_errors=True); time.sleep(1)
    R[f"canary_percall_{tag}"] = {"convert": res, "canary_hits": list(hits)}
    hits.clear(); res = listener_run(hardened, [canary], f"canary_listener_{tag}"); time.sleep(1)
    R[f"canary_listener_{tag}"] = {"convert": res["files"], "canary_hits": list(hits)}

srv.shutdown()
op = out / "per_call" / "order.pdf"
if op.exists():
    pymupdf.open(op)[0].get_pixmap(dpi=70).save(out / "per_call_order_p1.png")
lp = out / "listener_hardened" / "deck.pdf"
if lp.exists():
    dd = pymupdf.open(lp); dd[dd.page_count - 1].get_pixmap(dpi=70).save(out / "listener_deck_chart.png")
R["finished"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
(out / "lo_qualify_result.json").write_text(json.dumps(R, indent=1))
print(json.dumps(R, indent=1))
