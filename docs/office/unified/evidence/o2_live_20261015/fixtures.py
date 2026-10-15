"""Synthetic administrative fixtures for the O2 live run (plan §12 row J1; owner OD-U7: synthetic only).

Every name, number and address is invented. Nothing is copied from a real document.
"""
from __future__ import annotations

import io
from pathlib import Path


def official_letter(p: Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt
    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = "GOVERNMENT OF THE STATE OF SAMPLE · DEPARTMENT OF PUBLIC RECORDS"
    sec.footer.paragraphs[0].text = "Synthetic letter for testing · Page"
    t = d.add_paragraph("No. DPR/2026/0417")
    t.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("Dated: 15 October 2026").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    d.add_paragraph("To\nThe District Collector,\nSample District,\nSample City 100001")
    h = d.add_paragraph("Subject: Submission of the quarterly record-room inventory")
    h.runs[0].bold = True
    d.add_paragraph("Sir,")
    for k in range(4):
        par = d.add_paragraph(
            f"{k + 1}. With reference to the circular dated 2 September 2026, the department encloses the inventory of "
            "records held in the district record room for the quarter ending 30 September 2026. The inventory lists "
            "each register by series, year and shelf, and marks the registers sent for digitization.")
        par.paragraph_format.space_after = Pt(6)
    d.add_paragraph("Yours faithfully,\n\n(A. Sample)\nDeputy Secretary")
    d.save(p)


def office_order(p: Path):
    import docx
    d = docx.Document()
    d.add_heading("OFFICE ORDER No. 41 of 2026", 1)
    d.add_paragraph("The following officials are assigned to the record digitization drive with immediate effect.")
    tbl = d.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    for i, h in enumerate(("S. No.", "Name", "Designation", "Section")):
        tbl.rows[0].cells[i].text = h
    for n in range(1, 42):
        r = tbl.add_row().cells
        r[0].text, r[1].text = str(n), f"Official {n:02d}"
        r[2].text = ("Section Officer", "Assistant", "Clerk")[n % 3]
        r[3].text = f"Section {chr(65 + n % 6)}"
    d.add_paragraph("By order,\nJoint Secretary")
    d.save(p)


def notice_pair(notice: Path, corrigendum: Path):
    import docx
    for path, title, body in (
            (notice, "NOTICE", "Applications for the record assistant posts close on 31 October 2026."),
            (corrigendum, "CORRIGENDUM", "In the notice dated 1 October 2026, read 7 November 2026 for 31 October "
                                         "2026. Other conditions remain unchanged.")):
        d = docx.Document()
        d.add_heading(title, 1)
        d.add_paragraph("No. DPR/REC/2026/12")
        d.add_paragraph(body)
        d.save(path)


def memorandum(p: Path):
    import docx
    d = docx.Document()
    d.add_heading("MEMORANDUM", 1)
    d.add_paragraph("Sub: Delay in submission of attendance statements")
    for k in range(3):
        d.add_paragraph(f"{k + 1}. It has been observed that statements for September were received late from "
                        "three sections. Section heads shall ensure timely submission.", style="List Number")
    d.save(p)


def minutes(p: Path):
    import docx
    d = docx.Document()
    d.add_heading("Minutes of the review meeting held on 10 October 2026", 1)
    d.add_heading("Present", 2)
    for n in range(1, 7):
        d.add_paragraph(f"Member {n}", style="List Bullet")
    d.add_heading("Decisions", 2)
    for n in range(1, 6):
        d.add_paragraph(f"Item {n}: The committee approved the action plan for item {n}.", style="List Number")
    d.save(p)


def bilingual_letter(p: Path):
    """English and Hindi (Devanagari) text; checks complex-script rendering and font fallback."""
    import docx
    d = docx.Document()
    d.add_heading("Public notice / सार्वजनिक सूचना", 1)
    d.add_paragraph("The office will remain closed on 2 November 2026.")
    d.add_paragraph("कार्यालय 2 नवंबर 2026 को बंद रहेगा। सभी आवेदन अगले कार्य दिवस पर स्वीकार किए जाएंगे।")
    d.add_paragraph("आदेश से,\nउप सचिव")
    d.save(p)


def report20(p: Path):
    import docx
    d = docx.Document()
    d.add_heading("Annual Report 2025-26 (synthetic)", 0)
    for sec in range(1, 21):
        d.add_heading(f"{sec}. Chapter {sec}", 1)
        for k in range(5):
            d.add_paragraph(f"Paragraph {k + 1} of chapter {sec}. The department continued the digitization of "
                            "records, trained staff and reviewed its procedures during the year. " * 2)
        d.add_page_break()
    d.save(p)


def deck11(p: Path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[0])
    s.shapes.title.text = "Record digitization: progress review"
    s.placeholders[1].text = "Synthetic deck · October 2026"
    for n in range(2, 11):
        s = pr.slides.add_slide(pr.slide_layouts[1])
        s.shapes.title.text = f"Section {n - 1}: status"
        s.placeholders[1].text = f"Registers scanned: {n * 120}\nPending: {1000 - n * 80}\nIssues: none critical"
    s = pr.slides.add_slide(pr.slide_layouts[5])
    s.shapes.title.text = "Registers scanned per month"
    cd = CategoryChartData()
    cd.categories = ["Jun", "Jul", "Aug", "Sep"]
    cd.add_series("Scanned", (320, 410, 515, 640))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.5), Inches(8), Inches(5), cd)
    pr.save(p)


def budget_workbook(p: Path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Head", "Allocation 2026-27", "Spent to date"])
    rows = [("Salaries", 1250000, 610500), ("Digitization", 480000, 212750), ("Stationery", 35000, 18420),
            ("Travel", 60000, 22100), ("Maintenance", 90000, 41380)]
    for r in rows:
        ws.append(list(r))
    ws["A8"] = "Total"
    ws["B8"] = "=SUM(B2:B6)"
    ws["C8"] = "=SUM(C2:C6)"
    ws["B8"].value, ws["C8"].value = "=SUM(B2:B6)", "=SUM(C2:C6)"
    ws.print_area = "A1:C8"
    # Column widths wide enough for the text (run 4 used the default widths, so "Digitization" printed clipped,
    # exactly as the sheet shows it; that was a fixture defect, not a conversion defect).
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    wb.save(p)
    # Cached values for the formulas, as Excel would store them (LibreOffice never recalculates on load).
    import re
    import zipfile
    parts = {}
    with zipfile.ZipFile(p) as z:
        for n in z.namelist():
            parts[n] = z.read(n)
    x = parts["xl/worksheets/sheet1.xml"].decode()
    x = re.sub(r'(<c r="B8"[^>]*>)<f>SUM\(B2:B6\)</f>(<v></v>)?', r"\1<f>SUM(B2:B6)</f><v>1915000</v>", x)
    x = re.sub(r'(<c r="C8"[^>]*>)<f>SUM\(C2:C6\)</f>(<v></v>)?', r"\1<f>SUM(C2:C6)</f><v>905150</v>", x)
    parts["xl/worksheets/sheet1.xml"] = x.encode()
    with zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    return [("Salaries", "1250000"), ("Digitization", "480000"), ("Total", "1915000"), ("Spent to date", None)]


def scanned_pdf(p: Path, page_png: bytes):
    """A 'scanned' PDF: one page image and no text layer."""
    from PIL import Image
    im = Image.open(io.BytesIO(page_png)).convert("L")
    im.save(p, "PDF", resolution=150)
