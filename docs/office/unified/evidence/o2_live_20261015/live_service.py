"""O2 / O2P live run against the REAL Office service process on this Windows 11 PC (plan §12, A1 §A8).

    python -I live_service.py <out_dir>

What it drives: `python -m uri_office.service run` (a child process, loopback only, its own state folder, the pinned
LibreOffice, the WinRT raster worker, the pypdf worker, `--enable-printing` for the O2P rows), over HTTP with the
service's bearer token, exactly as the Flutter client calls it. Approvals go through `/a2ui/action` with the bound
hash, like a button press. Nothing bypasses an approval.

What it does NOT cover (reported, never claimed): the Flutter UI (separate UI run), the Brain (GPU window: BLOCKED),
the owner's own check (BLOCKED, owner away), physical printing (NOT VERIFIED, no physical printer; OD-U13 allows only
Microsoft Print to PDF), Microsoft Office fidelity (NOT VERIFIED, no Office licence).

Every document is synthetic (fixtures.py). Outputs go to <out_dir>; virtual-printer outputs go to a dedicated folder
that is emptied at the end (OD-U13).
"""
from __future__ import annotations

import hashlib
import http.client
import json
import os
import shutil
import statistics
import subprocess
import sys
import threading
import time
import traceback
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[4]
sys.path.insert(0, str(HERE))
import fixtures as FX  # noqa: E402

OUT = Path(sys.argv[1]).resolve()
shutil.rmtree(OUT, ignore_errors=True)
OUT.mkdir(parents=True)
FIXT, SAVED, SHOTS, STATE = OUT / "fixtures", OUT / "saved", OUT / "pages", OUT / "state"
for d in (FIXT, SAVED, SHOTS):
    d.mkdir()
PRINT_DIR = OUT / "virtual_print"            # OD-U13 dedicated folder, emptied at the end
PRINT_DIR.mkdir()
R: dict = {"started": time.strftime("%Y-%m-%dT%H:%M:%S%z"), "rows": {}, "timings": {}, "census": {}}


def sha(p: Path) -> str:
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def row(rid: str, verdict: str, **ev):
    R["rows"][rid] = {"verdict": verdict, **ev}
    print(f"[{verdict}] {rid} {json.dumps(ev, default=str)[:300]}", flush=True)


class Svc:
    def __init__(self):
        self.proc = None

    def start(self):
        env = {k: v for k, v in os.environ.items() if not k.upper().startswith("PYTHON")}
        self.log = open(OUT / "service.log", "ab")
        self.proc = subprocess.Popen([sys.executable, "-m", "uri_office.service", "run", "--state", str(STATE),
                                      "--port", "0", "--allow-root", str(OUT), "--enable-printing"],
                                     cwd=str(REPO), stdout=subprocess.PIPE, stderr=self.log, env=env)
        line = self.proc.stdout.readline().decode()
        info = json.loads(line)
        self.port = info["port"]
        self.token = (STATE / "api-token").read_text().strip()

    def stop(self, kill=False):
        if self.proc is None:
            return
        (self.proc.kill if kill else self.proc.terminate)()
        self.proc.wait(30)
        self.proc = None

    def call(self, method, path, body=None, raw=False, timeout=600):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=timeout)
        c.request(method, path, body=json.dumps(body).encode() if body is not None else None,
                  headers={"Authorization": f"Bearer {self.token}", "Content-Type": "application/json"})
        r = c.getresponse()
        data = r.read()
        c.close()
        if raw:
            return r.status, data
        return r.status, (json.loads(data) if data else {})


S = Svc()


def wait(fn, timeout=300, step=0.25):
    t0 = time.monotonic()
    while time.monotonic() - t0 < timeout:
        v = fn()
        if v:
            return v
        time.sleep(step)
    raise TimeoutError("wait")


def act(messages, button="apply"):
    sid = messages[0]["createSurface"]["surfaceId"]
    bound = messages[2]["updateDataModel"]["value"]["bound"]
    key = "plan_sha256" if "plan_sha256" in bound else "change_sha256"
    name = {("plan_sha256", "apply"): "office.run_plan", ("plan_sha256", "reject"): "office.discard_plan",
            ("change_sha256", "apply"): "office.apply_change", ("change_sha256", "reject"): "office.reject_change"}[
        (key, button)]
    return S.call("POST", "/a2ui/action", {"action": {"name": name, "surfaceId": sid, "sourceComponentId": button,
                                                      "context": {"surface_digest": bound["surface"]["digest"],
                                                                  key: bound[key]}}})


def open_doc(path: Path) -> dict:
    t0 = time.perf_counter()
    st, v = S.call("POST", "/open/document", {"path": str(path)})
    if st != 200:
        return {"error": v, "status": st}
    did = v["document_id"]
    v = wait(lambda: (lambda x: x if x["preview"]["state"] != "rendering" else None)(
        S.call("GET", f"/document/{did}")[1]), 300)
    v["first_ready_s"] = round(time.perf_counter() - t0, 2)
    return v


def page_png(did, n, w, name) -> tuple:
    t0 = time.perf_counter()
    st, data = S.call("GET", f"/document/{did}/page/{n}?w={w}", raw=True)
    ms = round((time.perf_counter() - t0) * 1000)
    if st == 200:
        (SHOTS / name).write_bytes(data)
    return st, ms


def save_as_pdf(source: dict, dest: Path) -> dict:
    t0 = time.perf_counter()
    st, wf = S.call("POST", "/workflow", {"capability": "save_as_pdf", "inputs": [source]})
    if st != 200:
        return {"error": wf, "status": st}
    wid = wf["workflow_id"]
    st, r = act(wf["messages"])
    v = wait(lambda: (lambda x: x if x["state"] not in ("running",) else None)(S.call("GET", f"/workflow/{wid}")[1]))
    out = {"workflow_id": wid, "state": v["state"], "findings": v.get("findings", []), "outputs": v.get("outputs"),
           "steps": [s["state"] for s in v["steps"]]}
    if v["state"] == "ready_to_save":
        if any(f["severity"] == "warning" for f in v["findings"]):
            S.call("POST", f"/workflow/{wid}/accept_warnings", {})
        st, card = S.call("POST", f"/workflow/{wid}/propose_save", {"save_as_path": str(dest)})
        if st == 200:
            st, r = act(card["messages"])
            out["save"] = r
    out["seconds"] = round(time.perf_counter() - t0, 2)
    return out


def proc_tree_pids():
    import psutil
    try:
        p = psutil.Process(S.proc.pid)
        return [p.pid] + [c.pid for c in p.children(recursive=True)]
    except psutil.Error:
        return []


def census(label):
    """RSS and CPU of the service tree; GPU dedicated memory and engine use of every process in the tree."""
    import psutil
    out = {"processes": []}
    for pid in proc_tree_pids():
        try:
            p = psutil.Process(pid)
            out["processes"].append({"pid": pid, "name": p.name(), "rss_mb": round(p.memory_info().rss / 2**20, 1),
                                     "cpu_s": round(sum(p.cpu_times()[:2]), 2)})
        except psutil.Error:
            pass
    pids = [x["pid"] for x in out["processes"]]
    ps = ("$ErrorActionPreference='SilentlyContinue'; $ids=@(" + ",".join(map(str, pids)) + "); "
          "$m=(Get-Counter '\\GPU Process Memory(*)\\Dedicated Usage').CounterSamples; "
          "$e=(Get-Counter '\\GPU Engine(*)\\Utilization Percentage').CounterSamples; "
          "$r=@(); foreach($i in $ids){ $mm=($m|?{$_.InstanceName -like \"pid_$($i)_*\"}|Measure-Object CookedValue -Sum).Sum;"
          " $ee=($e|?{$_.InstanceName -like \"pid_$($i)_*\"}|Measure-Object CookedValue -Maximum).Maximum;"
          " $r+=@{pid=$i;gpu_dedicated_bytes=[int64]$mm;gpu_engine_max=[double]$ee} }; $r|ConvertTo-Json -Compress")
    try:
        g = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=60)
        gj = json.loads(g.stdout) if g.stdout.strip() else []
        out["gpu"] = gj if isinstance(gj, list) else [gj]
    except Exception as e:  # noqa: BLE001
        out["gpu_error"] = repr(e)[:200]
    R["census"][label] = out
    return out


def temp_entries():
    t = Path(os.environ["TEMP"])
    try:
        return {p.name for p in t.iterdir()}
    except OSError:
        return set()


try:
    temp_before = temp_entries()
    # ------------------------------------------------------------------ fixtures (J1)
    F = {"letter": FIXT / "official_letter.docx", "order": FIXT / "office_order_41.docx",
         "notice": FIXT / "notice.docx", "corr": FIXT / "corrigendum.docx", "memo": FIXT / "memorandum.docx",
         "minutes": FIXT / "minutes.docx", "bilingual": FIXT / "bilingual_letter.docx",
         "report": FIXT / "annual_report_20.docx", "tender_docx": FIXT / "tender60.docx",
         "deck": FIXT / "progress_deck_11.pptx", "budget": FIXT / "budget.xlsx"}
    FX.official_letter(F["letter"])
    FX.office_order(F["order"])
    FX.notice_pair(F["notice"], F["corr"])
    FX.memorandum(F["memo"])
    FX.minutes(F["minutes"])
    FX.bilingual_letter(F["bilingual"])
    FX.report20(F["report"])
    FX.deck11(F["deck"])
    budget_cells = FX.budget_workbook(F["budget"])
    import docx
    d = docx.Document()
    for pg in range(1, 61):
        d.add_heading(f"Tender clause {pg}", 2)
        d.add_paragraph(f"Clause {pg}. The bidder shall comply with the conditions of the tender document. " * 25)
        if pg < 60:
            d.add_page_break()
    d.save(F["tender_docx"])
    hashes_before = {k: sha(v) for k, v in F.items()}

    # ------------------------------------------------------------------ service start, engine verification
    t0 = time.perf_counter()
    S.start()
    R["timings"]["service_start_s"] = round(time.perf_counter() - t0, 2)
    st, c = S.call("GET", "/contract")
    row("contract", "PASS" if c.get("contract") == "o2.1" else "FAIL", contract=c.get("contract"))
    t0 = time.perf_counter()
    eng = wait(lambda: (lambda e: e if e["engines"][0]["state"] in ("ready", "unavailable") else None)(
        S.call("GET", "/engines")[1]), 400, 1.0)
    R["timings"]["first_engine_verification_s"] = round(time.perf_counter() - t0, 2)
    row("AC-O2-6/19 engine verified at first start", "PASS" if eng["engines"][0]["state"] == "ready" else "FAIL",
        engine=eng["engines"][0])
    census("idle_after_start")

    # ------------------------------------------------------------------ open and preview every kind (AC-O2-1, K1)
    opened = {}
    for k in ("letter", "order", "notice", "corr", "memo", "minutes", "bilingual", "report", "deck"):
        v = open_doc(F[k])
        opened[k] = v
        ok = v.get("preview", {}).get("state") == "ready" and (v.get("pages") or 0) >= 1
        st, ms = page_png(v["document_id"], 1, 880, f"{k}_p1.png") if ok else (None, None)
        row(f"K1 open+preview {k}", "PASS" if ok and st == 200 else "FAIL", pages=v.get("pages"),
            first_ready_s=v.get("first_ready_s"), page1_ms=ms, findings=len(v.get("findings", [])),
            outline=len(v.get("outline", [])), error=v.get("error"))
        if ok and k in ("notice", "corr", "memo", "minutes"):
            S.call("POST", f"/document/{v['document_id']}/close", {})   # at most 8 documents are open at once
    # every page of the deck and the order, for visual inspection
    for k, n in (("deck", opened["deck"]["pages"]), ("order", opened["order"]["pages"])):
        for i in range(1, n + 1):
            page_png(opened[k]["document_id"], i, 880, f"{k}_p{i}.png")
    row("AC-O2 deck slide count", "PASS" if opened["deck"]["pages"] == 11 and opened["deck"].get("slides") == 11
        else "FAIL", pages=opened["deck"]["pages"], slides=opened["deck"].get("slides"))
    census("after_previews")

    # ------------------------------------------------------------------ U12 / H1 timings: warm 20-page DOCX
    rep = []
    for i in range(3):
        st, _ = S.call("POST", f"/document/{opened['report']['document_id']}/close", {})
        v = open_doc(F["report"])
        opened["report"] = v
        rep.append(v["first_ready_s"])
    row("U12/H1 warm 20-page DOCX first page visible", "PASS" if statistics.median(rep) <= 2.0 else "FAIL",
        seconds=rep, median=statistics.median(rep), threshold_s=2.0,
        note="cache hit after the first conversion of the same content (preview cache); see cold_conversion_s")

    # ------------------------------------------------------------------ Save as PDF (P-1, D1, I1, AC-O2-7)
    saved = {}
    for k in ("letter", "order", "bilingual", "deck", "report"):
        out = save_as_pdf({"document_id": opened[k]["document_id"]}, SAVED / f"{F[k].stem}.pdf")
        saved[k] = out
        ok = out.get("state") == "ready_to_save" and (out.get("save") or {}).get("status") == "applied" and \
            (SAVED / f"{F[k].stem}.pdf").exists()
        row(f"P-1/D1 Save as PDF {k}", "PASS" if ok else "FAIL", state=out.get("state"), seconds=out.get("seconds"),
            pages=(out.get("outputs") or [{}])[0].get("pages"),
            findings=[f"{f['severity']}:{f['check']}" for f in out.get("findings", [])])
    # no overwrite (D1)
    st, wf = S.call("POST", "/workflow", {"capability": "save_as_pdf",
                                         "inputs": [{"document_id": opened["letter"]["document_id"]}]})
    act(wf["messages"])
    wait(lambda: S.call("GET", f"/workflow/{wf['workflow_id']}")[1]["state"] != "running")
    S.call("POST", f"/workflow/{wf['workflow_id']}/accept_warnings", {})
    st, r = S.call("POST", f"/workflow/{wf['workflow_id']}/propose_save",
                   {"save_as_path": str(SAVED / "official_letter.pdf")})
    row("D1 no overwrite of an existing name", "PASS" if r.get("error") == "DESTINATION_EXISTS" else "FAIL",
        response=r.get("error"))
    # reopen the saved PDFs in Office (P-1 reopen)
    for k in ("letter", "deck"):
        v = open_doc(SAVED / f"{F[k].stem}.pdf")
        st, ms = page_png(v["document_id"], 1, 880, f"reopened_{k}_p1.png")
        row(f"P-1 reopen saved PDF {k}", "PASS" if v.get("pages") == saved[k]["outputs"][0]["pages"] and st == 200
            else "FAIL", pages=v.get("pages"), renderer=v.get("renderer"))
        S.call("POST", f"/document/{v['document_id']}/close", {})
    for k in ("bilingual", "report"):
        S.call("POST", f"/document/{opened[k]['document_id']}/close", {})   # stay within 8 open documents

    # ------------------------------------------------------------------ PDF viewing (P-2): 60-page tender + scanned
    tdoc = open_doc(F["tender_docx"])
    out = save_as_pdf({"document_id": tdoc["document_id"]}, SAVED / "tender60.pdf")
    S.call("POST", f"/document/{tdoc['document_id']}/close", {})
    tv = open_doc(SAVED / "tender60.pdf")
    did = tv["document_id"]
    widths = {}
    for w in (440, 880, 1240):
        times = []
        for n in (1, 30, 60):
            st, ms = page_png(did, n, w, f"tender_p{n}_w{w}.png")
            times.append(ms)
        widths[w] = times
    st, m = S.call("POST", f"/document/{did}/search", {"text": "Tender clause 37"})
    st2, m2 = S.call("POST", f"/document/{did}/search", {"text": "nonexistent phrase xyz"})
    row("P-2 PDF view: pages, zoom 440/880/1240, search", "PASS" if tv.get("pages") == 60 and
        [x["page"] for x in m.get("matches", [])][:1] == [37] and m2.get("matches") == [] else "FAIL",
        pages=tv.get("pages"), page_ms_by_width=widths, first_match=m.get("matches", [])[:1], no_match=m2)
    shot = (SHOTS / "letter_p1.png").read_bytes()
    FX.scanned_pdf(FIXT / "scanned_letter.pdf", shot)
    hashes_before["scanned"] = sha(FIXT / "scanned_letter.pdf")
    sv = open_doc(FIXT / "scanned_letter.pdf")
    st, sm = S.call("POST", f"/document/{sv['document_id']}/search", {"text": "record"})
    S.call("POST", f"/document/{sv['document_id']}/close", {})
    row("P-2 scanned PDF: preview and honest 'no searchable text'",
        "PASS" if sv.get("preview", {}).get("state") == "ready" and sm.get("error") == "NO_TEXT_LAYER" else "FAIL",
        search=sm)

    # ------------------------------------------------------------------ P-4 hostile and malformed PDFs
    hostile = {"zero": b"", "garbage": os.urandom(50000), "truncated": (SAVED / "tender60.pdf").read_bytes()[:20000],
               "deep_nesting": b"%PDF-1.4\n1 0 obj\n" + b"[" * 200000 + b"]" * 200000 + b"\nendobj\n%%EOF"}
    survived = True
    res = {}
    for name, data in hostile.items():
        p = FIXT / f"hostile_{name}.pdf"
        p.write_bytes(data)
        v = open_doc(p)
        res[name] = v.get("preview", {}).get("state") or v.get("error", {}).get("error")
        if v.get("document_id"):
            S.call("POST", f"/document/{v['document_id']}/close", {})
        survived = survived and S.call("GET", "/contract")[0] == 200
    row("P-4/AC-O2-20 hostile PDFs: worker fails alone, service survives",
        "PASS" if survived and all(x in ("unavailable", "ready") for x in res.values()) else "FAIL", results=res)

    # ------------------------------------------------------------------ P-6 Excel Save as PDF
    st, wbv = S.call("POST", "/open/local", {"path": str(F["budget"])})
    wid = wbv["workbook_id"]
    S.call("POST", f"/workbook/{wid}/edit", {"sheet": "Budget", "cell": "C2", "value": 999})
    st, r = S.call("POST", "/workflow", {"capability": "save_as_pdf", "inputs": [{"workbook_id": wid}]})
    blocked = r.get("error") == "UNSAVED_CHANGES"
    S.call("POST", f"/workbook/{wid}/undo", {})
    out = save_as_pdf({"workbook_id": wid}, SAVED / "budget.pdf")
    import pypdf
    text = "".join(pg.extract_text() or "" for pg in pypdf.PdfReader(str(SAVED / "budget.pdf")).pages) \
        if (SAVED / "budget.pdf").exists() else ""
    norm = text.replace(",", "").replace(" ", "")
    present = {label: (val is None or val in norm) and label.replace(" ", "") in norm for label, val in budget_cells}
    row("P-6 Excel: refused with unsaved changes; saved file -> PDF with the print-area values",
        "PASS" if blocked and all(present.values()) else "FAIL", refused=r.get("error"), values_present=present,
        pages=(out.get("outputs") or [{}])[0].get("pages"))

    # ------------------------------------------------------------------ F1 canaries through the live service
    sys.path.insert(0, str(REPO / "tests" / "office_u2"))
    import canaries as CN  # noqa: E402
    import http.server
    hits = []

    class H(http.server.BaseHTTPRequestHandler):
        def do_GET(self):
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()
        do_HEAD = do_POST = do_PROPFIND = do_OPTIONS = do_GET

        def log_message(self, *a):
            pass
    srv = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{srv.server_address[1]}/canary"
    can = {"docx": CN.docx_canary(FIXT / "canary.docx", base), "pptx": CN.pptx_canary(FIXT / "canary.pptx", base),
           "xlsx": CN.xlsx_canary(FIXT / "canary.xlsx", base)}
    import psutil
    smb = set()
    stop = threading.Event()

    def watch():
        while not stop.is_set():
            for cc in psutil.net_connections(kind="inet"):
                if cc.raddr and cc.raddr.ip == "127.0.0.2" and cc.raddr.port in (445, 139):
                    smb.add(cc.raddr.port)
            stop.wait(0.05)
    tw = threading.Thread(target=watch, daemon=True)
    tw.start()
    for k in ("docx", "pptx"):
        v = open_doc(can[k])
        save_as_pdf({"document_id": v["document_id"]}, SAVED / f"canary_{k}.pdf")
        S.call("POST", f"/document/{v['document_id']}/close", {})
    st, cwv = S.call("POST", "/open/local", {"path": str(can["xlsx"])})
    if st == 200:
        save_as_pdf({"workbook_id": cwv["workbook_id"]}, SAVED / "canary_xlsx.pdf")
    time.sleep(1.0)
    stop.set()
    srv.shutdown()
    row("F1 canaries (HTTP, SMB) through the live service", "PASS" if not hits and not smb else "FAIL",
        http_hits=hits, smb_ports=sorted(smb), xlsx_opened=st == 200)

    # ------------------------------------------------------------------ L2 repeats of D1
    reps = []
    for i in range(5):
        o = save_as_pdf({"document_id": opened["order"]["document_id"]}, SAVED / f"order_repeat_{i}.pdf")
        reps.append((o.get("state"), (o.get("outputs") or [{}])[0].get("pages"), o.get("seconds")))
    row("L2 five repeats of Save as PDF give identical outcomes",
        "PASS" if len({(a, b) for a, b, _ in reps}) == 1 and reps[0][0] == "ready_to_save" else "FAIL", runs=reps)

    # ------------------------------------------------------------------ O2P: discovery, preview, card, virtual job
    st, pr = S.call("GET", "/devices/printers")
    ps = subprocess.run(["powershell", "-NoProfile", "-Command",
                         "Get-Printer | Select-Object Name | ConvertTo-Json -Compress; "
                         "(Get-CimInstance Win32_Printer | Where-Object Default).Name"],
                        capture_output=True, text=True, timeout=60).stdout.strip().splitlines()
    win_names = {x["Name"] for x in json.loads(ps[0])} if ps else set()
    win_default = ps[1].strip() if len(ps) > 1 else ""
    ours = {p["name"] for p in pr.get("printers", [])}
    dflt = [p["name"] for p in pr.get("printers", []) if p["is_default"]]
    row("PR-1 discovery matches Windows (virtual printers only on this PC)",
        "PASS" if ours == win_names and dflt == [win_default] else "FAIL", ours=sorted(ours),
        windows=sorted(win_names), default=dflt, windows_default=win_default, physical=pr.get("physical"))
    vp = next((p for p in pr.get("printers", []) if p["name"] == "Microsoft Print to PDF"), None)
    if vp:
        caps = S.call("GET", f"/devices/printers/{vp['id']}")[1]["capabilities"]
        row("PR-2 capabilities read (virtual)", "PASS" if caps.get("papers") else "FAIL",
            duplex=caps.get("duplex"), color=caps.get("color"), papers=len(caps.get("papers", [])),
            max_copies=caps.get("max_copies"))
        src = {"document_id": opened["order"]["document_id"]}
        body = {"source": src, "printer_id": vp["id"], "pages": [1, 2], "output_path": str(PRINT_DIR / "order.pdf"),
                "settings": {"paper": "A4", "fit": "fit", "dpi": 300}}
        st, pv = S.call("POST", "/print/preview", body)
        row("PR-3 print preview: exactly the chosen pages, printable area", "PASS" if pv.get("pages") == [1, 2] and
            pv.get("printable_area", {}).get("ok") else "FAIL", pages=pv.get("pages"),
            printable=pv.get("printable_area"))
        st, card = S.call("POST", "/print/propose", body)
        written_before = (PRINT_DIR / "order.pdf").exists()
        # PR-4: a replayed / tampered approval is refused
        msgs = card["messages"]
        bad = json.loads(json.dumps(msgs))
        bad[2]["updateDataModel"]["value"]["bound"]["change_sha256"] = "0" * 64
        stb, rb = act(bad)
        st, r = act(msgs)
        st2, r2 = act(msgs)
        job = r.get("job")
        jv = wait(lambda: (lambda x: x if x["state"] not in ("queued", "printing") else None)(
            S.call("GET", f"/print/{job}")[1]), 180) if job else {}
        pages_out = len(pypdf.PdfReader(str(PRINT_DIR / "order.pdf")).pages) if (PRINT_DIR / "order.pdf").exists() \
            else None
        row("PR-4 tampered and replayed print approvals refused", "PASS" if stb != 200 and st2 != 200 else "FAIL",
            tampered=rb.get("error"), replay=r2.get("error"))
        row("PR-5 virtual printer job (OD-U13): approved pages only, no dialog, file written",
            "PASS" if not written_before and jv.get("state") == "sent" and pages_out == 2 else "FAIL",
            state=jv.get("state"), pages_out=pages_out, worker=jv.get("worker"), spooler=jv.get("spooler"),
            note="virtual printer evidence only; NOT physical qualification")
        st, again = S.call("POST", "/print/propose", body)
        row("PR-5b same output file refused (never replaced)", "PASS" if again.get("error") == "DESTINATION_EXISTS"
            else "FAIL", response=again.get("error"))
    for rid in ("PR-6 physical print", "PR-7 physical progress and cancel", "PR-8 real offline printer",
                "PR-9 printer error mid-job"):
        row(rid, "NOT VERIFIED", reason="no physical printer on this PC (2026-10-15); needs owner authorization")

    # ------------------------------------------------------------------ H2 census, then again after the idle stop
    census("after_conversions")
    time.sleep(75)
    census("after_70s_idle")
    import pypdf as _pp
    chart_text = _pp.PdfReader(str(SAVED / "progress_deck_11.pdf")).pages[10].extract_text() or ""
    row("B1 native chart drawn in the deck's PDF (slide 11)", "PASS" if "Jun" in chart_text and "Sep" in chart_text
        else "FAIL", chart_text=chart_text[:120])

    # ------------------------------------------------------------------ F3 kill mid-workflow, restart, resume
    v = open_doc(F["report"])
    st, wf = S.call("POST", "/workflow", {"capability": "save_as_pdf", "inputs": [{"document_id": v["document_id"]}]})
    act(wf["messages"])
    time.sleep(0.3)
    lo_pids_before = proc_tree_pids()
    S.stop(kill=True)
    time.sleep(3)
    import psutil
    orphans = [p for p in lo_pids_before if psutil.pid_exists(p)]
    S.start()
    rec = S.call("GET", f"/workflow/{wf['workflow_id']}")[1]
    st, rv = S.call("POST", f"/workflow/{wf['workflow_id']}/resume", {"inputs": [{"path": str(F["report"])}]})
    done = wait(lambda: (lambda x: x if x["state"] != "running" else None)(
        S.call("GET", f"/workflow/{wf['workflow_id']}")[1]), 300) if st == 200 else {}
    row("F3/AC-O2-5/AC-O2-9 kill mid-run: no orphan engine, interrupted, resumed", "PASS" if not orphans and
        rec.get("state") == "interrupted" and done.get("state") in ("ready_to_save", "needs_review") else "FAIL",
        orphans=orphans, after_restart=rec.get("state"), interrupted_at=rec.get("interrupted_at_step"),
        resumed=done.get("state"))

    # ------------------------------------------------------------------ G1 originals and temp
    after = {k: sha(v) for k, v in F.items()}
    after["scanned"] = sha(FIXT / "scanned_letter.pdf")
    changed = [k for k in hashes_before if hashes_before[k] != after[k]]
    row("G1 originals byte-identical after every item", "PASS" if not changed else "FAIL", changed=changed,
        files=len(hashes_before))
finally:
    try:
        S.stop()
    except Exception:  # noqa: BLE001
        pass
    temp_after = temp_entries()
    new_temp = sorted(x for x in temp_after - temp_before if not x.startswith(("pytest", "claude", "tmp")))
    R["temp_new_entries"] = new_temp
    # OD-U13 clean-up of the virtual printer outputs
    for p in PRINT_DIR.iterdir():
        p.unlink()
    R["virtual_print_dir_empty"] = not any(PRINT_DIR.iterdir())
    R["finished"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
    R["error"] = None if sys.exc_info()[0] is None else traceback.format_exc()[-2000:]
    (OUT / "live_result.json").write_text(json.dumps(R, indent=1, default=str), encoding="utf-8")
    print("RESULT", OUT / "live_result.json")
