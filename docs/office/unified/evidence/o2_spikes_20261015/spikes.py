"""O2 day-1 spikes (implementation research, G1 active). Nothing installed, nothing printed, no PyMuPDF.

S1  LibreOffice python.exe running a bridge script under -I: does `import uno` work and does a conversion succeed?
S3  ctypes Job Object (no pywin32): KILL_ON_JOB_CLOSE + process memory limit; closing the job kills soffice tree.
S4  Listener start-up with a persistent, re-seeded profile (2nd start) vs a fresh profile (1st start).
S6  Cost of a full-tree SHA-256 manifest of the LibreOffice extract; cost of executables + DLLs + python only.
S7  Listener timing: a 20+ page DOCX and a 60-page PDF-bound DOCX (load + export, 5 runs each).
S9  pypdf text positions and font sizes (visitor_text) on a LibreOffice PDF: are positions and sizes usable?
S10 PowerShell language mode; WinRT raster worker inside a Job Object with a 1.5 GB limit; malformed-PDF corpus
    (truncated, garbage, zero-byte, deep nesting); WinRT-raster OCR timing at 300 dpi.
Usage: python -I spikes.py <out_dir>
"""
import ctypes, ctypes.wintypes as W, hashlib, json, os, shutil, statistics, subprocess, sys, tempfile, time
from pathlib import Path
import docx, pypdf

HERE = Path(__file__).resolve().parent
QUAL = HERE.parents[0] / "lo_qualify_20261015"
PRINT = HERE.parents[0] / "pdf_print_probe_20261015"
PROG = Path(os.environ["LOCALAPPDATA"]) / "URI" / "engines" / "libreoffice-26.8.1"
SOFFICE, LOPY = PROG / "program" / "soffice.com", PROG / "program" / "python.exe"
TESS = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
out = Path(sys.argv[1]).resolve(); shutil.rmtree(out, ignore_errors=True); out.mkdir(parents=True)
R = {"started": time.strftime("%Y-%m-%dT%H:%M:%S%z")}

# ---------------- ctypes Job Object (S3) ----------------
k32 = ctypes.WinDLL("kernel32", use_last_error=True)


class IO_COUNTERS(ctypes.Structure):
    _fields_ = [(n, ctypes.c_ulonglong) for n in ("r", "w", "o", "rb", "wb", "ob")]


class BASIC(ctypes.Structure):
    _fields_ = [("PerProcessUserTimeLimit", ctypes.c_longlong), ("PerJobUserTimeLimit", ctypes.c_longlong),
                ("LimitFlags", W.DWORD), ("MinimumWorkingSetSize", ctypes.c_size_t),
                ("MaximumWorkingSetSize", ctypes.c_size_t), ("ActiveProcessLimit", W.DWORD),
                ("Affinity", ctypes.c_size_t), ("PriorityClass", W.DWORD), ("SchedulingClass", W.DWORD)]


class EXT(ctypes.Structure):
    _fields_ = [("Basic", BASIC), ("Io", IO_COUNTERS), ("ProcessMemoryLimit", ctypes.c_size_t),
                ("JobMemoryLimit", ctypes.c_size_t), ("PeakProcessMemoryUsed", ctypes.c_size_t),
                ("PeakJobMemoryUsed", ctypes.c_size_t)]


KILL_ON_CLOSE, PROC_MEM = 0x2000, 0x100
k32.CreateJobObjectW.restype = W.HANDLE
k32.OpenProcess.restype = W.HANDLE


def make_job(mem_limit=None):
    job = k32.CreateJobObjectW(None, None)
    info = EXT(); info.Basic.LimitFlags = KILL_ON_CLOSE | (PROC_MEM if mem_limit else 0)
    if mem_limit:
        info.ProcessMemoryLimit = mem_limit
    assert k32.SetInformationJobObject(job, 9, ctypes.byref(info), ctypes.sizeof(info)), ctypes.get_last_error()
    return job


def assign(job, pid):
    h = k32.OpenProcess(0x1F0FFF, False, pid)
    ok = k32.AssignProcessToJobObject(job, h)
    k32.CloseHandle(h)
    return bool(ok)


def peak(job):
    info = EXT(); k32.QueryInformationJobObject(job, 9, ctypes.byref(info), ctypes.sizeof(info), None)
    return info.PeakProcessMemoryUsed


def alive(pid):
    h = k32.OpenProcess(0x1000, False, pid)
    if not h:
        return False
    code = W.DWORD(); k32.GetExitCodeProcess(h, ctypes.byref(code)); k32.CloseHandle(h)
    return code.value == 259


# ---------------- fixtures ----------------
fx = out / "fixtures"; fx.mkdir()
big = fx / "report25.docx"
d = docx.Document()
for sec in range(1, 26):
    d.add_heading(f"{sec}. Section {sec}", 1)
    for k in range(9):
        d.add_paragraph(f"Paragraph {k + 1} of section {sec}. The department reviewed the files and recorded the "
                        "decisions taken at each meeting, with the actions due before the next review. " * 3)
d.save(big)
long60 = fx / "tender60.docx"
d = docx.Document()
for pg in range(1, 61):
    d.add_heading(f"Tender clause {pg}", 2)
    d.add_paragraph(f"Clause {pg}. " + "The bidder shall comply with the conditions of the tender document. " * 30)
    if pg < 60:
        d.add_page_break()
d.save(long60)

# ---------------- listener helpers ----------------
XCU = (QUAL / "lo_qualify.py").read_text(encoding="utf-8")
XCU = XCU[XCU.index('HARDENED_XCU = """') + 18:]; XCU = XCU[:XCU.index('"""')]


def start_listener(profile_dir, fresh):
    if fresh:
        shutil.rmtree(profile_dir, ignore_errors=True)
    (profile_dir / "user").mkdir(parents=True, exist_ok=True)
    (profile_dir / "user" / "registrymodifications.xcu").write_text(XCU, encoding="utf-8")   # re-seed every start
    job = make_job()
    pipe = f"uri_s_{os.getpid()}_{int(time.time() * 1000)}"
    p = subprocess.Popen([str(SOFFICE), f"-env:UserInstallation={profile_dir.as_uri()}", "--headless", "--invisible",
                          "--norestore", "--nolockcheck", "--nodefault", "--nologo",
                          f"--accept=pipe,name={pipe};urp;StarOffice.ComponentContext"],
                         stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    assigned = assign(job, p.pid)
    return job, p, pipe, assigned


def run_bridge(script, pipe, odir, files, isolated, extra=()):
    # isolated: -I plus bridge_isolated.py, which adds only the LibreOffice program folder (S1 fallback)
    pre = ["-I", str(HERE / "bridge_isolated.py")] if isolated else []
    args = [str(LOPY)] + pre + [str(script), pipe, str(odir)] + list(extra) + [str(f) for f in files]
    t0 = time.perf_counter(); cp = subprocess.run(args, capture_output=True, text=True, timeout=900)
    return cp, round(time.perf_counter() - t0, 3)


# ---------------- S1 + S4 + S3 kill ----------------
plain = subprocess.run([str(LOPY), "-I", "-c", "import uno"], capture_output=True, text=True)
R["S1_plain_-I_import_uno"] = {"rc": plain.returncode, "stderr_tail": plain.stderr.strip()[-90:]}
prof = out / "lo-profile"
od = out / "s1"; od.mkdir()
job, lo, pipe, assigned = start_listener(prof, fresh=True)
t0 = time.perf_counter()
cp, _ = run_bridge(QUAL / "uno_convert.py", pipe, od, [big], isolated=True)
first_ready = round(time.perf_counter() - t0, 2)
R["S1_uno_under_-I"] = {"rc": cp.returncode, "rows": [json.loads(l) for l in cp.stdout.splitlines() if l.startswith("{")],
                        "stderr_tail": cp.stderr[-400:], "first_result_s_fresh_profile": first_ready}
R["S3_assigned_with_ctypes"] = assigned
kids_before = [int(x) for x in subprocess.run(["powershell", "-NoProfile", "-Command",
               f"(Get-CimInstance Win32_Process -Filter 'ParentProcessId={lo.pid}').ProcessId"],
               capture_output=True, text=True).stdout.split()]
k32.CloseHandle(job); time.sleep(2)
R["S3_job_close"] = {"children_before": len(kids_before), "alive_after": [p for p in [lo.pid] + kids_before if alive(p)]}

# S4: second start with the persistent (re-seeded) profile
job, lo, pipe, _ = start_listener(prof, fresh=False)
t0 = time.perf_counter()
cp, _ = run_bridge(QUAL / "uno_convert.py", pipe, od, [big], isolated=True)
R["S4_persistent_profile_first_result_s"] = round(time.perf_counter() - t0, 2)
R["S4_rows"] = [json.loads(l) for l in cp.stdout.splitlines() if l.startswith("{")]

# ---------------- S7 timings on the warm listener ----------------
s7 = out / "s7"; s7.mkdir(); rows = {"report25.docx": [], "tender60.docx": []}
for _ in range(5):
    for f in (big, long60):
        cp, wall = run_bridge(QUAL / "uno_convert.py", pipe, s7, [f], isolated=True)
        r = [json.loads(l) for l in cp.stdout.splitlines() if l.startswith("{")]
        rows[f.name].append({"bridge_reported_s": r[0]["seconds"] if r else None, "wall_incl_bridge_start_s": wall})
pages = {f.name: len(pypdf.PdfReader(str(f)).pages) for f in s7.glob("*.pdf")}
R["S7"] = {"pages": pages, "runs": rows,
           "median_bridge_s": {k: statistics.median([x["bridge_reported_s"] for x in v if x["bridge_reported_s"] is not None] or [-1]) for k, v in rows.items()},
           "median_wall_s": {k: statistics.median(x["wall_incl_bridge_start_s"] for x in v) for k, v in rows.items()}}
k32.CloseHandle(job); time.sleep(1)

# ---------------- S6 manifest cost ----------------
t0 = time.perf_counter(); n = size = 0; h = hashlib.sha256()
for f in sorted(PROG.rglob("*")):
    if f.is_file():
        b = f.read_bytes(); n += 1; size += len(b); h.update(hashlib.sha256(b).digest())
full = round(time.perf_counter() - t0, 2)
t0 = time.perf_counter(); m = 0
for f in sorted(PROG.rglob("*")):
    if f.is_file() and f.suffix.lower() in (".exe", ".dll", ".com", ".bin", ".pyd"):
        hashlib.sha256(f.read_bytes()); m += 1
exe = round(time.perf_counter() - t0, 2)
R["S6_manifest"] = {"files": n, "bytes": size, "full_s": full, "exe_dll_files": m, "exe_dll_s": exe}

# ---------------- S9 pypdf positions and sizes ----------------
pdf = next(s7.glob("report25.pdf"))
rd = pypdf.PdfReader(str(pdf)); pg = rd.pages[0]; box = pg.mediabox
items = []


def vis(text, cm, tm, font, size):
    if text.strip():
        x, y = tm[4] * cm[0] + cm[4], tm[5] * cm[3] + cm[5]
        items.append({"t": text.strip()[:24], "x": round(x, 1), "y": round(y, 1), "size": round(size * tm[0], 2)})


pg.extract_text(visitor_text=vis)
R["S9"] = {"page_box": [float(v) for v in box], "spans": len(items), "sample": items[:6],
           "x_range": [min(i["x"] for i in items), max(i["x"] for i in items)] if items else None,
           "y_range": [min(i["y"] for i in items), max(i["y"] for i in items)] if items else None,
           "sizes": sorted({i["size"] for i in items})[:8]}

# ---------------- S10 ----------------
lm = subprocess.run(["powershell", "-NoProfile", "-Command", "$ExecutionContext.SessionState.LanguageMode"],
                    capture_output=True, text=True).stdout.strip()
corpus = out / "corpus"; corpus.mkdir()
good = (s7 / "tender60.pdf").read_bytes()
(corpus / "truncated.pdf").write_bytes(good[: len(good) // 3])
(corpus / "garbage.pdf").write_bytes(os.urandom(200_000))
(corpus / "zero.pdf").write_bytes(b"")
deep = b"%PDF-1.4\n1 0 obj\n" + b"[" * 200000 + b"]" * 200000 + b"\nendobj\ntrailer<</Root 1 0 R>>\n%%EOF"
(corpus / "deep_nesting.pdf").write_bytes(deep)
shutil.copy(s7 / "tender60.pdf", corpus / "valid60.pdf")
ras = PRINT / "winrt_pdf_render.ps1"
res = {}
for f in sorted(corpus.iterdir()):
    o = out / ("s10_" + f.stem); o.mkdir()
    job = make_job(mem_limit=1536 * 2**20)
    t0 = time.perf_counter()
    p = subprocess.Popen(["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                          str(ras), str(f), str(o), "1240"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    assign(job, p.pid)
    try:
        so, se = p.communicate(timeout=120); rc = p.returncode
    except subprocess.TimeoutExpired:
        k32.CloseHandle(job); job = None; so, se, rc = "", "TIMEOUT", "timeout"
    pk = peak(job) if job else None
    if job:
        k32.CloseHandle(job)
    res[f.name] = {"rc": rc, "seconds": round(time.perf_counter() - t0, 2), "pngs": len(list(o.glob("*.png"))),
                   "peak_process_mb": round(pk / 2**20, 1) if pk else None, "stderr_head": (se or "")[:160]}
R["S10"] = {"powershell_language_mode": lm, "corpus": res}

# WinRT-raster OCR at 300 dpi (A4 width 2480 px) on the scanned fixture from the print probe output, if present
scan = Path(os.environ["LOCALAPPDATA"]) / "Temp" / "uri_pdf_print_probe" / "fixtures" / "scan.pdf"
if scan.exists():
    o = out / "s10_ocr"; o.mkdir()
    t0 = time.perf_counter()
    subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ras), str(scan), str(o), "2480"],
                   capture_output=True, text=True, timeout=120)
    rr = round(time.perf_counter() - t0, 2); ocr = []
    for png in sorted(o.glob("*.png")):
        t1 = time.perf_counter()
        txt = subprocess.run([str(TESS), str(png), "stdout"], capture_output=True, text=True, timeout=120).stdout
        ocr.append({"page": png.name, "ocr_ms": round((time.perf_counter() - t1) * 1000), "head": txt.strip()[:70]})
    R["S10_ocr_winrt_300dpi"] = {"raster_wall_s": rr, "pages": ocr}
else:
    R["S10_ocr_winrt_300dpi"] = "NOT RUN: scan fixture not present"

R["pymupdf_imported"] = "fitz" in sys.modules or "pymupdf" in sys.modules
R["finished"] = time.strftime("%Y-%m-%dT%H:%M:%S%z")
(out / "spikes_result.json").write_text(json.dumps(R, indent=1))
print(json.dumps(R, indent=1))
